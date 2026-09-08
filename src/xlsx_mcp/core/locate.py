"""core/locate.py: the grid location resolver (DESIGN Section 5).

KS4W's location-object resolver adapted to the grid domain. Every positional
tool takes a ``location`` object carrying EXACTLY ONE selector key; this module
resolves it against an open openpyxl workbook to a normalized rectangle
``{sheet, min_row, min_col, max_row, max_col}`` before any tool acts.

Selectors (Phase 2):

    {"cell": "B7"}                                   single A1 cell
    {"range": "A1:C10"}                              rectangular A1 range
    {"a1": "B7"} / {"a1": "A1:C10"}                  either, one A1 key
    {"cell": "B7", "sheet": "Q3"}                    sheet-qualified sibling
    {"r1c1": "R7C2"} / {"r1c1": "R1C1:R10C3"}        R1C1 addressing
    {"name": "SalesTotal"}                           defined name
    {"named_range": "SalesTotal"}                    alias of name
    {"table": "Sales", "column": "Amount",
     "part": "data"}                                 structured reference
    {"used_range": "Q3"}                             the sheet's TRUE used range
    {"region": {"near": "B7"}}                        the contiguous data island
    {"search": {"text": "Q3 Total", "sheet": "Q3",
                "occurrence": 1, "match_case": false,
                "match": "exact"}}                    content-match addressing
    {"anchor": "gv1:..."}                             the rectangle a
                                                      get_grid_view showed,
                                                      content-verified

Discipline inherited from KS4W (DESIGN Section 5.2): no tool acts on first
match. A search or name that resolves to more than one target with no
disambiguator refuses with AmbiguousTarget (AMBIGUOUS_LOCATION) carrying every
match. Zero matches refuse with TargetNotFound (NOT_FOUND) plus nearest-miss
hints. Out-of-grid addresses refuse with RangeOutOfBounds. The used_range
selector computes the TRUE used range (value-bearing bounds), which differs
from openpyxl's dimension (see true_used_range); the incumbent gets this wrong
(demand_mining D7, incumbent_study defect 11).

The anchor selector (reserved since Phase 2, landed with the consolidation
phase) is view-stable addressing: get_grid_view returns one self-contained
token for the rectangle it SHOWED (sheet + A1 bounds + a content
fingerprint), and {"anchor": token} resolves to that rectangle only while
its content is unchanged, refusing StaleAnchor (STALE_ANCHOR) otherwise.
Cells already have stable A1 addresses, so per-cell/per-row anchors would be
ceremony; the one thing A1 cannot express is "the region as it was when I
looked", and that is all the anchor adds. The token is stateless (no
server-side session table), so it survives server restarts and works across
processes; the fingerprint is computed over raw stored values (formulas as
strings, data_only=False loads), the value space every mutating tool
resolves against.
"""

from __future__ import annotations

import base64
import hashlib
import re
from dataclasses import dataclass, field
from typing import Any

from openpyxl.utils import (
    coordinate_to_tuple,
    get_column_letter,
    range_boundaries,
)

from .errors import (
    AmbiguousTarget,
    RangeOutOfBounds,
    StaleAnchor,
    TargetNotFound,
    XlMcpError,
)

SELECTORS = (
    "cell", "range", "a1", "r1c1", "name", "named_range",
    "table", "used_range", "region", "search", "anchor",
)
#: nothing reserved anymore; the anchor selector landed with the
#: consolidation phase. Kept (empty) for import stability.
RESERVED_SELECTORS: tuple[str, ...] = ()

#: anchor token version tag; a token from a future incompatible scheme
#: refuses as malformed rather than mis-resolving.
_ANCHOR_TAG = "gv1"

#: Fingerprint field shape (blake2b, digest_size=5 -> 10 hex chars). A token
#: whose fingerprint field is not this shape is MALFORMED, not stale, and says
#: so (adversarial round: a truncated token used to report STALE_ANCHOR, which
#: sent the caller looking for a content change that never happened).
_FP_RE = re.compile(r"[0-9a-f]{10}")

#: Ceiling on the rectangle an anchor token may name. A view can only ever
#: anchor what it showed (at most 200 x 100 cells), but the token is caller-
#: supplied text, and the fingerprint walks every cell of the rectangle it
#: names: a hand-edited 'A1:XFD1048576' token costs ~17.2 BILLION cell reads
#: (measured: 20-40 CPU-minutes) inside a single-threaded stdio server, i.e. a
#: one-call hang. The rectangle is bounds-checked BEFORE any fingerprint work.
_MAX_ANCHOR_CELLS = 200_000

MAX_ROW = 1_048_576
MAX_COL = 16_384
MAX_CELL_CHARS = 32_767


def qualify_a1(sheet: str, a1: str) -> str:
    """'Data!B7', quoting the sheet name the way Excel does when it is not a
    bare identifier. One implementation, so a response that has to say WHICH
    sheet an address is on spells it the same way the formula bar does."""
    if re.search(r"[^A-Za-z0-9_]", sheet):
        sheet = "'" + sheet.replace("'", "''") + "'"
    return f"{sheet}!{a1}"


_MAX_LISTED = 25
_TABLE_PARTS = ("all", "data", "headers", "totals")
_MATCH_MODES = ("exact", "contains")


# ------------------------------------------------------------------ result


@dataclass(frozen=True)
class ResolvedGrid:
    """A normalized rectangle a tool can act on. Rows and columns are 1-based
    and inclusive. ``empty`` marks an empty used_range (an explicit empty
    result, not an error)."""

    sheet: str
    min_row: int
    min_col: int
    max_row: int
    max_col: int
    selector: str
    matched: dict[str, Any] = field(default_factory=dict)
    empty: bool = False

    @property
    def is_single(self) -> bool:
        return (self.min_row == self.max_row and self.min_col == self.max_col)

    @property
    def a1(self) -> str:
        tl = f"{get_column_letter(self.min_col)}{self.min_row}"
        if self.is_single:
            return tl
        br = f"{get_column_letter(self.max_col)}{self.max_row}"
        return f"{tl}:{br}"

    @property
    def sheet_qualified_a1(self) -> str:
        return qualify_a1(self.sheet, self.a1)

    def as_dict(self) -> dict:
        return {
            "sheet": self.sheet,
            "min_row": self.min_row, "min_col": self.min_col,
            "max_row": self.max_row, "max_col": self.max_col,
            "a1": self.a1, "selector": self.selector,
            "empty": self.empty, "matched": self.matched,
        }


# ------------------------------------------------------------------ helpers


def _clip(text: str, limit: int = 60) -> str:
    return text if len(text) <= limit else text[: limit - 3] + "..."


def _check_bounds(min_col: int, min_row: int, max_col: int, max_row: int,
                  what: str) -> None:
    if min_row < 1 or min_col < 1:
        raise RangeOutOfBounds(
            f"{what}: rows and columns are 1-based and must be >= 1")
    if max_row > MAX_ROW or max_col > MAX_COL:
        raise RangeOutOfBounds(
            f"{what} exceeds the grid limits (max {MAX_COL} columns x "
            f"{MAX_ROW} rows)")
    if max_row < min_row or max_col < min_col:
        raise RangeOutOfBounds(
            f"{what} is inverted (bottom-right precedes top-left)")


def _sheet_for(wb, sheet: str | None, *, selector: str):
    if sheet is None:
        ws = wb.active
        if ws is None:
            raise TargetNotFound("workbook has no active sheet")
        return ws
    if sheet not in wb.sheetnames:
        close = [s for s in wb.sheetnames if s.lower() == sheet.lower()]
        hint = ""
        if close:
            hint = f"; sheet names are case-sensitive, did you mean {close[0]!r}?"
        raise TargetNotFound(
            f"no sheet named {sheet!r} (sheets: "
            f"{', '.join(repr(s) for s in wb.sheetnames[:_MAX_LISTED])})"
            + hint)
    return wb[sheet]


def _ambiguous(message: str, matches: list[dict]) -> AmbiguousTarget:
    exc = AmbiguousTarget(message)
    exc.matches = matches
    return exc


# ------------------------------------------------------ true used range


def true_used_range(ws) -> tuple[int, int, int, int] | None:
    """The TRUE used range: the bounding box of VALUE-BEARING cells, 1-based
    (min_row, min_col, max_row, max_col). None when the sheet holds no values.

    This differs from openpyxl's ``ws.calculate_dimension()`` / ``ws.max_row``
    / ``ws.max_column`` on purpose. openpyxl's dimension is the min/max over
    every INSTANTIATED cell, which includes cells that were only formatted or
    merely touched, and it also trusts the stored ``<dimension>`` tag other
    tools can author wrong. Either over-reports (a sheet formatted out to
    column Z but with data only in A:C reports A:Z) or, when the tag is stale,
    mis-reports. Filtering on ``value is not None`` yields the range a human
    would call "used" and the range a read tool should page over. This is the
    incumbent's defect 11 fix (validate_excel_range rejecting legitimate
    targets by checking against a wrong used range).
    """
    cells = getattr(ws, "_cells", None)
    min_r = min_c = None
    max_r = max_c = 0
    if cells is not None:
        for (r, c), cell in cells.items():
            if cell.value is None:
                continue
            if min_r is None or r < min_r:
                min_r = r
            if min_c is None or c < min_c:
                min_c = c
            if r > max_r:
                max_r = r
            if c > max_c:
                max_c = c
    else:  # pragma: no cover - read-only worksheets
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                r, c = cell.row, cell.column
                if min_r is None or r < min_r:
                    min_r = r
                if min_c is None or c < min_c:
                    min_c = c
                max_r = max(max_r, r)
                max_c = max(max_c, c)
    if min_r is None:
        return None
    return min_r, min_c, max_r, max_c


# ------------------------------------------------------------------ selectors


def _resolve_a1(wb, ref: Any, sheet: str | None, selector: str) -> ResolvedGrid:
    if not isinstance(ref, str) or not ref.strip():
        raise XlMcpError(
            f"{selector} selector takes an A1 string like 'B7' or 'A1:C10'")
    ref = ref.strip()
    # A sheet embedded in the ref ("Q3!A1:B2") overrides / supplies the sheet.
    embedded_sheet = None
    if "!" in ref:
        pre, ref = ref.rsplit("!", 1)
        embedded_sheet = pre.strip()
        if embedded_sheet.startswith("'") and embedded_sheet.endswith("'"):
            embedded_sheet = embedded_sheet[1:-1].replace("''", "'")
        if sheet is not None and embedded_sheet.lower() != sheet.lower():
            raise XlMcpError(
                f"{selector} selector: the ref names sheet {embedded_sheet!r} "
                f"but the sibling sheet key says {sheet!r}; pass one")
    use_sheet = embedded_sheet if embedded_sheet is not None else sheet
    ws = _sheet_for(wb, use_sheet, selector=selector)
    try:
        min_col, min_row, max_col, max_row = range_boundaries(ref.upper())
    except Exception:
        raise XlMcpError(
            f"{selector} selector: {ref!r} is not a valid A1 reference")
    if None in (min_col, min_row, max_col, max_row):
        raise XlMcpError(
            f"{selector} selector: {ref!r} is a whole-row or whole-column "
            "reference; give a bounded cell or range")
    _check_bounds(min_col, min_row, max_col, max_row, f"{selector} {ref!r}")
    return ResolvedGrid(
        ws.title, min_row, min_col, max_row, max_col, selector,
        matched={"a1": ref.upper()})


def _resolve_r1c1(wb, ref: Any, sheet: str | None) -> ResolvedGrid:
    if not isinstance(ref, str) or not ref.strip():
        raise XlMcpError("r1c1 selector takes 'R7C2' or 'R1C1:R10C3'")
    ws = _sheet_for(wb, sheet, selector="r1c1")
    parts = ref.strip().upper().split(":")

    def one(tok: str) -> tuple[int, int]:
        m = re.fullmatch(r"R(\d+)C(\d+)", tok)
        if not m:
            raise XlMcpError(
                f"r1c1 selector: {tok!r} is not absolute R#C# form")
        return int(m.group(1)), int(m.group(2))

    if len(parts) == 1:
        r, c = one(parts[0])
        min_row = max_row = r
        min_col = max_col = c
    elif len(parts) == 2:
        r1, c1 = one(parts[0])
        r2, c2 = one(parts[1])
        min_row, max_row = min(r1, r2), max(r1, r2)
        min_col, max_col = min(c1, c2), max(c1, c2)
    else:
        raise XlMcpError(f"r1c1 selector: {ref!r} has too many ':' parts")
    _check_bounds(min_col, min_row, max_col, max_row, f"r1c1 {ref!r}")
    return ResolvedGrid(
        ws.title, min_row, min_col, max_row, max_col, "r1c1",
        matched={"r1c1": ref.strip().upper(),
                 "a1": f"{get_column_letter(min_col)}{min_row}"})


def _all_defined_names(wb) -> list[tuple[str, str, Any]]:
    """Every defined name as (scope, name, definition). scope is 'workbook' or
    a sheet title."""
    out: list[tuple[str, str, Any]] = []
    for name, defn in wb.defined_names.items():
        out.append(("workbook", name, defn))
    for ws in wb.worksheets:
        dn = getattr(ws, "defined_names", None)
        if dn is None:
            continue
        for name, defn in dn.items():
            out.append((ws.title, name, defn))
    return out


def _resolve_name(wb, value: Any, sheet: str | None,
                  scope_hint: str | None) -> ResolvedGrid:
    if not isinstance(value, str) or not value.strip():
        raise XlMcpError("name selector takes a defined-name string")
    target = value.strip()
    hits = [
        (scope, name, defn)
        for scope, name, defn in _all_defined_names(wb)
        if name == target
    ]
    if scope_hint is not None:
        hits = [h for h in hits if h[0].lower() == scope_hint.lower()
                or (scope_hint.lower() == "workbook" and h[0] == "workbook")]
    if not hits:
        allnames = sorted({n for _s, n, _d in _all_defined_names(wb)})
        raise TargetNotFound(
            f"no defined name {target!r}; names: "
            + (", ".join(repr(n) for n in allnames[:_MAX_LISTED])
               or "none defined"))
    if len(hits) > 1:
        raise _ambiguous(
            f"defined name {target!r} exists in {len(hits)} scopes "
            "(a sheet-scoped name shadows a workbook-scoped one); pass a "
            "'scope' modifier (a sheet title or 'workbook') to disambiguate",
            [{"scope": s, "name": n, "value": getattr(d, "value", None)}
             for s, n, d in hits])
    scope, name, defn = hits[0]
    dests = list(getattr(defn, "destinations", []) or [])
    if not dests:
        val = getattr(defn, "value", None)
        raise TargetNotFound(
            f"defined name {name!r} does not resolve to a cell range "
            f"(its definition is {val!r}); it may be a constant or a formula")
    if len(dests) > 1:
        raise _ambiguous(
            f"defined name {name!r} refers to {len(dests)} disjoint ranges; "
            "address the specific range you want by cell/range",
            [{"sheet": s, "range": c} for s, c in dests[:_MAX_LISTED]])
    dest_sheet, coord = dests[0]
    ws = _sheet_for(wb, dest_sheet, selector="name")
    min_col, min_row, max_col, max_row = range_boundaries(coord.replace("$", ""))
    _check_bounds(min_col, min_row, max_col, max_row, f"name {name!r}")
    return ResolvedGrid(
        ws.title, min_row, min_col, max_row, max_col, "name",
        matched={"name": name, "scope": scope, "definition": coord})


def _find_table(wb, name: str):
    for ws in wb.worksheets:
        tables = getattr(ws, "tables", {})
        if name in tables:
            return ws, tables[name]
        for tname in list(tables):
            # index by name: TableList.items() yields ref STRINGS, not Table
            # objects, so the old items() loop returned a str here
            if tname.lower() == name.lower():
                return ws, tables[tname]
    return None, None


def _resolve_table(wb, name: Any, spec: dict) -> ResolvedGrid:
    if not isinstance(name, str) or not name.strip():
        raise XlMcpError("table selector takes a table (ListObject) name")
    ws, table = _find_table(wb, name.strip())
    if table is None:
        allnames = [t for w in wb.worksheets for t in getattr(w, "tables", {})]
        raise TargetNotFound(
            f"no table named {name!r}; tables: "
            + (", ".join(repr(t) for t in allnames[:_MAX_LISTED]) or "none"))
    part = spec.get("part", "all")
    if part not in _TABLE_PARTS:
        raise XlMcpError(
            f"table selector: part must be one of {_TABLE_PARTS}, got {part!r}")
    column = spec.get("column")
    min_col, min_row, max_col, max_row = range_boundaries(table.ref)
    header_rows = table.headerRowCount if table.headerRowCount is not None else 1
    totals_rows = table.totalsRowCount if table.totalsRowCount is not None else 0

    top = min_row
    bottom = max_row
    if part == "headers":
        if header_rows == 0:
            raise TargetNotFound(f"table {name!r} has no header row")
        bottom = min_row + header_rows - 1
    elif part == "data":
        top = min_row + header_rows
        bottom = max_row - totals_rows
        if bottom < top:
            raise TargetNotFound(f"table {name!r} has no data rows")
    elif part == "totals":
        if totals_rows == 0:
            raise TargetNotFound(f"table {name!r} has no totals row")
        top = max_row - totals_rows + 1

    left = min_col
    right = max_col
    if column is not None:
        if not isinstance(column, str):
            raise XlMcpError("table selector: 'column' must be a column name")
        names = [c.name for c in table.tableColumns]
        if not names:  # tableColumns not materialized: read the header row
            names = [ws.cell(min_row, min_col + i).value
                     for i in range(max_col - min_col + 1)]
        idx = None
        for i, cn in enumerate(names):
            if cn is not None and str(cn).strip() == column:
                idx = i
                break
        if idx is None:
            raise TargetNotFound(
                f"table {name!r} has no column {column!r}; columns: "
                + ", ".join(repr(str(c)) for c in names if c is not None))
        left = right = min_col + idx
    _check_bounds(left, top, right, bottom, f"table {name!r}")
    return ResolvedGrid(
        ws.title, top, left, bottom, right, "table",
        matched={"table": table.displayName or table.name, "part": part,
                 "column": column, "ref": table.ref})


def _resolve_used_range(wb, value: Any,
                        sibling_sheet: str | None = None) -> ResolvedGrid:
    sheet = sibling_sheet
    if isinstance(value, str) and value.strip():
        sheet = value.strip()
    elif value not in (True, None) and not isinstance(value, str):
        raise XlMcpError(
            'used_range selector takes a sheet name string, or true for the '
            'active sheet (a sibling "sheet" key also works)')
    ws = _sheet_for(wb, sheet, selector="used_range")
    bounds = true_used_range(ws)
    if bounds is None:
        return ResolvedGrid(
            ws.title, 1, 1, 1, 1, "used_range",
            matched={"note": "sheet holds no values"}, empty=True)
    min_r, min_c, max_r, max_c = bounds
    return ResolvedGrid(
        ws.title, min_r, min_c, max_r, max_c, "used_range",
        matched={"note": "true used range (value-bearing bounds)",
                 "openpyxl_dimension": ws.calculate_dimension()})


def _resolve_region(wb, value: Any, sheet: str | None) -> ResolvedGrid:
    if not isinstance(value, dict) or "near" not in value:
        raise XlMcpError(
            'region selector takes {"near": "B7"} (the contiguous data island '
            "around a cell)")
    near = value["near"]
    if not isinstance(near, str):
        raise XlMcpError("region selector: 'near' must be an A1 cell")
    use_sheet = value.get("sheet", sheet)
    ws = _sheet_for(wb, use_sheet, selector="region")
    try:
        r0, c0 = coordinate_to_tuple(near.upper())
    except Exception:
        raise XlMcpError(f"region selector: {near!r} is not a valid A1 cell")

    # Probe without ws.cell(): that call INSTANTIATES an empty cell object for
    # every miss, polluting the model (and bloating a later save) with the
    # whole probed frontier. _cells.get() is a pure lookup.
    cells = getattr(ws, "_cells", None)

    def has_value(r: int, c: int) -> bool:
        if r < 1 or c < 1 or r > MAX_ROW or c > MAX_COL:
            return False
        if cells is not None:
            cell = cells.get((r, c))
            return cell is not None and cell.value is not None
        return ws.cell(r, c).value is not None  # read-only sheets

    if not has_value(r0, c0):
        # An empty anchor: the region is just that cell (Excel selects the
        # single cell when the current region is empty).
        return ResolvedGrid(
            ws.title, r0, c0, r0, c0, "region",
            matched={"near": near.upper(), "note": "anchor cell is empty"})

    top, bottom, left, right = r0, r0, c0, c0
    changed = True
    while changed:
        changed = False
        # grow up
        if any(has_value(top - 1, c) for c in range(left, right + 1)):
            top -= 1
            changed = True
        # grow down
        if any(has_value(bottom + 1, c) for c in range(left, right + 1)):
            bottom += 1
            changed = True
        # grow left
        if any(has_value(r, left - 1) for r in range(top, bottom + 1)):
            left -= 1
            changed = True
        # grow right
        if any(has_value(r, right + 1) for r in range(top, bottom + 1)):
            right += 1
            changed = True
    _check_bounds(left, top, right, bottom, f"region near {near!r}")
    return ResolvedGrid(
        ws.title, top, left, bottom, right, "region",
        matched={"near": near.upper()})


def _resolve_search(wb, spec: Any, sheet: str | None) -> ResolvedGrid:
    if isinstance(spec, str):
        spec = {"text": spec}
    if not isinstance(spec, dict) or "text" not in spec:
        raise XlMcpError(
            'search selector takes {"text": "...", "occurrence": 1, '
            '"match_case": false, "match": "exact", "sheet": "Q3"}')
    text = spec["text"]
    if not isinstance(text, str) or text == "":
        raise XlMcpError("search selector: 'text' must be a non-empty string")
    match_case = bool(spec.get("match_case", False))
    mode = spec.get("match", "exact")
    if mode not in _MATCH_MODES:
        raise XlMcpError(
            f"search selector: match must be one of {_MATCH_MODES}, "
            f"got {mode!r}")
    occurrence = spec.get("occurrence")
    if occurrence is not None:
        if isinstance(occurrence, bool) or not isinstance(occurrence, int) \
                or occurrence < 1:
            raise XlMcpError(
                "search selector: occurrence is a 1-based integer")
    search_sheet = spec.get("sheet", sheet)
    if search_sheet is not None:
        sheets = [_sheet_for(wb, search_sheet, selector="search")]
    else:
        sheets = list(wb.worksheets)

    needle = text if match_case else text.lower()

    def matches(cell_text: str) -> bool:
        hay = cell_text if match_case else cell_text.lower()
        return needle in hay if mode == "contains" else hay == needle

    hits: list[dict] = []
    for ws in sheets:
        cells = getattr(ws, "_cells", None)
        items = cells.items() if cells is not None else (
            ((c.row, c.column), c) for row in ws.iter_rows() for c in row)
        for (r, c), cell in items:
            val = cell.value
            if val is None:
                continue
            if matches(str(val)):
                hits.append({
                    "sheet": ws.title, "row": r, "col": c,
                    "cell": f"{get_column_letter(c)}{r}",
                    "context": _clip(str(val)),
                })
    hits.sort(key=lambda h: (h["sheet"], h["row"], h["col"]))

    if not hits:
        parts = [f"search {text!r} matched no cell"]
        if search_sheet is None:
            parts.append("searched every sheet")
        else:
            parts.append(f"on sheet {search_sheet!r}")
        if match_case:
            ci = []
            for ws in sheets:
                for row in ws.iter_rows():
                    for cell in row:
                        if cell.value is not None and \
                                text.lower() in str(cell.value).lower():
                            ci.append(f"{ws.title}!{cell.coordinate}")
                            break
                    if ci:
                        break
            if ci:
                parts.append(
                    f"matches exist ignoring case (e.g. {ci[0]}); pass "
                    "match_case:false")
        if mode == "exact":
            parts.append("this was an exact match; try match:'contains'")
        raise TargetNotFound(". ".join(parts) + ".")

    if occurrence is not None:
        if occurrence > len(hits):
            raise TargetNotFound(
                f"search {text!r}: occurrence {occurrence} out of range, only "
                f"{len(hits)} match(es)")
        chosen = hits[occurrence - 1]
    elif len(hits) > 1:
        raise _ambiguous(
            f"search {text!r} matched {len(hits)} cells across "
            f"{len({h['sheet'] for h in hits})} sheet(s). Pass occurrence "
            "(1-based), a 'sheet', or address by cell/name.",
            [{"sheet": h["sheet"], "cell": h["cell"], "context": h["context"]}
             for h in hits[:_MAX_LISTED]])
    else:
        chosen = hits[0]
    ws = wb[chosen["sheet"]]
    r, c = chosen["row"], chosen["col"]
    # Merged-cell anchor semantics: only the top-left of a merge holds a value,
    # so a value hit already resolves to the anchor. Note it when merged.
    matched = {"cell": chosen["cell"], "context": chosen["context"],
               "occurrence": (occurrence or (hits.index(chosen) + 1))}
    for mr in getattr(ws, "merged_cells", []).ranges if \
            getattr(ws, "merged_cells", None) else []:
        if (mr.min_row, mr.min_col) == (r, c) and \
                (mr.max_row, mr.max_col) != (r, c):
            matched["merged_anchor"] = str(mr)
            break
    return ResolvedGrid(ws.title, r, c, r, c, "search", matched=matched)


# ------------------------------------------------------------------ anchors


def grid_fingerprint(ws, min_row: int, min_col: int, max_row: int,
                     max_col: int) -> str:
    """A short content fingerprint of a rectangle: raw stored values
    (formula strings and literals, the data_only=False value space) keyed by
    coordinate, so any value change, move, insert, or delete inside the
    rectangle changes the digest. Formatting-only changes do not; the anchor
    guards content, not cosmetics.

    The BOUNDS are hashed in as well, so a token cannot be widened by hand:
    before that, an A1:C4 token re-pointed at A1:Z2000 still verified (the
    added cells were empty and empty cells contribute nothing), which let an
    edit land on cells the view never showed."""
    h = hashlib.blake2b(digest_size=5)
    h.update(f"{min_row},{min_col},{max_row},{max_col}|".encode("ascii"))
    cells = getattr(ws, "_cells", None)
    for r in range(min_row, max_row + 1):
        for c in range(min_col, max_col + 1):
            if cells is not None:
                cell = cells.get((r, c))
                v = None if cell is None else cell.value
            else:  # pragma: no cover - read-only worksheets
                v = ws.cell(r, c).value
            if v is not None:
                h.update(f"{r},{c},{v!r};".encode("utf-8", "surrogatepass"))
    return h.hexdigest()


def make_anchor(ws, min_row: int, min_col: int, max_row: int,
                max_col: int) -> str:
    """Build the self-contained view anchor token get_grid_view returns:
    gv1:<base64url(sheet)>:<A1 bounds>:<fingerprint>. Stateless by design;
    resolving it re-derives the fingerprint from the live workbook."""
    tl = f"{get_column_letter(min_col)}{min_row}"
    br = f"{get_column_letter(max_col)}{max_row}"
    a1_ref = tl if (min_row, min_col) == (max_row, max_col) else f"{tl}:{br}"
    b64 = base64.urlsafe_b64encode(
        ws.title.encode("utf-8")).decode("ascii").rstrip("=")
    fp = grid_fingerprint(ws, min_row, min_col, max_row, max_col)
    return f"{_ANCHOR_TAG}:{b64}:{a1_ref}:{fp}"


def _resolve_anchor(wb, token: Any) -> ResolvedGrid:
    bad = XlMcpError(
        "anchor selector takes the token a get_grid_view result carries "
        "(it looks like 'gv1:...'); re-run get_grid_view for a fresh one")
    if not isinstance(token, str) or not token.startswith(_ANCHOR_TAG + ":"):
        raise bad
    try:
        _tag, b64, rest = token.split(":", 2)
        a1_ref, fp = rest.rsplit(":", 1)
        pad = "=" * (-len(b64) % 4)
        sheet = base64.urlsafe_b64decode(b64 + pad).decode("utf-8")
        if not sheet or not a1_ref or not _FP_RE.fullmatch(fp):
            raise ValueError
    except (ValueError, UnicodeDecodeError):
        raise bad
    if sheet not in wb.sheetnames:
        raise StaleAnchor(
            f"the anchor points at sheet {sheet!r}, which no longer exists; "
            "re-run get_grid_view and resend with fresh addressing")
    ws = wb[sheet]
    try:
        min_col, min_row, max_col, max_row = range_boundaries(a1_ref)
    except Exception:
        raise bad
    if None in (min_col, min_row, max_col, max_row):
        raise bad
    _check_bounds(min_col, min_row, max_col, max_row, f"anchor {a1_ref!r}")
    cells = (max_row - min_row + 1) * (max_col - min_col + 1)
    if cells > _MAX_ANCHOR_CELLS:
        raise RangeOutOfBounds(
            f"anchor {a1_ref!r} names {cells:,} cells, over the "
            f"{_MAX_ANCHOR_CELLS:,}-cell anchor ceiling; a get_grid_view "
            "anchor never covers more than the rectangle it showed, so this "
            "token was edited by hand. Re-run get_grid_view for a fresh one")
    current = grid_fingerprint(ws, min_row, min_col, max_row, max_col)
    if current != fp:
        raise StaleAnchor(
            f"anchor {a1_ref} on sheet {sheet!r} is stale: the region's "
            "content changed after the view was taken. Re-run get_grid_view "
            "and resend with the fresh anchor (or address by cell/range if "
            "the change was your own intended edit).")
    return ResolvedGrid(
        sheet, min_row, min_col, max_row, max_col, "anchor",
        matched={"a1": a1_ref, "fingerprint": fp,
                 "note": "content verified unchanged since the view"})


# ------------------------------------------------------------------ public


def resolve_location(wb, location: Any, *,
                     default_sheet: str | None = None) -> ResolvedGrid:
    """Resolve one location object against an open openpyxl workbook.

    location: a dict with EXACTLY ONE selector key from SELECTORS, plus an
        optional sibling 'sheet' (the active sheet when omitted) and, for the
        name selector, an optional 'scope'. Zero or multiple selector keys
        refuse (BAD_PARAMS). A convenience: a bare string is treated as
        {"a1": <string>}.

    Returns a ResolvedGrid rectangle (1-based, inclusive). Raises XlMcpError
    (BAD_PARAMS), TargetNotFound (NOT_FOUND), AmbiguousTarget with .matches
    (AMBIGUOUS_LOCATION), RangeOutOfBounds (RANGE_OUT_OF_BOUNDS), or
    StaleAnchor (STALE_ANCHOR).
    """
    if isinstance(location, str):
        location = {"a1": location}
    if not isinstance(location, dict):
        raise XlMcpError(
            "location must be an object with exactly one selector key from "
            f"{list(SELECTORS)} plus an optional 'sheet'")
    modifiers = {"sheet", "scope", "column", "part"}
    present = [k for k in SELECTORS if k in location]
    unknown = sorted(set(location) - set(SELECTORS) - modifiers)
    if unknown:
        raise XlMcpError(
            f"unknown location key(s) {unknown}; selectors are "
            f"{list(SELECTORS)}, plus optional 'sheet' (and 'scope' for name, "
            "'column'/'part' for table)")
    if len(present) != 1:
        raise XlMcpError(
            "location needs exactly one selector key, got "
            f"{len(present)} ({present if present else 'none'}); selectors "
            f"are {list(SELECTORS)}")
    sel = present[0]
    value = location[sel]
    sheet = location.get("sheet", default_sheet)

    if sel in ("cell", "range", "a1"):
        return _resolve_a1(wb, value, sheet, sel)
    if sel == "r1c1":
        return _resolve_r1c1(wb, value, sheet)
    if sel in ("name", "named_range"):
        return _resolve_name(wb, value, sheet, location.get("scope"))
    if sel == "table":
        return _resolve_table(wb, value, location)
    if sel == "used_range":
        return _resolve_used_range(wb, value, sheet)
    if sel == "region":
        return _resolve_region(wb, value, sheet)
    if sel == "anchor":
        return _resolve_anchor(wb, value)
    return _resolve_search(wb, value, sheet)


__all__ = [
    "ResolvedGrid", "resolve_location", "true_used_range",
    "SELECTORS", "RESERVED_SELECTORS", "make_anchor", "grid_fingerprint",
    "MAX_ROW", "MAX_COL", "MAX_CELL_CHARS", "qualify_a1",
]
