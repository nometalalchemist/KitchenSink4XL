"""core/package.py: WorkbookPackage, the central open/mutate/save object.

Every file-based tool opens a workbook through this object, mutates it, and
saves through it, so the safety core runs end to end on every write without
per-tool wiring (DESIGN Sections 2.4, 3; PLAN reuse ledger 1.2). It integrates:

  - the HAZARD SCAN (core.hazard) on open, and the routing decision on save:
    a mutation of a workbook with drop-risk parts REFUSES with HAZARD_REFUSED,
    naming the at-risk parts and the remedy (the COM route when it exists, or
    an explicit allow_loss override), unless the mutation goes through a loss-
    safe path (a clean workbook, or a .xlsm whose only hazard is VBA, preserved
    via keep_vba). Drop-risk includes content with no part of its own: a
    worksheet extLst carrying x14 conditional formatting, sparklines or a
    slicer list refuses on the same terms, since openpyxl writes the worksheet
    back without it.
  - BACKUP-BEFORE-MUTATION (core.safesave two-slot backups) captured against
    the pre-mutation state immediately before promotion, plus the per-file
    write lock held across the whole read-modify-verify-save cycle.
  - ATOMIC VALIDATED SAVE with VERIFY-AFTER-WRITE (core.verify): the mutation
    is written to a temp path and verified (structural re-open, DEFAULT-FAIL
    part-inventory diff vs the pre-write scan so ANY unexplained part loss
    refuses, the fragile-part checks, content read-back against intent)
    BEFORE the original is touched. Ops that deliberately remove parts at the
    model level (sheet delete, table to_range, comment delete) declare it via
    expect_removal() so the inventory diff does not mistake the removal for
    silent loss. A failed verify leaves the original unmodified and refuses
    with VALIDATION_FAILED. After promotion a second verify runs; if it
    fails, the pre-mutation backup is restored.
  - FORMULA-WRITE SAFETY: any formula written through the package is
    normalized (core.calc._xlfn/_xlpm shim) and the workbook is marked
    fullCalcOnLoad so the next Excel/LibreOffice open recalculates.
  - STRUCTURAL-EDIT INTEGRITY: insert/delete rows/columns rewrite every
    reference (core.refs) so the workbook stays coherent, then pass verify.
  - DIRTY/LOCK HANDLING: a workbook held open in Excel surfaces as
    WORKBOOK_LOCKED (the PermissionError path proven in Phase 1) rather than a
    raw traceback.

The object addresses cells through core.locate (the grid location resolver).
"""

from __future__ import annotations

import os
import uuid
from pathlib import Path
from typing import Any, Callable

from . import calc as _calc
from . import hazard as _hazard
from . import limits as _limits
from . import locate as _locate
from . import refs as _refs
from . import safesave as _safesave
from . import verify as _verify
from .errors import (
    HazardRefused,
    ValidationFailed,
    WorkbookCorrupt,
    WorkbookLocked,
    WorkbookNotFound,
    XlMcpError,
)
from .sandbox import check_path


class WorkbookPackage:
    """An open workbook, its hazard report, its pending edits, and the save
    pipeline. Construct with ``WorkbookPackage.open(path)``; mutate via the
    edit methods; persist with ``save()``. Hold one instance per mutation and
    let ``save()`` (which takes the write lock) serialize concurrent writers.
    """

    def __init__(self, path: str, *, com_manager: Any = None):
        self.path = os.fspath(path)
        self.com_manager = com_manager
        self._workbook = None            # lazy openpyxl load
        self._keep_vba = self.path.lower().endswith(".xlsm")
        self._hazard: _hazard.HazardReport | None = None
        self._pre_parts: list[str] = []
        self._pre_sizes: dict[str, int] = {}
        self._intended: dict[tuple[str, str], tuple[str, Any]] = {}
        self._formula_written = False
        self._structural: list[_refs.RefEdit] = []
        self._changed: dict[str, Any] = {}
        self._expected_removals: set[str] = set()
        self._expected_preserved: set[str] = set()
        #: (mtime_ns, size) of the file the model and the hazard scan were
        #: read from. save() re-stats under the write lock and refuses if
        #: another writer changed the file in between (see _check_unchanged).
        self._stamp: tuple[int, int] | None = None
        #: extensions openpyxl announced it was discarding at load time
        #: (x14 conditional formatting, sparkline and slicer-list extLst
        #: blocks); they live INSIDE surviving parts, so no part-level verify
        #: can see them. hazard.scan_path detects them up front by reading the
        #: worksheet extLst; this list is the independent second detector, and
        #: the gate refuses on either unless allow_loss covers it.
        self._dropped_extensions: list[str] = []
        #: hazard keys the caller was explicitly warned would be dropped when
        #: they passed allow_loss; the verify gate excuses these and nothing
        #: else (allow_loss is not a blanket amnesty for fragile parts).
        self._warned_loss_keys: set[str] = set()
        #: Worksheet parts whose ca="1" always-calculate flags and cached
        #: values the save put back (core.calc.restore_always_calc_cache).
        self._restored_calc_cache = 0

    # ------------------------------------------------------------- open

    @classmethod
    def open(cls, path: str, *, com_manager: Any = None) -> "WorkbookPackage":
        """Open a workbook for reading/mutation. Runs the hazard scan and
        captures the pre-write part list. Does not load the openpyxl model
        until a mutation or an explicit ``workbook`` access needs it."""
        p = check_path(path, "open workbook")
        if not os.path.exists(p):
            raise WorkbookNotFound(f"no such workbook: {p}")
        if os.path.isdir(p):
            raise XlMcpError(
                f"{p} is a directory, not a workbook file; pass the path of "
                "an .xlsx/.xlsm file")
        pkg = cls(p, com_manager=com_manager)
        rep = _hazard.scan_path(p)
        if rep.error is not None:
            # A lock is not corruption: the wrong word sends a panicked
            # owner toward a destructive "repair" (destroyer round, M-2).
            if rep.error_kind == "locked":
                raise WorkbookLocked(f"{p}: {rep.error}")
            raise WorkbookCorrupt(f"{p}: {rep.error}")
        pkg._hazard = rep
        pkg._pre_parts = list(rep.parts)
        pkg._pre_sizes = dict(rep.sizes)
        pkg._stamp = _file_stamp(p)
        return pkg

    def _check_unchanged(self) -> None:
        """Refuse if the file changed since this package read it.

        The write lock serializes THIS server's writers, but Excel, another
        tool, or a second process can land a write between open() (which took
        the hazard scan and the part inventory) and save() (which writes the
        model loaded from the older bytes). Without this check that write is
        silently clobbered, and the pre-write inventory the verify gate
        compares against describes a file that no longer exists. Adversarial
        round: reproduced end to end, the concurrent cell simply vanished."""
        if self._stamp is None:
            return
        now = _file_stamp(self.path)
        if now is None or now == self._stamp:
            return
        exc = XlMcpError(
            f"{Path(self.path).name} changed on disk after this operation "
            "read it (another process or Excel wrote to it). Refusing rather "
            "than overwriting that change with a stale copy; re-run the "
            "operation against the current file.")
        exc.code = "CONFLICT"
        raise exc

    @property
    def hazard(self) -> _hazard.HazardReport:
        if self._hazard is None:
            self._hazard = _hazard.scan_path(self.path)
        return self._hazard

    @property
    def workbook(self):
        """The openpyxl workbook, loaded on first access. keep_vba is set for
        .xlsm so the VBA project survives the round-trip (Phase 1 finding 4)."""
        if self._workbook is None:
            import zipfile as _zipfile
            import openpyxl
            import warnings as _warnings
            try:
                with _warnings.catch_warnings(record=True) as caught:
                    _warnings.simplefilter("always")
                    self._workbook = openpyxl.load_workbook(
                        self.path, keep_vba=self._keep_vba, data_only=False,
                        rich_text=False)
            except _zipfile.BadZipFile:
                # The promised conversion: a mutation-path load must refuse
                # as WorkbookCorrupt, never a raw BadZipFile (destroyer
                # round, L-4) -- unless a byte-range lock is the real cause.
                if _hazard.file_read_blocked(self.path):
                    raise WorkbookLocked(
                        f"{self.path}: {_hazard._LOCKED_ERROR}") from None
                raise WorkbookCorrupt(
                    f"{self.path} is not a valid .xlsx/.xlsm package (the "
                    "zip structure is damaged); the file was NOT modified"
                ) from None
            self._dropped_extensions = _extension_drops(caught)
        return self._workbook

    # ------------------------------------------------------------- address

    def resolve(self, location: Any, *,
                default_sheet: str | None = None) -> _locate.ResolvedGrid:
        """Resolve a location object to a normalized rectangle (core.locate)."""
        return _locate.resolve_location(
            self.workbook, location, default_sheet=default_sheet)

    def _ws(self, sheet: str | None):
        wb = self.workbook
        if sheet is None:
            return wb.active
        return wb[sheet]

    # ------------------------------------------------------------- mutate

    def set_cell(self, sheet: str | None, coord: str, value: Any) -> None:
        """Write one cell. A string starting with '=' is treated as a formula
        (normalized through the _xlfn shim); anything else is a literal."""
        if isinstance(value, str) and value.startswith("=") and value != "=":
            self.set_formula(sheet, coord, value)
            return
        if value == "=":
            # A bare "=" is not a formula; the file stores it as an inline
            # string either way, but the response used to REPORT it as
            # kind:"formula" (fresh-eyes round, L-1). Store and report it
            # honestly as text.
            ws = self._ws(sheet)
            cell = ws[coord]
            cell.value = value
            cell.data_type = "s"
            self._intended[(ws.title, coord.upper())] = ("value", value)
            return
        # Control characters, lone surrogates, and text past Excel's cell
        # ceiling: refuse in-envelope here rather than let openpyxl's
        # IllegalCharacterError escape the Section 7 shape (and rather than
        # accept a value the reply cannot encode, or one that only the
        # post-write verify would notice).
        _limits.check_cell_text(value, what=f"the value for {coord}")
        ws = self._ws(sheet)
        ws[coord] = value
        self._intended[(ws.title, coord.upper())] = ("value", value)

    def set_formula(self, sheet: str | None, coord: str, formula: str) -> None:
        """Write a formula, normalized so modern functions do not land as
        #NAME? and the workbook recalculates on open."""
        ws = self._ws(sheet)
        _limits.check_text_storable(formula, what=f"the formula for {coord}")
        normalized, _prefixed = _calc.normalize_formula(formula)
        ws[coord] = normalized
        self._intended[(ws.title, coord.upper())] = ("formula", normalized)
        self._formula_written = True

    def modify_structure(self, edit: _refs.RefEdit) -> _refs.RewriteReport:
        """Insert or delete rows/columns and rewrite every reference so the
        workbook stays coherent (core.refs). openpyxl shifts the cell grid;
        core.refs repairs formulas, names, conditional formats, data
        validations, table refs, and merged ranges. The change is verified on
        save like any other mutation."""
        wb = self.workbook
        ws = wb[edit.sheet]
        if edit.kind == _refs.INSERT_ROWS:
            ws.insert_rows(edit.index, edit.count)
        elif edit.kind == _refs.DELETE_ROWS:
            ws.delete_rows(edit.index, edit.count)
        elif edit.kind == _refs.INSERT_COLS:
            ws.insert_cols(edit.index, edit.count)
        elif edit.kind == _refs.DELETE_COLS:
            ws.delete_cols(edit.index, edit.count)
        elif edit.kind == _refs.MOVE:
            raise ValueError(
                "modify_structure handles insert/delete; MOVE is a copy_range "
                "concern that rewrites refs directly via core.refs")
        report = _refs.rewrite_workbook(wb, edit)
        self._structural.append(edit)
        if report.formulas:
            self._formula_written = True
        # Structural edits move cells, so the tracked intents no longer point
        # at the right coordinates; content read-back is by report, not cells.
        self._intended.clear()
        self._changed.setdefault("structural", []).append(
            {"kind": edit.kind, "sheet": edit.sheet, "index": edit.index,
             "count": edit.count, "rewrites": report.as_dict()})
        return report

    def expect_removal(self, *prefixes: str) -> None:
        """Register part-name prefixes an op DELIBERATELY removes at the model
        level (sheet delete, table to_range, comment delete), so the default-
        fail inventory check in verify-after-write does not flag the removal
        as silent loss. Scoped to the next successful save only; cleared with
        the other pending-edit state. This is the narrow, per-operation escape
        hatch; nothing else excuses a non-fragile part from the inventory
        diff."""
        self._expected_removals.update(p.lower() for p in prefixes)

    def expect_preserved(self, *hazard_keys: str) -> None:
        """Register hazard KEYS (core.hazard spec keys, e.g. 'media') whose
        SEV_DROPS gate refusal this op downgrades to a warning, because the
        op's own save path has been OBSERVED to preserve those parts and the
        default-fail part checks in verify-after-write still refuse the save
        if they are in fact lost. The sanctioned user: ops/objects.py
        manage_image, whose Pillow-backed load round-trips picture drawings
        and xl/media/ intact (observed on this machine for Excel-authored
        and openpyxl-authored images alike; the op verifies the picture-only
        precondition before registering). This NEVER weakens verification:
        part_loss_check runs with allow_loss=False for these keys, so an
        actual loss still fails the write and the original stays untouched.
        Scoped to the next successful save; cleared with pending state."""
        self._expected_preserved.update(k.lower() for k in hazard_keys)

    # ------------------------------------------------------------- save

    def _lock_state(self) -> tuple[bool, bool]:
        """(lockfile_present, writable). The ~$ owner lockfile is a SIGNAL,
        not authority: a stale one survives an Excel crash (com_ground_truth
        exp 4 caveat, verified in the COM-tier gate), so the write-probe
        decides. A real Excel hold denies the r+b open; a stale lockfile
        leaves the file writable."""
        p = Path(self.path)
        owner = p.parent / ("~$" + p.name)
        present = owner.exists()
        writable = True
        try:
            fh = open(p, "r+b")
            fh.close()
        except OSError:
            writable = False
        return present, writable

    def _extension_warnings(self) -> list[str]:
        """The announced-loss line for in-part extension blocks openpyxl
        discarded at load. Reported only on the allow_loss path, because
        without allow_loss the gate now REFUSES instead (see _hazard_gate).

        These blocks live INSIDE xl/worksheets/sheetN.xml, which survives the
        save at full size, so no part-level check can catch them: x14
        conditional formatting (every modern data bar and icon set),
        sparkline groups, slicer lists. The hazard scan reads the worksheet
        extLst and flags them up front; openpyxl's own load-time
        announcement, captured here, is the independent second detector and
        the one that names the exact extensions it threw away."""
        if not self._dropped_extensions:
            return []
        return ["openpyxl discarded in-part extension blocks on load, so "
                "they are NOT in the saved file: "
                + "; ".join(self._dropped_extensions)
                + ". These live inside the worksheet part (x14 conditional "
                  "formatting, sparklines, slicer lists), so no part-level "
                  "check can catch them; use the com pack to edit this "
                  "workbook with full fidelity."]

    def _refuse_extension_loss(self) -> None:
        """Backstop refusal for extension blocks openpyxl announced dropping
        that the hazard scan did not flag.

        The scan is the primary detector and normally fires first, so this
        path is reached only when the two disagree: a worksheet extLst the
        scan could not read, or an extension openpyxl drops from somewhere
        the scan does not look. Drop-class content must not go through on a
        detector disagreement, so the announcement alone refuses."""
        exc = HazardRefused(
            "openpyxl discarded in-part extension blocks when it loaded this "
            "workbook, so a file-based save would write it back without them: "
            + "; ".join(self._dropped_extensions)
            + ". These live inside the worksheet part (x14 conditional "
              "formatting, sparklines, slicer lists), so no part-level check "
              "can catch the loss after the fact. Refusing the mutation "
              "rather than destroy them. Remedy: use the com pack (Excel "
              "saves with everything intact), or pass allow_loss:true to "
              "proceed with a backup and accept the loss.")
        exc.detail = {"dropped_extensions": list(self._dropped_extensions),
                      "routes": ["enable COM (com pack)",
                                 "allow_loss:true with backup"]}
        raise exc

    def _hazard_gate(self, allow_loss: bool) -> list[str]:
        """Apply the routing decision. Returns advisory warnings; raises
        HazardRefused when a drop-risk mutation has no loss-safe path.

        In-part extension blocks are drop-class under the ratified package
        policy (drop-risk refuses without allow_loss, degrade-risk warns and
        proceeds), so they raise here like any other SEV_DROPS hazard. They
        used to warn and proceed, which announced the loss but still let a
        data bar disappear from a workbook the user never agreed to sacrifice.
        With allow_loss the announcement survives as a warning."""
        rep = self.hazard
        warnings: list[str] = []
        if rep.clean:
            if self._dropped_extensions:
                # The scan saw nothing but openpyxl still threw extensions
                # away; the backstop governs.
                if not allow_loss:
                    self._refuse_extension_loss()
                warnings.extend(self._extension_warnings())
            return warnings
        drop_keys = list(rep.lossy_keys)  # SEV_DROPS
        degradable = [h.label for h in rep.hazards
                      if h.severity == _hazard.SEV_DEGRADES]
        degrade_keys = [h.key for h in rep.hazards
                        if h.severity == _hazard.SEV_DEGRADES]
        conditional = [h for h in rep.hazards
                       if h.severity == _hazard.SEV_CONDITIONAL]

        # A conditional hazard (VBA) survives ONLY because the openpyxl load
        # passes keep_vba, and keep_vba is set from the .xlsm extension. A
        # vbaProject.bin inside a plain .xlsx (renamed or mis-authored file)
        # gets NO keep_vba, so the save drops it: treat it as a hard drop,
        # never claim it is preserved.
        if conditional and not self._keep_vba:
            drop_keys.extend(h.key for h in conditional)
            conditional = []

        # Keys the op registered via expect_preserved: the gate downgrades
        # them to a warning because the op's save path preserves them and
        # verify-after-write still hard-fails if they are actually lost.
        preserved = [k for k in drop_keys if k in self._expected_preserved]
        if preserved:
            drop_keys = [k for k in drop_keys if k not in preserved]
            labels = [next(h.label for h in rep.hazards if h.key == k)
                      for k in preserved]
            warnings.append(
                ", ".join(labels) + " are expected to survive this "
                "operation's save path (observed round-trip); "
                "verify-after-write refuses the save if they are lost")

        if drop_keys and not allow_loss:
            labels = [next(h.label for h in rep.hazards if h.key == k)
                      for k in drop_keys]
            parts = [p for h in rep.hazards if h.key in drop_keys
                     for p in h.parts]
            exc = HazardRefused(
                "this workbook holds " + ", ".join(labels)
                + " that a file-based (openpyxl) save drops silently. Refusing "
                "the mutation rather than destroy them. Remedy: enable the COM "
                "route once available (Excel saves with everything intact), or "
                "pass allow_loss:true to proceed with a backup and accept the "
                "loss.")
            exc.detail = {"parts": parts, "labels": labels,
                          "routes": ["enable COM (com pack)",
                                     "allow_loss:true with backup"]}
            raise exc
        if allow_loss and (drop_keys or degrade_keys):
            # Record exactly which hazard families the caller was told they
            # were sacrificing; verify-after-write excuses those and only
            # those (see verify.part_loss_check loss_keys). DEGRADE-warned
            # families are included: "may degrade" concretely means chart
            # styling sub-parts (colors1.xml, style1.xml) and anchors can
            # vanish, and scoping the amnesty to DROP families alone made
            # the refusal's advertised remedy a dead end on any file that
            # held both classes, while a file with ONLY degrade families
            # fell through to the blanket-None amnesty: same loss, opposite
            # outcomes depending on unrelated content (destroyer round,
            # M-3). One consistent rule now: allow_loss excuses exactly the
            # families the gate warned about, of either class.
            self._warned_loss_keys = set(drop_keys) | set(degrade_keys)
        if drop_keys and allow_loss:
            labels = [next(h.label for h in rep.hazards if h.key == k)
                      for k in drop_keys]
            warnings.append(
                "allow_loss: proceeding despite " + ", ".join(labels)
                + " which the file-based save drops; the backup is the safety "
                "net")
        # The in-part extension backstop, same rule as the clean path. Without
        # allow_loss, reaching this line means drop_keys was empty (a
        # populated one already raised above), so an announcement the scan
        # missed is the only thing standing between the caller and a lost data
        # bar: refuse on it. With allow_loss it becomes the itemized warning.
        if self._dropped_extensions:
            if not allow_loss:
                self._refuse_extension_loss()
            warnings.extend(self._extension_warnings())
        if degradable:
            warnings.append(
                ", ".join(degradable) + " are re-serialized through openpyxl's "
                "model; sub-features the model does not know may degrade"
                + (" (with allow_loss, sub-parts these families lose on the "
                   "round-trip, such as Excel-authored chart styling parts, "
                   "count as accepted losses)" if allow_loss else ""))
        if conditional:
            warnings.append(
                "VBA project preserved via keep_vba; note that form and "
                "ActiveX controls still break on a file-based edit")
        return warnings

    def _default_saver(self, tmp: str) -> None:
        wb = self.workbook
        if self._formula_written:
            try:
                wb.calculation.fullCalcOnLoad = True
            except Exception:
                pass
        wb.save(tmp)
        # openpyxl leaves an EMPTY <v></v> on every formula cell it writes.
        # Excel reads it as an error value, which a normal recalculation
        # overwrites but an ITERATIVE one seeds from, leaving a circular
        # formula stuck at #VALUE! (numbers-safety gate). Strip it before the
        # verify gate sees the package.
        try:
            _calc.strip_empty_cached_values(tmp)
        except Exception:  # noqa: BLE001
            pass  # cosmetic-plus: never fail a save over this
        # The ADJACENT case: openpyxl also drops the always-calculate flag
        # and any REAL cached value from every formula cell. Harmless under
        # normal calculation (fullCalcOnLoad fills it back in), fatal under
        # ITERATIVE calculation, where the iteration seeds from the current
        # value and a seedless circular formula sticks on #VALUE! forever.
        # Untouched formulas get their ca="1" and their <v> back; an edited
        # one never does (insane round, M-4).
        try:
            self._restored_calc_cache = _calc.restore_always_calc_cache(
                self.path, tmp)
        except Exception:  # noqa: BLE001
            self._restored_calc_cache = 0

    def _run_verify(self, target: str, allow_loss: bool) -> _verify.VerifyResult:
        """Run the verify gate and NEVER let it raise.

        Every check inside verify reads a package the server just produced,
        and a package that is broken in the right way makes those readers
        throw rather than return a verdict (adversarial round: a dangling
        chart relationship made content_readback's openpyxl load raise
        KeyError straight out of save). A raising verify is worse than a
        failing one: in the POST-promote position the exception skips the
        restore-from-backup branch entirely and leaves the broken file in
        place. An exception is therefore a verification FAILURE, reported as
        one."""
        try:
            return _verify.verify_after_write(
                target, pre_parts=self._pre_parts, intended=self._intended,
                allow_loss=allow_loss, pre_sizes=self._pre_sizes,
                expected_removals=frozenset(self._expected_removals),
                loss_keys=self._warned_loss_keys or None)
        except Exception as exc:  # noqa: BLE001
            return _verify.VerifyResult(
                ok=False,
                reasons=[f"the verification pass could not read the produced "
                         f"package ({type(exc).__name__}: {exc}); treating "
                         "that as a failed verify"])

    def _restore_from_backup(self) -> bool:
        import shutil
        slot = _safesave.slot_dir(self.path) / _safesave.PREV_SLOT
        if slot.exists():
            shutil.copy2(str(slot), self.path)
            return True
        return False

    def save(self, *, allow_loss: bool = False, backup: bool = True,
             saver: Callable[[str], None] | None = None,
             verify_com: bool | None = None) -> dict:
        """Persist pending mutations through the full safety pipeline and
        return the mutation-success envelope. Raises HazardRefused,
        ValidationFailed, or WorkbookLocked (all mapped to closed codes by the
        envelope). On any refusal the original file is left as it was (a failed
        pre-promote verify never touches it; a failed post-promote verify
        restores it from the backup).

        verify_com is the DEEP, authoritative check: after promotion the file
        is opened in a private hidden Excel worker and a repair-free open is
        required. It is the third layer of the "Excel will refuse this" gate
        (core.limits is the first, honest reporting the second), and it stays
        OPT-IN because it costs a COM round trip and needs Excel installed.
        None means "take the default", which is False unless KS4XL_VERIFY_COM
        is set."""
        if verify_com is None:
            verify_com = com_verify_default()
        path = self.path
        with _safesave.write_lock(path):
            self._check_unchanged()
            lock_present, writable = self._lock_state()
            if not writable:
                raise WorkbookLocked(
                    f"{Path(path).name} is open in Excel (or another process "
                    "holds it); close it, or use a live/COM route, then retry")
            warnings = self._hazard_gate(allow_loss)
            if lock_present:
                warnings.append(
                    f"a stale owner lockfile (~${Path(path).name}) is "
                    "present but the file is writable; a prior Excel "
                    "session likely crashed. Proceeding; the ~$ file can "
                    "be deleted safely")

            _sweep_stale_write_tmps(os.path.dirname(os.path.abspath(path)))
            tmp = os.path.join(
                os.path.dirname(os.path.abspath(path)) or ".",
                f".ks4xl-write-{uuid.uuid4().hex}{Path(path).suffix}")
            try:
                (saver or self._default_saver)(tmp)
            except PermissionError as exc:
                _silent_remove(tmp)
                raise WorkbookLocked(
                    f"{Path(path).name}: cannot write (it may be open in "
                    f"Excel). {exc}")
            except OSError as exc:
                # Disk-full and friends: the refusal is environmental, not
                # the caller's parameters, and the half-written temp must
                # not linger (destroyer round, L-2).
                _silent_remove(tmp)
                raise _os_write_error(Path(path).name, exc)
            except BaseException:
                _silent_remove(tmp)
                raise

            pre = self._run_verify(tmp, allow_loss)
            if not pre.ok:
                _silent_remove(tmp)
                raise ValidationFailed(
                    "verify-after-write failed; the original file was NOT "
                    "modified. " + "; ".join(pre.reasons))

            # Two-phase backup rotation: STAGE the pre-mutation content now,
            # COMMIT it onto the slots only after a successful promote. A
            # promote that fails must leave the slots untouched; rotating
            # first burned the prev undo of the previous successful mutation
            # even though nothing was written (destroyer round, M-1).
            backup_slot = None
            ticket = None
            if backup:
                ticket = _safesave.prepare_rotation(path)
            try:
                _safesave.replace_with_retry(tmp, path)
            except PermissionError as exc:
                if ticket is not None:
                    ticket.abort()
                _silent_remove(tmp)
                raise WorkbookLocked(
                    f"{Path(path).name}: cannot replace the file (it may be "
                    f"open in Excel). {exc}")
            except OSError as exc:
                if ticket is not None:
                    ticket.abort()
                _silent_remove(tmp)
                raise _os_write_error(Path(path).name, exc)
            except BaseException:
                if ticket is not None:
                    ticket.abort()
                _silent_remove(tmp)
                raise
            if ticket is not None:
                try:
                    ticket.commit()
                    backup_slot = "prev"
                except OSError as exc:
                    ticket.abort()
                    warnings.append(
                        "the save succeeded but the backup slot rotation "
                        f"failed ({type(exc).__name__}: {exc}); the prev "
                        "slot still holds the state before the PREVIOUS "
                        "mutation, not this one")

            post = self._run_verify(path, allow_loss)
            if not post.ok:
                restored = False
                if backup:
                    restored = self._restore_from_backup()
                raise ValidationFailed(
                    "post-promote verify failed"
                    + (" and the file was restored from the backup"
                       if restored else "")
                    + ". " + "; ".join(post.reasons))

            verified_com = None
            if verify_com:
                verified_com = self._verify_com_after_promote(
                    path, backup=backup, warnings=warnings)

            # The on-disk package changed; the cached model and scan are stale.
            self._workbook = None
            self._hazard = _hazard.scan_path(path)
            self._pre_parts = list(self._hazard.parts)
            self._pre_sizes = dict(self._hazard.sizes)
            changed = dict(self._changed)
            if self._intended:
                changed["cells"] = [
                    {"sheet": s, "cell": c, "kind": k}
                    for (s, c), (k, _v) in self._intended.items()]
            self._intended = {}
            self._formula_written = False
            self._structural = []
            self._changed = {}
            self._expected_removals = set()
            self._expected_preserved = set()
            self._warned_loss_keys = set()
            self._dropped_extensions = []
            # The file we just wrote is now the baseline a second save on this
            # package must compare against.
            self._stamp = _file_stamp(path)
            result = {
                "ok": True,
                "file": path,
                "changed": changed,
                "saved": True,
                "backup": backup_slot,
                "verified": True,
                "warnings": warnings,
            }
            if verified_com is not None:
                result["verified_com"] = verified_com
            return result

    def _verify_com_after_promote(self, path: str, *, backup: bool,
                                  warnings: list[str]) -> bool | None:
        """The optional DEEP verification (verify_com:true): after promotion,
        open the produced file in a private hidden Excel worker and require a
        clean, repair-free open (com.session.opens_clean). On failure the
        backup is restored and the save refuses. When COM is unavailable the
        save stands on the file-tier verify and says so in warnings (None)."""
        from ..com import session as _com_session
        ok, why = _com_session.com_available()
        if not ok:
            warnings.append(
                f"verify_com was requested but Excel/COM is unavailable "
                f"({why}); the save stands on the file-tier verification "
                "only")
            return None
        try:
            res = _com_session.opens_clean(path)
        except Exception as exc:  # noqa: BLE001
            warnings.append(
                "verify_com could not run (" + f"{type(exc).__name__}: {exc}"
                + "); the save stands on the file-tier verification only")
            return None
        if not res.get("opens_clean"):
            restored = self._restore_from_backup() if backup else False
            raise ValidationFailed(
                "deep verification failed: Excel refused the produced file "
                "or demanded a repair (" + str(res.get("excel_says", ""))
                + ")" + (" and the file was restored from the backup"
                         if restored else ""))
        return True

    def close(self) -> None:
        if self._workbook is not None:
            try:
                self._workbook.close()
            except Exception:
                pass
            self._workbook = None

    def __enter__(self) -> "WorkbookPackage":
        return self

    def __exit__(self, *exc) -> None:
        self.close()


#: Environment switch making the deep Excel verification the DEFAULT for
#: every save on this server, rather than a per-call opt-in.
#:
#: The insane round's headline was that verify-after-write has no notion of
#: "Excel will refuse this": it proves the produced package is valid OOXML
#: that reads back as intended, and a file Excel throws out can be exactly
#: that. core.limits closes the known routes cheaply, but only Excel can
#: answer for the unknown ones, and com_validate_opens_clean answers
#: correctly (it got all ten of the round's unopenable files right). This is
#: the switch that puts it on the write path. It stays OFF by default because
#: it costs a COM round trip per save and needs Excel installed; a headless
#: CI run and a Linux install must keep working. Per-call verify_com:true
#: always wins over the default, and since the live COM stress round (M-4) it
#: is actually reachable: every mutating tool takes verify_com, so a single
#: risky save can ask for Excel's verdict without turning it on server-wide.
#: The claim was true of this API and false of the surface for a whole wave.
VERIFY_COM_ENV = "KS4XL_VERIFY_COM"


def com_verify_default() -> bool:
    """Whether saves deep-verify through Excel unless told otherwise."""
    return os.environ.get(VERIFY_COM_ENV, "").strip().lower() in (
        "1", "true", "yes", "on")


def _file_stamp(path: str) -> tuple[int, int] | None:
    """(mtime_ns, size) identity of the file on disk, or None if it is gone."""
    try:
        st = os.stat(path)
    except OSError:
        return None
    return (st.st_mtime_ns, st.st_size)


#: openpyxl's own wording when it discards an extLst block it cannot model
#: ("Conditional Formatting extension is not supported and will be removed",
#: "Slicer List extension ...", "Unknown extension ...").
_EXT_DROP_MARKER = "is not supported and will be removed"


def _extension_drops(caught) -> list[str]:
    """De-duplicated extension-drop notices from a captured warning list."""
    out: list[str] = []
    for w in caught or ():
        text = str(getattr(w, "message", w))
        if _EXT_DROP_MARKER in text and text not in out:
            out.append(text)
    return out


def _silent_remove(path: str) -> None:
    try:
        os.remove(path)
    except OSError:
        pass


def _os_write_error(name: str, exc: OSError) -> XlMcpError:
    """An honest refusal for OS-level write failures. Disk-full used to map
    to BAD_PARAMS ("the caller's mistake") with the raw errno text
    (destroyer round, L-2); it is the environment's state, so it carries
    CONFLICT and says so in plain words. The original file is intact in
    every case: the failure happened on the temp or at the promote, never
    mid-way through the workbook's own bytes."""
    err = XlMcpError(
        (f"{name}: the disk is full (no space left to write the new "
         "version). The original file is intact and was NOT modified; free "
         "some disk space and retry.")
        if getattr(exc, "errno", None) == 28 else
        (f"{name}: the operating system refused the write "
         f"({type(exc).__name__}: {exc}). The original file is intact and "
         "was NOT modified."))
    err.code = "CONFLICT"
    return err


#: A .ks4xl-write-* temp older than this is an orphan from a killed process
#: (no mutation legitimately runs this long; the write lock's stale window
#: is 10 minutes) and is swept on the next save in the same folder. Kill
#: storms left these forever; the write.lock self-heals, the temps did not
#: (destroyer round, L-2).
_STALE_TMP_SECONDS = 60 * 60


def _sweep_stale_write_tmps(directory: str) -> None:
    """Best-effort janitor for aged .ks4xl-write-* temp files. Runs under
    the write lock on the save path; a FRESH temp (a concurrent writer's
    live work) is never touched."""
    try:
        now = __import__("time").time()
        with os.scandir(directory or ".") as it:
            for entry in it:
                if not entry.name.startswith(".ks4xl-write-"):
                    continue
                try:
                    if not entry.is_file():
                        continue
                    if now - entry.stat().st_mtime > _STALE_TMP_SECONDS:
                        os.remove(entry.path)
                except OSError:
                    continue
    except OSError:
        pass


__all__ = ["WorkbookPackage", "com_verify_default", "VERIFY_COM_ENV"]
