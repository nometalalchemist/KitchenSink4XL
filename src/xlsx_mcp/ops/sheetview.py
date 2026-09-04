"""ops/sheetview.py: set_view, the sheet-view settings tool (lite).

Freeze panes, split panes, gridline and heading visibility, zoom, the
active cell / selection, and the sheet tab color, all stored in the sheet
view and sheet properties parts openpyxl round-trips natively. View
settings are workbook state (what a user sees on open), so the mutation
runs through WorkbookPackage like any other write: hazard gate, backup,
atomic verified save.

Freeze and split are mutually exclusive on a sheet's pane (Excel's own
rule); asking for both in one call refuses. Split positions are taken in
POINTS (1/72 inch) and stored in the file's native twentieths of a point.
"""

from __future__ import annotations

import re
from typing import Any

from openpyxl.utils.cell import range_boundaries
from openpyxl.worksheet.views import Pane, Selection

from ..core.errors import XlMcpError
from ..core.package import WorkbookPackage

_CELL_RE = re.compile(r"^[A-Za-z]{1,3}[1-9][0-9]{0,6}$")
_HEX_RE = re.compile(r"^#?([0-9A-Fa-f]{6})$")


def _check_range(text: str, what: str) -> tuple[int, int, int, int]:
    try:
        return range_boundaries(text.upper())
    except Exception:  # noqa: BLE001
        raise XlMcpError(
            f"{what} must be an A1 cell or range, got {text!r}") from None


def set_view(path: str, sheet: str | None = None,
             freeze: str | None = None, split: dict | None = None,
             gridlines: bool | None = None, headings: bool | None = None,
             zoom: int | None = None, selection: str | None = None,
             tab_color: str | None = None, allow_loss: bool = False,
             backup: bool = True,
             verify_com: bool | None = None) -> dict:
    """Set sheet-view state: freeze panes, split panes, gridlines/headings
    visibility, zoom, active selection, tab color. One backup + one
    verified save."""
    settings = {"freeze": freeze, "split": split, "gridlines": gridlines,
                "headings": headings, "zoom": zoom, "selection": selection,
                "tab_color": tab_color}
    if all(v is None for v in settings.values()):
        raise XlMcpError(
            "nothing to set; give at least one of freeze, split, gridlines, "
            "headings, zoom, selection, tab_color")
    if freeze is not None and split is not None:
        raise XlMcpError(
            "freeze and split are mutually exclusive (one pane per sheet); "
            "set one of them per call")

    pkg = WorkbookPackage.open(path)
    wb = pkg.workbook
    if sheet is not None and sheet not in wb.sheetnames:
        raise XlMcpError(
            f"no sheet named {sheet!r}; sheets: {wb.sheetnames}")
    ws = wb[sheet] if sheet is not None else wb.active
    view = ws.sheet_view
    detail: dict[str, Any] = {"sheet": ws.title}

    if freeze is not None:
        if str(freeze).lower() in ("clear", "none", ""):
            ws.freeze_panes = None
            detail["freeze"] = None
        else:
            cell = str(freeze).upper()
            if not _CELL_RE.match(cell):
                raise XlMcpError(
                    f"freeze must be a single cell like 'B2' (freezes rows "
                    f"above and columns left of it) or 'clear', got "
                    f"{freeze!r}")
            if cell == "A1":
                raise XlMcpError(
                    "freeze='A1' freezes nothing (no rows above, no columns "
                    "left); use 'clear' to remove panes")
            ws.freeze_panes = cell
            detail["freeze"] = cell

    if split is not None:
        if not isinstance(split, dict) or not (
                set(split) and set(split) <= {"x", "y"}):
            raise XlMcpError(
                "split must be a dict with 'x' and/or 'y' positions in "
                "points, e.g. {'x': 200} or {'y': 150, 'x': 300}")
        x = split.get("x")
        y = split.get("y")
        for name, v in (("x", x), ("y", y)):
            if v is not None and (
                    isinstance(v, bool) or not isinstance(v, (int, float))
                    or v <= 0):
                raise XlMcpError(f"split {name} must be a positive number "
                                 "of points")
        active_pane = ("bottomRight" if x and y
                       else "topRight" if x else "bottomLeft")
        view.pane = Pane(
            xSplit=float(x) * 20 if x else None,
            ySplit=float(y) * 20 if y else None,
            state="split", activePane=active_pane)
        detail["split"] = {k: v for k, v in (("x", x), ("y", y))
                           if v is not None}

    if gridlines is not None:
        view.showGridLines = bool(gridlines)
        detail["gridlines"] = bool(gridlines)
    if headings is not None:
        view.showRowColHeaders = bool(headings)
        detail["headings"] = bool(headings)

    if zoom is not None:
        if isinstance(zoom, bool) or not isinstance(zoom, int) \
                or not (10 <= zoom <= 400):
            raise XlMcpError("zoom must be an integer percent from 10 to 400")
        view.zoomScale = zoom
        detail["zoom"] = zoom

    if selection is not None:
        ref = str(selection).upper().replace(" ", "")
        min_col, min_row, _mc, _mr = _check_range(ref, "selection")
        from openpyxl.utils import get_column_letter
        active_cell = f"{get_column_letter(min_col)}{min_row}"
        pane_attr = None
        if view.pane is not None and view.pane.activePane:
            pane_attr = view.pane.activePane
        view.selection = [Selection(pane=pane_attr, activeCell=active_cell,
                                    sqref=ref)]
        detail["selection"] = ref

    if tab_color is not None:
        if str(tab_color).lower() in ("clear", "none", ""):
            ws.sheet_properties.tabColor = None
            detail["tab_color"] = None
        else:
            m = _HEX_RE.match(str(tab_color))
            if not m:
                raise XlMcpError(
                    f"tab_color must be a 6-digit hex color like 'FF9900' "
                    f"or 'clear', got {tab_color!r}")
            ws.sheet_properties.tabColor = "00" + m.group(1).upper()
            detail["tab_color"] = m.group(1).upper()

    pkg._changed["view"] = detail
    return pkg.save(allow_loss=allow_loss, backup=backup,
                    verify_com=verify_com)


__all__ = ["set_view"]
