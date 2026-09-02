"""ops/tables.py: Excel tables (ListObjects), DESIGN Section 11.

create_table turns a range into a ListObject with a header row, a table style,
and an optional totals row. get_table reads a table back as records or a matrix
with column addressing and the honest cached-value story. manage_table is the
advanced lifecycle (add/delete rows and columns, rename, resize, toggle totals,
convert to a plain range, restyle).

Structural row/column edits route through WorkbookPackage.modify_structure, so
core.refs rewrites every formula, name, conditional-format, data-validation, and
merge that pointed across the change, and the table ref is then set explicitly
to its intended bounds (the boundary-insert case openpyxl's transposer cannot
disambiguate). Reads load the workbook directly and never mutate.
"""

from __future__ import annotations

import re
from typing import Any

from openpyxl.utils import (
    column_index_from_string,
    get_column_letter,
    range_boundaries,
)

from ..core import refs as _refs
from ..core.errors import TargetNotFound, XlMcpError
from ..core.package import WorkbookPackage
from . import gridio

#: Excel table (ListObject) name grammar: a letter or underscore, then letters,
#: digits, periods, underscores; never a cell-reference shape like "A1".
_NAME_RE = re.compile(r"^[A-Za-z_\\][A-Za-z0-9_.]*$")
_CELLREF_RE = re.compile(r"^[A-Za-z]{1,3}\$?\d+$")

TABLE_ACTIONS = (
    "add_row", "delete_row", "add_column", "delete_column", "rename",
    "resize", "toggle_totals", "to_range", "set_style",
)

_TOTALS_FUNCS = frozenset({
    "sum", "average", "count", "countNums", "min", "max", "stdDev", "var",
    "none", "custom",
})


def _validate_name(wb, name: str) -> str:
    if not isinstance(name, str) or not name.strip():
        raise XlMcpError("a table needs a non-empty name")
    name = name.strip()
    if _CELLREF_RE.match(name) or not _NAME_RE.match(name):
        raise XlMcpError(
            f"table name {name!r} is not a legal Excel table name (start with "
            "a letter or underscore, no spaces, not a cell reference)")
    if len(name) > 255:
        raise XlMcpError("table name exceeds 255 characters")
    existing = {t.lower() for w in wb.worksheets for t in getattr(w, "tables", {})}
    if name.lower() in existing:
        raise XlMcpError(f"a table named {name!r} already exists in the workbook")
    return name


def _find_table(wb, name: str):
    for ws in wb.worksheets:
        tables = getattr(ws, "tables", {})
        if name in tables:
            return ws, tables[name]
        for tname, t in tables.items():
            if tname.lower() == name.lower():
                return ws, t
    raise TargetNotFound(
        f"no table named {name!r}; tables: "
        + (", ".join(repr(t) for w in wb.worksheets
                     for t in getattr(w, "tables", {})) or "none"))


def _bounds(ref: str) -> tuple[int, int, int, int]:
    min_col, min_row, max_col, max_row = range_boundaries(ref)
    return min_row, min_col, max_row, max_col


def _ref(min_row: int, min_col: int, max_row: int, max_col: int) -> str:
    return (f"{get_column_letter(min_col)}{min_row}:"
            f"{get_column_letter(max_col)}{max_row}")


def _header_names(ws, min_row: int, min_col: int, max_col: int) -> list[str]:
    return [ws.cell(min_row, c).value
            for c in range(min_col, max_col + 1)]


def _sync_columns(ws, table, min_row: int, min_col: int, max_col: int) -> None:
    """Rebuild the table's column list from the current header row, but only if
    the columns were materialized; an empty column list lets openpyxl re-derive
    the columns from the header cells on save. Existing totals functions and
    labels are preserved by column name."""
    if not table.tableColumns:
        return
    from openpyxl.worksheet.table import TableColumn
    prior = {tc.name: tc for tc in table.tableColumns}
    names = _header_names(ws, min_row, min_col, max_col)
    cols = []
    for i, hn in enumerate(names):
        cn = str(hn) if hn is not None else f"Column{i + 1}"
        tc = TableColumn(id=i + 1, name=cn)
        old = prior.get(cn)
        if old is not None:
            tc.totalsRowFunction = old.totalsRowFunction
            tc.totalsRowLabel = old.totalsRowLabel
        cols.append(tc)
    table.tableColumns = cols


# --------------------------------------------------------------- create_table


def create_table(path: str, location: Any, name: str, header: bool = True,
                 style: str = "TableStyleMedium9", row_stripes: bool = True,
                 col_stripes: bool = False, totals_row: bool = False,
                 totals: dict | None = None, sheet: str | None = None,
                 allow_loss: bool = False, backup: bool = True) -> dict:
    """Turn a range into a table (ListObject). One backup + one verified save."""
    from openpyxl.worksheet.table import Table, TableColumn, TableStyleInfo

    pkg = WorkbookPackage.open(path)
    wb = pkg.workbook
    name = _validate_name(wb, name)
    grid = pkg.resolve(location, default_sheet=sheet)
    if grid.empty:
        raise XlMcpError("cannot build a table on an empty range")
    ws = wb[grid.sheet]
    min_row, min_col, max_row, max_col = (
        grid.min_row, grid.min_col, grid.max_row, grid.max_col)
    ncols = max_col - min_col + 1

    if header:
        raw = _header_names(ws, min_row, min_col, max_col)
        seen: set[str] = set()
        col_names: list[str] = []
        for i, h in enumerate(raw):
            hn = str(h).strip() if h is not None and str(h).strip() else \
                f"Column{i + 1}"
            base, k = hn, 1
            while hn.lower() in seen:
                k += 1
                hn = f"{base}{k}"
            seen.add(hn.lower())
            col_names.append(hn)
        header_row_count = 1
    else:
        col_names = [f"Column{i + 1}" for i in range(ncols)]
        header_row_count = 0

    if totals_row and max_row - min_row + 1 - header_row_count < 1:
        raise XlMcpError("a totals row needs at least one data row above it")

    table = Table(displayName=name, ref=_ref(min_row, min_col, max_row, max_col))
    table.headerRowCount = header_row_count
    need_columns = (not header) or totals_row or bool(totals)
    if need_columns:
        cols = []
        for i, cn in enumerate(col_names):
            tc = TableColumn(id=i + 1, name=cn)
            cols.append(tc)
        table.tableColumns = cols
    if totals_row:
        table.totalsRowCount = 1
        table.ref = _ref(min_row, min_col, max_row + 1, max_col)
        funcs = totals or {}
        for i, tc in enumerate(table.tableColumns):
            fn = funcs.get(tc.name)
            if fn is not None:
                if fn not in _TOTALS_FUNCS:
                    raise XlMcpError(
                        f"totals function {fn!r} must be one of "
                        f"{sorted(_TOTALS_FUNCS)}")
                tc.totalsRowFunction = fn
            elif i == 0:
                tc.totalsRowLabel = "Total"
    table.tableStyleInfo = TableStyleInfo(
        name=style, showRowStripes=row_stripes, showColumnStripes=col_stripes,
        showFirstColumn=False, showLastColumn=False)
    ws.add_table(table)

    pkg._changed["table"] = {
        "created": name, "sheet": ws.title, "ref": table.ref,
        "columns": col_names, "header": header, "totals_row": totals_row}
    return pkg.save(allow_loss=allow_loss, backup=backup)


# ------------------------------------------------------------------ get_table


def get_table(path: str, name: str, columns: list | None = None,
              values: str = "cached", records: bool = False) -> dict:
    """Read a table's data by name. columns projects a subset; values is
    cached | formula | both. Read-only; nothing is written."""
    if values not in gridio.VALUE_MODES:
        raise XlMcpError(f"values must be one of {gridio.VALUE_MODES}")
    formula_wb = gridio.open_wb(path, data_only=False) \
        if values in ("formula", "both") else None
    cached_wb = gridio.open_wb(path, data_only=True) \
        if values in ("cached", "both") else None
    base = formula_wb if formula_wb is not None else cached_wb
    try:
        ws, table = _find_table(base, name)
        min_row, min_col, max_row, max_col = _bounds(table.ref)
        hrc = table.headerRowCount if table.headerRowCount is not None else 1
        trc = table.totalsRowCount if table.totalsRowCount is not None else 0
        all_names = [str(base[ws.title].cell(min_row, c).value)
                     if hrc else f"Column{c - min_col + 1}"
                     for c in range(min_col, max_col + 1)]
        if columns:
            sel = []
            for c in columns:
                if str(c) not in all_names:
                    raise TargetNotFound(
                        f"table {name!r} has no column {c!r}; columns: "
                        f"{all_names}")
                sel.append(all_names.index(str(c)))
        else:
            sel = list(range(len(all_names)))
        proj_names = [all_names[i] for i in sel]
        data_top = min_row + hrc
        data_bottom = max_row - trc
        rws = base[ws.title]
        cws = cached_wb[ws.title] if cached_wb is not None else None
        fws = formula_wb[ws.title] if formula_wb is not None else None
        out_rows: list[Any] = []
        for r in range(data_top, data_bottom + 1):
            row = []
            for i in sel:
                c = min_col + i
                if values == "formula":
                    v = fws.cell(r, c).value
                elif values == "cached":
                    v = cws.cell(r, c).value
                else:
                    fv = fws.cell(r, c).value
                    v = cws.cell(r, c).value if not (
                        isinstance(fv, str) and fv.startswith("=")) else fv
                row.append(gridio.compact_value(v))
            out_rows.append(dict(zip(proj_names, row)) if records else row)
        return {
            "sheet": ws.title, "table": table.displayName or table.name,
            "ref": table.ref, "columns": proj_names,
            "rows": out_rows, "row_count": len(out_rows),
            "has_totals_row": trc > 0, "value_mode": values,
        }
    finally:
        for wb in (formula_wb, cached_wb):
            if wb is not None:
                wb.close()


# --------------------------------------------------------------- manage_table


def manage_table(path: str, name: str, action: str, values: list | None = None,
                 index: int | None = None, column: str | None = None,
                 new_name: str | None = None, new_ref: str | None = None,
                 style: str | None = None, on: bool = True,
                 totals: dict | None = None, row_stripes: bool | None = None,
                 col_stripes: bool | None = None, sheet: str | None = None,
                 allow_loss: bool = False, backup: bool = True) -> dict:
    """Advanced table lifecycle. One backup + one verified save per call."""
    if action not in TABLE_ACTIONS:
        raise XlMcpError(f"action must be one of {TABLE_ACTIONS}, got {action!r}")
    pkg = WorkbookPackage.open(path)
    wb = pkg.workbook
    ws, table = _find_table(wb, name)
    min_row, min_col, max_row, max_col = _bounds(table.ref)
    hrc = table.headerRowCount if table.headerRowCount is not None else 1
    trc = table.totalsRowCount if table.totalsRowCount is not None else 0
    detail: dict[str, Any] = {"action": action, "table": table.displayName}

    if action == "add_row":
        rows = _as_rows(values)
        n = len(rows)
        insert_at = max_row - trc + 1
        pkg.modify_structure(_refs.RefEdit(
            ws.title, _refs.INSERT_ROWS, index=insert_at, count=n))
        ws = pkg.workbook[ws.title]
        table = ws.tables[table.name]
        table.ref = _ref(min_row, min_col, max_row + n, max_col)
        for i, row in enumerate(rows):
            for j, val in enumerate(row[:max_col - min_col + 1]):
                pkg.set_cell(ws.title,
                             gridio.a1(insert_at + i, min_col + j), val)
        detail.update(added_rows=n, ref=table.ref)

    elif action == "delete_row":
        if index is None or int(index) < 1:
            raise XlMcpError("delete_row needs index (1-based data row)")
        data_rows = max_row - trc - (min_row + hrc) + 1
        if int(index) > data_rows:
            raise XlMcpError(
                f"row index {index} out of range; the table has {data_rows} "
                "data row(s)")
        sheet_row = min_row + hrc + int(index) - 1
        pkg.modify_structure(_refs.RefEdit(
            ws.title, _refs.DELETE_ROWS, index=sheet_row, count=1))
        ws = pkg.workbook[ws.title]
        table = ws.tables[table.name]
        table.ref = _ref(min_row, min_col, max_row - 1, max_col)
        detail.update(deleted_row=int(index), ref=table.ref)

    elif action == "add_column":
        if not column:
            raise XlMcpError("add_column needs column (the new header name)")
        new_col = max_col + 1
        if new_col > _refs.MAX_COL:
            raise XlMcpError("the table already reaches the last column")
        if hrc:
            pkg.set_cell(ws.title, gridio.a1(min_row, new_col), column)
        col_vals = _as_flat(values)
        data_top = min_row + hrc
        for i, val in enumerate(col_vals):
            pkg.set_cell(ws.title, gridio.a1(data_top + i, new_col), val)
        table.ref = _ref(min_row, min_col, max_row, new_col)
        _sync_columns(ws, table, min_row, min_col, new_col)
        detail.update(added_column=column, ref=table.ref)

    elif action == "delete_column":
        if not column:
            raise XlMcpError("delete_column needs column (the header name)")
        names = _header_names(ws, min_row, min_col, max_col)
        idx = next((i for i, h in enumerate(names)
                    if h is not None and str(h).strip() == column), None)
        if idx is None:
            raise TargetNotFound(
                f"table {name!r} has no column {column!r}; columns: {names}")
        if max_col - min_col + 1 == 1:
            raise XlMcpError("cannot delete the only column of a table")
        sheet_col = min_col + idx
        pkg.modify_structure(_refs.RefEdit(
            ws.title, _refs.DELETE_COLS, index=sheet_col, count=1))
        ws = pkg.workbook[ws.title]
        table = ws.tables[table.name]
        table.ref = _ref(min_row, min_col, max_row, max_col - 1)
        _sync_columns(ws, table, min_row, min_col, max_col - 1)
        detail.update(deleted_column=column, ref=table.ref)

    elif action == "rename":
        if not new_name:
            raise XlMcpError("rename needs new_name")
        new = _validate_name(wb, new_name)
        old = table.name
        del ws.tables[old]
        table.displayName = new
        table.name = new
        ws.add_table(table)
        detail.update(renamed_from=old, renamed_to=new)

    elif action == "resize":
        if not new_ref:
            raise XlMcpError("resize needs new_ref (an A1 range)")
        nb = _bounds(new_ref.upper())
        if nb[0] != min_row or nb[1] != min_col:
            raise XlMcpError(
                "resize must keep the table's top-left header anchor "
                f"({get_column_letter(min_col)}{min_row}); move it with a "
                "separate edit if needed")
        table.ref = new_ref.upper()
        detail.update(ref=table.ref)

    elif action == "toggle_totals":
        if on and trc == 0:
            if max_row + 1 > _refs.MAX_ROW:
                raise XlMcpError("no room below the table for a totals row")
            table.totalsRowCount = 1
            table.ref = _ref(min_row, min_col, max_row + 1, max_col)
            names = _header_names(ws, min_row, min_col, max_col)
            from openpyxl.worksheet.table import TableColumn
            cols = []
            funcs = totals or {}
            for i, hn in enumerate(names):
                cn = str(hn) if hn is not None else f"Column{i + 1}"
                tc = TableColumn(id=i + 1, name=cn)
                fn = funcs.get(cn)
                if fn is not None:
                    if fn not in _TOTALS_FUNCS:
                        raise XlMcpError(
                            f"totals function {fn!r} must be one of "
                            f"{sorted(_TOTALS_FUNCS)}")
                    tc.totalsRowFunction = fn
                elif i == 0:
                    tc.totalsRowLabel = "Total"
                cols.append(tc)
            table.tableColumns = cols
            detail.update(totals_row=True, ref=table.ref)
        elif not on and trc == 1:
            table.totalsRowCount = 0
            table.ref = _ref(min_row, min_col, max_row - 1, max_col)
            for tc in table.tableColumns or []:
                tc.totalsRowFunction = None
                tc.totalsRowLabel = None
            detail.update(totals_row=False, ref=table.ref)
        else:
            detail.update(totals_row=bool(trc), note="already in that state")

    elif action == "to_range":
        del ws.tables[table.name]
        detail.update(converted_to_range=True)

    else:  # set_style
        from openpyxl.worksheet.table import TableStyleInfo
        cur = table.tableStyleInfo
        table.tableStyleInfo = TableStyleInfo(
            name=style if style is not None else (
                cur.name if cur else "TableStyleMedium9"),
            showRowStripes=row_stripes if row_stripes is not None else (
                cur.showRowStripes if cur else True),
            showColumnStripes=col_stripes if col_stripes is not None else (
                cur.showColumnStripes if cur else False),
            showFirstColumn=cur.showFirstColumn if cur else False,
            showLastColumn=cur.showLastColumn if cur else False)
        detail.update(style=table.tableStyleInfo.name)

    pkg._changed["table"] = detail
    return pkg.save(allow_loss=allow_loss, backup=backup)


def _as_rows(values) -> list[list]:
    if not values:
        raise XlMcpError("add_row needs values (a row list, or a list of rows)")
    if all(isinstance(v, list) for v in values):
        return values
    return [values]


def _as_flat(values) -> list:
    if not values:
        return []
    if all(isinstance(v, list) for v in values):
        return [v[0] if v else None for v in values]
    return list(values)


__all__ = ["create_table", "get_table", "manage_table", "TABLE_ACTIONS"]
