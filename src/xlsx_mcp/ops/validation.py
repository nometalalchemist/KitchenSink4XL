"""ops/validation.py: the multiplex read-only checker (DESIGN Section 11).

Grammar-matched to KitchenSink4Word's validate: checks=[...] picks the
batteries, the result is {passed, results: {check: {passed, findings}}},
and passed=false means FINDINGS, never a failed call. Where an underlying
op already owns the report shape (audit_styles, the hazard scan,
get_external_links, audit_formulas' sections), the findings are that op's
verbatim result, so a caller who knows one shape knows both.

Checks: structure (package opens clean, sheet integrity), references
(#REF!/#NAME? and kin), names (broken defined names), merges (overlapping
or orphaned merged ranges), tables (duplicates, broken refs, overlaps),
formatting_bloat (audit_styles' counters as a check), hazards (the
round-trip scan as a check), external_links, calc_staleness (formulas
lacking cached values, which read as blank to every non-Excel consumer).

Everything here is read-only; repairs live in the editing tools.
"""

from __future__ import annotations

import zipfile
from typing import Any, Callable

from openpyxl.utils.cell import range_boundaries

from ..core import hazard as _hazard
from ..core import locate as _locate
from ..core.errors import XlMcpError
from ..core.sandbox import check_path
from . import format as _format
from . import formulas as _formulas
from . import gridio
from . import inspectors as _inspectors

DEFAULT_CHECKS = ("structure", "references", "calc_staleness")


# --------------------------------------------------------------- the checks


#: What the structure check can honestly say about ``opens_clean`` when
#: nothing asked Excel. openpyxl loading a package is NOT Excel accepting it:
#: the insane round produced ten files that openpyxl read happily and Excel
#: answered "Open method of Workbooks class failed" for, and every one of
#: them was reported here as ``opens_clean: true`` because the field was
#: hard-coded (M-1). The field now carries this string instead of a claim,
#: and becomes a real boolean only when the COM check actually ran.
NOT_CHECKED = "not checked (no Excel verdict was requested)"

#: Set KS4XL_VALIDATE_COM=1 to make the structure check route to
#: com_validate_opens_clean whenever the com pack is available.
_COM_ENV = "KS4XL_VALIDATE_COM"


def _com_opens_clean(path: str):
    """Excel's own verdict, or None when it was not asked / not available.
    com_validate_opens_clean is the authoritative corruption smoke test and
    got every one of the round's ten unopenable files right; this is the
    wire from the file-tier check to it."""
    import os

    if os.environ.get(_COM_ENV, "").strip().lower() not in ("1", "true",
                                                            "yes", "on"):
        return None
    try:
        from ..com import session as _com_session
        ok, _why = _com_session.com_available()
        if not ok:
            return None
        return bool(_com_session.opens_clean(path).get("opens_clean"))
    except Exception:  # noqa: BLE001 - an unavailable check is 'not checked'
        return None


def _check_structure(path: str) -> tuple[bool, dict]:
    issues: list[str] = []
    kind = _hazard.ole_container_kind(check_path(path, "validate workbook"))
    if kind is not None:
        # An OLE container (legacy .xls or an encrypted package): the zip
        # check below would either fail with the wrong word ("not a zip")
        # or, for a modern-Excel .xls with its embedded theme fragment,
        # "succeed" and report a missing workbook part (geriatric H-1/L-1).
        return False, {
            "opens_clean": False,
            "opens_clean_source": "OLE compound-file signature",
            "issues": [_hazard.ole_refusal_text(kind, path)]}
    try:
        with zipfile.ZipFile(check_path(path, "validate workbook")) as zf:
            names = set(zf.namelist())
        for required in ("[Content_Types].xml", "xl/workbook.xml"):
            if required not in names:
                issues.append(f"package is missing {required}")
    except zipfile.BadZipFile:
        return False, {"opens_clean": False,
                       "opens_clean_source": "the file is not a zip at all",
                       "issues": ["not a valid zip / OOXML package"]}
    try:
        wb = gridio.open_wb(path)
    except Exception as exc:  # noqa: BLE001
        return False, {"opens_clean": False,
                       "opens_clean_source": "openpyxl could not load it",
                       "issues": issues + [f"openpyxl cannot load it: {exc}"]}
    try:
        visible = [ws for ws in wb.worksheets
                   if ws.sheet_state == "visible"]
        if not visible:
            issues.append("no visible sheet (Excel refuses such a file)")
        seen: dict[str, str] = {}
        for name in wb.sheetnames:
            low = name.lower()
            if low in seen:
                issues.append(
                    f"sheet names collide case-insensitively: "
                    f"{seen[low]!r} and {name!r}")
            seen[low] = name
        excel_verdict = _com_opens_clean(path)
        if excel_verdict is False:
            issues.append(
                "Excel refuses to open this file or demands a repair "
                "(com_validate_opens_clean)")
        findings = {
            "opens_clean": (NOT_CHECKED if excel_verdict is None
                            else excel_verdict),
            "opens_clean_source": (
                "Excel (com_validate_opens_clean)" if excel_verdict is not None
                else "openpyxl loaded the package; Excel was NOT asked. Run "
                     "com_validate_opens_clean, or set "
                     f"{_COM_ENV}=1, for an Excel verdict"),
            "openpyxl_loads": True,
            "sheet_count": len(wb.sheetnames),
            "visible_sheets": len(visible),
            "issues": issues,
        }
        return not issues, findings
    finally:
        wb.close()


def _audit_sections(path: str, wanted: set[str]) -> dict[str, dict]:
    """One audit_formulas pass serving both formula-backed checks."""
    audit = _formulas.audit_formulas(path)
    out: dict[str, dict] = {}
    if "references" in wanted:
        findings = dict(audit["error_cells"])
        if audit.get("unknown_sheet_references"):
            findings["unknown_sheet_references"] = \
                audit["unknown_sheet_references"]
        out["references"] = findings
    if "calc_staleness" in wanted:
        findings = dict(audit["missing_cached_values"])
        if findings["count"]:
            findings["note"] = (
                "these formula cells have no cached value and read as "
                "blank to every non-Excel consumer until a recalculation")
        out["calc_staleness"] = findings
    return out


def _check_names(path: str) -> tuple[bool, dict]:
    wb = gridio.open_wb(path)
    try:
        known = {n.lower() for n in wb.sheetnames}
        broken: list[dict] = []
        total = 0
        for scope, name, defn in _locate._all_defined_names(wb):
            total += 1
            value = getattr(defn, "value", None) or ""
            reasons: list[str] = []
            if "#REF!" in value:
                reasons.append("contains #REF!")
            for m in _formulas._SHEET_REF_RE.finditer(
                    _formulas._strip_strings(value)):
                target = (m.group(1) or m.group(2) or "").replace("''", "'")
                if "[" in target:
                    continue  # external workbook reference
                if target.lower() not in known:
                    reasons.append(f"references missing sheet {target!r}")
            if not value.strip():
                reasons.append("empty definition")
            if reasons:
                broken.append({"name": name, "scope": scope,
                               "value": value, "reasons": reasons})
        return not broken, {"defined_names": total, "broken": broken,
                            "count": len(broken)}
    finally:
        wb.close()


def _overlap(a: tuple[int, int, int, int],
             b: tuple[int, int, int, int]) -> bool:
    return not (a[2] < b[0] or b[2] < a[0] or a[3] < b[1] or b[3] < a[1])


def _check_merges(path: str) -> tuple[bool, dict]:
    wb = gridio.open_wb(path)
    try:
        overlaps: list[dict] = []
        orphans: list[dict] = []
        total = 0
        for ws in wb.worksheets:
            ranges = [str(r) for r in ws.merged_cells.ranges]
            total += len(ranges)
            bounds = [range_boundaries(r) for r in ranges]
            for i in range(len(bounds)):
                for j in range(i + 1, len(bounds)):
                    if _overlap(bounds[i], bounds[j]):
                        overlaps.append({"sheet": ws.title,
                                         "ranges": [ranges[i], ranges[j]]})
            used = _locate.true_used_range(ws)
            for ref, (c0, r0, c1, r1) in zip(ranges, bounds):
                if used is None:
                    orphans.append({"sheet": ws.title, "range": ref,
                                    "reason": "sheet holds no values"})
                    continue
                ur0, uc0, ur1, uc1 = used
                if r1 < ur0 or r0 > ur1 or c1 < uc0 or c0 > uc1:
                    orphans.append({
                        "sheet": ws.title, "range": ref,
                        "reason": "entirely outside the used range"})
        findings = {"merged_ranges": total, "overlaps": overlaps,
                    "orphans": orphans,
                    "count": len(overlaps) + len(orphans)}
        return not overlaps and not orphans, findings
    finally:
        wb.close()


def _check_tables(path: str) -> tuple[bool, dict]:
    wb = gridio.open_wb(path)
    try:
        problems: list[dict] = []
        seen: dict[str, str] = {}
        per_sheet: dict[str, list[tuple[str, tuple]]] = {}
        total = 0
        for ws in wb.worksheets:
            tmap = getattr(ws, "tables", {})
            for tname in list(tmap):
                total += 1
                ref = tmap[tname].ref
                low = tname.lower()
                if low in seen:
                    problems.append({
                        "table": tname, "sheet": ws.title,
                        "problem": f"name collides with {seen[low]!r} "
                                   "(table names are workbook-global)"})
                seen[low] = tname
                try:
                    b = range_boundaries(ref)
                except Exception:  # noqa: BLE001
                    problems.append({"table": tname, "sheet": ws.title,
                                     "problem": f"broken ref {ref!r}"})
                    continue
                if b[0] > b[2] or b[1] > b[3] or \
                        b[2] > _locate.MAX_COL or b[3] > _locate.MAX_ROW:
                    problems.append({"table": tname, "sheet": ws.title,
                                     "problem": f"ref {ref!r} is out of "
                                                "bounds"})
                    continue
                for other, ob in per_sheet.setdefault(ws.title, []):
                    if _overlap(b, ob):
                        problems.append({
                            "table": tname, "sheet": ws.title,
                            "problem": f"overlaps table {other!r}"})
                per_sheet[ws.title].append((tname, b))
        return not problems, {"tables": total, "problems": problems,
                              "count": len(problems)}
    finally:
        wb.close()


def _check_formatting_bloat(path: str) -> tuple[bool, dict]:
    findings = _format.audit_styles(path)
    return findings["risk"] == "ok", findings


def _check_hazards(path: str) -> tuple[bool, dict]:
    rep = _hazard.scan_path(check_path(path, "scan workbook"))
    return (rep.error is None and not rep.would_lose), rep.as_dict()


def _check_external_links(path: str) -> tuple[bool, dict]:
    findings = _inspectors.get_external_links(path)
    return findings["count"] == 0, findings


_CHECKS: dict[str, Callable[[str], tuple[bool, dict]] | None] = {
    "structure": _check_structure,
    "references": None,       # served by the shared audit pass
    "names": _check_names,
    "merges": _check_merges,
    "tables": _check_tables,
    "formatting_bloat": _check_formatting_bloat,
    "hazards": _check_hazards,
    "external_links": _check_external_links,
    "calc_staleness": None,   # served by the shared audit pass
}

CHECK_NAMES = tuple(_CHECKS)


# ------------------------------------------------------------------ dispatch


def validate(path: str, checks: list[str] | None = None) -> dict:
    """Run the requested read-only checks and return one report:
    {passed, results: {check: {passed, findings}}}."""
    wanted: list[str] = []
    for check in (checks if checks is not None else list(DEFAULT_CHECKS)):
        if check not in _CHECKS:
            raise XlMcpError(
                f"unknown check {check!r}; checks: " + " | ".join(CHECK_NAMES))
        if check not in wanted:
            wanted.append(check)
    if not wanted:
        raise XlMcpError("checks must name at least one check")

    results: dict[str, dict[str, Any]] = {}
    audit_wanted = {c for c in wanted if _CHECKS[c] is None}
    audit_results = _audit_sections(path, audit_wanted) if audit_wanted \
        else {}
    for check in wanted:
        runner = _CHECKS[check]
        if runner is None:
            findings = audit_results[check]
            # count covers the error cells. The references check carries a
            # second finding, unknown_sheet_references, and a verdict that
            # ignored it reported passed:true over its own non-empty payload.
            passed = (findings["count"] == 0
                      and not findings.get("unknown_sheet_references"))
        else:
            passed, findings = runner(path)
        results[check] = {"passed": passed, "findings": findings}
    # "passed", not "ok": the envelope's top-level ok means THE CALL
    # SUCCEEDED; a battery that ran fine but found problems is a passed=false
    # result, not an error.
    return {"passed": all(r["passed"] for r in results.values()),
            "checks_run": wanted, "results": results}


__all__ = ["validate", "CHECK_NAMES", "DEFAULT_CHECKS"]
