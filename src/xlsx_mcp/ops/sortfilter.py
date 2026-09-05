"""ops/sortfilter.py: sort and filter (DESIGN Section 11, lite).

sort_range reorders whole rows of a range (or table body) by one or more keys,
writing the reordered values and styles back. A moved formula has its relative
references shifted by its row displacement (core.refs.offset_formula), Excel's
own sort semantics, so a formula that referred to its own row keeps referring to
its own row after the move; absolute ($) anchors stay put.

set_filter is the honest answer to a file-format landmine: an .xlsx autofilter
stores the filter CRITERIA, not which rows are hidden. A criteria-only write
would leave every row visible until Excel re-applies the filter. So set_filter
EVALUATES the criteria here and hides the non-matching data rows (row-level
hidden flag), and also records the autofilter range so Excel shows the filter
UI. The hidden state is a snapshot over the current cached / literal values;
Excel re-evaluates on its own when the user touches the filter. Formula cells
with no cached value cannot be evaluated and are reported rather than guessed.
clear_filter removes the autofilter and unhides the rows.

All three route mutations through WorkbookPackage for the backup and verify.
"""

from __future__ import annotations

from copy import copy as _copy
from typing import Any

from openpyxl.utils import get_column_letter

from ..core import arrays as _arrays
from ..core import calc as _calc
from ..core import refs as _refs
from ..core.errors import TargetNotFound, UnsupportedStructure, XlMcpError
from ..core.package import WorkbookPackage
from . import cells as _cells
from . import gridio


def _column_index(col: Any, header_names: list[str], min_col: int) -> int:
    """Resolve a key/criteria column to a 0-based offset within the block."""
    if isinstance(col, int):
        return col - 1
    s = str(col)
    if s in header_names:
        return header_names.index(s)
    from openpyxl.utils import column_index_from_string
    try:
        return column_index_from_string(s.upper()) - min_col
    except Exception:
        raise XlMcpError(f"no column {col!r} in the range; headers: "
                         f"{header_names}")


def _header_names(ws, row: int, min_col: int, max_col: int) -> list[str]:
    return [str(ws.cell(row, c).value) if ws.cell(row, c).value is not None
            else get_column_letter(c)
            for c in range(min_col, max_col + 1)]


def _refuse_merges(ws, grid, what: str) -> None:
    """Refuse when the target rectangle overlaps a merged range.

    Rewriting rows under a merge is not a thing openpyxl can do: only the
    top-left of a merge holds a value and the rest are read-only MergedCell
    objects, so the write half of a sort died mid-flight with
    "'MergedCell' object attribute 'value' is read-only" (adversarial round)
    after the reads had already run. Refuse up front, name the merges, and
    say the remedy."""
    overlapping = []
    for mr in list(getattr(ws, "merged_cells", None).ranges
                   if getattr(ws, "merged_cells", None) else []):
        if (mr.min_row <= grid.max_row and mr.max_row >= grid.min_row
                and mr.min_col <= grid.max_col and mr.max_col >= grid.min_col):
            overlapping.append(str(mr))
    if overlapping:
        raise UnsupportedStructure(
            f"cannot {what} {grid.a1} on {grid.sheet!r}: it overlaps merged "
            f"range(s) {', '.join(sorted(overlapping))}, and merged cells "
            f"cannot be rewritten row by row. Unmerge them first "
            f"(set_merge action='unmerge'), then {what}.")


def sort_range(path: str, location: Any, keys: list, has_header: bool = True,
               sheet: str | None = None, allow_loss: bool = False,
               backup: bool = True,
               verify_com: bool | None = None) -> dict:
    """Sort a range or table body by one or more keys. Backup + verify."""
    if not keys or not isinstance(keys, list):
        raise XlMcpError("keys must be a non-empty list of {column, order}")
    pkg = WorkbookPackage.open(path)
    grid = pkg.resolve(location, default_sheet=sheet)
    if grid.empty:
        raise XlMcpError("cannot sort an empty range")
    gridio.guard_cell_count(grid)
    ws = pkg.workbook[grid.sheet]
    _refuse_merges(ws, grid, "sort")
    # An array formula cannot be reordered row by row: its ref anchor stays
    # behind and the workbook stops opening. Excel refuses the identical Sort
    # with "You can't change part of an array" (core.arrays).
    _arrays.refuse_if_touched(ws, grid, "sort")
    cached = gridio.open_wb(path, data_only=True)
    cws = cached[grid.sheet]

    min_row, min_col, max_row, max_col = (
        grid.min_row, grid.min_col, grid.max_row, grid.max_col)
    data_top = min_row + 1 if has_header else min_row
    header_names = _header_names(ws, min_row, min_col, max_col) if has_header \
        else [get_column_letter(c) for c in range(min_col, max_col + 1)]

    key_specs = []
    for k in keys:
        if isinstance(k, dict):
            col = k.get("column")
            order = str(k.get("order", "asc")).lower()
        else:
            col, order = k, "asc"
        if col is None:
            raise XlMcpError("each key needs a column")
        key_specs.append((_column_index(col, header_names, min_col),
                          order in ("desc", "descending")))

    # Hidden rows, Excel's way (COM ground truth 2026-09-05): sorting a
    # range with an ACTIVE AUTOFILTER sorts only the visible rows and
    # leaves filter-hidden rows pinned in place with their data; a sort
    # over MANUALLY hidden rows (no filter) moves every row while the
    # hidden flags stay with the row POSITIONS. The old code always did
    # the latter, so a sort after set_filter left the filter hiding the
    # WRONG rows with zero warnings (fresh-eyes round, M-2).
    hidden_rows = [r for r in range(data_top, max_row + 1)
                   if ws.row_dimensions[r].hidden]
    pin_hidden = bool(hidden_rows) and bool(ws.auto_filter.ref)
    if pin_hidden:
        hidden_set = set(hidden_rows)
        sortable_rows = [r for r in range(data_top, max_row + 1)
                         if r not in hidden_set]
    else:
        sortable_rows = list(range(data_top, max_row + 1))

    key_offsets = {o for o, _rev in key_specs}
    uncached_keys = 0
    #: Formula cells in the sorted block that DID have a cached value. The
    #: openpyxl round-trip does not carry a cache through, so every one of
    #: them reads back 'absent' after this save. uncached_keys counts the
    #: opposite population (cells with no cache at read time), which is why a
    #: fully-computed workbook produced zero warnings while losing all of its
    #: numbers (insane round, M-5).
    dropped_cache = 0
    rows = []
    for r in sortable_rows:
        vals, styles, sortvals, texts = [], [], [], []
        for c in range(min_col, max_col + 1):
            cell = ws.cell(r, c)
            vals.append(cell.value)
            styles.append(_copy(cell._style))
            # A TEXT cell that merely starts with '=' is data, not a formula
            # (import_data's neutralized injection text); it must keep its
            # string type through the rewrite instead of going live.
            texts.append(_calc.looks_like_formula_text(cell))
            cv = cws.cell(r, c).value
            if _calc.is_formula_cell(cell):
                if cv is None:
                    if (c - min_col) in key_offsets:
                        uncached_keys += 1  # sorts by formula TEXT, not value
                else:
                    dropped_cache += 1
            sortvals.append(cv if cv is not None else cell.value)
        rows.append({"src": r, "vals": vals, "styles": styles,
                     "sort": sortvals, "text": texts})
    cached.close()

    for offset, reverse in reversed(key_specs):
        rows.sort(key=lambda row, o=offset, rv=reverse: _cells._sort_key(
            row["sort"][o], reverse=rv), reverse=reverse)

    wrote_formula = False
    for i, row in enumerate(rows):
        dest_r = sortable_rows[i]
        dr = dest_r - row["src"]
        for j, (val, style, is_text) in enumerate(
                zip(row["vals"], row["styles"], row["text"])):
            cell = ws.cell(dest_r, min_col + j)
            out = val
            if is_text:
                cell.value = out
                cell.data_type = "s"   # stays TEXT, never re-armed
                cell._style = _copy(style)
                continue
            if isinstance(val, str) and val.startswith("="):
                if dr != 0:
                    out = _refs.offset_formula(val, dr, 0, grid.sheet)
                wrote_formula = True
            cell.value = out
            cell._style = _copy(style)
    if wrote_formula:
        pkg._formula_written = True
    pkg._changed["sorted"] = {
        "sheet": grid.sheet, "range": grid.a1,
        "rows_sorted": len(rows),
        **({"rows_pinned_hidden": len(hidden_rows)} if pin_hidden else {}),
        "keys": [
            {"column": header_names[o] if o < len(header_names) else o,
             "order": "desc" if rev else "asc"} for o, rev in key_specs]}
    result = pkg.save(allow_loss=allow_loss, backup=backup,
                      verify_com=verify_com)
    notes: list[str] = []
    if pin_hidden:
        notes.append(
            f"{len(hidden_rows)} filter-hidden row(s) were left in place "
            "and excluded from the sort, matching Excel's semantics for "
            "sorting a filtered range; clear_filter first to sort every "
            "row.")
    elif hidden_rows:
        notes.append(
            f"{len(hidden_rows)} manually hidden row(s) were sorted along "
            "with the visible rows (Excel's own behavior without a "
            "filter); hidden flags stay with the row POSITIONS, so "
            "different rows may be hidden now.")
    if uncached_keys:
        notes.append(
            f"{uncached_keys} sort-key cell(s) hold a formula with no cached "
            "value and were ordered by their formula TEXT, not their result. "
            "Run recalculate or open in Excel first for a value-accurate "
            "sort.")
    if dropped_cache:
        notes.append(
            f"{dropped_cache} formula cell(s) in the sorted range had a "
            "cached value before this sort and have none after it (label "
            "'absent'): a file-based save cannot carry Excel's cache through "
            "a formula rewrite. " + gridio.ABSENT_WARNING)
    if notes:
        result["warnings"] = list(result.get("warnings", [])) + notes
    return result


def set_filter(path: str, location: Any, criteria: list | None = None,
               sheet: str | None = None, allow_loss: bool = False,
               backup: bool = True,
               verify_com: bool | None = None) -> dict:
    """Apply an autofilter and hide the non-matching rows. Backup + verify."""
    pkg = WorkbookPackage.open(path)
    grid = pkg.resolve(location, default_sheet=sheet)
    if grid.empty:
        raise XlMcpError("cannot filter an empty range")
    ws = pkg.workbook[grid.sheet]
    min_row, min_col, max_row, max_col = (
        grid.min_row, grid.min_col, grid.max_row, grid.max_col)
    header_names = _header_names(ws, min_row, min_col, max_col)

    preds = []
    for cr in (criteria or []):
        if not isinstance(cr, dict) or "column" not in cr:
            raise XlMcpError(
                "each criterion is {column, op, value}; op defaults to eq")
        # Unknown keys used to be IGNORED: the natural LLM shape
        # {"column": "Category", "equals": "Freight"} parsed as op=eq
        # value=None and hid every data row with ok:true (destroyer
        # round, L-1). A stray key or a value-less binary op refuses like
        # a wrong op does.
        unknown = set(cr) - {"column", "op", "value"}
        if unknown:
            raise XlMcpError(
                f"unknown criterion key(s) {sorted(unknown)}; each "
                "criterion is {column, op, value} (op defaults to eq; "
                "for equality put the wanted value under 'value')")
        op = cr.get("op", "eq")
        if op not in (_cells._OPS_BINARY | _cells._OPS_SET | _cells._OPS_UNARY):
            raise XlMcpError(f"unknown filter op {op!r}")
        if op not in _cells._OPS_UNARY and "value" not in cr:
            raise XlMcpError(
                f"criterion for column {cr['column']!r} has op {op!r} but "
                "no value; only is_blank/not_blank need none")
        offset = _column_index(cr["column"], header_names, min_col)
        preds.append((offset, op, cr.get("value")))

    cached = gridio.open_wb(path, data_only=True)
    cws = cached[grid.sheet]
    hidden = shown = uneval = 0
    warnings: list[str] = []
    for r in range(min_row + 1, max_row + 1):
        keep = True
        for offset, op, value in preds:
            c = min_col + offset
            cv = cws.cell(r, c).value
            fv = ws.cell(r, c).value
            if cv is None and _calc.is_formula_cell(ws.cell(r, c)):
                uneval += 1
                continue  # cannot evaluate a formula with no cached value
            if not _cells._cmp(cv if cv is not None else fv, op, value):
                keep = False
                break
        ws.row_dimensions[r].hidden = not keep
        if keep:
            shown += 1
        else:
            hidden += 1
    cached.close()
    if uneval:
        warnings.append(
            f"{uneval} filtered cell(s) hold a formula with no cached value; "
            "those rows were kept visible. Run recalculate or open in Excel "
            "first for an exact filter.")

    ws.auto_filter.ref = grid.a1
    for offset, op, value in preds:
        if op in ("eq", "in"):
            vals = value if isinstance(value, list) else [value]
            ws.auto_filter.add_filter_column(
                offset, [str(v) for v in vals], blank=False)

    pkg._changed["filter"] = {
        "sheet": grid.sheet, "range": grid.a1,
        "rows_shown": shown, "rows_hidden": hidden,
        "evaluated": "cached-and-literal values, non-matching rows hidden"}
    result = pkg.save(allow_loss=allow_loss, backup=backup,
                      verify_com=verify_com)
    if warnings:
        result["warnings"] = list(result.get("warnings", [])) + warnings
    return result


def clear_filter(path: str, location: Any = None, sheet: str | None = None,
                 allow_loss: bool = False, backup: bool = True,
                 verify_com: bool | None = None) -> dict:
    """Remove the autofilter and unhide the rows it hid. Backup + verify."""
    pkg = WorkbookPackage.open(path)
    wb = pkg.workbook
    ws = wb[sheet] if sheet is not None else wb.active
    ref = ws.auto_filter.ref
    if location is not None:
        grid = pkg.resolve(location, default_sheet=sheet)
        ws = wb[grid.sheet]
        ref = ws.auto_filter.ref or grid.a1
    if not ref:
        # NOT_FOUND, matching the documented contract (the tool description
        # says "refuses (NOT_FOUND)"); a bare XlMcpError mapped to BAD_PARAMS
        # and contradicted the docs (author field testing, lite finding #26).
        raise TargetNotFound(
            f"no autofilter on sheet {ws.title!r} to clear")
    from openpyxl.utils import range_boundaries
    _c0, r0, _c1, r1 = range_boundaries(str(ref))
    unhidden = 0
    for r in range(r0, r1 + 1):
        if ws.row_dimensions[r].hidden:
            ws.row_dimensions[r].hidden = False
            unhidden += 1
    ws.auto_filter.ref = None
    ws.auto_filter.filterColumn = []
    pkg._changed["filter_cleared"] = {
        "sheet": ws.title, "range": str(ref), "rows_unhidden": unhidden}
    return pkg.save(allow_loss=allow_loss, backup=backup,
                    verify_com=verify_com)


__all__ = ["sort_range", "set_filter", "clear_filter"]
