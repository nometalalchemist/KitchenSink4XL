"""ops/lifecycle.py: workbook and worksheet lifecycle + discovery.

The orient-before-editing family (DESIGN Section 11, lite): create and copy
workbooks, read workbook metadata (sheets, the TRUE used range per sheet, named
ranges, defined tables, and the round-trip hazard summary), manage the sheet
lifecycle (add / delete / rename / copy / reorder / hide), and surface the
hazard scan as a user-facing health readout.

Every mutation routes through WorkbookPackage, so the hazard gate, the
backup-before-mutation, and verify-after-write all run without per-op wiring.
create_workbook writes a brand-new file (no existing content to endanger, so no
scan) but still honors the sandbox and never clobbers an existing path unless
told to. Reads load the workbook directly (no mutation, no backup).
"""

from __future__ import annotations

import os
from typing import Any

from ..core import hazard as _hazard
from ..core import locate as _locate
from ..core import calc as _calc
from ..core.errors import TargetNotFound, XlMcpError
from ..core.package import WorkbookPackage
from ..core.sandbox import check_path
from . import format as _format
from . import gridio

WS_ACTIONS = ("add", "delete", "rename", "copy", "reorder", "hide", "unhide")
_HIDE_STATES = {"hidden": "hidden", "very_hidden": "veryHidden",
                "veryhidden": "veryHidden"}


def _check_sheet_name(name: str) -> None:
    """Excel's sheet-name grammar, applied BEFORE openpyxl sees the title.

    This is the single entry point for every sheet title the file tier
    writes (create_workbook, and manage_worksheet add / rename / copy), and
    it delegates to core.limits.check_sheet_title, which owns the measured
    grammar: 1-31 characters, none of : \\ / ? * [ ], no apostrophe at
    either end, and no character Excel cannot store.

    The delegation is the mutation round's D1 finding (2026-09-06):
    check_sheet_title was implemented, exported, docstringed and wired to
    NOTHING, so every one of its mutants survived the full suite and sheet
    naming rode on whatever openpyxl happened to enforce. openpyxl
    validates only length and the banned charset, and it raises a bare
    ValueError doing it; the apostrophe rule it has no notion of at all,
    which is how a name starting with one used to produce a workbook Excel
    refuses to OPEN while verify-after-write still passed (fresh-eyes
    round, M-1). Refusing here means one message, one exception type, and
    a refusal before any part of the package is touched."""
    from ..core.limits import check_sheet_title
    check_sheet_title(name)


# ----------------------------------------------------------------- create/copy


def _refuse_backup_store_target(p: str, what: str) -> None:
    """A workbook written INTO a .ks4xl-backups store overwrites somebody's
    undo; no creation tool may target it (destroyer/fresh-eyes H-1 class)."""
    from pathlib import Path as _Path
    from ..core.safesave import BACKUP_DIR_NAME
    if any(part.lower() == BACKUP_DIR_NAME
           for part in _Path(os.path.abspath(p)).parts):
        raise XlMcpError(
            f"refusing to {what} inside a {BACKUP_DIR_NAME} backup store: "
            "overwriting a backup slot destroys the undo it holds. Choose "
            "a path outside the backup store.")


def _write_atomic(p: str, writer) -> None:
    """Write via a sibling temp file + os.replace. This matters beyond crash
    atomicity: the backup slots are HARDLINKS of the pre-overwrite file, so
    an in-place wb.save/copy2 onto the same inode would rewrite the backup
    it just took. os.replace points the path at a NEW inode and the slot
    keeps the old bytes."""
    import uuid
    tmp = os.path.join(
        os.path.dirname(os.path.abspath(p)) or ".",
        f".ks4xl-write-{uuid.uuid4().hex}{os.path.splitext(p)[1]}")
    try:
        writer(tmp)
        from ..core.safesave import replace_with_retry
        replace_with_retry(tmp, p)
    except BaseException:
        try:
            os.remove(tmp)
        except OSError:
            pass
        raise


def _rotate_before_clobber(p: str) -> dict:
    """Backup-before-mutation for the overwrite=true paths: the workbook
    about to be replaced rotates into its prev/anchor slots FIRST, so one
    habitual overwrite:true is no longer unrecoverable ruin (destroyer
    H-2). Returns the extra result fields."""
    from ..core import safesave as _safesave
    _safesave.rotate_slots(p)
    return {
        "backup": "prev",
        "replaced_existing": True,
        "note": ("the replaced workbook was backed up first: "
                 "manage_backups(action='restore', path=..., source='prev') "
                 "brings it back"),
    }


def create_workbook(path: str, sheets: list[str] | None = None,
                    overwrite: bool = False) -> dict:
    """Create a new .xlsx at `path` with the named sheets (default one sheet
    'Sheet1'). Refuses an existing path unless overwrite=True; overwrite
    rotates the existing workbook into its backup slots before replacing
    it (backup-before-mutation covers this last bypass too)."""
    p = check_path(path, "create workbook")
    _refuse_backup_store_target(p, "create a workbook")
    if os.path.exists(p) and not overwrite:
        raise FileExistsError(
            f"{p} already exists; pass overwrite=true to replace it (the "
            "existing file is backed up to its prev slot first)")
    names = [str(s) for s in (sheets or ["Sheet1"]) if str(s).strip()]
    if not names:
        names = ["Sheet1"]
    seen: set[str] = set()
    for n in names:
        low = n.lower()
        if low in seen:
            raise XlMcpError(f"duplicate sheet name {n!r}")
        seen.add(low)
        _check_sheet_name(n)
    extra: dict = {}
    if os.path.exists(p) and overwrite:
        extra = _rotate_before_clobber(p)
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.title = names[0]
    for n in names[1:]:
        wb.create_sheet(title=n)
    _write_atomic(p, wb.save)
    return {"file": p, "created": True, "sheets": names, "saved": True,
            **extra}


def copy_workbook(src: str, dst: str, overwrite: bool = False) -> dict:
    """Copy a workbook file byte-for-byte to a new path (no round-trip through
    openpyxl, so nothing is degraded). Refuses an existing dst unless
    overwrite=True; overwrite rotates the existing dst into its backup
    slots before replacing it."""
    import shutil
    s = check_path(src, "read source workbook")
    d = check_path(dst, "write copy")
    if not os.path.exists(s):
        raise TargetNotFound(f"no such workbook: {s}")
    _refuse_backup_store_target(d, "copy a workbook")
    from ..core.safesave import canonical_key as _ck
    if os.path.exists(d) and _ck(s) == _ck(d):
        raise XlMcpError(
            "src and dst are the same file; nothing to copy")
    if os.path.exists(d) and not overwrite:
        raise FileExistsError(
            f"{d} already exists; pass overwrite=true to replace it (the "
            "existing file is backed up to its prev slot first)")
    extra: dict = {}
    if os.path.exists(d) and overwrite:
        extra = _rotate_before_clobber(d)
    _write_atomic(d, lambda tmp: shutil.copy2(s, tmp))
    return {"file": d, "copied_from": s, "saved": True, **extra}


# -------------------------------------------------------------- read metadata


def get_workbook_metadata(path: str) -> dict:
    """Read structural metadata without mutating: every sheet with its state and
    TRUE used range, defined names, tables, and the hazard summary."""
    wb = gridio.open_wb(path, data_only=False)
    try:
        sheets = []
        for ws in wb.worksheets:
            bounds = _locate.true_used_range(ws)
            if bounds is None:
                used = None
                dims = {"rows": 0, "cols": 0}
            else:
                r0, c0, r1, c1 = bounds
                used = gridio.a1(r0, c0) + ":" + gridio.a1(r1, c1)
                dims = {"rows": r1 - r0 + 1, "cols": c1 - c0 + 1}
            entry = {
                "name": ws.title,
                "state": ws.sheet_state,
                "used_range": used,
                "dimensions": dims,
                "merged_cells": len(list(getattr(
                    ws, "merged_cells", []).ranges)) if getattr(
                    ws, "merged_cells", None) else 0,
            }
            # Row and column grouping, the read side of the outline
            # surface. Reported only when the sheet actually has groups, so
            # the ordinary workbook's metadata does not grow a key that is
            # always empty; a budget outline is invisible to a caller that
            # only reads values, and a collapsed group hides rows for a
            # reason that has nothing to do with a filter.
            outline = _format.sheet_outline(ws)
            if outline:
                entry["outline"] = outline
            sheets.append(entry)
        names = []
        for scope, name, defn in _locate._all_defined_names(wb):
            names.append({"name": name, "scope": scope,
                          "value": getattr(defn, "value", None)})
        tables = []
        for ws in wb.worksheets:
            tmap = getattr(ws, "tables", {})
            for tname in list(tmap):
                # index by name: TableList.items() yields ref STRINGS, not
                # Table objects (a latent AttributeError the re-audit caught;
                # this path crashed on any workbook containing a table)
                t = tmap[tname]
                tables.append({"name": tname, "sheet": ws.title, "ref": t.ref})
        rep = _hazard.scan_path(check_path(path, "scan workbook"))
        return {
            "file": os.path.abspath(path),
            "sheet_count": len(sheets),
            "sheets": sheets,
            "defined_names": names,
            "tables": tables,
            "active_sheet": wb.active.title if wb.active is not None else None,
            "hazards": {
                "clean": rep.clean,
                "would_lose": rep.would_lose,
                "labels": rep.labels(),
            },
        }
    finally:
        wb.close()


def diagnose_workbook(path: str) -> dict:
    """The hazard scan surfaced as a health readout: which fragile parts the
    file holds, whether an openpyxl edit would drop any of them, the routing
    recommendation, and a light integrity summary."""
    scan_path = check_path(path, "scan workbook")
    rep = _hazard.scan_path(scan_path)
    if rep.error is not None:
        return {"file": os.path.abspath(path), "readable": False,
                "error": rep.error}
    route_surgical, reason_s = _hazard.route(rep, edit="surgical")
    route_struct, reason_st = _hazard.route(rep, edit="structural")
    wb = gridio.open_wb(path, data_only=False)
    try:
        visible = [ws for ws in wb.worksheets
                   if ws.sheet_state == "visible"]
        formula_cells = 0
        for ws in wb.worksheets:
            cells = getattr(ws, "_cells", None)
            if cells is None:
                continue
            for cell in cells.values():
                if _calc.is_formula_cell(cell):
                    formula_cells += 1
        health = {
            "sheet_count": len(wb.worksheets),
            "visible_sheets": len(visible),
            "formula_cells": formula_cells,
            "keep_vba": scan_path.lower().endswith(".xlsm"),
        }
    finally:
        wb.close()
    return {
        "file": os.path.abspath(path),
        "readable": True,
        "hazards": rep.as_dict(),
        "routing": {
            "surgical_edit": {"route": route_surgical, "reason": reason_s},
            "structural_edit": {"route": route_struct, "reason": reason_st},
        },
        "health": health,
    }


# --------------------------------------------------------------- sheet manage


def manage_worksheet(path: str, action: str, sheet: str | None = None,
                     new_name: str | None = None, index: int | None = None,
                     state: str | None = None, allow_loss: bool = False,
                     backup: bool = True,
                     verify_com: bool | None = None) -> dict:
    """Add / delete / rename / copy / reorder / hide a worksheet. One backup +
    one verified save per call through WorkbookPackage."""
    if action not in WS_ACTIONS:
        raise XlMcpError(
            f"action must be one of {WS_ACTIONS}, got {action!r}")
    pkg = WorkbookPackage.open(path)
    wb = pkg.workbook

    def _need(name: str | None, what: str) -> str:
        if not name:
            raise XlMcpError(f"{action} needs {what}")
        if name not in wb.sheetnames:
            raise TargetNotFound(
                f"no sheet named {name!r} (sheets: "
                f"{', '.join(wb.sheetnames)})")
        return name

    detail: dict[str, Any] = {"action": action}
    if action == "add":
        name = new_name or sheet
        if not name:
            # It used to default to "Sheet", so a call that forgot new_name
            # silently created a differently named sheet than the caller
            # meant, and only the SECOND such call ever complained.
            raise XlMcpError(
                "add needs new_name (the title for the new sheet)")
        if name in wb.sheetnames:
            raise XlMcpError(f"a sheet named {name!r} already exists")
        _check_sheet_name(name)
        pos = index if index is not None else len(wb.sheetnames)
        ws = wb.create_sheet(title=name, index=pos)
        detail.update(sheet=name, index=wb.sheetnames.index(ws.title))
    elif action == "delete":
        name = _need(sheet, "the sheet to delete")
        if len(wb.sheetnames) == 1:
            raise XlMcpError(
                "cannot delete the only sheet; a workbook needs one visible "
                "sheet")
        visible = [w for w in wb.worksheets if w.sheet_state == "visible"]
        if len(visible) == 1 and visible[0].title == name:
            raise XlMcpError(
                f"cannot delete {name!r}: it is the only VISIBLE sheet (the "
                "others are hidden) and a workbook needs one visible sheet; "
                "unhide another sheet first")
        del wb[name]
        # A deleted sheet takes its own part plus any per-sheet companion
        # parts (comments, comment VML anchors, tables) with it; declare the
        # removal so the default-fail inventory verify does not read it as
        # silent loss. Fragile per-sheet parts (charts, pivots) stay under
        # the hazard gate and still need allow_loss or COM.
        pkg.expect_removal("xl/worksheets/", "xl/chartsheets/",
                           "xl/comments", "xl/drawings/commentsdrawing",
                           "xl/tables/")
        detail.update(sheet=name)
    elif action == "rename":
        name = _need(sheet, "the sheet to rename")
        if not new_name:
            raise XlMcpError("rename needs new_name")
        _check_sheet_name(new_name)
        if new_name in wb.sheetnames and new_name != name:
            raise XlMcpError(f"a sheet named {new_name!r} already exists")
        wb[name].title = new_name
        detail.update(sheet=name, new_name=new_name)
    elif action == "copy":
        name = _need(sheet, "the sheet to copy")
        src = wb[name]
        dup = wb.copy_worksheet(src)
        if new_name:
            _check_sheet_name(new_name)
            dup.title = new_name
        detail.update(sheet=name, new_name=dup.title)
    elif action == "reorder":
        name = _need(sheet, "the sheet to move")
        if index is None:
            raise XlMcpError("reorder needs index (0-based target position)")
        target = max(0, min(int(index), len(wb.sheetnames) - 1))
        cur = wb.sheetnames.index(name)
        wb.move_sheet(name, offset=target - cur)
        detail.update(sheet=name, index=target)
    elif action == "hide":
        name = _need(sheet, "the sheet to hide")
        st = _HIDE_STATES.get((state or "hidden").lower())
        if st is None:
            raise XlMcpError("state must be 'hidden' or 'very_hidden'")
        visible = [w for w in wb.worksheets if w.sheet_state == "visible"]
        if len(visible) == 1 and visible[0].title == name:
            raise XlMcpError(
                "cannot hide the only visible sheet; a workbook needs one")
        wb[name].sheet_state = st
        detail.update(sheet=name, state=st)
    else:  # unhide
        name = _need(sheet, "the sheet to unhide")
        wb[name].sheet_state = "visible"
        detail.update(sheet=name, state="visible")

    pkg._changed["worksheet"] = detail
    result = pkg.save(allow_loss=allow_loss, backup=backup,
                      verify_com=verify_com)
    return result


__all__ = [
    "create_workbook", "copy_workbook", "get_workbook_metadata",
    "diagnose_workbook", "manage_worksheet", "WS_ACTIONS",
]
