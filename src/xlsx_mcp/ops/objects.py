"""ops/objects.py: images and charts (DESIGN Section 11; design pack as re-cut).

manage_image inserts, lists, deletes, and extracts cell-anchored images.
openpyxl only carries images through a load+save when Pillow is installed,
so insert requires Pillow (a declared dependency of this pack) and the
mutating actions register expect_preserved for the media/drawings hazard
keys ONLY after confirming every drawing in the workbook is picture-only
(rels all target images, drawing XML holds no sp/grpSp/cxnSp). The gate
refusal becomes a warning on that observed-safe path; verify-after-write
still refuses the save if a media part is actually lost, so the safety net
never weakens. A workbook whose drawings hold real shapes keeps the loud
HAZARD_REFUSED, because those shapes genuinely die on a file-based save.

manage_chart creates basic chart types from a data range, lists, and
deletes. Chart fidelity is model-mediated: openpyxl re-serializes every
chart through its own model, so a complex Excel-authored chart may lose
sub-features it does not know; the hazard scan tracks charts as a DEGRADES
hazard and the save warns rather than refuses. Deliberate delete registers
expect_removal so the default-fail part checks recognize the intent.

Sparklines are NOT here: they live in an x14 extension inside the sheet XML
that openpyxl does not model; the deferral is recorded in DESIGN.md.

Both mutators run a post-save object-count read-back (the object analog of
the cell content read-back): a mismatch restores the backup and refuses.
"""

from __future__ import annotations

import re
import zipfile
from typing import Any

from openpyxl.utils import get_column_letter

from ..core import hazard as _hazard
from ..core.errors import (
    TargetNotFound,
    ValidationFailed,
    WorkbookNotFound,
    XlMcpError,
)
from ..core.package import WorkbookPackage
from . import gridio

IMAGE_ACTIONS = ("insert", "list", "delete", "extract")
CHART_ACTIONS = ("create", "list", "delete")

CHART_TYPES = ("bar", "bar_horizontal", "line", "pie", "doughnut", "area",
               "scatter")

_IMG_REL_SUFFIX = "/image"
_TYPE_RE = re.compile(r'Type="([^"]+)"')
#: shape locals that mean real shape content (a picture, 'pic', is fine here)
_REAL_SHAPE_LOCALS = frozenset({"sp", "grpSp", "cxnSp"})


# ------------------------------------------------------------- shared helpers


def _pillow_available() -> bool:
    try:
        import PIL  # noqa: F401
        return True
    except ImportError:
        return False


def _picture_only_drawings(path: str, report) -> bool:
    """True when every part of the drawings hazard is a picture-only drawing:
    its rels exist and target only images, and the drawing XML carries no
    real shape element (sp / grpSp / cxnSp; pic is the picture itself).
    OBSERVED on this machine (openpyxl 3.1.5 + Pillow 12.3): such drawings
    and their xl/media/ payloads round-trip intact through load+save, for
    Excel-authored and openpyxl-authored images alike. Anything else stays
    conservatively unsafe."""
    hz = next((h for h in report.hazards if h.key == "drawings"), None)
    if hz is None:
        return False
    drawing_xmls = [p for p in hz.parts
                    if "/_rels/" not in p.lower()
                    and p.lower().endswith(".xml")]
    if not drawing_xmls:
        return False
    try:
        with zipfile.ZipFile(path) as zf:
            names = set(zf.namelist())
            for part in drawing_xmls:
                head, base = part.rsplit("/", 1)
                rels = f"{head}/_rels/{base}.rels"
                if rels not in names:
                    return False
                types = _TYPE_RE.findall(
                    zf.read(rels).decode("utf-8", "replace"))
                if not types or not all(
                        t.rstrip("/").lower().endswith(_IMG_REL_SUFFIX)
                        for t in types):
                    return False
                from xml.etree.ElementTree import fromstring
                root = fromstring(zf.read(part))
                for el in root.iter():
                    tag = el.tag
                    local = tag.rsplit("}", 1)[-1] if isinstance(tag, str) \
                        else ""
                    if local in _REAL_SHAPE_LOCALS:
                        return False
    except Exception:  # noqa: BLE001
        return False
    return True


def _arm_image_preservation(pkg: WorkbookPackage) -> None:
    """Downgrade the media/drawings gate refusal to a warning when the
    observed-safe preconditions hold (Pillow present, picture-only
    drawings). Otherwise leave the gate alone so it refuses honestly."""
    rep = pkg.hazard
    keys = set(rep.lossy_keys)
    if not ({"media", "drawings"} & keys):
        return
    if not _pillow_available():
        return
    if "drawings" in keys and not _picture_only_drawings(pkg.path, rep):
        return
    if "drawings" not in keys and "media" in keys:
        # media with no picture drawing means it is anchored some other way
        # (header/footer VML, themes); no observation covers that, refuse.
        return
    pkg.expect_preserved("media", "drawings")


def _anchor_cell(obj) -> str | None:
    """Best-effort A1 anchor of an image or chart from its anchor object."""
    anchor = getattr(obj, "anchor", None)
    frm = getattr(anchor, "_from", None)
    if frm is None:
        return None
    try:
        return f"{get_column_letter(frm.col + 1)}{frm.row + 1}"
    except Exception:  # noqa: BLE001
        return None


def _restore_and_refuse(pkg: WorkbookPackage, backup: bool, msg: str):
    restored = False
    if backup:
        restored = pkg._restore_from_backup()
    raise ValidationFailed(
        msg + (" The file was restored from the backup." if restored
               else " No backup was taken (backup=false), so the saved "
                    "file stands; inspect it before further edits."))


# ------------------------------------------------------------------ images


def manage_image(path: str, action: str, image_file: str | None = None,
                 location: Any = None, sheet: str | None = None,
                 width: int | None = None, height: int | None = None,
                 index: int | None = None, out_dir: str | None = None,
                 allow_loss: bool = False, backup: bool = True) -> dict:
    """Insert / list / delete / extract images. Backup + verify on write."""
    if action not in IMAGE_ACTIONS:
        raise XlMcpError(
            f"action must be one of {IMAGE_ACTIONS}, got {action!r}")

    if action == "list":
        wb = gridio.open_wb(path)
        try:
            out = []
            for ws in (wb.worksheets if sheet is None else [wb[sheet]]):
                for i, img in enumerate(getattr(ws, "_images", []), start=1):
                    # display size comes from the anchor extent (EMU, 9525
                    # per pixel) when present; openpyxl's img.width is the
                    # source file's pixel size, not the placed size
                    w = getattr(img, "width", None)
                    h = getattr(img, "height", None)
                    ext = getattr(getattr(img, "anchor", None), "ext", None)
                    if ext is not None and ext.cx and ext.cy:
                        w = round(ext.cx / 9525)
                        h = round(ext.cy / 9525)
                    out.append({
                        "sheet": ws.title, "index": i,
                        "anchor": _anchor_cell(img),
                        "width": w, "height": h,
                        "format": getattr(img, "format", None)})
            return {"images": out, "count": len(out)}
        finally:
            wb.close()

    if action == "extract":
        if not out_dir:
            raise XlMcpError("extract needs out_dir (a directory to write "
                             "the image files into)")
        import os
        from ..core.sandbox import check_path
        od = check_path(out_dir, "write extracted images")
        if not os.path.isdir(od):
            raise XlMcpError(f"out_dir is not an existing directory: {od}")
        wp = check_path(path, "open workbook")
        if not os.path.exists(wp):
            raise WorkbookNotFound(f"no such workbook: {wp}")
        written = []
        with zipfile.ZipFile(wp) as zf:
            media = [n for n in zf.namelist()
                     if n.lower().startswith("xl/media/")]
            for name in sorted(media):
                base = name.rsplit("/", 1)[-1]
                dest = os.path.join(od, base)
                with open(dest, "wb") as fh:
                    fh.write(zf.read(name))
                written.append(dest)
        return {"extracted": written, "count": len(written)}

    # ------------------------------------------------------------- mutations
    if not _pillow_available():
        raise XlMcpError(
            "image mutations need the Pillow library, which is not "
            "importable in this environment")

    pkg = WorkbookPackage.open(path)
    _arm_image_preservation(pkg)

    if action == "insert":
        if not image_file:
            raise XlMcpError("insert needs image_file (path to the image)")
        from ..core.sandbox import check_path
        ip = check_path(image_file, "read image file")
        import os
        if not os.path.exists(ip):
            raise TargetNotFound(f"no such image file: {ip}")
        from openpyxl.drawing.image import Image as _XLImage
        try:
            img = _XLImage(ip)
        except Exception as exc:  # noqa: BLE001
            raise XlMcpError(
                f"{ip} could not be read as an image ({exc})")
        grid = pkg.resolve(location if location is not None
                           else {"cell": "A1"}, default_sheet=sheet)
        if width is not None:
            img.width = int(width)
        if height is not None:
            img.height = int(height)
        ws = pkg.workbook[grid.sheet]
        pre_count = len(ws._images)
        anchor = f"{get_column_letter(grid.min_col)}{grid.min_row}"
        ws.add_image(img, anchor)
        pkg._changed["image"] = {"sheet": grid.sheet, "inserted": anchor,
                                 "source": ip}
        result = pkg.save(allow_loss=allow_loss, backup=backup)
        expected, verb = pre_count + 1, "inserted"
        target_sheet = grid.sheet
    else:  # delete
        if sheet is None:
            raise XlMcpError("delete needs sheet (and index from list, or "
                             "location as the image's anchor cell)")
        ws = pkg.workbook[sheet]
        images = list(ws._images)
        if not images:
            raise TargetNotFound(f"sheet {sheet!r} has no images")
        target = None
        if index is not None:
            if not 1 <= index <= len(images):
                raise TargetNotFound(
                    f"index {index} out of range; sheet {sheet!r} has "
                    f"{len(images)} image(s)")
            target = images[index - 1]
        elif location is not None:
            grid = pkg.resolve(location, default_sheet=sheet)
            want = f"{get_column_letter(grid.min_col)}{grid.min_row}"
            hits = [im for im in images if _anchor_cell(im) == want]
            if not hits:
                raise TargetNotFound(
                    f"no image anchored at {want} on {sheet!r}; anchors: "
                    + ", ".join(str(_anchor_cell(im)) for im in images))
            if len(hits) > 1:
                raise XlMcpError(
                    f"{len(hits)} images share the anchor {want}; delete "
                    "by index instead (see action list)")
            target = hits[0]
        else:
            raise XlMcpError("delete needs index or location")
        ws._images.remove(target)
        pre_count = len(images)
        pkg.expect_removal("xl/media/", "xl/drawings/")
        pkg._changed["image"] = {"sheet": sheet,
                                 "deleted": _anchor_cell(target)}
        result = pkg.save(allow_loss=allow_loss, backup=backup)
        expected, verb = pre_count - 1, "deleted"
        target_sheet = sheet

    # post-save object read-back (the image analog of cell read-back)
    wb = gridio.open_wb(path)
    try:
        got = len(wb[target_sheet]._images)
    finally:
        wb.close()
    if got != expected:
        _restore_and_refuse(
            pkg, backup,
            f"after the save {target_sheet!r} holds {got} image(s) where "
            f"{expected} were expected (one {verb}).")
    result["changed"]["images_on_sheet"] = got
    return result


# ------------------------------------------------------------------- charts


def _title_text(title) -> str | None:
    if title is None:
        return None
    try:
        rich = title.tx.rich
        return "".join(
            (r.t or "") for p in rich.p for r in (p.r or [])) or None
    except Exception:  # noqa: BLE001
        return None


def _build_chart(chart_type: str):
    from openpyxl import chart as _c
    if chart_type == "bar":
        ch = _c.BarChart()
        ch.type = "col"
    elif chart_type == "bar_horizontal":
        ch = _c.BarChart()
        ch.type = "bar"
    elif chart_type == "line":
        ch = _c.LineChart()
    elif chart_type == "pie":
        ch = _c.PieChart()
    elif chart_type == "doughnut":
        ch = _c.DoughnutChart()
    elif chart_type == "area":
        ch = _c.AreaChart()
    else:  # scatter
        ch = _c.ScatterChart()
        ch.style = 13
    return ch


def manage_chart(path: str, action: str, chart_type: str | None = None,
                 data: Any = None, categories: Any = None,
                 title: str | None = None, x_title: str | None = None,
                 y_title: str | None = None, anchor: str | None = None,
                 sheet: str | None = None, index: int | None = None,
                 titles_from_data: bool = True, allow_loss: bool = False,
                 backup: bool = True) -> dict:
    """Create / list / delete charts (openpyxl chart model). Backup + verify
    on write."""
    if action not in CHART_ACTIONS:
        raise XlMcpError(
            f"action must be one of {CHART_ACTIONS}, got {action!r}")

    if action == "list":
        wb = gridio.open_wb(path)
        try:
            out = []
            for ws in (wb.worksheets if sheet is None else [wb[sheet]]):
                for i, ch in enumerate(getattr(ws, "_charts", []), start=1):
                    out.append({
                        "sheet": ws.title, "index": i,
                        "type": type(ch).__name__,
                        "title": _title_text(getattr(ch, "title", None)),
                        "anchor": _anchor_cell(ch)})
            return {"charts": out, "count": len(out)}
        finally:
            wb.close()

    pkg = WorkbookPackage.open(path)

    if action == "create":
        if chart_type not in CHART_TYPES:
            raise XlMcpError(
                f"chart_type must be one of {CHART_TYPES}, got "
                f"{chart_type!r}")
        if data is None:
            raise XlMcpError("create needs data (a location object for the "
                             "data range, headers in the first row)")
        grid = pkg.resolve(data, default_sheet=sheet)
        if grid.is_single:
            raise XlMcpError(
                f"the data range resolved to the single cell {grid.a1}; a "
                "chart needs a rectangle of data")
        from openpyxl.chart import Reference, Series
        ws = pkg.workbook[grid.sheet]
        ch = _build_chart(chart_type)
        if chart_type == "scatter":
            # first data column is X, each further column is a Y series
            if grid.max_col - grid.min_col < 1:
                raise XlMcpError(
                    "scatter needs at least two columns (x values, then "
                    "one or more y series)")
            xref = Reference(ws, min_col=grid.min_col,
                             min_row=grid.min_row + (1 if titles_from_data
                                                     else 0),
                             max_row=grid.max_row)
            for c in range(grid.min_col + 1, grid.max_col + 1):
                yref = Reference(ws, min_col=c, min_row=grid.min_row,
                                 max_row=grid.max_row)
                ser = Series(yref, xref,
                             title_from_data=titles_from_data)
                ch.series.append(ser)
        else:
            dref = Reference(ws, min_col=grid.min_col, min_row=grid.min_row,
                             max_col=grid.max_col, max_row=grid.max_row)
            ch.add_data(dref, titles_from_data=titles_from_data)
            if categories is not None:
                cgrid = pkg.resolve(categories, default_sheet=grid.sheet)
                cref = Reference(
                    pkg.workbook[cgrid.sheet], min_col=cgrid.min_col,
                    min_row=cgrid.min_row, max_col=cgrid.max_col,
                    max_row=cgrid.max_row)
                ch.set_categories(cref)
        if title:
            ch.title = title
        if x_title:
            ch.x_axis.title = x_title
        if y_title:
            ch.y_axis.title = y_title
        at = anchor or (f"{get_column_letter(min(grid.max_col + 2, 16384))}"
                        f"{grid.min_row}")
        pre_count = len(ws._charts)
        ws.add_chart(ch, at)
        pkg._changed["chart"] = {"sheet": grid.sheet, "created": chart_type,
                                 "anchor": at, "data": grid.a1}
        result = pkg.save(allow_loss=allow_loss, backup=backup)
        expected, verb = pre_count + 1, "created"
        target_sheet = grid.sheet
    else:  # delete
        if sheet is None:
            raise XlMcpError("delete needs sheet (and index or title from "
                             "action list)")
        ws = pkg.workbook[sheet]
        charts = list(getattr(ws, "_charts", []))
        if not charts:
            raise TargetNotFound(f"sheet {sheet!r} has no charts")
        target = None
        if index is not None:
            if not 1 <= index <= len(charts):
                raise TargetNotFound(
                    f"index {index} out of range; sheet {sheet!r} has "
                    f"{len(charts)} chart(s)")
            target = charts[index - 1]
        elif title is not None:
            hits = [c for c in charts
                    if _title_text(getattr(c, "title", None)) == title]
            if not hits:
                raise TargetNotFound(
                    f"no chart titled {title!r} on {sheet!r}; titles: "
                    + ", ".join(str(_title_text(getattr(c, "title", None)))
                                for c in charts))
            if len(hits) > 1:
                raise XlMcpError(
                    f"{len(hits)} charts share the title {title!r}; delete "
                    "by index instead")
            target = hits[0]
        else:
            raise XlMcpError("delete needs index or title")
        ws._charts.remove(target)
        pre_count = len(charts)
        pkg.expect_removal("xl/charts/", "xl/drawings/")
        pkg._changed["chart"] = {"sheet": sheet,
                                 "deleted": type(target).__name__}
        result = pkg.save(allow_loss=allow_loss, backup=backup)
        expected, verb = pre_count - 1, "deleted"
        target_sheet = sheet

    # post-save object read-back (the chart analog of cell read-back)
    wb = gridio.open_wb(path)
    try:
        got = len(wb[target_sheet]._charts)
    finally:
        wb.close()
    if got != expected:
        _restore_and_refuse(
            pkg, backup,
            f"after the save {target_sheet!r} holds {got} chart(s) where "
            f"{expected} were expected (one {verb}).")
    result["changed"]["charts_on_sheet"] = got
    return result


__all__ = ["manage_image", "manage_chart", "IMAGE_ACTIONS", "CHART_ACTIONS",
           "CHART_TYPES"]
