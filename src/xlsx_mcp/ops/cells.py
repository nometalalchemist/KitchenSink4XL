"""ops/cells.py: cell and range read / write / copy / move, plus the
token-shaped read (query_range), the headline capability.

Reads (read_range, query_range) load the workbook directly and never mutate.
Writes (set_cell, write_range, clear_range, copy_range, move_range) route their
mutation through WorkbookPackage, so the hazard gate, the backup, and
verify-after-write run automatically. Formula writes are normalized through the
_xlfn shim and flag fullCalcOnLoad, so a modern function does not land as
#NAME? and the next Excel open recalculates.

query_range is the context-safety answer (DESIGN 1.5, demand D1): rather than
pull a whole sheet into the agent's context, it filters rows, projects a column
subset, sorts, paginates, and aggregates SERVER-SIDE, returning only the slice
or summary the caller asked for.

Phase 3c additions: get_cells / set_cells, the SCATTER pair complementing the
rectangular read_range / write_range (many individually addressed cells in one
call, one save, one verify), and set_merge (merge / unmerge / list with the
Excel semantics: merge keeps the top-left value, and absorbing cells that hold
values refuses without an explicit confirm flag rather than discarding data
silently).
"""

from __future__ import annotations

from copy import copy as _copy
from typing import Any

from ..core import calc as _calc
from ..core import refs as _refs
from ..core.errors import RangeOutOfBounds, TargetNotFound, XlMcpError
from ..core.package import WorkbookPackage
from . import gridio

CLEAR_WHAT = ("contents", "formats", "all")
COPY_WHAT = ("all", "values", "formulas", "formats")
MERGE_ACTIONS = ("merge", "unmerge", "list")

#: A single scatter call over more cells than this refuses and asks for
#: read_range / write_range / apply_edits, which are shaped for bulk.
MAX_SCATTER_CELLS = 1_000


# ------------------------------------------------------------------- reads


def read_range(path: str, location: Any, values: str = "cached",
               sheet: str | None = None) -> dict:
    """Read a cell or range. values='cached' returns last calculated values,
    'formula' the formula strings, 'both' pairs each cell with its label."""
    if values not in gridio.VALUE_MODES:
        raise XlMcpError(
            f"values must be one of {gridio.VALUE_MODES}, got {values!r}")
    formula_wb = gridio.open_wb(path, data_only=False) \
        if values in ("formula", "both") else None
    cached_wb = gridio.open_wb(path, data_only=True) \
        if values in ("cached", "both") else None
    try:
        base = formula_wb if formula_wb is not None else cached_wb
        grid = gridio.resolve(base, location, default_sheet=sheet,
                              path=path)
        if grid.empty:
            return {"sheet": grid.sheet, "range": grid.a1, "empty": True,
                    "values": [], "value_mode": values}
        gridio.guard_cell_count(grid)
        vals, labels, has_formula = gridio.read_matrix(
            grid, mode=values, formula_wb=formula_wb, cached_wb=cached_wb,
            path=path)
        vals = [[gridio.compact_value(v) for v in row] for row in vals]
        out: dict[str, Any] = {
            "sheet": grid.sheet, "range": grid.a1,
            "rows": len(vals), "cols": len(vals[0]) if vals else 0,
            "value_mode": values, "values": vals,
        }
        if has_formula and values != "formula":
            out["labels"] = labels
            note = gridio.absent_note(labels)
            if note:
                out["warning"] = note
        return out
    finally:
        for wb in (formula_wb, cached_wb):
            if wb is not None:
                wb.close()


# --------------------------------------------------------------- basic writes


def set_cell(path: str, location: Any, value: Any, sheet: str | None = None,
             allow_loss: bool = False, backup: bool = True) -> dict:
    """Write ONE cell. A string beginning with '=' is stored as a formula
    (normalized); anything else is a literal. One backup + one verified save."""
    pkg = WorkbookPackage.open(path)
    grid = pkg.resolve(location, default_sheet=sheet)
    if not grid.is_single:
        raise XlMcpError(
            f"set_cell needs a single cell; {grid.a1} is a range. Use "
            "write_range for a block.")
    pkg.set_cell(grid.sheet, gridio.a1(grid.min_row, grid.min_col), value)
    return pkg.save(allow_loss=allow_loss, backup=backup)


def write_range(path: str, location: Any, data: list[list[Any]],
                sheet: str | None = None, allow_loss: bool = False,
                backup: bool = True) -> dict:
    """Write a 2D block of values/formulas anchored at the location's top-left.
    Formula strings ('=...') are normalized. One backup + one verified save."""
    if not isinstance(data, list) or (data and not all(
            isinstance(r, list) for r in data)):
        raise XlMcpError("data must be a 2D array (list of row lists)")
    widths = {len(r) for r in data}
    if len(widths) > 1:
        # A ragged block used to be accepted and written row by row, so the
        # short rows quietly left whatever was underneath them in place.
        raise XlMcpError(
            f"data rows have different lengths {sorted(widths)}; a written "
            "block must be rectangular (pad the short rows with null to "
            "clear those cells)")
    pkg = WorkbookPackage.open(path)
    grid = pkg.resolve(location, default_sheet=sheet)
    rows = len(data)
    cols = max((len(r) for r in data), default=0)
    if rows == 0 or cols == 0:
        raise XlMcpError("data is empty; nothing to write")
    top, left = grid.min_row, grid.min_col
    if (top + rows - 1) > gridio._locate.MAX_ROW or \
            (left + cols - 1) > gridio._locate.MAX_COL:
        raise RangeOutOfBounds(
            "the block would extend past the grid limits from its anchor")
    if rows * cols > gridio.MAX_READ_CELLS:
        raise RangeOutOfBounds(
            f"the block is {rows * cols:,} cells, over the "
            f"{gridio.MAX_READ_CELLS:,}-cell write ceiling; split it")
    for i, row in enumerate(data):
        for j, val in enumerate(row):
            pkg.set_cell(grid.sheet, gridio.a1(top + i, left + j), val)
    result = pkg.save(allow_loss=allow_loss, backup=backup)
    result["changed"]["anchor"] = gridio.a1(top, left)
    result["changed"]["shape"] = {"rows": rows, "cols": cols}
    return result


def clear_range(path: str, location: Any, what: str = "contents",
                sheet: str | None = None, allow_loss: bool = False,
                backup: bool = True) -> dict:
    """Clear a cell or range: 'contents' (values/formulas), 'formats' (styles),
    or 'all'. One backup + one verified save."""
    if what not in CLEAR_WHAT:
        raise XlMcpError(f"what must be one of {CLEAR_WHAT}, got {what!r}")
    pkg = WorkbookPackage.open(path)
    grid = pkg.resolve(location, default_sheet=sheet)
    gridio.guard_cell_count(grid)
    ws = pkg.workbook[grid.sheet]
    from openpyxl.styles import Alignment, Border, Font, PatternFill
    n = 0
    for r in range(grid.min_row, grid.max_row + 1):
        for c in range(grid.min_col, grid.max_col + 1):
            cell = ws.cell(r, c)
            if what in ("contents", "all"):
                if cell.value is not None:
                    cell.value = None
                    pkg._intended[(grid.sheet, gridio.a1(r, c))] = (
                        "value", None)
                    n += 1
            if what in ("formats", "all"):
                cell.font = Font()
                cell.fill = PatternFill()
                cell.border = Border()
                cell.alignment = Alignment()
                cell.number_format = "General"
    result = pkg.save(allow_loss=allow_loss, backup=backup)
    result["changed"]["cleared"] = {"range": grid.a1, "what": what}
    return result


# --------------------------------------------------------------- copy / move


def _read_block(ws, grid, data_only_ws=None):
    """Buffer (value, style, cached, text_lookalike) for each cell of a
    rectangle. The last flag marks a TEXT cell whose string starts with '='
    (import_data's neutralized injection text), which must be written back as
    text rather than re-armed into a live formula."""
    block = []
    for r in range(grid.min_row, grid.max_row + 1):
        row = []
        for c in range(grid.min_col, grid.max_col + 1):
            cell = ws.cell(r, c)
            cached = data_only_ws.cell(r, c).value if data_only_ws else None
            row.append((cell.value, _copy(cell._style), cached,
                        _calc.looks_like_formula_text(cell)))
        block.append(row)
    return block


def _place(cell, value, *, as_text: bool) -> None:
    """Write a buffered value into a cell, keeping neutralized text TEXT."""
    cell.value = value
    if as_text and isinstance(value, str):
        cell.data_type = "s"


def copy_range(path: str, source: Any, dest: Any, what: str = "all",
               adjust_formulas: bool = True, sheet: str | None = None,
               allow_loss: bool = False, backup: bool = True) -> dict:
    """Copy a source rectangle to a destination anchor: 'all', 'values',
    'formulas', or 'formats'. Relative references in copied formulas shift by
    the paste offset (Excel semantics) unless adjust_formulas=false. One backup
    + one verified save."""
    if what not in COPY_WHAT:
        raise XlMcpError(f"what must be one of {COPY_WHAT}, got {what!r}")
    pkg = WorkbookPackage.open(path)
    src = pkg.resolve(source, default_sheet=sheet)
    dst = pkg.resolve(dest, default_sheet=sheet)
    gridio.guard_cell_count(src)
    sws = pkg.workbook[src.sheet]
    dws = pkg.workbook[dst.sheet]
    dows = None
    if what == "values":
        dob = gridio.open_wb(path, data_only=True)
        dows = dob[src.sheet]
    block = _read_block(sws, src, dows)
    dr = dst.min_row - src.min_row
    dc = dst.min_col - src.min_col
    if (dst.min_row + src.max_row - src.min_row) > gridio._locate.MAX_ROW or \
            (dst.min_col + src.max_col - src.min_col) > gridio._locate.MAX_COL:
        raise RangeOutOfBounds("the paste would extend past the grid limits")
    wrote_formula = False
    for i, row in enumerate(block):
        for j, (val, style, cached, is_text) in enumerate(row):
            tcell = dws.cell(dst.min_row + i, dst.min_col + j)
            if what in ("all", "values", "formulas"):
                out = val
                as_text = is_text
                if what == "values":
                    out = cached if (isinstance(val, str)
                                     and val.startswith("=")
                                     and not is_text) else val
                elif isinstance(val, str) and val.startswith("=") \
                        and not is_text:
                    if adjust_formulas:
                        out = _refs.offset_formula(val, dr, dc, src.sheet)
                    wrote_formula = True
                _place(tcell, out, as_text=as_text)
            if what in ("all", "formats"):
                tcell._style = _copy(style)
    if dows is not None:
        dob.close()
    if wrote_formula:
        pkg._formula_written = True
    result = pkg.save(allow_loss=allow_loss, backup=backup)
    result["changed"]["copied"] = {
        "from": f"{src.sheet}!{src.a1}",
        "to": f"{dst.sheet}!{gridio.a1(dst.min_row, dst.min_col)}",
        "what": what}
    return result


def move_range(path: str, source: Any, dest: Any, sheet: str | None = None,
               allow_loss: bool = False, backup: bool = True) -> dict:
    """Move a rectangle to a new anchor on the SAME sheet, rewriting every
    reference that pointed into the source so the workbook stays coherent
    (references follow the cells, Excel move semantics). One backup + one
    verified save."""
    pkg = WorkbookPackage.open(path)
    src = pkg.resolve(source, default_sheet=sheet)
    dst = pkg.resolve(dest, default_sheet=sheet)
    if src.sheet != dst.sheet:
        raise XlMcpError(
            "move_range is single-sheet; use copy_range across sheets, then "
            "clear_range the source")
    gridio.guard_cell_count(src)
    ws = pkg.workbook[src.sheet]
    block = _read_block(ws, src)
    dr = dst.min_row - src.min_row
    dc = dst.min_col - src.min_col
    if dr == 0 and dc == 0:
        raise XlMcpError("the destination equals the source; nothing to move")
    if (dst.min_row + src.max_row - src.min_row) > gridio._locate.MAX_ROW or \
            (dst.min_col + src.max_col - src.min_col) > gridio._locate.MAX_COL:
        raise RangeOutOfBounds("the move would extend past the grid limits")
    from openpyxl.styles import Alignment, Border, Font, PatternFill
    # Clear the source first (buffered), then place at the destination.
    for r in range(src.min_row, src.max_row + 1):
        for c in range(src.min_col, src.max_col + 1):
            cell = ws.cell(r, c)
            cell.value = None
            cell.font = Font(); cell.fill = PatternFill()
            cell.border = Border(); cell.alignment = Alignment()
            cell.number_format = "General"
    wrote_formula = False
    for i, row in enumerate(block):
        for j, (val, style, _cached, is_text) in enumerate(row):
            tcell = ws.cell(dst.min_row + i, dst.min_col + j)
            _place(tcell, val, as_text=is_text)
            tcell._style = _copy(style)
            if isinstance(val, str) and val.startswith("=") and not is_text:
                wrote_formula = True
    # References that pointed into the source rectangle now follow it to dest.
    edit = _refs.RefEdit(
        src.sheet, _refs.MOVE,
        src=(src.min_row, src.min_col, src.max_row, src.max_col),
        dst=(dst.min_row, dst.min_col))
    report = _refs.rewrite_workbook(pkg.workbook, edit)
    if wrote_formula or report.formulas:
        pkg._formula_written = True
    result = pkg.save(allow_loss=allow_loss, backup=backup)
    result["changed"]["moved"] = {
        "from": f"{src.sheet}!{src.a1}",
        "to": f"{src.sheet}!{gridio.a1(dst.min_row, dst.min_col)}",
        "reference_rewrites": report.as_dict()}
    return result


# ------------------------------------------------ query_range (headline read)


_OPS_BINARY = {"eq", "ne", "gt", "ge", "lt", "le",
               "contains", "startswith", "endswith", "regex"}
_OPS_SET = {"in", "not_in"}
_OPS_UNARY = {"is_blank", "not_blank"}
_AGG_FUNCS = {"count", "count_nonblank", "count_distinct", "sum", "avg",
              "mean", "min", "max", "first", "last"}


def _num(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return v
    if isinstance(v, str):
        try:
            return float(v)
        except ValueError:
            return None
    return None


def _cmp(cell, op: str, target) -> bool:
    if op in _OPS_UNARY:
        blank = cell is None or (isinstance(cell, str) and cell == "")
        return blank if op == "is_blank" else not blank
    if op in _OPS_SET:
        seq = target if isinstance(target, list) else [target]
        hit = any(_scalar_eq(cell, t) for t in seq)
        return hit if op == "in" else not hit
    if op in ("contains", "startswith", "endswith"):
        s = "" if cell is None else str(cell)
        t = "" if target is None else str(target)
        s, t = s.lower(), t.lower()
        return {"contains": t in s, "startswith": s.startswith(t),
                "endswith": s.endswith(t)}[op]
    if op == "regex":
        from ..core import _regex
        s = "" if cell is None else str(cell)
        return bool(_regex.finditer(str(target), s))
    # ordered / equality comparisons: numeric when both coerce, else string
    cn, tn = _num(cell), _num(target)
    if cn is not None and tn is not None:
        a, b = cn, tn
    elif op in ("eq", "ne"):
        return _scalar_eq(cell, target) if op == "eq" \
            else not _scalar_eq(cell, target)
    else:
        # Blank cells never satisfy an ordered comparison (Excel's filter
        # behavior: "less than 5" excludes blanks). Before the re-audit a
        # blank coerced to "" and "" <= "5" let blank rows through lt/le.
        if cell is None or cell == "":
            return False
        a = str(cell)
        b = "" if target is None else str(target)
    return {"eq": a == b, "ne": a != b, "gt": a > b, "ge": a >= b,
            "lt": a < b, "le": a <= b}[op]


def _scalar_eq(a, b) -> bool:
    an, bn = _num(a), _num(b)
    if an is not None and bn is not None:
        return an == bn
    return ("" if a is None else str(a)) == ("" if b is None else str(b))


def query_range(path: str, location: Any = None, sheet: str | None = None,
                header: bool = True, columns: list | None = None,
                where: list | None = None, match: str = "all",
                order_by: list | None = None, aggregate: list | None = None,
                group_by: Any = None, limit: int | None = None,
                offset: int = 0, distinct: bool = False,
                records: bool = False, values: str = "cached") -> dict:
    """Server-side filter / project / sort / paginate / aggregate over a range,
    so an agent reads only the rows and columns it needs.

    location defaults to the sheet's TRUE used range. With header=true the first
    row names the columns (referenced by name in columns/where/order_by/
    aggregate/group_by); otherwise columns are the A1 letters. where is a list
    of {column, op, value} predicates combined by match ('all'|'any'); ops:
    eq ne gt ge lt le contains startswith endswith regex in not_in is_blank
    not_blank. aggregate is a list of {column, func} (count count_nonblank
    count_distinct sum avg min max first last), optionally per group_by. Returns
    a compact projection (arrays by default, records=true for objects) with the
    matched / returned / scanned counts. Read-only; nothing is written.
    """
    if match not in ("all", "any"):
        raise XlMcpError("match must be 'all' or 'any'")
    if values not in gridio.VALUE_MODES:
        raise XlMcpError(f"values must be one of {gridio.VALUE_MODES}")
    data_only = values != "formula"
    wb = gridio.open_wb(path, data_only=data_only)
    try:
        loc = location if location is not None else {
            "used_range": sheet if sheet is not None else True}
        grid = gridio.resolve(wb, loc, default_sheet=sheet, path=path)
        if grid.empty:
            return {"sheet": grid.sheet, "source": grid.a1, "mode": "rows",
                    "columns": [], "rows": [], "matched": 0, "returned": 0,
                    "scanned": 0, "truncated": False}
        gridio.guard_cell_count(grid)
        ws = wb[grid.sheet]
        matrix = [[gridio.compact_value(ws.cell(r, c).value)
                   for c in range(grid.min_col, grid.max_col + 1)]
                  for r in range(grid.min_row, grid.max_row + 1)]
        # STALENESS: on the cached load an uncalculated formula cell is
        # indistinguishable from a blank one, so a filter skipped it and an
        # aggregate summed around it while the caller saw a clean number.
        # Probe only when a blank actually appeared in the rectangle.
        stale_cells: list[str] = []
        if data_only and any(v is None for row in matrix for v in row):
            mask = gridio.formula_mask(path, grid)
            if mask:
                for (r, c) in sorted(mask):
                    if matrix[r - grid.min_row][c - grid.min_col] is None:
                        stale_cells.append(gridio.a1(r, c))
        stale_note = (
            f"{len(stale_cells)} cell(s) in {grid.a1} hold formulas with NO "
            "cached value; they were read as BLANK, so filters skipped them "
            "and aggregates were computed without them. Run recalculate (com "
            "pack) or open in Excel before trusting these numbers"
        ) if stale_cells else None
        # column names + data rows
        from openpyxl.utils import get_column_letter
        letters = [get_column_letter(c)
                   for c in range(grid.min_col, grid.max_col + 1)]
        if header and matrix:
            col_names = [str(h) if h is not None else letters[i]
                         for i, h in enumerate(matrix[0])]
            body = matrix[1:]
        else:
            col_names = letters
            body = matrix
        idx = {name: i for i, name in enumerate(col_names)}
        # allow addressing header columns by letter too
        for i, lt in enumerate(letters):
            idx.setdefault(lt, i)

        def col_i(name) -> int:
            if isinstance(name, int):
                if 1 <= name <= len(col_names):
                    return name - 1
                raise XlMcpError(f"column index {name} out of range")
            if name in idx:
                return idx[name]
            raise XlMcpError(
                f"no column {name!r}; columns are {col_names}")

        scanned = len(body)
        # filter
        preds = where or []
        for p in preds:
            if not isinstance(p, dict) or "column" not in p or "op" not in p:
                raise XlMcpError(
                    "each where clause is {column, op, value}")
            if p["op"] not in (_OPS_BINARY | _OPS_SET | _OPS_UNARY):
                raise XlMcpError(f"unknown op {p['op']!r}")

        def keep(row) -> bool:
            results = []
            for p in preds:
                ci = col_i(p["column"])
                cell = row[ci] if ci < len(row) else None
                results.append(_cmp(cell, p["op"], p.get("value")))
            if not results:
                return True
            return all(results) if match == "all" else any(results)

        matched_rows = [r for r in body if keep(r)]
        matched = len(matched_rows)

        if group_by and not aggregate:
            # Grouping only exists inside aggregate mode. A group_by on its
            # own used to be dropped on the floor: the caller asked for a
            # grouped summary, got ungrouped rows, and nothing said so (the
            # column name was not even validated).
            raise XlMcpError(
                "group_by only applies in aggregate mode; pass aggregate "
                "(for example [{'func': 'count'}] or "
                "[{'func': 'sum', 'column': 'Amount'}]) alongside it, or use "
                "distinct=true to get unique rows")

        # aggregate mode
        if aggregate:
            for a in aggregate:
                if not isinstance(a, dict) or "func" not in a:
                    raise XlMcpError("each aggregate is {column, func}")
                if a["func"] not in _AGG_FUNCS:
                    raise XlMcpError(f"unknown aggregate func {a['func']!r}")
            gbs = ([group_by] if isinstance(group_by, (str, int))
                   else list(group_by or []))
            gb_idx = [col_i(g) for g in gbs]
            groups: dict[tuple, list] = {}
            order: list[tuple] = []
            for row in matched_rows:
                key = tuple(row[i] if i < len(row) else None for i in gb_idx)
                if key not in groups:
                    groups[key] = []
                    order.append(key)
                groups[key].append(row)
            out_groups = []
            for key in order:
                rows = groups[key]
                aggs = {}
                for a in aggregate:
                    ci = col_i(a["column"]) if "column" in a else None
                    vals = [row[ci] if ci is not None and ci < len(row)
                            else None for row in rows]
                    aggs[_agg_label(a)] = _apply_agg(a["func"], vals)
                entry = {"aggregates": aggs}
                if gbs:
                    entry["group"] = {gbs[i]: key[i] for i in range(len(gbs))}
                out_groups.append(entry)
            agg_out = {
                "sheet": grid.sheet, "source": grid.a1, "mode": "aggregate",
                "group_by": gbs, "groups": out_groups,
                "matched": matched, "scanned": scanned,
            }
            if stale_note:
                agg_out["warning"] = stale_note
                agg_out["uncalculated_cells"] = stale_cells[:100]
            return agg_out

        # sort BEFORE projection, so order_by works on any source column,
        # projected or not (before the re-audit an order_by column missing
        # from `columns` was silently ignored).
        if order_by:
            for spec in reversed(order_by):
                ci = col_i(spec["column"]) if isinstance(spec, dict) \
                    else col_i(spec)
                desc = isinstance(spec, dict) and \
                    str(spec.get("dir", "asc")).lower() in ("desc", "descending")
                matched_rows.sort(
                    key=lambda r, k=ci, d=desc: _sort_key(
                        r[k] if k < len(r) else None, reverse=d),
                    reverse=desc)
        # projection
        proj_idx = ([col_i(c) for c in columns] if columns
                    else list(range(len(col_names))))
        proj_names = [col_names[i] for i in proj_idx]
        rows = [[row[i] if i < len(row) else None for i in proj_idx]
                for row in matched_rows]
        if distinct:
            seen = set()
            uniq = []
            for r in rows:
                sig = tuple(str(x) for x in r)
                if sig not in seen:
                    seen.add(sig)
                    uniq.append(r)
            rows = uniq
        total_after_filter = len(rows)
        if offset:
            rows = rows[max(0, int(offset)):]
        truncated = False
        if limit is not None:
            if int(limit) < 0:
                raise XlMcpError("limit must be >= 0")
            truncated = len(rows) > int(limit)
            rows = rows[: int(limit)]
        payload_rows = ([dict(zip(proj_names, r)) for r in rows]
                        if records else rows)
        row_out = {
            "sheet": grid.sheet, "source": grid.a1, "mode": "rows",
            "header": header, "columns": proj_names,
            "rows": payload_rows,
            "matched": matched, "returned": len(rows),
            "scanned": scanned, "distinct_after_filter": total_after_filter,
            "truncated": truncated, "value_mode": values,
        }
        if stale_note:
            row_out["warning"] = stale_note
            row_out["uncalculated_cells"] = stale_cells[:100]
        return row_out
    finally:
        wb.close()


def _agg_label(a: dict) -> str:
    return a.get("as") or (f"{a['func']}_{a['column']}" if "column" in a
                           else a["func"])


def _apply_agg(func: str, vals: list):
    nonblank = [v for v in vals if v is not None and v != ""]
    if func == "count":
        return len(vals)
    if func == "count_nonblank":
        return len(nonblank)
    if func == "count_distinct":
        return len({str(v) for v in nonblank})
    if func == "first":
        return nonblank[0] if nonblank else None
    if func == "last":
        return nonblank[-1] if nonblank else None
    nums = [n for n in (_num(v) for v in nonblank) if n is not None]
    if func in ("sum", "avg", "mean") and not nums:
        return 0 if func == "sum" else None
    if func == "sum":
        return _round(sum(nums))
    if func in ("avg", "mean"):
        return _round(sum(nums) / len(nums)) if nums else None
    if func in ("min", "max"):
        pool = nums if nums else nonblank
        if not pool:
            return None
        return min(pool) if func == "min" else max(pool)
    return None


def _round(x):
    if isinstance(x, float) and x == int(x):
        return int(x)
    return round(x, 10) if isinstance(x, float) else x


#: Excel's cached error literals, which sort after booleans.
_ERROR_LITERALS = frozenset({
    "#REF!", "#NAME?", "#VALUE!", "#DIV/0!", "#N/A", "#NULL!", "#NUM!",
    "#SPILL!", "#CALC!", "#GETTING_DATA",
})


def _excel_serial(v) -> float:
    """A date/time as Excel's own serial number, so a date column sorts with
    numbers exactly as Excel sorts it."""
    import datetime as _dt
    if isinstance(v, _dt.datetime):
        base = _dt.datetime(1899, 12, 30)
        return (v - base).total_seconds() / 86400.0
    if isinstance(v, _dt.date):
        return float((v - _dt.date(1899, 12, 30)).days)
    if isinstance(v, _dt.time):
        return (v.hour * 3600 + v.minute * 60 + v.second) / 86400.0
    if isinstance(v, _dt.timedelta):
        return v.total_seconds() / 86400.0
    return 0.0


def _sort_key(v, *, reverse: bool = False):
    """Excel's own sort order for a mixed-type column.

    Excel ranks ascending: numbers, then text (case-insensitively), then
    FALSE, then TRUE, then error values, and BLANKS ALWAYS LAST -- in both
    directions, which is why the blank rank flips when the caller sorts
    descending. The pre-gate key sorted blanks first on a descending sort,
    compared text case-SENSITIVELY (so 'Zebra' preceded 'apple'), coerced
    numeric TEXT like '10' into a number (Excel keeps it text, after every
    real number), and dropped booleans and error literals into the text run.
    """
    import datetime as _dt
    if v is None or v == "":
        return (-1 if reverse else 5, 0.0, "")
    if isinstance(v, bool):
        return (2, 1.0 if v else 0.0, "")
    if isinstance(v, (int, float)):
        return (0, float(v), "")
    if isinstance(v, (_dt.datetime, _dt.date, _dt.time, _dt.timedelta)):
        return (0, _excel_serial(v), "")
    s = str(v)
    if s in _ERROR_LITERALS:
        return (3, 0.0, s)
    return (1, 0.0, s.casefold())


# ------------------------------------------------------- scatter read / write


def _single_cell_grid(base_wb, item: Any, i: int, sheet: str | None,
                      path: str | None = None):
    """Resolve one scatter item's address to a single-cell grid, or refuse
    naming the offending item."""
    grid = gridio.resolve(base_wb, item, default_sheet=sheet, path=path)
    if not grid.is_single:
        raise XlMcpError(
            f"cells[{i}] resolves to the range {grid.a1}, not a single cell; "
            "the scatter tools take individual cells (use read_range / "
            "write_range for rectangles)")
    return grid


def get_cells(path: str, cells: list, values: str = "cached",
              sheet: str | None = None) -> dict:
    """Read many individually addressed cells in one call. Each item is a
    location object or A1 string resolving to ONE cell; values is cached |
    formula | both, labelled like read_range."""
    if values not in gridio.VALUE_MODES:
        raise XlMcpError(
            f"values must be one of {gridio.VALUE_MODES}, got {values!r}")
    if not isinstance(cells, list) or not cells:
        raise XlMcpError(
            "cells must be a non-empty list of cell addresses "
            "(location objects or A1 strings)")
    if len(cells) > MAX_SCATTER_CELLS:
        raise RangeOutOfBounds(
            f"{len(cells):,} cells is over the {MAX_SCATTER_CELLS:,}-cell "
            "scatter ceiling; use read_range or query_range for bulk reads")
    formula_wb = gridio.open_wb(path, data_only=False) \
        if values in ("formula", "both") else None
    cached_wb = gridio.open_wb(path, data_only=True) \
        if values in ("cached", "both") else None
    try:
        base = formula_wb if formula_wb is not None else cached_wb
        out = []
        absent = 0
        for i, item in enumerate(cells):
            grid = _single_cell_grid(base, item, i, sheet, path=path)
            vals, labels, _hf = gridio.read_matrix(
                grid, mode=values, formula_wb=formula_wb, cached_wb=cached_wb,
                path=path)
            label = labels[0][0]
            if label == "absent":
                absent += 1
            out.append({
                "sheet": grid.sheet, "cell": grid.a1,
                "value": gridio.compact_value(vals[0][0]), "label": label})
        result: dict[str, Any] = {
            "count": len(out), "value_mode": values, "cells": out}
        if absent:
            result["warning"] = f"{absent} of them: " + gridio.ABSENT_WARNING
        return result
    finally:
        for wb in (formula_wb, cached_wb):
            if wb is not None:
                wb.close()


def set_cells(path: str, cells: list, sheet: str | None = None,
              allow_loss: bool = False, backup: bool = True) -> dict:
    """Write many individually addressed cells as ONE batch: every address is
    resolved before anything is written, then one backup + one verified save.
    Each item is {cell (or location), value}; '=' strings become formulas."""
    if not isinstance(cells, list) or not cells:
        raise XlMcpError(
            "cells must be a non-empty list of {cell, value} objects")
    if len(cells) > MAX_SCATTER_CELLS:
        raise RangeOutOfBounds(
            f"{len(cells):,} cells is over the {MAX_SCATTER_CELLS:,}-cell "
            "scatter ceiling; use write_range or apply_edits for bulk writes")
    pkg = WorkbookPackage.open(path)
    # Validation pass: resolve every address and check every item's shape
    # BEFORE mutating, so one bad item refuses the whole batch untouched.
    plan: list[tuple[Any, Any]] = []
    for i, item in enumerate(cells):
        if not isinstance(item, dict):
            raise XlMcpError(f"cells[{i}] is not an object")
        if "value" not in item:
            raise XlMcpError(f"cells[{i}] is missing 'value'")
        loc = item.get("cell", item.get("location"))
        if loc is None:
            raise XlMcpError(f"cells[{i}] is missing 'cell' (or 'location')")
        grid = _single_cell_grid(pkg.workbook, loc, i,
                                 item.get("sheet", sheet))
        plan.append((grid, item["value"]))
    for grid, value in plan:
        pkg.set_cell(grid.sheet, grid.a1, value)
    result = pkg.save(allow_loss=allow_loss, backup=backup)
    result["changed"]["cells_written"] = len(plan)
    return result


# -------------------------------------------------------------------- merges


def _overlaps(a, b) -> bool:
    return not (a.max_row < b.min_row or b.max_row < a.min_row
                or a.max_col < b.min_col or b.max_col < a.min_col)


def set_merge(path: str, action: str, location: Any = None,
              sheet: str | None = None, confirm_data_loss: bool = False,
              allow_loss: bool = False, backup: bool = True) -> dict:
    """Merge or unmerge a cell range, or list merges. Excel semantics: a merge
    keeps only the top-left value; absorbing cells that hold values refuses
    without confirm_data_loss=true rather than discarding them silently."""
    if action not in MERGE_ACTIONS:
        raise XlMcpError(
            f"action must be one of {MERGE_ACTIONS}, got {action!r}")

    if action == "list":
        wb = gridio.open_wb(path)
        try:
            if sheet is not None:
                # Resolve through locate for its missing-sheet error paths.
                title = gridio.resolve(
                    wb, {"used_range": True}, default_sheet=sheet).sheet
                sheets = [wb[title]]
            else:
                sheets = wb.worksheets
            out = {}
            total = 0
            for ws in sheets:
                ranges = sorted(
                    str(r) for r in getattr(ws, "merged_cells", []).ranges) \
                    if getattr(ws, "merged_cells", None) else []
                out[ws.title] = ranges
                total += len(ranges)
            return {"merges": out, "count": total}
        finally:
            wb.close()

    if location is None:
        raise XlMcpError(f"action {action!r} needs a 'location' range")
    pkg = WorkbookPackage.open(path)
    grid = pkg.resolve(location, default_sheet=sheet)
    ws = pkg.workbook[grid.sheet]
    existing = list(getattr(ws, "merged_cells", []).ranges) \
        if getattr(ws, "merged_cells", None) else []

    if action == "unmerge":
        hit = next((r for r in existing
                    if (r.min_row, r.min_col, r.max_row, r.max_col)
                    == (grid.min_row, grid.min_col, grid.max_row,
                        grid.max_col)), None)
        if hit is None:
            listed = ", ".join(str(r) for r in existing[:25]) or "none"
            raise TargetNotFound(
                f"{grid.a1} on {grid.sheet!r} is not a merged range "
                f"(unmerge needs the exact stored range; merges: {listed})")
        ws.unmerge_cells(str(hit))
        result = pkg.save(allow_loss=allow_loss, backup=backup)
        result["changed"]["unmerged"] = {"sheet": grid.sheet,
                                         "range": grid.a1}
        return result

    # merge
    if grid.is_single:
        raise XlMcpError(
            f"{grid.a1} is a single cell; a merge needs a multi-cell range")
    clashes = [str(r) for r in existing if _overlaps(r, grid)]
    if clashes:
        raise XlMcpError(
            f"the range {grid.a1} on {grid.sheet!r} overlaps existing "
            f"merge(s) {', '.join(clashes[:25])}; unmerge them first")
    absorbed = []
    for r in range(grid.min_row, grid.max_row + 1):
        for c in range(grid.min_col, grid.max_col + 1):
            if (r, c) == (grid.min_row, grid.min_col):
                continue
            if ws.cell(r, c).value is not None:
                absorbed.append(gridio.a1(r, c))
    if absorbed and not confirm_data_loss:
        exc = XlMcpError(
            f"merging {grid.a1} would DISCARD the values in "
            f"{', '.join(absorbed[:25])}"
            + (f" (+{len(absorbed) - 25} more)" if len(absorbed) > 25 else "")
            + " (Excel keeps only the top-left value). Pass "
            "confirm_data_loss:true to proceed; the prev backup slot is "
            "the undo")
        exc.code = "CONFLICT"
        raise exc
    for coord in absorbed:
        ws[coord] = None
        pkg._intended[(grid.sheet, coord)] = ("value", None)
    ws.merge_cells(grid.a1)
    result = pkg.save(allow_loss=allow_loss, backup=backup)
    result["changed"]["merged"] = {"sheet": grid.sheet, "range": grid.a1,
                                  "absorbed_cleared": absorbed}
    if absorbed:
        warnings = list(result.get("warnings", []))
        warnings.append(
            f"confirm_data_loss: {len(absorbed)} absorbed cell value(s) were "
            "discarded (Excel merge semantics keep only the top-left value)")
        result["warnings"] = warnings
    return result


__all__ = [
    "read_range", "set_cell", "write_range", "clear_range",
    "copy_range", "move_range", "query_range",
    "get_cells", "set_cells", "set_merge",
    "CLEAR_WHAT", "COPY_WHAT", "MERGE_ACTIONS", "MAX_SCATTER_CELLS",
]
