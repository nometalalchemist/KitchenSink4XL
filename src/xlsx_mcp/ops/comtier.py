"""ops/comtier.py: the COM tier, driving real Excel for what files cannot do.

Every operation here runs through com.session.ComExecutor: process-wide
serialization (one COM worker thread), DisplayAlerts/EnableEvents suppressed
and restored around every operation, bounded timeouts with clean structured
refusals, save/export retry with backoff, and PID-journaled private workers
only (NEVER the user's Excel session; a file open in the user's Excel refuses
with WORKBOOK_LOCKED semantics).

Rulings in force:
- CALC: every recalc issues an explicit Application.CalculateFull; the Phase 1
  spike proved open-time calc does NOT fire under manual calculation mode.
- PIVOTS: REAL pivot tables only, created and refreshed through Excel. Never
  a fake static table.
- VBA: preserve/inspect only. No macro authoring, and macro EXECUTION
  (com_run_macro) is deferred entirely from v1 (author-decision item).

Mutating operations rotate the safesave backup slots BEFORE Excel touches the
file and run a structural verify after; a failed verify restores the backup.
Excel's own save replaces the whole package (part names and order change
legitimately), so the file-tier part-inventory diff does not apply here; the
structural check plus Excel-as-writer is the honesty basis, and
com_validate_opens_clean is available as the deep check.
"""

from __future__ import annotations

import os
import re
import time
from pathlib import Path
from typing import Any, Callable

from ..com import session as _session
from ..core import calc as _calc
from ..core import hazard as _hazard
from ..core import safesave as _safesave
from ..core import verify as _verify
from ..core.errors import (
    ExcelDisconnected,
    TargetNotFound,
    ValidationFailed,
    XlMcpError,
)
from ..core.sandbox import check_path

# ------------------------------------------------------------ Excel constants

XL_DATABASE = 1
XL_ROW_FIELD = 1
XL_COLUMN_FIELD = 2
XL_PAGE_FIELD = 3
XL_TYPE_PDF = 0
XL_SCREEN = 1     # CopyPicture Appearance
XL_BITMAP = 2     # CopyPicture Format

AGG_FUNCS: dict[str, int] = {
    "sum": -4157, "count": -4112, "average": -4106, "max": -4136,
    "min": -4139, "product": -4149, "count_numbers": -4113,
    "stdev": -4155, "var": -4164,
}

FILE_FORMATS: dict[str, int] = {
    "xlsx": 51, "xlsm": 52, "xlsb": 50, "xls": 56,
    "csv": 62,      # xlCSVUTF8 (Excel 2016+; UTF-8 with BOM)
    "ods": 60,
}

SPARK_TYPES: dict[str, int] = {"line": 1, "column": 2, "win_loss": 3}

_PDF_SCOPES = ("workbook", "sheet", "range")


# ---------------------------------------------------------------- helpers


def _norm_path(path: str, label: str) -> str:
    p = check_path(path, label)
    if not os.path.exists(p):
        from ..core.errors import WorkbookNotFound
        raise WorkbookNotFound(f"no such workbook: {p}")
    return p


def _out_path(path: str, label: str, *, overwrite: bool,
              source: str | None = None,
              workbook_target: bool = False) -> tuple[str, dict]:
    """Output-target guard for the COM export/render/convert family: the
    same non-mutating-writer covenant as the file-tier exports
    (core.outguard): never the source workbook, never inside the backup
    store, an existing target refuses without overwrite:true, and
    overwrite first keeps a timestamped .bak of what it replaces."""
    from ..core.outguard import guard_out_file
    p, info = guard_out_file(path, source=source, overwrite=overwrite,
                             what=label, workbook_target=workbook_target)
    parent = os.path.dirname(os.path.abspath(p))
    if parent and not os.path.isdir(parent):
        raise TargetNotFound(f"output directory does not exist: {parent}")
    return p, info


def _ws(wb, sheet: str | None):
    """Worksheet by name (or the first sheet), with an honest not-found."""
    if sheet is None:
        return wb.Worksheets(1)
    try:
        return wb.Worksheets(sheet)
    except Exception:
        names = [str(wb.Worksheets(i + 1).Name)
                 for i in range(int(wb.Worksheets.Count))]
        raise TargetNotFound(
            f"no sheet named {sheet!r}; sheets: {names}") from None


def _op_of(label: str) -> str:
    """The operation name out of an executor label like
    ``recalculate(book.xlsx)``, for use in a refusal message."""
    return str(label).split("(", 1)[0].strip() or "the operation"


def _wrap_com_error(exc: Exception, doing: str) -> XlMcpError:
    """Turn any COM failure into one of this server's errors, with Excel's
    own words in it and never a pywin32 tuple repr.

    Two outcomes, and the split is the live COM stress round's M-2. A COM
    failure that means EXCEL DIED (the taskkilled-mid-write case, HRESULT
    0x800706BE) is not a bad parameter and must not be coded as one: an
    unattended orchestrator reading BAD_PARAMS will not retry, it will go
    "fix" a call that was correct. Those raise ExcelDisconnected, which the
    envelope maps to CONFLICT. Everything else stays the plain refusal."""
    msg = _session.excel_error_message(exc)
    if _session.excel_process_died(exc):
        return ExcelDisconnected(
            f"the Excel worker process stopped answering while {doing}: "
            f"{msg}. The operation did not complete. Retry: the COM layer "
            "re-arms with a fresh Excel on the next call.")
    return XlMcpError(f"Excel refused while {doing}: {msg}")


def _guarded(body: Callable[[Any, Any], dict], doing: str
             ) -> Callable[[Any, Any], dict]:
    """Run an operation body with the same COM-error translation the open and
    save steps get. Before this, only open_workbook and wb.Save were wrapped,
    so a com_error raised by any COM call inside a body (CalculateFull, a
    Range assignment, a Worksheets lookup) propagated untranslated and
    envelope.refusal() rendered it with str() as the raw pywin32 tuple
    ``(-2147023170, 'The remote procedure call failed.', None, None)``
    (live COM stress, M-2)."""
    def run(app, wb) -> dict:
        try:
            return body(app, wb)
        except XlMcpError:
            raise
        except Exception as exc:  # noqa: BLE001
            raise _wrap_com_error(exc, doing) from exc
    return run


def _refuse_encrypted_without_password(path: str,
                                       password: str | None) -> None:
    """An encrypted (CFB) file opened with NO password HANGS Excel on a modal
    prompt (probe-proven: no supplied argument suppresses it), so it is
    refused by signature BEFORE any COM call.

    A WRONG password is a different case and this message used to describe it
    wrongly. It said a wrong password surfaces as the operation timeout,
    which taught callers to budget 60 seconds for a mistyped password. The
    positional-password delivery in session.open_workbook fixed that: a wrong
    password now comes back in about a tenth of a second as Excel's own
    "The password you supplied is not correct" (live COM stress, L-2)."""
    if password is None and _session.is_encrypted_package(path):
        from ..core.errors import WorkbookProtected
        raise WorkbookProtected(
            f"{Path(path).name} is password-protected (encrypted package); "
            "this operation needs the password. "
            "com_validate_opens_clean and com_save_with_password accept "
            "one; a WRONG password comes back at once with Excel's own "
            "wrong-password message, so there is nothing to wait out.")


def _run_readonly(label: str, path: str,
                  body: Callable[[Any, Any], dict],
                  timeout: float | None = None,
                  password: str | None = None) -> dict:
    """Open the workbook read-only in the pooled worker, run body(app, wb),
    close without saving. Serialized, alert-suppressed, timeout-bounded."""
    _refuse_encrypted_without_password(path, password)
    body = _guarded(body, f"running {_op_of(label)} in the workbook")

    def job(manager) -> dict:
        w = manager.acquire()
        wb = None
        prior = _session._assert_hygiene(w.app)
        try:
            try:
                wb = _session.open_workbook(w.app, path, read_only=True,
                                            password=password)
            except XlMcpError:
                raise
            except Exception as exc:  # noqa: BLE001
                raise _wrap_com_error(exc, "opening the workbook") from exc
            out = body(w.app, wb)
            out.setdefault("instance_pid", w.pid)
            return out
        finally:
            if wb is not None:
                try:
                    wb.Close(False)
                except Exception:
                    pass
            _session._restore_hygiene(w.app, *prior)
            manager.release_to_pool(w)

    return _session.run_com(label, job, timeout=timeout)


def _run_mutation(label: str, path: str,
                  body: Callable[[Any, Any], dict],
                  *, backup: bool = True, timeout: float | None = None,
                  save: bool = True, password: str | None = None) -> dict:
    """The mutating pattern: guard (not open elsewhere), rotate backup slots,
    open in the pooled worker, run body(app, wb), save with retry, close,
    structural verify (restore the backup on failure)."""
    _refuse_encrypted_without_password(path, password)
    if _hazard.ole_container_kind(path) == _hazard.OLE_KIND_BIFF:
        # Excel would happily edit and re-save the BIFF file, but the
        # post-save structural verify requires an OOXML package, so the
        # mutation would be applied and then rolled back. Refuse up front
        # with the actual remedy instead (geriatric round, H-1). Read-only
        # COM ops (com_convert_format among them) take .xls fine.
        from ..core.errors import UnsupportedStructure
        exc = UnsupportedStructure(
            f"{Path(path).name} is a legacy Excel 97-2003 (BIFF) workbook; "
            "COM mutations on this server write and verify OOXML packages "
            "only. Convert it first: com_convert_format(path=..., "
            "output='...xlsx'), then edit the converted file.")
        exc.hint_tools = ("com_convert_format",)
        raise exc
    body = _guarded(body, f"running {_op_of(label)} in the workbook")
    warnings = _session.guard_target_closed(path)
    pre_report = None
    try:
        pre_report = _hazard.scan_path(path)
    except Exception:
        pre_report = None
    # Two-phase rotation (destroyer M-1, same covenant as the file tier):
    # stage the pre-mutation content now, commit onto the slots only after
    # Excel's save came back. A COM op that fails before or during save
    # leaves the slots exactly as they were instead of burning prev.
    backup_slot = None
    ticket = None
    if backup:
        ticket = _safesave.prepare_rotation(path)

    def job(manager) -> dict:
        w = manager.acquire()
        wb = None
        prior = _session._assert_hygiene(w.app)
        try:
            try:
                wb = _session.open_workbook(w.app, path, password=password)
            except XlMcpError:
                raise
            except Exception as exc:  # noqa: BLE001
                raise _wrap_com_error(exc, "opening the workbook") from exc
            changed = body(w.app, wb)
            if save:
                try:
                    _session.com_retry(wb.Save,
                                       label=f"save {Path(path).name}")
                except XlMcpError:
                    raise
                except Exception as exc:  # noqa: BLE001
                    raise _wrap_com_error(exc, "saving the workbook") from exc
            try:
                wb.Close(False)
            except XlMcpError:
                raise
            except Exception as exc:  # noqa: BLE001
                raise _wrap_com_error(exc, "closing the workbook") from exc
            wb = None
            return {"changed": changed, "instance_pid": w.pid}
        finally:
            if wb is not None:
                try:
                    wb.Close(False)
                except Exception:
                    pass
            _session._restore_hygiene(w.app, *prior)
            manager.release_to_pool(w)

    try:
        out = _session.run_com(label, job, timeout=timeout)
    except BaseException:
        if ticket is not None:
            ticket.abort()  # nothing promoted; the slots stay as they were
        raise
    if ticket is not None:
        try:
            ticket.commit()
            backup_slot = "prev"
        except OSError as exc:
            ticket.abort()
            warnings.append(
                "the save succeeded but the backup slot rotation failed "
                f"({type(exc).__name__}: {exc}); the prev slot still holds "
                "the state before the PREVIOUS mutation, not this one")

    verified = False
    if save and password is None:
        ok, reasons = _verify.structural_check(path)
        if not ok:
            restored = _restore_backup(path) if backup_slot else False
            raise ValidationFailed(
                "the Excel-saved file failed the structural verify"
                + (" and was restored from the backup" if restored else "")
                + ": " + "; ".join(reasons))
        verified = True
    elif save and password is not None:
        # An encrypted package is a CFB container, not a plain zip; the
        # structural zip check cannot read it. Excel-as-writer plus a COM
        # reopen (com_validate_opens_clean with the password) is the check.
        verified = False

    result: dict[str, Any] = {
        "ok": True,
        "file": path,
        "changed": out.get("changed", {}),
        "saved": bool(save),
        "backup": backup_slot,
        "verified": verified,
        "engine": "com",
        "instance_pid": out.get("instance_pid"),
        "warnings": warnings,
    }
    if pre_report is not None and pre_report.hazards:
        result["preserved"] = (
            "Excel saved the workbook itself, so fragile parts "
            f"({', '.join(pre_report.labels())}) are preserved at full "
            "fidelity")
    return result


def _restore_backup(path: str) -> bool:
    import shutil
    slot = _safesave.slot_dir(path) / _safesave.PREV_SLOT
    if slot.exists():
        shutil.copy2(str(slot), path)
        return True
    return False


# ============================================================== recalculate


def recalculate(path: str, engine: str = "auto",
                timeout_seconds: float | None = None,
                backup: bool = True) -> dict:
    """COM-first fidelity recalculation. engine='com' opens the workbook in a
    private hidden Excel, issues an explicit CalculateFull (open-time calc
    does NOT fire under manual mode; spike-proven), saves, and reports how
    many formula cells gained cached values. engine='formulas' is the
    best-effort pure-Python compute: values are RETURNED, the file is NOT
    modified. engine='auto' prefers COM."""
    p = _norm_path(path, "recalculate")
    _refuse_encrypted_without_password(p, None)
    engine = str(engine).strip().lower()
    if engine not in ("auto", "com", "formulas"):
        raise XlMcpError(
            f"engine must be 'auto', 'com', or 'formulas', got {engine!r}")
    ok, why = _session.com_available()
    if engine == "auto":
        engine = "com" if ok else "formulas"
    if engine == "com" and not ok:
        from ..core.errors import CalcUnavailable
        raise CalcUnavailable(
            f"a fidelity recalc needs Excel and COM is unavailable here "
            f"({why}); engine='formulas' computes classic functions "
            "best-effort without touching the file")

    if engine == "formulas":
        values = _formulas_compute_only(p)
        return {
            "ok": True, "file": p, "engine": "formulas", "saved": False,
            "recalculated": False,
            "computed_values_sample": values,
            "note": ("best-effort pure-Python compute of classic functions; "
                     "the FILE is unchanged and its caches are as before. "
                     "Dynamic arrays, LAMBDA, and modern functions are not "
                     "covered; engine='com' is the fidelity path"),
        }

    absent_before = _count_absent_cached(p)
    seeded: list[int] = [0]

    def body(app, wb) -> dict:
        app.CalculateFull()
        try:
            if bool(app.Iteration):
                seeded[0] = _reseed_iterative_errors(app, wb)
        except Exception:  # noqa: BLE001
            pass
        return {"calculated": "CalculateFull"}

    result = _run_mutation(f"recalculate({Path(p).name})", p, body,
                           backup=backup, timeout=timeout_seconds)
    absent_after = _count_absent_cached(p)
    result["recalculated"] = True
    result["changed"] = {
        "calculation": "explicit CalculateFull in Excel, cache saved",
        "formula_cells_without_cached_value_before": absent_before,
        "formula_cells_without_cached_value_after": absent_after,
    }
    result["freshness"] = (
        "cached values were computed by Excel just now (label: computed); "
        "reads through any engine will see current results")
    if seeded[0]:
        result["changed"]["iterative_cells_reseeded"] = seeded[0]
        result["warnings"] = list(result.get("warnings", [])) + [
            f"{seeded[0]} formula cell(s) under iterative calculation had no "
            "seed value and were loading as errors; their formulas were "
            "re-entered so Excel could iterate them"]
    return result


def _reseed_iterative_errors(app, wb) -> int:
    """Re-enter error-valued formulas so iterative calculation can converge.

    Under iterative calc Excel seeds every pass from the cell's CURRENT
    value. A formula that arrived from a file-tier write has NO cached value,
    so Excel loads it as an error and each iteration propagates that error:
    the numbers-safety gate caught =B1+1 stuck at #VALUE! forever in a
    workbook where Excel's own authoring of the identical formula converges.
    Re-entering the formula through Excel restores a numeric seed. Cells whose
    error is genuine (#REF!, #DIV/0!) are unharmed: re-entering the same
    formula recomputes the same error. Legacy CSE array cells are skipped,
    since assigning .Formula would break the array.

    The re-entry goes through .Formula2, NOT .Formula. HasArray is False for a
    DYNAMIC array, so the CSE skip above never covered them, and assigning
    legacy .Formula re-enters in implicit-intersection mode: an Excel probe
    (edge audit follow-up, 2026-09-04) took an error-valued
    =FILTER(A1:A3,A1:A3>99), assigned .Formula = .Formula, and Excel then
    reported the formula as =@FILTER(A1:A3,A1:A3>99) and stored it as a plain
    <f> with the t="array" ref= spill attributes GONE. The formula had been
    silently converted. .Formula2 round-trips it unchanged. Pre-2019 Excel has
    no .Formula2 (and no dynamic arrays either), so the fallback below is the
    old path where it is still the correct one."""
    xl_cell_type_formulas, xl_errors = -4123, 16
    total = 0
    for ws in wb.Worksheets:
        try:
            errors = ws.Cells.SpecialCells(xl_cell_type_formulas, xl_errors)
        except Exception:  # noqa: BLE001
            continue  # SpecialCells raises when the sheet has no error cells
        try:
            for cell in errors:
                try:
                    if cell.HasArray:
                        continue
                    try:
                        cell.Formula2 = cell.Formula2
                    except Exception:  # noqa: BLE001 - pre-2019 Excel
                        cell.Formula = cell.Formula
                    total += 1
                except Exception:  # noqa: BLE001
                    continue
        except Exception:  # noqa: BLE001
            continue
    if total:
        app.CalculateFull()
    return total


def _formulas_compute_only(path: str) -> dict:
    """Best-effort pure-Python compute via the `formulas` library, WITHOUT
    writing anything (the library's own writer rebuilds the package, which
    is the exact fidelity loss this server exists to prevent). Returns a
    capped sample of computed cell values."""
    try:
        import formulas  # type: ignore
    except Exception as exc:  # noqa: BLE001
        from ..core.errors import CalcUnavailable
        raise CalcUnavailable(
            f"the pure-Python fallback engine is not installed ({exc}); "
            "engine='com' is the fidelity path") from exc
    try:
        xl_model = formulas.ExcelModel().loads(path).finish()
        solution = xl_model.calculate()
    except Exception as exc:  # noqa: BLE001
        from ..core.errors import CalcUnavailable
        raise CalcUnavailable(
            f"the pure-Python fallback could not compute: "
            f"{type(exc).__name__}: {exc}") from exc
    return {k: _calc._clean(v) for k, v in list(solution.items())[:50]}


def _count_absent_cached(path: str) -> int:
    """Formula cells with NO cached value (read as blank outside Excel)."""
    import openpyxl
    keep_vba = path.lower().endswith(".xlsm")
    fwb = openpyxl.load_workbook(path, keep_vba=keep_vba, data_only=False)
    cwb = openpyxl.load_workbook(path, keep_vba=keep_vba, data_only=True)
    try:
        n = 0
        for ws in fwb.worksheets:
            cws = cwb[ws.title]
            for (r, c), cell in getattr(ws, "_cells", {}).items():
                # cell TYPE, not a leading '=': neutralized injection text is
                # data, and counting it as an uncalculated formula made the
                # recalculate report claim work it never had to do.
                if (_calc.is_formula_cell(cell)
                        and cws.cell(r, c).value is None):
                    n += 1
        return n
    finally:
        fwb.close()
        cwb.close()


# ============================================================== real pivots


def com_manage_pivot(path: str, action: str, name: str | None = None,
                     source_sheet: str | None = None,
                     source_range: str | None = None,
                     dest_sheet: str | None = None,
                     dest_cell: str = "A3",
                     rows: list[str] | None = None,
                     columns: list[str] | None = None,
                     filters: list[str] | None = None,
                     values: list[dict] | None = None,
                     timeout_seconds: float | None = None,
                     backup: bool = True) -> dict:
    """Create, refresh, delete, or list REAL pivot tables through Excel."""
    p = _norm_path(path, "com_manage_pivot")
    action = str(action).strip().lower()
    if action not in ("create", "refresh", "delete", "list"):
        raise XlMcpError(
            f"action must be create, refresh, delete, or list; got {action!r}")

    if action == "list":
        def body(app, wb) -> dict:
            return {"pivot_tables": _list_pivots(wb)}
        return _run_readonly(f"pivot_list({Path(p).name})", p, body,
                             timeout=timeout_seconds)

    if action == "create":
        if not source_range:
            raise XlMcpError("create needs source_range (an A1 range with a "
                             "header row, e.g. 'A1:D200')")
        if not values:
            raise XlMcpError(
                "create needs values: a list like "
                "[{'field': 'Amount', 'func': 'sum'}]")
        for spec in values:
            func = str(spec.get("func", "sum")).lower()
            if func not in AGG_FUNCS:
                raise XlMcpError(
                    f"unknown aggregation {func!r}; one of "
                    f"{sorted(AGG_FUNCS)}")
            if not spec.get("field"):
                raise XlMcpError("every values entry needs a 'field'")

        def body(app, wb) -> dict:
            try:
                src_ws = _ws(wb, source_sheet)
                src = src_ws.Range(source_range)
                if dest_sheet is not None:
                    try:
                        dst_ws = wb.Worksheets(dest_sheet)
                    except Exception:
                        dst_ws = wb.Worksheets.Add()
                        dst_ws.Name = dest_sheet
                else:
                    dst_ws = wb.Worksheets.Add()
                pt_name = name or f"KS4XLPivot{int(time.time()) % 100000}"
                cache = wb.PivotCaches().Create(
                    XL_DATABASE, f"'{src_ws.Name}'!{src.Address}")
                pt = cache.CreatePivotTable(
                    dst_ws.Range(dest_cell), pt_name)
                for f in rows or []:
                    pt.PivotFields(f).Orientation = XL_ROW_FIELD
                for f in columns or []:
                    pt.PivotFields(f).Orientation = XL_COLUMN_FIELD
                for f in filters or []:
                    pt.PivotFields(f).Orientation = XL_PAGE_FIELD
                added = []
                for spec in values:
                    func = str(spec.get("func", "sum")).lower()
                    caption = spec.get("caption") or \
                        f"{func.capitalize()} of {spec['field']}"
                    pt.AddDataField(pt.PivotFields(spec["field"]),
                                    caption, AGG_FUNCS[func])
                    added.append(caption)
                return {"pivot_created": pt_name,
                        "sheet": str(dst_ws.Name),
                        "location": str(pt.TableRange2.Address),
                        "rows": rows or [], "columns": columns or [],
                        "filters": filters or [], "values": added}
            except XlMcpError:
                raise
            except Exception as exc:  # noqa: BLE001
                raise _wrap_com_error(
                    exc, "creating the pivot table (check that source_range "
                    "includes a header row and every field name matches a "
                    "header exactly)") from exc

        return _run_mutation(f"pivot_create({Path(p).name})", p, body,
                             backup=backup, timeout=timeout_seconds)

    if action == "refresh":
        def body(app, wb) -> dict:
            found = []
            for i in range(1, int(wb.Worksheets.Count) + 1):
                ws = wb.Worksheets(i)
                for j in range(1, int(ws.PivotTables().Count) + 1):
                    pt = ws.PivotTables(j)
                    if name is None or str(pt.Name) == name:
                        try:
                            pt.RefreshTable()
                        except Exception as exc:  # noqa: BLE001
                            raise _wrap_com_error(
                                exc, f"refreshing pivot {pt.Name!r}") from exc
                        found.append(str(pt.Name))
            if name is not None and not found:
                raise TargetNotFound(
                    f"no pivot table named {name!r}; use action='list'")
            if not found:
                raise TargetNotFound("this workbook has no pivot tables")
            return {"refreshed": found}
        return _run_mutation(f"pivot_refresh({Path(p).name})", p, body,
                             backup=backup, timeout=timeout_seconds)

    # delete
    if not name:
        raise XlMcpError("delete needs the pivot table name; "
                         "use action='list' to see them")

    def body(app, wb) -> dict:
        for i in range(1, int(wb.Worksheets.Count) + 1):
            ws = wb.Worksheets(i)
            for j in range(1, int(ws.PivotTables().Count) + 1):
                pt = ws.PivotTables(j)
                if str(pt.Name) == name:
                    addr = str(pt.TableRange2.Address)
                    sheet_name = str(ws.Name)
                    pt.TableRange2.Clear()
                    return {"pivot_deleted": name, "sheet": sheet_name,
                            "cleared_range": addr}
        raise TargetNotFound(
            f"no pivot table named {name!r}; use action='list'")

    return _run_mutation(f"pivot_delete({Path(p).name})", p, body,
                         backup=backup, timeout=timeout_seconds)


def _list_pivots(wb) -> list[dict]:
    out: list[dict] = []
    for i in range(1, int(wb.Worksheets.Count) + 1):
        ws = wb.Worksheets(i)
        for j in range(1, int(ws.PivotTables().Count) + 1):
            pt = ws.PivotTables(j)
            entry = {"name": str(pt.Name), "sheet": str(ws.Name),
                     "location": str(pt.TableRange2.Address)}
            try:
                entry["source"] = str(pt.PivotCache().SourceData)
            except Exception:
                pass
            try:
                entry["refreshed"] = str(pt.RefreshDate)
            except Exception:
                pass
            out.append(entry)
    return out


# ============================================================ export / render


def com_export_pdf(path: str, output: str, scope: str = "workbook",
                   sheet: str | None = None, range_a1: str | None = None,
                   overwrite: bool = False,
                   timeout_seconds: float | None = None) -> dict:
    """Export the workbook, one sheet, or a range to PDF via Excel."""
    p = _norm_path(path, "com_export_pdf")
    out, out_info = _out_path(output, "com_export_pdf output",
                              overwrite=overwrite, source=p)
    scope = str(scope).strip().lower()
    if scope not in _PDF_SCOPES:
        raise XlMcpError(f"scope must be one of {_PDF_SCOPES}, got {scope!r}")
    if scope == "range" and not range_a1:
        raise XlMcpError("scope='range' needs range_a1 (e.g. 'A1:F40')")

    def body(app, wb) -> dict:
        try:
            if scope == "workbook":
                target = wb
            elif scope == "sheet":
                target = _ws(wb, sheet)
            else:
                target = _ws(wb, sheet).Range(range_a1)
            _session.com_retry(
                lambda: target.ExportAsFixedFormat(XL_TYPE_PDF,
                                                   os.path.abspath(out)),
                label="ExportAsFixedFormat")
        except XlMcpError:
            raise
        except Exception as exc:  # noqa: BLE001
            raise _wrap_com_error(exc, "exporting to PDF") from exc
        size = os.path.getsize(out) if os.path.exists(out) else 0
        return {"exported": out, "scope": scope, "bytes": size}

    result = _run_readonly(f"export_pdf({Path(p).name})", p, body,
                           timeout=timeout_seconds)
    if not os.path.exists(out) or os.path.getsize(out) == 0:
        raise ValidationFailed(
            "Excel reported success but the PDF is missing or empty")
    result.update(out_info)
    result["ok"] = True
    result["file"] = p
    return result


def com_render_sheet(path: str, output: str, sheet: str | None = None,
                     range_a1: str | None = None, overwrite: bool = False,
                     timeout_seconds: float | None = None) -> dict:
    """Render a sheet's used range (or a given range) to a PNG image for
    visual verification, via CopyPicture into a temporary chart canvas."""
    p = _norm_path(path, "com_render_sheet")
    out, out_info = _out_path(output, "com_render_sheet output",
                              overwrite=overwrite, source=p)
    if not out.lower().endswith(".png"):
        raise XlMcpError("output must be a .png path")

    def body(app, wb) -> dict:
        try:
            ws = _ws(wb, sheet)
            rng = ws.Range(range_a1) if range_a1 else ws.UsedRange
            rng.CopyPicture(XL_SCREEN, XL_BITMAP)
            width, height = float(rng.Width), float(rng.Height)
            if width <= 0 or height <= 0:
                raise XlMcpError("the target range has no visible size")
            co = ws.ChartObjects().Add(0, 0, width, height)
            try:
                co.Chart.Paste()
                _session.com_retry(
                    lambda: co.Chart.Export(os.path.abspath(out), "PNG"),
                    label="Chart.Export")
            finally:
                co.Delete()
            return {"rendered": out, "range": str(rng.Address),
                    "sheet": str(ws.Name)}
        except XlMcpError:
            raise
        except Exception as exc:  # noqa: BLE001
            raise _wrap_com_error(exc, "rendering the range to an "
                                  "image") from exc

    result = _run_readonly(f"render_sheet({Path(p).name})", p, body,
                           timeout=timeout_seconds)
    if not os.path.exists(out) or os.path.getsize(out) == 0:
        raise ValidationFailed(
            "Excel reported success but the PNG is missing or empty")
    result.update(out_info)
    result["ok"] = True
    result["file"] = p
    return result


def com_convert_format(path: str, output: str, format: str | None = None,
                       overwrite: bool = False,
                       timeout_seconds: float | None = None) -> dict:
    """Convert between workbook formats via Excel's own SaveAs."""
    p = _norm_path(path, "com_convert_format")
    out, out_info = _out_path(output, "com_convert_format output",
                              overwrite=overwrite, source=p,
                              workbook_target=True)
    fmt = (format or Path(out).suffix.lstrip(".")).strip().lower()
    if fmt not in FILE_FORMATS:
        raise XlMcpError(
            f"format must be one of {sorted(FILE_FORMATS)}, got {fmt!r}")
    if os.path.abspath(out) == os.path.abspath(p):
        raise XlMcpError("output must differ from the source path; "
                         "conversion never overwrites its own source")

    def body(app, wb) -> dict:
        try:
            if overwrite and os.path.exists(out):
                os.remove(out)
            _session.com_retry(
                lambda: wb.SaveAs(os.path.abspath(out),
                                  FILE_FORMATS[fmt]),
                label="SaveAs")
        except XlMcpError:
            raise
        except Exception as exc:  # noqa: BLE001
            raise _wrap_com_error(exc, f"converting to {fmt}") from exc
        note = None
        if fmt == "csv":
            note = "CSV carries the FIRST worksheet's cells only (values, " \
                   "no formulas or formatting); that is the format, not a bug"
        return {"converted": out, "format": fmt,
                **({"note": note} if note else {})}

    def job_wrapper(app, wb):
        return body(app, wb)

    result = _run_readonly(f"convert({Path(p).name} -> {fmt})", p,
                           job_wrapper, timeout=timeout_seconds)
    if not os.path.exists(out) or os.path.getsize(out) == 0:
        raise ValidationFailed(
            "Excel reported success but the converted file is missing or "
            "empty")
    result.update(out_info)
    result["ok"] = True
    result["file"] = p
    return result


# =========================================================== encrypt / misc


def com_save_with_password(path: str, password: str,
                           current_password: str | None = None,
                           timeout_seconds: float | None = None,
                           backup: bool = True) -> dict:
    """Password-protect (or re-key / decrypt) a workbook using Excel's real
    encryption. password='' removes the protection (current_password then
    required)."""
    p = _norm_path(path, "com_save_with_password")
    if not isinstance(password, str):
        raise XlMcpError("password must be a string ('' removes protection)")
    if password == "" and not current_password:
        raise XlMcpError(
            "removing a password needs current_password to open the file")
    _refuse_encrypted_without_password(
        p, current_password if current_password else None)
    warnings = _session.guard_target_closed(p)
    backup_slot = None
    ticket = None
    if backup:
        ticket = _safesave.prepare_rotation(p)

    def job(manager) -> dict:
        w = manager.acquire()
        wb = None
        prior = _session._assert_hygiene(w.app)
        try:
            try:
                wb = _session.open_workbook(w.app, p,
                                            password=current_password)
            except XlMcpError:
                raise
            except Exception as exc:  # noqa: BLE001
                raise _wrap_com_error(
                    exc, "opening the workbook (a wrong current_password "
                    "surfaces here)") from exc
            try:
                wb.Password = password
                _session.com_retry(wb.Save, label="encrypted save")
            except Exception as exc:  # noqa: BLE001
                raise _wrap_com_error(exc, "applying the password") from exc
            wb.Close(False)
            wb = None
            return {"instance_pid": w.pid}
        finally:
            if wb is not None:
                try:
                    wb.Close(False)
                except Exception:
                    pass
            _session._restore_hygiene(w.app, *prior)
            manager.release_to_pool(w)

    try:
        out = _session.run_com(f"encrypt({Path(p).name})", job,
                               timeout=timeout_seconds)
    except BaseException:
        if ticket is not None:
            ticket.abort()
        raise
    if ticket is not None:
        try:
            ticket.commit()
            backup_slot = "prev"
        except OSError as exc:
            ticket.abort()
            warnings.append(
                "the save succeeded but the backup slot rotation failed "
                f"({type(exc).__name__}: {exc}); the prev slot still holds "
                "the state before the PREVIOUS mutation, not this one")
    encrypted = password != ""
    if encrypted:
        # An encrypted package is a CFB container; a plain zip open failing
        # is the expected signature of success, not corruption.
        import zipfile
        try:
            with zipfile.ZipFile(p):
                zip_openable = True
        except Exception:
            zip_openable = False
        if zip_openable:
            raise ValidationFailed(
                "the file is still a plain (unencrypted) package after the "
                "password save; the backup holds the pre-call state")
    else:
        ok, reasons = _verify.structural_check(p)
        if not ok:
            restored = _restore_backup(p) if backup_slot else False
            raise ValidationFailed(
                "the decrypted save failed the structural verify"
                + (" and was restored from the backup" if restored else "")
                + ": " + "; ".join(reasons))
    return {
        "ok": True, "file": p, "saved": True, "backup": backup_slot,
        "engine": "com", "instance_pid": out.get("instance_pid"),
        "changed": {"encryption": "applied" if encrypted else "removed"},
        "verified": not encrypted,
        "warnings": warnings + ([
            "the file is now encrypted; every file-based tool on this "
            "server will refuse it until the password is removed. Keep the "
            "password safe; it is NOT recoverable"] if encrypted else []),
    }


def com_autofit(path: str, sheet: str | None = None,
                columns: str | None = None, rows: str | None = None,
                timeout_seconds: float | None = None,
                backup: bool = True) -> dict:
    """True column/row autofit using Excel's real text metrics."""
    p = _norm_path(path, "com_autofit")

    def body(app, wb) -> dict:
        try:
            ws = _ws(wb, sheet)
            done: dict[str, str] = {"sheet": str(ws.Name)}
            if columns is None and rows is None:
                ws.UsedRange.EntireColumn.AutoFit()
                done["columns"] = "all used columns"
            if columns is not None:
                ws.Columns(columns).AutoFit()
                done["columns"] = columns
            if rows is not None:
                ws.Rows(rows).AutoFit()
                done["rows"] = rows
            return done
        except XlMcpError:
            raise
        except Exception as exc:  # noqa: BLE001
            raise _wrap_com_error(exc, "autofitting (columns like 'A:D', "
                                  "rows like '1:20')") from exc

    return _run_mutation(f"autofit({Path(p).name})", p, body,
                         backup=backup, timeout=timeout_seconds)


def com_goal_seek(path: str, target_cell: str, target_value: float,
                  changing_cell: str, sheet: str | None = None,
                  save: bool = True, timeout_seconds: float | None = None,
                  backup: bool = True) -> dict:
    """Excel's Goal Seek: adjust changing_cell until target_cell (a formula)
    reaches target_value; convergence-checked."""
    p = _norm_path(path, "com_goal_seek")
    try:
        goal = float(target_value)
    except (TypeError, ValueError):
        raise XlMcpError("target_value must be a number") from None

    def body(app, wb) -> dict:
        try:
            ws = _ws(wb, sheet)
            tgt = ws.Range(target_cell)
            chg = ws.Range(changing_cell)
            if not str(tgt.Formula or "").startswith("="):
                raise XlMcpError(
                    f"target_cell {target_cell} holds no formula; Goal Seek "
                    "adjusts an input until a FORMULA reaches the goal")
            converged = bool(tgt.GoalSeek(goal, chg))
            result_value = tgt.Value
            input_value = chg.Value
            if not converged:
                raise XlMcpError(
                    f"Goal Seek did not converge: {target_cell} reached "
                    f"{result_value!r} against the goal {goal!r}. The file "
                    "was not saved.")
            return {"converged": True, "target_cell": target_cell,
                    "target_value": goal, "achieved": result_value,
                    "changing_cell": changing_cell,
                    "solution": input_value}
        except XlMcpError:
            raise
        except Exception as exc:  # noqa: BLE001
            raise _wrap_com_error(exc, "running Goal Seek") from exc

    if save:
        return _run_mutation(f"goal_seek({Path(p).name})", p, body,
                             backup=backup, timeout=timeout_seconds)
    result = _run_readonly(f"goal_seek({Path(p).name})", p, body,
                           timeout=timeout_seconds)
    result["ok"] = True
    result["file"] = p
    result["saved"] = False
    result["note"] = "save:false; the solution is reported but not persisted"
    return result


def _qualify_sparkline_source(source: str, sheet_name: str) -> str:
    """Bind an UNQUALIFIED sparkline source to the group's OWN sheet.

    Excel's ``SparklineGroups.Add`` resolves a bare "B2:F10" against the ACTIVE
    sheet, not the sheet the group lives on. Under the pooled hidden worker
    (whose active sheet is whatever the file stored) an unqualified source
    silently binds the sparkline to the WRONG sheet's data -- the field
    report's "list reports Pivot2!" was that mis-binding, not a display glitch
    (the stored xm:f was literally Pivot2!). Qualifying with the group's sheet
    binds it where the caller meant and makes ``.SourceData`` read back the
    right sheet. A caller-supplied "Sheet!ref" (a deliberate cross-sheet
    sparkline) already carries its sheet and is left untouched.
    """
    if "!" in source:
        return source
    nm = sheet_name
    if re.search(r"[^A-Za-z0-9_]", nm):
        nm = "'" + nm.replace("'", "''") + "'"
    return f"{nm}!{source}"


def com_set_sparkline(path: str, action: str = "create",
                      location: str | None = None,
                      source: str | None = None,
                      type: str = "line", sheet: str | None = None,
                      timeout_seconds: float | None = None,
                      backup: bool = True) -> dict:
    """Create, clear, or list sparkline groups through Excel (the file tier
    defers sparklines here: they live in x14 extLst XML openpyxl cannot
    round-trip)."""
    p = _norm_path(path, "com_set_sparkline")
    action = str(action).strip().lower()
    if action not in ("create", "clear", "list"):
        raise XlMcpError(
            f"action must be create, clear, or list; got {action!r}")

    if action == "list":
        def body(app, wb) -> dict:
            groups: list[dict] = []
            for i in range(1, int(wb.Worksheets.Count) + 1):
                ws = wb.Worksheets(i)
                try:
                    # ws.Cells, not UsedRange: sparkline cells hold no
                    # values, so they can sit OUTSIDE the used range.
                    scope = ws.Cells
                    cnt = int(scope.SparklineGroups.Count)
                except Exception:
                    cnt = 0
                for j in range(1, cnt + 1):
                    g = scope.SparklineGroups.Item(j)
                    groups.append({
                        "sheet": str(ws.Name),
                        "location": str(g.Location.Address),
                        "source": str(g.SourceData),
                        "type": {v: k for k, v in SPARK_TYPES.items()}.get(
                            int(g.Type), str(g.Type)),
                    })
            return {"sparkline_groups": groups}
        return _run_readonly(f"sparkline_list({Path(p).name})", p, body,
                             timeout=timeout_seconds)

    if not location:
        raise XlMcpError("location is required (the cell/range that shows "
                         "the sparklines, e.g. 'G2:G10')")

    if action == "clear":
        def body(app, wb) -> dict:
            try:
                ws = _ws(wb, sheet)
                ws.Range(location).SparklineGroups.Clear()
                return {"sparklines_cleared": location,
                        "sheet": str(ws.Name)}
            except XlMcpError:
                raise
            except Exception as exc:  # noqa: BLE001
                raise _wrap_com_error(exc, "clearing sparklines") from exc
        return _run_mutation(f"sparkline_clear({Path(p).name})", p, body,
                             backup=backup, timeout=timeout_seconds)

    # create
    kind = str(type).strip().lower()
    if kind not in SPARK_TYPES:
        raise XlMcpError(
            f"type must be one of {sorted(SPARK_TYPES)}, got {type!r}")
    if not source:
        raise XlMcpError("create needs source: the data range the "
                         "sparklines chart, e.g. 'A2:F10'")

    def body(app, wb) -> dict:
        try:
            ws = _ws(wb, sheet)
            bound = _qualify_sparkline_source(source, str(ws.Name))
            group = ws.Range(location).SparklineGroups.Add(SPARK_TYPES[kind], bound)
            return {"sparklines_created": location, "source": bound,
                    "type": kind, "sheet": str(ws.Name),
                    "count": int(group.Count) if hasattr(group, "Count")
                    else None}
        except XlMcpError:
            raise
        except Exception as exc:  # noqa: BLE001
            raise _wrap_com_error(
                exc, "creating sparklines (location and source must be "
                "compatible shapes: one sparkline cell per source row or "
                "column)") from exc

    return _run_mutation(f"sparkline_create({Path(p).name})", p, body,
                         backup=backup, timeout=timeout_seconds)


# ======================================================= validate / status


def com_validate_opens_clean(path: str, password: str | None = None,
                             timeout_seconds: float | None = None) -> dict:
    """Open the file in a private hidden Excel and report whether Excel
    accepts it without a repair prompt (the authoritative corruption smoke
    test). Read-only; never saves."""
    p = check_path(path, "com_validate_opens_clean")
    if password is not None:
        def body(app, wb) -> dict:
            return {"opens_clean": True, "workbook": str(wb.Name),
                    "worksheets": int(wb.Worksheets.Count)}
        try:
            result = _run_readonly(f"opens_clean({Path(p).name})", p, body,
                                   timeout=timeout_seconds,
                                   password=password)
        except XlMcpError:
            raise
        except Exception as exc:  # noqa: BLE001
            return {"ok": True, "file": p, "opens_clean": False,
                    "excel_says": _session.excel_error_message(exc)}
        result["ok"] = True
        result["file"] = p
        return result
    result = _session.opens_clean(p, timeout=timeout_seconds)
    result["ok"] = True
    result["file"] = p
    return result


def com_status() -> dict:
    """Honest COM layer status: availability, worker/pool state, journaled
    PIDs, contention and timing. Never spawns Excel."""
    status = _session.get_executor().status()
    status["ok"] = True
    status["policy"] = (
        "private pooled hidden workers only; this server never attaches to "
        "or edits your open Excel session, and refuses files Excel holds "
        "open")
    return status


__all__ = [
    "recalculate", "com_manage_pivot", "com_export_pdf", "com_render_sheet",
    "com_convert_format", "com_save_with_password", "com_autofit",
    "com_goal_seek", "com_set_sparkline", "com_validate_opens_clean",
    "com_status", "AGG_FUNCS", "FILE_FORMATS", "SPARK_TYPES",
]
