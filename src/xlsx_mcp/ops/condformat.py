"""ops/condformat.py: conditional formatting (DESIGN Section 11; design pack as re-cut).

manage_conditional_format adds, lists, and deletes conditional-formatting rules
over a range: cell-value comparisons, color scales, data bars, icon sets,
formula rules, and top/bottom rules.

Honest note on evaluation: a conditional-format rule is stored DECLARATIVELY in
the file. This tool writes the rule; Excel (or LibreOffice) evaluates it and
paints the cells when the workbook opens. The saved file carries the rule, not a
rendered result, exactly like the calc story: file-tier writes the instruction,
the application computes the appearance. core.refs already rewrites CF ranges on
structural edits, so a rule stays anchored to its cells. Mutations route through
WorkbookPackage for the backup and verify.
"""

from __future__ import annotations

from typing import Any

from openpyxl.utils import get_column_letter

from ..core.errors import TargetNotFound, XlMcpError
from ..core.package import WorkbookPackage

CF_ACTIONS = ("add", "list", "delete")
#: The six base rule types 1.0 shipped, plus the six it did not. The gap was
#: never an openpyxl ceiling: formatting/rule.py accepts all twelve, and
#: uniqueValues, duplicateValues, aboveAverage, timePeriod, containsBlanks
#: and containsErrors are all buttons a user can point at in Excel's own
#: Highlight Cells and Top/Bottom menus. Only the x14 extension rules
#: (negative-fill data bars, custom icon sets) are a real library ceiling,
#: and those wait for the preservation work, because writing one through the
#: file tier means writing a worksheet extLst the writer then drops.
CF_TYPES = ("cell_is", "color_scale", "data_bar", "icon_set", "formula",
            "top_bottom", "unique", "duplicate", "above_average",
            "time_period", "blanks", "errors")

#: Excel's timePeriod vocabulary, verbatim: a rule stores one of these
#: strings and Excel evaluates it against the system date on open, the same
#: deal the rest of the conditional-format surface makes.
_TIME_PERIODS = ("today", "yesterday", "tomorrow", "last7Days", "thisWeek",
                 "lastWeek", "nextWeek", "thisMonth", "lastMonth",
                 "nextMonth")
_OPERATORS = ("greaterThan", "lessThan", "greaterThanOrEqual",
              "lessThanOrEqual", "equal", "notEqual", "between", "notBetween",
              "containsText", "notContains", "beginsWith", "endsWith")


def _hex(c: str | None, default: str | None = None) -> str | None:
    if c is None:
        return default
    s = str(c).lstrip("#").upper()
    if len(s) == 6:
        s = "FF" + s
    if len(s) != 8 or any(ch not in "0123456789ABCDEF" for ch in s):
        raise XlMcpError(f"color {c!r} is not a valid hex color")
    return s


def _fill(color: str | None):
    from openpyxl.styles import PatternFill
    rgb = _hex(color, "FFFFFF00")
    return PatternFill(start_color=rgb, end_color=rgb, fill_type="solid")


def _build_rule(cf_type: str, params: dict):
    from openpyxl.formatting.rule import (
        ColorScaleRule, DataBarRule, FormulaRule, IconSetRule, Rule)
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Font

    if cf_type == "cell_is":
        op = params.get("operator", "greaterThan")
        if op not in _OPERATORS:
            raise XlMcpError(f"operator must be one of {_OPERATORS}")
        formula = params.get("formula")
        if formula is None:
            raise XlMcpError("cell_is needs formula (a value or [lo, hi])")
        formula = formula if isinstance(formula, list) else [formula]
        return CellIsRule(
            operator=op, formula=[str(f) for f in formula],
            fill=_fill(params.get("fill")),
            font=Font(color=_hex(params["font_color"]))
            if params.get("font_color") else None)

    if cf_type == "formula":
        formula = params.get("formula")
        if not formula:
            raise XlMcpError("formula rule needs formula (an Excel expression)")
        f = formula if isinstance(formula, list) else [formula]
        return FormulaRule(formula=[str(x) for x in f],
                           fill=_fill(params.get("fill")))

    if cf_type == "color_scale":
        colors = params.get("colors") or ["FFF8696B", "FFFFEB84", "FF63BE7B"]
        if len(colors) == 2:
            return ColorScaleRule(
                start_type="min", start_color=_hex(colors[0]),
                end_type="max", end_color=_hex(colors[1]))
        return ColorScaleRule(
            start_type="min", start_color=_hex(colors[0]),
            mid_type="percentile", mid_value=50, mid_color=_hex(colors[1]),
            end_type="max", end_color=_hex(colors[2]))

    if cf_type == "data_bar":
        return DataBarRule(
            start_type="min", end_type="max",
            color=_hex(params.get("color"), "FF638EC6"),
            showValue=params.get("show_value", True))

    if cf_type == "icon_set":
        return IconSetRule(
            icon_style=params.get("icon_style", "3TrafficLights1"),
            type="percent", values=params.get("values", [0, 33, 67]),
            showValue=params.get("show_value", True),
            reverse=params.get("reverse", False))

    from openpyxl.styles.differential import DifferentialStyle

    if cf_type == "top_bottom":
        rank = int(params.get("rank", 10))
        rule = Rule(type="top10", rank=rank,
                    percent=bool(params.get("percent", False)),
                    bottom=bool(params.get("bottom", False)))
        rule.dxf = DifferentialStyle(fill=_fill(params.get("fill")))
        return rule

    # The six that take no parameter but a format: Excel decides what
    # qualifies (a duplicate, a blank, an error, a date in the period), and
    # this side only says which rule and how to paint it. dxf carries the
    # font when one is asked for, the same way cell_is does.
    font = (Font(color=_hex(params["font_color"]))
            if params.get("font_color") else None)
    style = DifferentialStyle(fill=_fill(params.get("fill")), font=font)

    if cf_type in ("unique", "duplicate"):
        rule = Rule(type="uniqueValues" if cf_type == "unique"
                    else "duplicateValues")
        rule.dxf = style
        return rule

    if cf_type == "above_average":
        # Excel models below-average and the standard-deviation bands as the
        # same rule with different attributes, so one type covers six of the
        # menu's entries.
        below = bool(params.get("below", False))
        rule = Rule(type="aboveAverage", aboveAverage=not below,
                    equalAverage=bool(params.get("equal_average", False)))
        std_dev = params.get("std_dev")
        if std_dev is not None:
            if not isinstance(std_dev, int) or isinstance(std_dev, bool) \
                    or not 1 <= std_dev <= 3:
                raise XlMcpError(
                    "std_dev must be 1, 2, or 3 standard deviations, or "
                    "omitted for a plain average comparison")
            rule.stdDev = std_dev
        rule.dxf = style
        return rule

    if cf_type == "time_period":
        period = params.get("period")
        if period not in _TIME_PERIODS:
            raise XlMcpError(
                f"time_period needs period, one of {_TIME_PERIODS}; got "
                f"{period!r}")
        rule = Rule(type="timePeriod", timePeriod=period)
        rule.dxf = style
        return rule

    # blanks / errors. Excel writes containsBlanks and containsErrors with a
    # formula that reproduces the test, and its own UI does the same, so the
    # rule is written that way rather than left for a reader to infer.
    first = str(params.get("first_cell", "A1")).upper()
    if cf_type == "blanks":
        rule = Rule(type="containsBlanks",
                    formula=[f"LEN(TRIM({first}))=0"])
    else:
        rule = Rule(type="containsErrors", formula=[f"ISERROR({first})"])
    rule.dxf = style
    return rule


def manage_conditional_format(path: str, action: str, location: Any = None,
                              cf_type: str | None = None,
                              params: dict | None = None,
                              sheet: str | None = None, index: int | None = None,
                              allow_loss: bool = False,
                              backup: bool = True,
                              verify_com: bool | None = None) -> dict:
    """Add / list / delete conditional-format rules. Backup + verify on write."""
    if action not in CF_ACTIONS:
        raise XlMcpError(f"action must be one of {CF_ACTIONS}, got {action!r}")

    if action == "list":
        from . import gridio
        wb = gridio.open_wb(path, data_only=False)
        try:
            target = sheet
            out = []
            for ws in (wb.worksheets if target is None else [wb[target]]):
                cfc = getattr(ws, "conditional_formatting", None)
                rules_map = getattr(cfc, "_cf_rules", {}) or {}
                for cf_obj, rules in rules_map.items():
                    for i, rule in enumerate(rules):
                        out.append({
                            "sheet": ws.title, "range": str(cf_obj.sqref),
                            "index": i, "type": rule.type,
                            "priority": rule.priority,
                            "operator": getattr(rule, "operator", None)})
            return {"rules": out, "count": len(out)}
        finally:
            wb.close()

    pkg = WorkbookPackage.open(path)
    if action == "add":
        if cf_type not in CF_TYPES:
            raise XlMcpError(f"cf_type must be one of {CF_TYPES}")
        grid = pkg.resolve(location, default_sheet=sheet)
        ws = pkg.workbook[grid.sheet]
        # containsBlanks and containsErrors carry a formula written against
        # the range's TOP-LEFT cell, which Excel then applies relatively
        # across the rest. Defaulting that to A1 would work only for ranges
        # that happen to start there, so it comes from the resolved
        # rectangle unless the caller names a cell.
        rule_params = dict(params or {})
        rule_params.setdefault(
            "first_cell",
            f"{get_column_letter(grid.min_col)}{grid.min_row}")
        rule = _build_rule(cf_type, rule_params)
        ws.conditional_formatting.add(grid.a1, rule)
        pkg._changed["conditional_format"] = {
            "added": cf_type, "sheet": ws.title, "range": grid.a1}
    else:  # delete
        if location is None:
            raise XlMcpError("delete needs location (the rule's range)")
        grid = pkg.resolve(location, default_sheet=sheet)
        ws = pkg.workbook[grid.sheet]
        cfc = ws.conditional_formatting
        rules_map = getattr(cfc, "_cf_rules", {}) or {}
        # Exact match on the whole sqref or on one of its member ranges.
        # A substring test is a trap ("A1" is a substring of "A10:A20" and
        # "A1:B2" of "A1:B20"), so member strings are compared whole.
        target_key = None
        for cf_obj in list(rules_map):
            members = [str(x) for x in cf_obj.sqref.ranges]
            if str(cf_obj.sqref) == grid.a1 or grid.a1 in members:
                target_key = cf_obj
                break
        if target_key is None:
            raise TargetNotFound(
                f"no conditional-format rule on range {grid.a1} of "
                f"{ws.title!r}")
        rules = rules_map[target_key]
        if index is not None:
            if not 0 <= int(index) < len(rules):
                raise XlMcpError(
                    f"rule index {index} out of range (0..{len(rules) - 1})")
            del rules[int(index)]
            removed = 1
            if not rules:
                del rules_map[target_key]
        else:
            removed = len(rules)
            del rules_map[target_key]
        pkg._changed["conditional_format"] = {
            "deleted_rules": removed, "sheet": ws.title, "range": grid.a1}

    return pkg.save(allow_loss=allow_loss, backup=backup,
                    verify_com=verify_com)


__all__ = ["manage_conditional_format", "CF_ACTIONS", "CF_TYPES"]
