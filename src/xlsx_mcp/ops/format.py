"""ops/format.py: cell formatting and dimensions (DESIGN Section 11, lite).

format_cells writes number format, font, fill, border, and alignment; it MERGES
each attribute onto the cell's existing style, so setting one property never
clobbers the others (curing the incumbent's format-clobber defect). openpyxl
deduplicates styles into its shared style table automatically, so repeated
formatting does not bloat toward the 64,000-format ceiling.

set_dimensions writes column widths and row heights, hides rows/columns, and
services an autofit REQUEST as a best-effort width approximation (true autofit
needs Excel's text metrics via the COM tier, so this is marked approximate,
never claimed as exact). Both route their mutation through WorkbookPackage.

Phase 3c additions (the style layer, design pack as re-cut): apply_style (named cell
styles, builtin or defined in-call), copy_format (the format painter: one
source cell's style painted onto a range), and audit_styles (the read-only
style-bloat audit toward the 64,000-cell-format ceiling the research flagged:
distinct format counts from styles.xml, per-sheet usage, and the heaviest
offenders). Audit reads; apply and copy write through WorkbookPackage.
"""

from __future__ import annotations

from typing import Any

from ..core.errors import TargetNotFound, XlMcpError
from ..core.package import WorkbookPackage
from ..core.sandbox import check_path
from . import gridio

_H_ALIGN = {"left", "center", "right", "fill", "justify", "centerContinuous",
            "distributed", "general"}
_V_ALIGN = {"top", "center", "bottom", "justify", "distributed"}
#: Excel stores text rotation as 0-180 degrees, plus the sentinel 255 for
#: vertically stacked text. openpyxl validates against the same set but
#: reports it by printing all 182 members, which buries the one fact the
#: caller needs. checked_rotation below states the range instead.
_ROTATION_STACKED = 255


def _checked_rotation(value):
    """Validate a text_rotation before openpyxl prints its 182-value set."""
    if value is None:
        return value
    if isinstance(value, bool) or not isinstance(value, (int, float)) \
            or int(value) != value:
        raise XlMcpError(
            f"text_rotation must be a whole number of degrees; got "
            f"{value!r}")
    value = int(value)
    if 0 <= value <= 180 or value == _ROTATION_STACKED:
        return value
    raise XlMcpError(
        f"text_rotation must be 0 to 180 degrees, or {_ROTATION_STACKED} for "
        f"vertically stacked text; got {value}")


_BORDER_STYLES = {"thin", "medium", "thick", "dashed", "dotted", "double",
                  "hair", "mediumDashed", "dashDot", "mediumDashDot",
                  "dashDotDot", "slantDashDot", "none"}
_SIDES = ("left", "right", "top", "bottom")


def _norm_color(c: Any) -> str | None:
    if c is None:
        return None
    if not isinstance(c, str):
        raise XlMcpError("color must be a hex string like 'FF0000' or "
                         "'FFFF0000'")
    s = c.lstrip("#").upper()
    if len(s) == 6:
        s = "FF" + s
    if len(s) != 8 or any(ch not in "0123456789ABCDEF" for ch in s):
        raise XlMcpError(f"color {c!r} is not a valid hex color")
    return s


def format_cells(path: str, location: Any, number_format: str | None = None,
                 font: dict | None = None, fill: dict | None = None,
                 border: dict | None = None, alignment: dict | None = None,
                 sheet: str | None = None, allow_loss: bool = False,
                 backup: bool = True,
                 verify_com: bool | None = None) -> dict:
    """Apply number format, font, fill, border, and/or alignment to a range,
    merging onto the existing style so unspecified attributes are preserved.
    One backup + one verified save."""
    if not any([number_format, font, fill, border, alignment]):
        raise XlMcpError(
            "pass at least one of number_format, font, fill, border, "
            "alignment")
    pkg = WorkbookPackage.open(path)
    grid = pkg.resolve(location, default_sheet=sheet)
    gridio.guard_cell_count(grid)
    ws = pkg.workbook[grid.sheet]
    from copy import copy
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    def new_font(existing):
        f = font
        return Font(
            name=f.get("name", existing.name),
            size=f.get("size", existing.size),
            bold=f.get("bold", existing.bold),
            italic=f.get("italic", existing.italic),
            underline=("single" if f["underline"] else None)
            if "underline" in f else existing.underline,
            strike=f.get("strike", existing.strike),
            color=_norm_color(f["color"]) if f.get("color") else existing.color)

    def new_fill(existing):
        color = _norm_color(fill.get("color") or fill.get("fg"))
        if color is None:
            return existing
        return PatternFill(fill_type=fill.get("pattern", "solid"),
                           start_color=color,
                           end_color=_norm_color(fill.get("bg")) or color)

    def new_border(existing):
        style = border.get("style", "thin")
        if style not in _BORDER_STYLES:
            raise XlMcpError(f"border style must be one of {sorted(_BORDER_STYLES)}")
        col = _norm_color(border.get("color")) or "FF000000"
        side = Side(style=None if style == "none" else style, color=col)
        sides = border.get("sides") or _SIDES
        kw = {s: side if s in sides else getattr(existing, s) for s in _SIDES}
        return Border(**kw)

    def new_align(existing):
        a = alignment
        h = a.get("horizontal", existing.horizontal)
        v = a.get("vertical", existing.vertical)
        if h is not None and h not in _H_ALIGN:
            raise XlMcpError(f"horizontal must be one of {sorted(_H_ALIGN)}")
        if v is not None and v not in _V_ALIGN:
            raise XlMcpError(f"vertical must be one of {sorted(_V_ALIGN)}")
        return Alignment(
            horizontal=h, vertical=v,
            wrap_text=a.get("wrap_text", existing.wrap_text),
            text_rotation=_checked_rotation(
                a.get("text_rotation", existing.text_rotation)),
            indent=a.get("indent", existing.indent))

    n = 0
    for r in range(grid.min_row, grid.max_row + 1):
        for c in range(grid.min_col, grid.max_col + 1):
            cell = ws.cell(r, c)
            if number_format is not None:
                cell.number_format = number_format
            if font is not None:
                cell.font = new_font(cell.font)
            if fill is not None:
                cell.fill = new_fill(cell.fill)
            if border is not None:
                cell.border = new_border(cell.border)
            if alignment is not None:
                cell.alignment = new_align(cell.alignment)
            n += 1
    result = pkg.save(allow_loss=allow_loss, backup=backup,
                      verify_com=verify_com)
    result["changed"]["formatted"] = {"range": grid.a1, "cells": n}
    return result


#: Outline levels Excel supports: 1 through 7, plus 0 meaning ungrouped.
MAX_OUTLINE_LEVEL = 7


def _outline_spans(specs, *, axis: str) -> list[tuple[int, int, int, bool]]:
    """Normalize group specs to (start, end, level, collapsed).

    A span is {start, end} with an optional level (default 1) and collapsed
    (default false). Rows take numbers; columns take letters or numbers, and
    both are resolved to indices before anything is written."""
    from openpyxl.utils import column_index_from_string

    def index(value, what: str) -> int:
        if isinstance(value, bool):
            raise XlMcpError(f"{what} must be a {axis} reference, not a bool")
        if isinstance(value, int):
            return value
        if isinstance(value, str) and axis == "column":
            try:
                return column_index_from_string(value.strip().upper())
            except Exception:
                raise XlMcpError(
                    f"{what}: {value!r} is not a column letter") from None
        if isinstance(value, str) and value.strip().isdigit():
            return int(value)
        raise XlMcpError(
            f"{what} must be a {'row number' if axis == 'row' else 'column '
            'letter or number'}; got {value!r}")

    out = []
    for spec in specs:
        if not isinstance(spec, dict) or "start" not in spec:
            raise XlMcpError(
                f"each {axis} group is {{start, end}} with optional level "
                f"(1-{MAX_OUTLINE_LEVEL}) and collapsed; got {spec!r}")
        start = index(spec["start"], "start")
        end = index(spec.get("end", spec["start"]), "end")
        if end < start:
            start, end = end, start
        level = spec.get("level", 1)
        if isinstance(level, bool) or not isinstance(level, int) \
                or not 1 <= level <= MAX_OUTLINE_LEVEL:
            raise XlMcpError(
                f"level must be 1 to {MAX_OUTLINE_LEVEL} (Excel's outline "
                f"depth); got {level!r}")
        out.append((start, end, level, bool(spec.get("collapsed", False))))
    return out


def sheet_outline(ws) -> dict | None:
    """The sheet's row and column groups, or None when it has none.

    Read side of the grouping surface. Contiguous runs at the same outline
    level are reported as one span, which is how a person sees them in
    Excel: the bracket in the margin, not the per-row attribute underneath
    it. A sheet with no grouping returns None so the ordinary workbook pays
    nothing for this."""
    from openpyxl.utils import column_index_from_string, get_column_letter

    def spans(dims, *, ranged: bool):
        """Per-index (level, hidden), coalesced into contiguous spans.

        Rows are stored one dimension per row. COLUMNS are not: openpyxl
        keeps one ColumnDimension spanning min..max, keyed by the first
        letter, so reading column_dimensions['C'] on a B:C group creates a
        fresh empty dimension and reports level 0. The min/max attributes
        are the truth for columns and must be walked, not the keys."""
        levels: dict[int, tuple[int, bool]] = {}
        for key, dim in dims.items():
            lvl = int(getattr(dim, "outlineLevel", 0) or 0)
            if not lvl:
                continue
            hidden = bool(getattr(dim, "hidden", False))
            if ranged:
                lo = int(getattr(dim, "min", 0) or 0)
                hi = int(getattr(dim, "max", 0) or 0)
                if not lo or not hi:
                    lo = hi = (key if isinstance(key, int)
                               else column_index_from_string(str(key)))
            else:
                lo = hi = int(key)
            for idx in range(lo, hi + 1):
                levels[idx] = (lvl, hidden)
        out: list[dict] = []
        for idx in sorted(levels):
            lvl, hidden = levels[idx]
            if out and out[-1]["level"] == lvl and out[-1]["_end"] == idx - 1:
                out[-1]["_end"] = idx
                out[-1]["collapsed"] = out[-1]["collapsed"] and hidden
            else:
                out.append({"level": lvl, "_start": idx, "_end": idx,
                            "collapsed": hidden})
        return out

    rows = spans(ws.row_dimensions, ranged=False)
    cols = spans(ws.column_dimensions, ranged=True)
    if not rows and not cols:
        return None
    out: dict[str, Any] = {}
    if rows:
        out["row_groups"] = [
            {"start": s["_start"], "end": s["_end"], "level": s["level"],
             "collapsed": s["collapsed"]} for s in rows]
    if cols:
        out["column_groups"] = [
            {"start": get_column_letter(s["_start"]),
             "end": get_column_letter(s["_end"]), "level": s["level"],
             "collapsed": s["collapsed"]} for s in cols]
    props = getattr(ws, "sheet_properties", None)
    outline_pr = getattr(props, "outlinePr", None) if props else None
    if outline_pr is not None:
        out["summary_below"] = bool(
            getattr(outline_pr, "summaryBelow", True))
        out["summary_right"] = bool(
            getattr(outline_pr, "summaryRight", True))
    return out


def set_dimensions(path: str, sheet: str | None = None,
                   column_widths: dict | None = None,
                   row_heights: dict | None = None,
                   autofit_columns: list | None = None,
                   hide_columns: list | None = None,
                   hide_rows: list | None = None,
                   group_rows: list | None = None,
                   group_columns: list | None = None,
                   ungroup_rows: list | None = None,
                   ungroup_columns: list | None = None,
                   outline_summary: dict | None = None,
                   allow_loss: bool = False, backup: bool = True,
                   verify_com: bool | None = None) -> dict:
    """Set column widths and row heights, hide rows/columns, group and
    ungroup them into outline levels, and service an autofit request as a
    best-effort width approximation (true autofit needs Excel via the com
    pack). One backup + one verified save."""
    if not any([column_widths, row_heights, autofit_columns, hide_columns,
                hide_rows, group_rows, group_columns, ungroup_rows,
                ungroup_columns, outline_summary]):
        raise XlMcpError(
            "pass at least one of column_widths, row_heights, "
            "autofit_columns, hide_columns, hide_rows, group_rows, "
            "group_columns, ungroup_rows, ungroup_columns, outline_summary")
    from openpyxl.utils import get_column_letter
    pkg = WorkbookPackage.open(path)
    wb = pkg.workbook
    ws = wb[sheet] if sheet is not None else wb.active

    def col_letter(key) -> str:
        if isinstance(key, int):
            return get_column_letter(key)
        return str(key).upper()

    changed: dict[str, Any] = {}
    if column_widths:
        for key, width in column_widths.items():
            ws.column_dimensions[col_letter(key)].width = float(width)
        changed["column_widths"] = len(column_widths)
    if row_heights:
        for key, height in row_heights.items():
            ws.row_dimensions[int(key)].height = float(height)
        changed["row_heights"] = len(row_heights)
    if autofit_columns:
        approx = {}
        for key in autofit_columns:
            letter = col_letter(key)
            from openpyxl.utils import column_index_from_string
            ci = column_index_from_string(letter)
            longest = 0
            for r in range(1, min(ws.max_row, 5000) + 1):
                v = ws.cell(r, ci).value
                if v is not None:
                    longest = max(longest, len(str(v)))
            width = min(max(longest + 2, 8), 80) * 1.05
            ws.column_dimensions[letter].width = round(width, 2)
            approx[letter] = ws.column_dimensions[letter].width
        changed["autofit_columns"] = {"approximate": True, "widths": approx}
    if hide_columns:
        for key in hide_columns:
            ws.column_dimensions[col_letter(key)].hidden = True
        changed["hidden_columns"] = [col_letter(k) for k in hide_columns]
    if hide_rows:
        for key in hide_rows:
            ws.row_dimensions[int(key)].hidden = True
        changed["hidden_rows"] = [int(k) for k in hide_rows]

    # Grouping. openpyxl models the whole thing (dimensions.group sets
    # outlineLevel on each member and hidden on a collapsed span), so this
    # is a wiring job, not a new format: the tree simply never called it.
    # Excel draws the bracket from outlineLevel and decides collapsed state
    # from the members' hidden flags, so ungroup clears both.
    if group_rows:
        spans = _outline_spans(group_rows, axis="row")
        for start, end, level, collapsed in spans:
            ws.row_dimensions.group(start, end, outline_level=level,
                                    hidden=collapsed)
        changed["grouped_rows"] = [
            {"start": s, "end": e, "level": lv, "collapsed": c}
            for s, e, lv, c in spans]
    if group_columns:
        spans = _outline_spans(group_columns, axis="column")
        for start, end, level, collapsed in spans:
            ws.column_dimensions.group(
                get_column_letter(start), get_column_letter(end),
                outline_level=level, hidden=collapsed)
        changed["grouped_columns"] = [
            {"start": get_column_letter(s), "end": get_column_letter(e),
             "level": lv, "collapsed": c} for s, e, lv, c in spans]
    if ungroup_rows:
        spans = _outline_spans(ungroup_rows, axis="row")
        for start, end, _lv, _c in spans:
            for r in range(start, end + 1):
                dim = ws.row_dimensions[r]
                dim.outlineLevel = 0
                dim.hidden = False
        changed["ungrouped_rows"] = [{"start": s, "end": e}
                                     for s, e, _lv, _c in spans]
    if ungroup_columns:
        spans = _outline_spans(ungroup_columns, axis="column")
        for start, end, _lv, _c in spans:
            # A column group is ONE dimension spanning min..max, keyed by
            # its first letter, so clearing per-letter would miss it and
            # silently create empty dimensions instead. Every dimension
            # whose span intersects the request is cleared.
            for dim in list(ws.column_dimensions.values()):
                lo = int(getattr(dim, "min", 0) or 0)
                hi = int(getattr(dim, "max", 0) or 0)
                if lo and hi and lo <= end and hi >= start:
                    dim.outlineLevel = 0
                    dim.hidden = False
        changed["ungrouped_columns"] = [
            {"start": get_column_letter(s), "end": get_column_letter(e)}
            for s, e, _lv, _c in spans]
    if outline_summary:
        unknown = sorted(set(outline_summary) - {"below", "right"})
        if unknown:
            raise XlMcpError(
                f"outline_summary takes below and right; got {unknown}")
        pr = ws.sheet_properties.outlinePr
        if "below" in outline_summary:
            pr.summaryBelow = bool(outline_summary["below"])
        if "right" in outline_summary:
            pr.summaryRight = bool(outline_summary["right"])
        changed["outline_summary"] = {"below": bool(pr.summaryBelow),
                                      "right": bool(pr.summaryRight)}

    pkg._changed["dimensions"] = {"sheet": ws.title, **changed}
    return pkg.save(allow_loss=allow_loss, backup=backup,
                    verify_com=verify_com)


# ----------------------------------------------------------- named styles


def _build_named_style(name: str, define: dict):
    """Build a NamedStyle from a define dict ({number_format, font, fill,
    border, alignment}, the format_cells vocabularies)."""
    from openpyxl.styles import (
        Alignment, Border, Font, NamedStyle, PatternFill, Side)
    known = {"number_format", "font", "fill", "border", "alignment"}
    unknown = sorted(set(define) - known)
    if unknown:
        raise XlMcpError(
            f"define has unknown key(s) {unknown}; it takes {sorted(known)}")
    if not any(define.get(k) for k in known):
        raise XlMcpError(
            "define must set at least one of number_format, font, fill, "
            "border, alignment")
    ns = NamedStyle(name=name)
    if define.get("number_format"):
        ns.number_format = str(define["number_format"])
    f = define.get("font")
    if f:
        ns.font = Font(
            name=f.get("name", "Calibri"), size=f.get("size", 11),
            bold=f.get("bold", False), italic=f.get("italic", False),
            underline="single" if f.get("underline") else None,
            strike=f.get("strike", False),
            color=_norm_color(f["color"]) if f.get("color") else None)
    fl = define.get("fill")
    if fl:
        color = _norm_color(fl.get("color") or fl.get("fg"))
        if color is None:
            raise XlMcpError("define.fill needs a 'color'")
        ns.fill = PatternFill(fill_type=fl.get("pattern", "solid"),
                              start_color=color,
                              end_color=_norm_color(fl.get("bg")) or color)
    b = define.get("border")
    if b:
        style = b.get("style", "thin")
        if style not in _BORDER_STYLES:
            raise XlMcpError(
                f"border style must be one of {sorted(_BORDER_STYLES)}")
        col = _norm_color(b.get("color")) or "FF000000"
        side = Side(style=None if style == "none" else style, color=col)
        sides = b.get("sides") or _SIDES
        ns.border = Border(**{s: side if s in sides else Side()
                              for s in _SIDES})
    a = define.get("alignment")
    if a:
        h = a.get("horizontal")
        v = a.get("vertical")
        if h is not None and h not in _H_ALIGN:
            raise XlMcpError(f"horizontal must be one of {sorted(_H_ALIGN)}")
        if v is not None and v not in _V_ALIGN:
            raise XlMcpError(f"vertical must be one of {sorted(_V_ALIGN)}")
        ns.alignment = Alignment(
            horizontal=h, vertical=v, wrap_text=a.get("wrap_text"),
            text_rotation=_checked_rotation(a.get("text_rotation", 0)),
            indent=a.get("indent", 0))
    return ns


def apply_style(path: str, style: str, location: Any = None,
                define: dict | None = None, sheet: str | None = None,
                allow_loss: bool = False, backup: bool = True,
                verify_com: bool | None = None) -> dict:
    """Apply a named cell style to a range, defining it first when `define`
    is given. Builtin Excel style names (Good, Bad, Input, Title...) work
    out of the box; define with no location registers the style only."""
    if not isinstance(style, str) or not style.strip():
        raise XlMcpError("style must be a style name")
    style = style.strip()
    if location is None and define is None:
        raise XlMcpError(
            "pass a location to apply the style, a define object to create "
            "it, or both")
    pkg = WorkbookPackage.open(path)
    wb = pkg.workbook
    existing = list(wb.named_styles)
    from openpyxl.styles.builtins import styles as _builtins

    if define is not None:
        if style in existing:
            raise XlMcpError(
                f"a named style {style!r} already exists in this workbook; "
                "pick another name (styles cannot be redefined in place)")
        wb.add_named_style(_build_named_style(style, define))
    elif style not in existing and style not in _builtins:
        candidates = sorted(set(existing) | set(_builtins))
        close = [n for n in candidates if n.lower() == style.lower()]
        hint = f"; did you mean {close[0]!r}?" if close else ""
        raise TargetNotFound(
            f"no style named {style!r}{hint} Workbook styles: "
            + (", ".join(repr(n) for n in existing) or "none")
            + ". Builtins include: Normal, Good, Bad, Neutral, Input, "
            "Output, Calculation, Check Cell, Warning Text, Title, "
            "Headline 1-4, Total, Note, and the Accent1-6 families.")
    elif style not in existing:
        # A builtin used for the first time: register a private copy so the
        # shared module-level object is never bound to this workbook.
        from copy import deepcopy
        wb.add_named_style(deepcopy(_builtins[style]))

    applied = 0
    grid = None
    if location is not None:
        grid = pkg.resolve(location, default_sheet=sheet)
        gridio.guard_cell_count(grid)
        ws = wb[grid.sheet]
        for r in range(grid.min_row, grid.max_row + 1):
            for c in range(grid.min_col, grid.max_col + 1):
                ws.cell(r, c).style = style
                applied += 1
    result = pkg.save(allow_loss=allow_loss, backup=backup,
                      verify_com=verify_com)
    result["changed"]["style"] = {
        "name": style, "defined": define is not None,
        "range": grid.a1 if grid else None, "cells": applied}
    return result


def copy_format(path: str, source: Any, dest: Any, sheet: str | None = None,
                allow_loss: bool = False, backup: bool = True,
                verify_com: bool | None = None) -> dict:
    """The format painter: copy ONE source cell's complete format (font, fill,
    border, alignment, number format) onto every cell of a target range.
    Values are never touched."""
    from copy import copy as _copy
    pkg = WorkbookPackage.open(path)
    src = pkg.resolve(source, default_sheet=sheet)
    if not src.is_single:
        raise XlMcpError(
            f"source must be a single cell ({src.a1} is a range); the "
            "painter copies one cell's format, like Excel's")
    dst = pkg.resolve(dest, default_sheet=sheet)
    gridio.guard_cell_count(dst)
    sws = pkg.workbook[src.sheet]
    dws = pkg.workbook[dst.sheet]
    style = sws.cell(src.min_row, src.min_col)._style
    painted = 0
    for r in range(dst.min_row, dst.max_row + 1):
        for c in range(dst.min_col, dst.max_col + 1):
            dws.cell(r, c)._style = _copy(style)
            painted += 1
    result = pkg.save(allow_loss=allow_loss, backup=backup,
                      verify_com=verify_com)
    result["changed"]["painted"] = {
        "from": f"{src.sheet}!{src.a1}", "to": f"{dst.sheet}!{dst.a1}",
        "cells": painted}
    return result


# ----------------------------------------------------------- style audit

#: Excel's hard ceiling on distinct cell formats per workbook.
STYLE_CEILING = 64_000
_RISK_ELEVATED = 4_000
_RISK_CRITICAL = 32_000


def _styles_xml_counts(path: str) -> dict:
    """Part-level truth: element counts straight from xl/styles.xml (the
    count= attributes can lie, so children are counted)."""
    import zipfile
    from xml.etree import ElementTree as ET
    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    with zipfile.ZipFile(path) as zf:
        try:
            root = ET.fromstring(zf.read("xl/styles.xml"))
        except KeyError:
            return {"cell_formats": 0, "fonts": 0, "fills": 0, "borders": 0,
                    "custom_number_formats": 0, "named_styles": 0}

    def n(tag: str) -> int:
        el = root.find(ns + tag)
        return len(list(el)) if el is not None else 0

    return {
        "cell_formats": n("cellXfs"),
        "fonts": n("fonts"),
        "fills": n("fills"),
        "borders": n("borders"),
        "custom_number_formats": n("numFmts"),
        "named_styles": n("cellStyles"),
    }


def _describe_style(cell) -> dict:
    font = cell.font
    fdesc = font.name or "default font"
    if isinstance(font.size, (int, float)):
        fdesc += f" {font.size:g}"
    for flag, label in ((font.bold, "bold"), (font.italic, "italic"),
                        (font.underline, "underline"), (font.strike, "strike")):
        if flag:
            fdesc += f" {label}"
    fill = None
    if cell.fill is not None and cell.fill.fill_type:
        rgb = getattr(cell.fill.start_color, "rgb", None)
        fill = rgb if isinstance(rgb, str) else cell.fill.fill_type
    border = cell.border
    bordered = border is not None and any(
        side is not None and side.style
        for side in (getattr(border, s, None) for s in _SIDES))
    return {"font": fdesc, "fill": fill, "bordered": bordered,
            "number_format": cell.number_format}


def audit_styles(path: str, top: int = 10) -> dict:
    """Read-only style audit: distinct-format counts from styles.xml, usage
    per sheet, the heaviest formats by cell count, and the bloat verdict
    against the 64,000-format ceiling."""
    if isinstance(top, bool) or not isinstance(top, int) or top < 1:
        raise XlMcpError("top must be an integer >= 1")
    p = check_path(path, "audit styles")
    counts = _styles_xml_counts(p)
    wb = gridio.open_wb(path)
    try:
        usage: dict[tuple, int] = {}
        rep: dict[tuple, tuple[str, Any]] = {}
        per_sheet: dict[str, dict] = {}
        default_key = None
        for ws in wb.worksheets:
            cells = getattr(ws, "_cells", None)
            if cells is None:  # pragma: no cover - read-only worksheets
                continue
            sheet_keys = set()
            for (r, c), cell in cells.items():
                key = tuple(cell._style) if cell.has_style else "default"
                if key == "default":
                    default_key = key
                sheet_keys.add(key)
                usage[key] = usage.get(key, 0) + 1
                if key not in rep:
                    rep[key] = (f"{ws.title}!{gridio.a1(r, c)}", cell)
            per_sheet[ws.title] = {
                "cells_scanned": len(cells),
                "distinct_formats": len(sheet_keys)}
        heaviest = []
        for key, n in sorted(usage.items(), key=lambda kv: -kv[1])[:top]:
            addr, cell = rep[key]
            entry = {"cells": n, "example": addr}
            if key == default_key:
                entry["default"] = True
            else:
                entry.update(_describe_style(cell))
            heaviest.append(entry)
        n_xfs = counts["cell_formats"]
        risk = ("critical" if n_xfs >= _RISK_CRITICAL
                else "elevated" if n_xfs >= _RISK_ELEVATED else "ok")
        return {
            **counts,
            "ceiling": STYLE_CEILING,
            "risk": risk,
            "risk_thresholds": {"elevated": _RISK_ELEVATED,
                                "critical": _RISK_CRITICAL},
            "distinct_formats_in_use": len(usage),
            "per_sheet": per_sheet,
            "heaviest_formats": heaviest,
            "note": (
                "cell_formats counts the styles.xml cellXfs registry, the "
                "number Excel checks against its 64,000 ceiling; "
                "distinct_formats_in_use counts formats actually carried by "
                "cells. A large gap means orphaned registry entries "
                "(style bloat from past edits)."),
        }
    finally:
        wb.close()


__all__ = ["format_cells", "set_dimensions", "apply_style", "copy_format",
           "audit_styles", "STYLE_CEILING"]
