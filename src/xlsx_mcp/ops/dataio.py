"""ops/dataio.py: import and export (DESIGN Section 11, lite).

import_data writes CSV / TSV / JSON into a sheet at an anchor. export_range
reads a range, table, or sheet out to CSV / TSV / JSON. With query_range these
are the interop layer the demand research flagged.

import_data runs the formula-injection lint the security posture requires: a
cell whose text begins with =, +, -, or @ is data that a spreadsheet would
execute as a formula, so by default such cells are written as TEXT (never as a
live formula) and the count is reported; pass formulas=true to write them as
formulas deliberately. Truncation is loud, never silent: an import past the
write ceiling refuses rather than dropping rows.

export_range states which value mode produced the output (cached, formula, or
both), so a cell holding a formula with no cached value is never passed off as
blank. import_data routes its write through WorkbookPackage for the backup and
verify; export_range is read-only.
"""

from __future__ import annotations

import csv as _csv
import io as _io
import json as _json
import os
import re
from typing import Any

from openpyxl.utils import get_column_letter

from ..core.errors import RangeOutOfBounds, XlMcpError
from ..core.package import WorkbookPackage
from ..core.sandbox import check_path
from . import gridio

_FORMATS = ("auto", "csv", "tsv", "json")
_EXPORT_FORMATS = ("csv", "tsv", "json")
_INJECTION = ("=", "+", "-", "@")


def _detect_format(fmt: str, source_file: str | None) -> str:
    if fmt != "auto":
        return fmt
    if source_file:
        low = source_file.lower()
        if low.endswith(".json"):
            return "json"
        if low.endswith(".tsv"):
            return "tsv"
    return "csv"


def _parse(text: str, fmt: str, delimiter: str | None,
           header: bool) -> list[list[Any]]:
    if fmt == "json":
        obj = _json.loads(text)
        if isinstance(obj, dict):
            obj = [obj]
        if not isinstance(obj, list):
            raise XlMcpError("JSON import expects a list of objects or rows")
        if obj and isinstance(obj[0], dict):
            cols: list[str] = []
            for rec in obj:
                for k in rec:
                    if k not in cols:
                        cols.append(k)
            rows = [[rec.get(c) for c in cols] for rec in obj]
            return [cols] + rows if header else rows
        return [list(r) if isinstance(r, list) else [r] for r in obj]
    delim = delimiter or ("\t" if fmt == "tsv" else ",")
    reader = _csv.reader(_io.StringIO(text), delimiter=delim)
    return [[_coerce(v) for v in row] for row in reader]


def _coerce(v: str) -> Any:
    if v == "":
        return None
    try:
        if v.strip().lstrip("-").isdigit():
            return int(v)
        f = float(v)
    except ValueError:
        return v
    # float() accepts "nan"/"inf"/"infinity"; Excel has no such numbers and
    # a NaN cell value serializes as garbage, so those stay text.
    if f != f or f in (float("inf"), float("-inf")):
        return v
    return f


def import_data(path: str, source: str | None = None,
                source_file: str | None = None, fmt: str = "auto",
                location: Any = None, sheet: str | None = None,
                header: bool = True, delimiter: str | None = None,
                encoding: str = "utf-8", formulas: bool = False,
                allow_loss: bool = False, backup: bool = True) -> dict:
    """Import CSV / TSV / JSON into a sheet at an anchor. Backup + verify."""
    if fmt not in _FORMATS:
        raise XlMcpError(f"fmt must be one of {_FORMATS}")
    if not source and not source_file:
        raise XlMcpError("pass source (inline text) or source_file (a path)")
    if source_file:
        sp = check_path(source_file, "read import file")
        with open(sp, encoding=encoding) as fh:
            text = fh.read()
    else:
        text = source
    fmt = _detect_format(fmt, source_file)
    matrix = _parse(text, fmt, delimiter, header)
    rows = len(matrix)
    ncols = max((len(r) for r in matrix), default=0)
    if rows == 0 or ncols == 0:
        raise XlMcpError("the import produced no cells")
    if rows * ncols > gridio.MAX_READ_CELLS:
        raise RangeOutOfBounds(
            f"the import is {rows * ncols:,} cells, over the "
            f"{gridio.MAX_READ_CELLS:,}-cell write ceiling; split the source")

    pkg = WorkbookPackage.open(path)
    grid = pkg.resolve(location if location is not None else {"cell": "A1"},
                       default_sheet=sheet)
    top, left = grid.min_row, grid.min_col
    if top + rows - 1 > gridio._locate.MAX_ROW or \
            left + ncols - 1 > gridio._locate.MAX_COL:
        raise RangeOutOfBounds("the import would extend past the grid limits")
    ws = pkg.workbook[grid.sheet]
    neutralized = 0
    for i, row in enumerate(matrix):
        for j, val in enumerate(row):
            r, c = top + i, left + j
            if isinstance(val, str) and val[:1] in _INJECTION and not formulas:
                cell = ws.cell(r, c)
                cell.value = val
                cell.data_type = "s"
                neutralized += 1
                pkg._intended[(grid.sheet, gridio.a1(r, c))] = ("value", val)
            else:
                pkg.set_cell(grid.sheet, gridio.a1(r, c), val)
    result = pkg.save(allow_loss=allow_loss, backup=backup)
    result["changed"]["imported"] = {
        "sheet": grid.sheet, "anchor": gridio.a1(top, left),
        "rows": rows, "cols": ncols, "format": fmt}
    warnings = list(result.get("warnings", []))
    if neutralized:
        warnings.append(
            f"{neutralized} imported cell(s) began with a formula character "
            "(=, +, -, @) and were written as TEXT to block formula injection; "
            "pass formulas=true to import them as live formulas")
    result["warnings"] = warnings
    return result


def export_range(path: str, location: Any = None, sheet: str | None = None,
                 fmt: str = "csv", header: bool = True, values: str = "cached",
                 records: bool = False, out_file: str | None = None) -> dict:
    """Export a range / table / sheet to CSV / TSV / JSON. Read-only.

    location defaults to the sheet's true used range. values is cached |
    formula | both. out_file, when given, writes the serialized text to a
    sandboxed path; otherwise the text is returned inline."""
    if fmt not in _EXPORT_FORMATS:
        raise XlMcpError(f"fmt must be one of {_EXPORT_FORMATS}")
    if values not in gridio.VALUE_MODES:
        raise XlMcpError(f"values must be one of {gridio.VALUE_MODES}")
    formula_wb = gridio.open_wb(path, data_only=False) \
        if values in ("formula", "both") else None
    cached_wb = gridio.open_wb(path, data_only=True) \
        if values in ("cached", "both") else None
    base = formula_wb if formula_wb is not None else cached_wb
    try:
        loc = location if location is not None else {
            "used_range": sheet if sheet is not None else True}
        grid = gridio.resolve(base, loc, default_sheet=sheet, path=path)
        if grid.empty:
            content = "[]" if fmt == "json" else ""
            return {"sheet": grid.sheet, "range": grid.a1, "format": fmt,
                    "rows": 0, "content": content, "value_mode": values}
        gridio.guard_cell_count(grid)
        vals, _labels, _hf = gridio.read_matrix(
            grid, mode=values, formula_wb=formula_wb, cached_wb=cached_wb)
        vals = [[gridio.compact_value(v) for v in row] for row in vals]
        if fmt == "json":
            if header and vals:
                cols = [str(h) if h is not None else get_column_letter(
                    grid.min_col + i) for i, h in enumerate(vals[0])]
                body = vals[1:]
                payload = ([dict(zip(cols, r)) for r in body] if records
                           else {"columns": cols, "rows": body})
            else:
                payload = vals
            content = _json.dumps(payload, ensure_ascii=False, default=str,
                                  indent=2)
        else:
            delim = "\t" if fmt == "tsv" else ","
            buf = _io.StringIO()
            writer = _csv.writer(buf, delimiter=delim, lineterminator="\n")
            for row in vals:
                writer.writerow(["" if v is None else v for v in row])
            content = buf.getvalue()
        out: dict[str, Any] = {
            "sheet": grid.sheet, "range": grid.a1, "format": fmt,
            "rows": len(vals), "cols": len(vals[0]) if vals else 0,
            "value_mode": values}
        if out_file:
            op = check_path(out_file, "write export file")
            with open(op, "w", encoding="utf-8", newline="") as fh:
                fh.write(content)
            out["written_to"] = op
        else:
            out["content"] = content
        return out
    finally:
        for wb in (formula_wb, cached_wb):
            if wb is not None:
                wb.close()


_FILE_FORMATS = ("csv", "tsv", "json")
_UNSAFE_NAME = re.compile(r'[\\/:*?"<>|]+')


def _serialize_matrix(vals, fmt: str, header: bool, records: bool,
                      min_col: int) -> Any:
    """Shared range-to-text/payload serialization (export_range's logic,
    factored so export_file emits identical output per sheet). For json the
    PAYLOAD is returned (the caller decides bundling); csv/tsv return
    text."""
    if fmt == "json":
        if header and vals:
            cols = [str(h) if h is not None else get_column_letter(
                min_col + i) for i, h in enumerate(vals[0])]
            body = vals[1:]
            return ([dict(zip(cols, r)) for r in body] if records
                    else {"columns": cols, "rows": body})
        return vals
    delim = "\t" if fmt == "tsv" else ","
    buf = _io.StringIO()
    writer = _csv.writer(buf, delimiter=delim, lineterminator="\n")
    for row in vals:
        writer.writerow(["" if v is None else v for v in row])
    return buf.getvalue()


def export_file(path: str, fmt: str = "csv", sheets: list | None = None,
                out_dir: str | None = None, out_file: str | None = None,
                header: bool = True, values: str = "cached",
                records: bool = False) -> dict:
    """Export whole sheets (or the whole workbook) to a CSV/TSV set or one
    JSON bundle. Read-only.

    csv/tsv: one file per sheet into out_dir (named <stem>_<sheet>.csv), or
    inline per-sheet text when out_dir is omitted. json: a single bundle
    keyed by sheet name, written to out_file or returned inline."""
    if fmt not in _FILE_FORMATS:
        raise XlMcpError(f"fmt must be one of {_FILE_FORMATS}")
    if values not in gridio.VALUE_MODES:
        raise XlMcpError(f"values must be one of {gridio.VALUE_MODES}")
    if out_dir and fmt == "json":
        raise XlMcpError("json produces ONE bundle; pass out_file, not "
                         "out_dir")
    if out_file and fmt != "json":
        raise XlMcpError(f"{fmt} produces one file per sheet; pass "
                         "out_dir, not out_file")

    formula_wb = gridio.open_wb(path, data_only=False) \
        if values in ("formula", "both") else None
    cached_wb = gridio.open_wb(path, data_only=True) \
        if values in ("cached", "both") else None
    base = formula_wb if formula_wb is not None else cached_wb
    try:
        wanted = sheets if sheets else base.sheetnames
        missing = [s for s in wanted if s not in base.sheetnames]
        if missing:
            raise XlMcpError(
                f"no sheet(s) named {missing}; sheets: {base.sheetnames}")
        per_sheet: dict[str, Any] = {}
        exported: list[dict] = []
        for name in wanted:
            grid = gridio.resolve(base, {"used_range": name})
            if grid.empty:
                per_sheet[name] = [] if fmt == "json" else ""
                exported.append({"sheet": name, "rows": 0, "empty": True})
                continue
            gridio.guard_cell_count(grid)
            vals, _labels, _hf = gridio.read_matrix(
                grid, mode=values, formula_wb=formula_wb,
                cached_wb=cached_wb)
            vals = [[gridio.compact_value(v) for v in row] for row in vals]
            per_sheet[name] = _serialize_matrix(
                vals, fmt, header, records, grid.min_col)
            exported.append({"sheet": name, "range": grid.a1,
                             "rows": len(vals)})

        out: dict[str, Any] = {"format": fmt, "value_mode": values,
                               "sheets": exported,
                               "sheet_count": len(exported)}
        if fmt == "json":
            bundle = {"workbook": os.path.basename(path),
                      "value_mode": values, "sheets": per_sheet}
            content = _json.dumps(bundle, ensure_ascii=False, default=str,
                                  indent=2)
            if out_file:
                op = check_path(out_file, "write export file")
                with open(op, "w", encoding="utf-8", newline="") as fh:
                    fh.write(content)
                out["written_to"] = op
            else:
                out["content"] = content
        else:
            if out_dir:
                od = check_path(out_dir, "write export files")
                if not os.path.isdir(od):
                    raise XlMcpError(
                        f"out_dir is not an existing directory: {od}")
                stem = os.path.splitext(os.path.basename(path))[0]
                written = []
                for name, text in per_sheet.items():
                    safe = _UNSAFE_NAME.sub("_", name)
                    dest = os.path.join(od, f"{stem}_{safe}.{fmt}")
                    with open(dest, "w", encoding="utf-8",
                              newline="") as fh:
                        fh.write(text)
                    written.append(dest)
                out["written_to"] = written
            else:
                out["content"] = per_sheet
        return out
    finally:
        for wb in (formula_wb, cached_wb):
            if wb is not None:
                wb.close()


__all__ = ["import_data", "export_range", "export_file"]
