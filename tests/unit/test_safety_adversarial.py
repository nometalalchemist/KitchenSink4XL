"""Adversarial tests for the safety core, added by the re-audit.

Every test here encodes a hole the re-audit found (or probed and ruled out):
fragile part types the original HAZARD_SPECS missed (all empirically confirmed
to drop on an openpyxl round-trip), the chart-drawing heuristic bypass (a
textbox beside a chart needs no relationship, so a rels-only test misses it),
loss by REPLACEMENT rather than omission (a fragile part written back empty or
as garbage XML), the lazy read_only load letting a truncated sheet part pass
structural verify, and the keep_vba gate claiming preservation for a
vbaProject.bin inside a plain .xlsx where keep_vba never applies.
"""

from __future__ import annotations

import shutil
import zipfile
from pathlib import Path

import openpyxl
import pytest

from xlsx_mcp.core import hazard, verify
from xlsx_mcp.core.errors import HazardRefused, ValidationFailed
from xlsx_mcp.core.package import WorkbookPackage


# ------------------------------------------------------------ fixtures


def _make_clean(path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A1"] = 1
    ws["A2"] = 2
    wb.save(path)
    wb.close()
    return path


def _inject_part(src: Path, dst: Path, member: str, data: bytes,
                 content_type: str | None = None) -> Path:
    """Copy src to dst adding one extra part (and a content-type override so
    the fixture is realistic enough for openpyxl to load)."""
    with zipfile.ZipFile(src) as zin, \
            zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            payload = zin.read(item.filename)
            if item.filename == "[Content_Types].xml" and content_type:
                payload = payload.replace(
                    b"</Types>",
                    ('<Override PartName="/' + member + '" ContentType="'
                     + content_type + '"/></Types>').encode())
            zout.writestr(item, payload)
        zout.writestr(member, data)
    return dst


def _replace_part(src: Path, dst: Path, member: str, data: bytes) -> Path:
    with zipfile.ZipFile(src) as zin, \
            zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            payload = data if item.filename == member else zin.read(
                item.filename)
            zout.writestr(item, payload)
    return dst


# --------------------------------------- missing fragile-part classes (fixed)

# Every one of these was empirically confirmed by the re-audit to be silently
# DROPPED by an openpyxl load+save while the pre-audit scan reported clean.
_NEW_DROP_PARTS = [
    ("xl/embeddings/oleObject1.bin", "embeddings"),
    ("xl/persons/person.xml", "threaded_comments"),
    ("xl/richData/rdrichvalue.xml", "rich_metadata"),
    ("xl/namedSheetViews/namedSheetView1.xml", "named_sheet_views"),
    ("xl/queryTables/queryTable1.xml", "query_tables"),
    ("xl/model/item.data", "data_model"),
    ("customXml/item1.xml", "power_query"),
]


@pytest.mark.parametrize("part,key", _NEW_DROP_PARTS)
def test_previously_missed_drop_parts_are_flagged(part, key):
    r = hazard.scan_names(["xl/workbook.xml", part])
    assert key in [h.key for h in r.hazards], (
        f"{part} drops on an openpyxl round-trip and must be flagged")
    assert r.would_lose is True
    assert hazard.route(r, edit="structural",
                        com_available=False)[0] == hazard.ROUTE_REFUSE


def test_external_links_deliberately_not_flagged():
    """openpyxl MODELS external links and preserves them on round-trip
    (verified empirically by the re-audit), so they carry no spec."""
    r = hazard.scan_names(["xl/workbook.xml",
                           "xl/externalLinks/externalLink1.xml"])
    assert r.clean is True


def test_printer_settings_deliberately_not_flagged():
    """printerSettings drops on round-trip, but flagging it would refuse
    nearly every workbook that ever printed, over a cosmetic loss. Policy
    documented in the hazard module docstring."""
    r = hazard.scan_names(["xl/workbook.xml",
                           "xl/printerSettings/printerSettings1.bin"])
    assert r.clean is True


def test_custom_xml_label_is_honest_about_scope():
    """The customXml flag covers MORE than Power Query (add-in stores,
    SharePoint property sets); the label must not claim PQ specifically."""
    r = hazard.scan_names(["customXml/itemProps1.xml"])
    assert r.hazards and "customXml" in r.hazards[0].label


# ------------------------------------- chart-vs-shape heuristic bypass (fixed)

_CHART_REL = (b'<Relationships><Relationship Type="http://schemas.'
              b'openxmlformats.org/officeDocument/2006/relationships/chart" '
              b'Target="/xl/charts/chart1.xml" Id="rId1"/></Relationships>')

_DRAWING_CHART_ONLY = (
    b'<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/'
    b'drawingml/2006/spreadsheetDrawing">'
    b'<xdr:twoCellAnchor><xdr:graphicFrame/></xdr:twoCellAnchor>'
    b'</xdr:wsDr>')

_DRAWING_CHART_PLUS_TEXTBOX = (
    b'<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/'
    b'drawingml/2006/spreadsheetDrawing">'
    b'<xdr:twoCellAnchor><xdr:graphicFrame/></xdr:twoCellAnchor>'
    b'<xdr:twoCellAnchor><xdr:sp/></xdr:twoCellAnchor>'
    b'</xdr:wsDr>')


def _drawing_reader(rels: bytes, drawing: bytes):
    def reader(member: str) -> bytes:
        return rels if member.endswith(".rels") else drawing
    return reader


def test_chart_plus_inline_textbox_is_still_a_drop():
    """THE BYPASS: a textbox drawn beside a chart in the same drawing has no
    relationship of its own, so its rels are all-chart. The pre-audit rels-only
    heuristic called this chart-only and would have silently dropped the
    textbox. The drawing XML content check closes it."""
    names = ["xl/drawings/drawing1.xml",
             "xl/drawings/_rels/drawing1.xml.rels",
             "xl/charts/chart1.xml"]
    r = hazard.scan_names(
        names, rels_reader=_drawing_reader(_CHART_REL,
                                           _DRAWING_CHART_PLUS_TEXTBOX))
    assert "drawings" in [h.key for h in r.hazards]
    assert r.would_lose is True


def test_pure_chart_drawing_still_passes_content_check():
    names = ["xl/drawings/drawing1.xml",
             "xl/drawings/_rels/drawing1.xml.rels",
             "xl/charts/chart1.xml"]
    r = hazard.scan_names(
        names, rels_reader=_drawing_reader(_CHART_REL, _DRAWING_CHART_ONLY))
    assert "drawings" not in [h.key for h in r.hazards]
    assert "charts" in [h.key for h in r.hazards]


def test_unparseable_drawing_is_conservatively_flagged():
    names = ["xl/drawings/drawing1.xml",
             "xl/drawings/_rels/drawing1.xml.rels"]
    r = hazard.scan_names(
        names, rels_reader=_drawing_reader(_CHART_REL, b"<not xml"))
    assert "drawings" in [h.key for h in r.hazards]


# ------------------------------------------- loss by replacement (fixed)


def test_fragile_part_replaced_with_empty_fails_verify(tmp_path):
    src = _make_clean(tmp_path / "src.xlsx")
    withpart = _inject_part(
        src, tmp_path / "withpart.xlsx", "xl/slicers/slicer1.xml",
        b"<slicers>real content</slicers>",
        "application/vnd.ms-excel.slicer+xml")
    pre = hazard.scan_path(str(withpart))
    emptied = _replace_part(withpart, tmp_path / "emptied.xlsx",
                            "xl/slicers/slicer1.xml", b"")
    result = verify.verify_after_write(
        str(emptied), pre_parts=pre.parts, pre_sizes=pre.sizes)
    assert result.ok is False
    assert any("replaced with empty" in r for r in result.reasons)


def test_fragile_xml_part_replaced_with_garbage_fails_verify(tmp_path):
    src = _make_clean(tmp_path / "src.xlsx")
    withpart = _inject_part(
        src, tmp_path / "withpart.xlsx", "xl/slicers/slicer1.xml",
        b"<slicers>real content</slicers>",
        "application/vnd.ms-excel.slicer+xml")
    pre = hazard.scan_path(str(withpart))
    garbled = _replace_part(withpart, tmp_path / "garbled.xlsx",
                            "xl/slicers/slicer1.xml", b"\x00garbage not xml")
    result = verify.verify_after_write(
        str(garbled), pre_parts=pre.parts, pre_sizes=pre.sizes)
    assert result.ok is False
    assert any("no longer well-formed" in r for r in result.reasons)


def test_untouched_fragile_part_passes_content_check(tmp_path):
    src = _make_clean(tmp_path / "src.xlsx")
    withpart = _inject_part(
        src, tmp_path / "withpart.xlsx", "xl/slicers/slicer1.xml",
        b"<slicers>real content</slicers>",
        "application/vnd.ms-excel.slicer+xml")
    pre = hazard.scan_path(str(withpart))
    ok, problems = verify.part_content_check(pre.sizes, str(withpart))
    assert ok is True and problems == []


# ---------------------------- lazy-load blind spot in structural verify (fixed)


def test_truncated_sheet_xml_fails_structural_check(tmp_path):
    """openpyxl's read_only load is lazy and never parses sheet XML, so before
    the re-audit a truncated/garbage sheet part passed structural verify
    whenever no cell read-back ran (exactly the structural-edit saves, which
    clear their cell intents). The streaming XML parse closes this."""
    src = _make_clean(tmp_path / "src.xlsx")
    with zipfile.ZipFile(src) as zf:
        sheet = next(n for n in zf.namelist()
                     if n.startswith("xl/worksheets/") and n.endswith(".xml"))
        good = zf.read(sheet)
    truncated = _replace_part(src, tmp_path / "trunc.xlsx", sheet,
                              good[: len(good) // 2])
    ok, reasons = verify.structural_check(str(truncated))
    assert ok is False
    assert any("not well-formed XML" in r for r in reasons)


def test_structural_verify_failure_leaves_original_untouched(tmp_path):
    """End to end through the save pipeline: a saver that produces a garbage
    sheet (a raw-surgery failure mode) must refuse and leave the original
    byte-for-byte unchanged."""
    path = _make_clean(tmp_path / "wb.xlsx")
    original = path.read_bytes()
    pkg = WorkbookPackage.open(str(path))
    pkg.set_cell("S", "A3", 3)

    def bad_saver(tmp: str) -> None:
        with zipfile.ZipFile(path) as zin, \
                zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename.startswith("xl/worksheets/"):
                    data = data[: len(data) // 2]
                zout.writestr(item, data)

    with pytest.raises(ValidationFailed):
        pkg.save(saver=bad_saver)
    assert path.read_bytes() == original


# --------------------------------------- keep_vba gate on a .xlsx path (fixed)


def _make_xlsx_with_vba(tmp_path: Path) -> Path:
    src = _make_clean(tmp_path / "plain.xlsx")
    return _inject_part(
        src, tmp_path / "renamed_macro.xlsx", "xl/vbaProject.bin",
        b"\xd0\xcf\x11\xe0fake-vba", "application/vnd.ms-office.vbaProject")


def test_vba_in_plain_xlsx_refuses_instead_of_claiming_keep_vba(tmp_path):
    """keep_vba only applies to .xlsm paths. A vbaProject.bin inside a plain
    .xlsx (renamed file) would be DROPPED by the save, but the pre-audit gate
    warned 'preserved via keep_vba'. It must refuse as a hard drop."""
    path = _make_xlsx_with_vba(tmp_path)
    pkg = WorkbookPackage.open(str(path))
    pkg.set_cell("S", "A3", 3)
    with pytest.raises(HazardRefused) as exc_info:
        pkg.save()
    assert "VBA" in str(exc_info.value)


def test_vba_in_plain_xlsx_allow_loss_proceeds_with_honest_warning(tmp_path):
    path = _make_xlsx_with_vba(tmp_path)
    pkg = WorkbookPackage.open(str(path))
    pkg.set_cell("S", "A3", 3)
    result = pkg.save(allow_loss=True)
    assert result["ok"] is True
    assert any("drops" in w for w in result["warnings"])
    assert not any("preserved via keep_vba" in w for w in result["warnings"])
    # and the part really is gone (which is exactly why the refusal exists)
    with zipfile.ZipFile(path) as zf:
        assert "xl/vbaProject.bin" not in zf.namelist()


def test_xlsm_keep_vba_warning_still_accurate(tmp_path):
    """For a real .xlsm the conditional warning stays (keep_vba IS applied)."""
    src = _make_clean(tmp_path / "m.xlsx")
    macro = _inject_part(
        src, tmp_path / "m.xlsm", "xl/vbaProject.bin",
        b"\xd0\xcf\x11\xe0fake-vba", "application/vnd.ms-office.vbaProject")
    pkg = WorkbookPackage.open(str(macro))
    pkg.set_cell("S", "A3", 3)
    result = pkg.save()
    assert result["ok"] is True
    assert any("keep_vba" in w for w in result["warnings"])


# --------------------------------------------------- scan sizes plumbing


def test_scan_path_reports_part_sizes(tmp_path):
    path = _make_clean(tmp_path / "wb.xlsx")
    rep = hazard.scan_path(str(path))
    assert rep.sizes, "scan_path must fill uncompressed part sizes"
    assert set(rep.sizes) == set(rep.parts)
    assert all(isinstance(v, int) for v in rep.sizes.values())


def test_package_save_end_to_end_still_green(tmp_path):
    """Regression guard: the new verify stages must not break a normal save."""
    path = _make_clean(tmp_path / "wb.xlsx")
    pkg = WorkbookPackage.open(str(path))
    pkg.set_cell("S", "A3", "=SUM(A1:A2)")
    result = pkg.save()
    assert result["ok"] is True and result["verified"] is True
    shutil.rmtree(tmp_path / ".ks4xl-backups", ignore_errors=True)
