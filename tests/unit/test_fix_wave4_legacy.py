"""Regressions for fix wave 4 -- geriatric round H-1, L-1, M-1
(2026-09-05).

H-1: a genuine .xls escaped the envelope as a raw InvalidFileException
(advice pointing at xlrd); every file entry point now sniffs the 8-byte OLE
magic and refuses honestly, naming the format and com_convert_format, with
the encrypted-package kind told apart by stream names. L-1: renamed BIFF
under .xlsx got the same misdiagnosis; the sniff is extension-blind. M-1:
Excel names comment anchors vmlDrawing*.vml, so every real old file with a
sticky note tripped HAZARD_REFUSED on plain edits; a VML content classifier
now demotes content-verified Note-only anchors while form controls, drawn
shapes, orphans, and unparseable parts stay flagged.
"""

from __future__ import annotations

import shutil
import struct
import time
import zipfile
from pathlib import Path

import openpyxl
import pytest

from xlsx_mcp.core import calc as core_calc
from xlsx_mcp.core import hazard as core_hazard
from xlsx_mcp.core.errors import (
    HazardRefused,
    UnsupportedStructure,
    ValidationFailed,
    WorkbookProtected,
)
from xlsx_mcp.ops import cells as cells_ops
from xlsx_mcp.ops import lifecycle
from xlsx_mcp.ops import structure as structure_ops

FULL = "0.30000000000000004"          # 0.1 + 0.2 as Excel stores it


def _patch_part(path: Path, part: str, transform) -> None:
    """Rewrite one zip member in place (transform: bytes -> bytes)."""
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(path) as zin,             zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == part:
                data = transform(data)
            zout.writestr(item, data)
    tmp.replace(path)


def _sheet_xml(path: Path) -> str:
    with zipfile.ZipFile(path) as z:
        return z.read("xl/worksheets/sheet1.xml").decode("utf-8")


def _rename_part(path: Path, old: str, new: str,
                 patch_refs: dict[str, tuple[bytes, bytes]] | None = None
                 ) -> None:
    """Rename a zip member and patch references in other members."""
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(path) as zin,             zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            name = new if item.filename == old else item.filename
            if patch_refs and item.filename in patch_refs:
                a, b = patch_refs[item.filename]
                data = data.replace(a, b)
            zout.writestr(name, data)
    tmp.replace(path)


# ======================================================================
# Geriatric H-1 (HIGH) + L-1: legacy OLE containers refuse honestly
# ======================================================================


def _cfb_bytes(streams: list[str]) -> bytes:
    """A minimal OLE compound file: valid header + one directory sector
    holding a root entry plus the named streams."""
    header = bytearray(512)
    header[0:8] = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
    struct.pack_into("<H", header, 26, 3)      # major version 3
    struct.pack_into("<H", header, 28, 0xFFFE)  # little-endian
    struct.pack_into("<H", header, 30, 9)      # 512-byte sectors
    struct.pack_into("<H", header, 32, 6)      # mini sector shift
    struct.pack_into("<i", header, 44, 1)      # one FAT sector
    struct.pack_into("<i", header, 48, 0)      # first directory sector = 0
    for off in range(76, 512, 4):
        struct.pack_into("<i", header, off, -1)

    def entry(name: str, etype: int) -> bytes:
        e = bytearray(128)
        raw = name.encode("utf-16-le")
        e[0:len(raw)] = raw
        struct.pack_into("<H", e, 64, len(raw) + 2)
        e[66] = etype
        e[67] = 1
        struct.pack_into("<i", e, 68, -1)
        struct.pack_into("<i", e, 72, -1)
        struct.pack_into("<i", e, 76, -1)
        return bytes(e)

    sector = bytearray(512)
    entries = [entry("Root Entry", 5)] + [entry(s, 2) for s in streams]
    for i, e in enumerate(entries[:4]):
        sector[i * 128:(i + 1) * 128] = e
    return bytes(header) + bytes(sector)


def _biff_file(tmp_path: Path, name: str) -> Path:
    p = tmp_path / name
    p.write_bytes(_cfb_bytes(["Workbook"]))
    return p


class TestLegacyXlsHonestRefusal:
    def test_kind_classifier(self, tmp_path):
        biff = _biff_file(tmp_path, "old.xls")
        assert core_hazard.ole_container_kind(str(biff)) == "biff"
        enc = tmp_path / "enc.xlsx"
        enc.write_bytes(_cfb_bytes(["EncryptionInfo", "EncryptedPackage"]))
        assert core_hazard.ole_container_kind(str(enc)) == "encrypted"
        bare = tmp_path / "bare.xls"
        bare.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 504)
        assert core_hazard.ole_container_kind(str(bare)) == "ole"
        trunc = tmp_path / "trunc.xls"
        trunc.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64)
        assert core_hazard.ole_container_kind(str(trunc)) == "ole"
        wb = openpyxl.Workbook()
        real = tmp_path / "real.xlsx"
        wb.save(real)
        assert core_hazard.ole_container_kind(str(real)) is None

    @pytest.mark.parametrize("fname", ["old.xls", "old_lying.xlsx"])
    def test_reads_refuse_naming_format_and_remedy(self, tmp_path, fname):
        # The matrix rows: same honest refusal whatever the extension says.
        book = _biff_file(tmp_path, fname)
        from xlsx_mcp.ops import gridio
        with pytest.raises(UnsupportedStructure) as ei:
            gridio.open_wb(str(book))
        msg = str(ei.value)
        assert "97-2003" in msg and "BIFF" in msg
        assert "com_convert_format" in msg
        assert "xlrd" not in msg

    @pytest.mark.parametrize("fname", ["old.xls", "old_lying.xlsx"])
    def test_mutations_refuse_the_same_way(self, tmp_path, fname):
        book = _biff_file(tmp_path, fname)
        before = book.read_bytes()
        with pytest.raises(UnsupportedStructure) as ei:
            cells_ops.set_cell(str(book), {"cell": "A1"}, 1)
        assert "com_convert_format" in str(ei.value)
        assert book.read_bytes() == before

    def test_metadata_and_diagnose_refuse_or_report_honestly(self, tmp_path):
        book = _biff_file(tmp_path, "meta.xls")
        with pytest.raises(UnsupportedStructure):
            lifecycle.get_workbook_metadata(str(book))
        d = lifecycle.diagnose_workbook(str(book))
        assert d["readable"] is False
        assert "com_convert_format" in d["error"]

    def test_pure_ole_refuses_with_both_possibilities(self, tmp_path):
        bare = tmp_path / "mystery.xls"
        bare.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 504)
        from xlsx_mcp.ops import gridio
        with pytest.raises(UnsupportedStructure) as ei:
            gridio.open_wb(str(bare))
        msg = str(ei.value)
        assert "com_convert_format" in msg and "password" in msg

    def test_encrypted_package_refuses_as_protected(self, tmp_path):
        enc = tmp_path / "enc.xlsx"
        enc.write_bytes(_cfb_bytes(["EncryptionInfo", "EncryptedPackage"]))
        from xlsx_mcp.ops import gridio
        with pytest.raises(WorkbookProtected) as ei:
            gridio.open_wb(str(enc))
        assert "password" in str(ei.value)

    def test_invalid_file_exception_backstop_in_envelope(self):
        from openpyxl.utils.exceptions import InvalidFileException
        from xlsx_mcp import envelope
        assert any(issubclass(InvalidFileException, t)
                   for t in envelope.CATCHABLE)
        assert envelope.classify(InvalidFileException("x")) == \
            "UNSUPPORTED_CONTENT"

    def test_validate_structure_check_names_the_format(self, tmp_path):
        from xlsx_mcp.ops import validation
        book = _biff_file(tmp_path, "val.xls")
        ok, findings = validation._check_structure(str(book))
        assert ok is False
        assert any("BIFF" in i for i in findings["issues"])

    def test_scan_path_reports_ole_error_kind(self, tmp_path):
        book = _biff_file(tmp_path, "scan.xls")
        rep = core_hazard.scan_path(str(book))
        assert rep.error_kind == "ole"
        assert "com_convert_format" in rep.error

    def test_com_tier_no_longer_calls_biff_encrypted(self, tmp_path):
        # H-1 corollary: is_encrypted_package used to be a bare CFB-magic
        # check, so the ONE useful route (com_convert_format) refused every
        # genuine .xls as "password-protected".
        from xlsx_mcp.com import session as com_session
        biff = _biff_file(tmp_path, "conv.xls")
        assert com_session.is_encrypted_package(str(biff)) is False
        enc = tmp_path / "enc2.xlsx"
        enc.write_bytes(_cfb_bytes(["EncryptionInfo", "EncryptedPackage"]))
        assert com_session.is_encrypted_package(str(enc)) is True
        # unclassifiable OLE stays possibly-encrypted (modal-prompt risk)
        trunc = tmp_path / "t.xls"
        trunc.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64)
        assert com_session.is_encrypted_package(str(trunc)) is True


# ======================================================================
# Geriatric M-1 (MEDIUM): Excel-named comment anchors stop over-refusing
# ======================================================================


def _excel_named_comment_book(tmp_path: Path, name: str) -> Path:
    """A workbook with a legacy comment whose VML anchor carries EXCEL's
    naming (vmlDrawing1.vml), the way every real old file arrives."""
    p = tmp_path / name
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A1"] = "data"
    from openpyxl.comments import Comment
    ws["B2"].comment = Comment("note text", "Author")
    wb.save(p)
    _rename_part(
        p, "xl/drawings/commentsDrawing1.vml", "xl/drawings/vmlDrawing1.vml",
        patch_refs={"xl/worksheets/_rels/sheet1.xml.rels":
                    (b"commentsDrawing1.vml", b"vmlDrawing1.vml")})
    return p


class TestExcelNamedCommentAnchors:
    def test_scan_demotes_content_verified_comment_anchor(self, tmp_path):
        book = _excel_named_comment_book(tmp_path, "note.xlsx")
        rep = core_hazard.scan_path(str(book))
        assert "drawings" not in {h.key for h in rep.hazards}
        assert rep.comment_anchor_vml == ["xl/drawings/vmlDrawing1.vml"]

    def test_plain_edit_no_longer_needs_allow_loss(self, tmp_path):
        book = _excel_named_comment_book(tmp_path, "note2.xlsx")
        r = cells_ops.set_cell(str(book), {"cell": "A1"}, "edited")
        assert r["ok"] and r["verified"]
        wb = openpyxl.load_workbook(book)
        c = wb["S"]["B2"].comment
        assert c is not None and c.text == "note text" \
            and c.author == "Author"
        wb.close()

    def test_form_control_vml_stays_flagged(self, tmp_path):
        book = _excel_named_comment_book(tmp_path, "ctrl.xlsx")
        _patch_part(book, "xl/drawings/vmlDrawing1.vml",
                    lambda b: b.replace(b'ObjectType="Note"',
                                        b'ObjectType="Checkbox"'))
        rep = core_hazard.scan_path(str(book))
        assert "drawings" in {h.key for h in rep.hazards}
        with pytest.raises(HazardRefused):
            cells_ops.set_cell(str(book), {"cell": "A1"}, "edited")

    def test_unparseable_vml_stays_flagged(self, tmp_path):
        book = _excel_named_comment_book(tmp_path, "junk.xlsx")
        _patch_part(book, "xl/drawings/vmlDrawing1.vml",
                    lambda b: b"<xml><v:shape unclosed</xml>")
        rep = core_hazard.scan_path(str(book))
        assert "drawings" in {h.key for h in rep.hazards}

    def test_drawn_shape_beside_note_stays_flagged(self, tmp_path):
        book = _excel_named_comment_book(tmp_path, "mix.xlsx")
        _patch_part(book, "xl/drawings/vmlDrawing1.vml",
                    lambda b: b.replace(
                        b"</xml>",
                        b'<ns9:oval xmlns:ns9="urn:schemas-microsoft-com:'
                        b'vml"/></xml>'))
        rep = core_hazard.scan_path(str(book))
        assert "drawings" in {h.key for h in rep.hazards}

    def test_orphan_note_anchor_without_comments_stays_flagged(self, tmp_path):
        # No comments part anywhere: openpyxl will not regenerate the
        # anchor, so the demotion must not fire.
        book = _excel_named_comment_book(tmp_path, "orphan.xlsx")
        tmp = book.with_suffix(".tmp")
        with zipfile.ZipFile(book) as zin, \
                zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                if item.filename == "xl/comments/comment1.xml":
                    continue
                zout.writestr(item, zin.read(item.filename))
        tmp.replace(book)
        rep = core_hazard.scan_path(str(book))
        assert "drawings" in {h.key for h in rep.hazards}

    def test_namelist_only_scan_stays_conservative(self):
        rep = core_hazard.scan_names(
            ["xl/workbook.xml", "xl/comments1.xml",
             "xl/drawings/vmlDrawing1.vml"])
        assert "drawings" in {h.key for h in rep.hazards}
