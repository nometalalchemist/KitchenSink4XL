"""Re-audit tests for core/refs.py: whole-column / whole-row span
transposition (the previously documented gap, closed because it was
load-bearing: =SUM(B:B) went stale on every column edit and print-title
names like $1:$2 went stale on row inserts), copy-offset semantics for
spans, and ArrayFormula cell rewriting. All expectations hand-computed
against Excel's own adjustment behavior.
"""

from __future__ import annotations

import openpyxl
from openpyxl.worksheet.formula import ArrayFormula

from xlsx_mcp.core import refs

R = refs.RefEdit


def tf(formula, edit, sheet="S"):
    return refs.transpose_formula(formula, edit, sheet)


# ------------------------------------------------- whole-column spans


def test_whole_col_shifts_on_col_insert():
    e = R("S", "insert_cols", index=2, count=1)
    assert tf("=SUM(B:B)", e) == "=SUM(C:C)"
    assert tf("=SUM(A:A)", e) == "=SUM(A:A)"      # before the insert
    assert tf("=SUM(A:C)", e) == "=SUM(A:D)"      # insert inside expands


def test_whole_col_delete_semantics():
    e = R("S", "delete_cols", index=2, count=1)
    assert tf("=SUM(B:B)", e) == "=SUM(#REF!)"    # wholly deleted
    assert tf("=SUM(A:C)", e) == "=SUM(A:B)"      # partial shrinks
    assert tf("=SUM(C:E)", e) == "=SUM(B:D)"      # after the band shifts


def test_whole_col_untouched_by_row_edits_and_move():
    assert tf("=SUM(B:B)", R("S", "insert_rows", index=1, count=5)) \
        == "=SUM(B:B)"
    assert tf("=SUM(B:B)", R("S", "delete_rows", index=1, count=5)) \
        == "=SUM(B:B)"
    assert tf("=SUM(B:B)", R("S", "move", src=(1, 1, 5, 5), dst=(10, 10))) \
        == "=SUM(B:B)"


def test_whole_col_absolute_anchors_preserved():
    e = R("S", "insert_cols", index=1, count=2)
    assert tf("=SUM($B:$B)+SUM($A:C)", e) == "=SUM($D:$D)+SUM($C:E)"


def test_whole_col_sheet_scope():
    e = R("Data", "insert_cols", index=1, count=1)
    assert refs.transpose_formula("=SUM(Data!B:B)+SUM(B:B)", e, "Summary") \
        == "=SUM(Data!C:C)+SUM(B:B)"


# ------------------------------------------------- whole-row spans


def test_whole_row_shifts_on_row_insert():
    e = R("S", "insert_rows", index=2, count=1)
    assert tf("=SUM(4:6)", e) == "=SUM(5:7)"
    assert tf("=SUM(1:1)", e) == "=SUM(1:1)"


def test_whole_row_delete_semantics():
    e = R("S", "delete_rows", index=2, count=2)
    assert tf("=SUM(2:3)", e) == "=SUM(#REF!)"
    assert tf("=SUM(1:5)", e) == "=SUM(1:3)"
    assert tf("=SUM(4:6)", e) == "=SUM(2:4)"


def test_whole_row_untouched_by_col_edits():
    assert tf("=SUM(2:3)", R("S", "insert_cols", index=1, count=2)) \
        == "=SUM(2:3)"


def test_print_titles_style_bare_ref():
    # _xlnm.Print_Titles definitions carry $1:$2 style refs
    assert refs.transpose_ref("$1:$2", R("S", "insert_rows", index=1,
                                         count=1), "S") == "$2:$3"
    assert refs.transpose_ref("$1:$2", R("S", "delete_rows", index=1,
                                         count=2), "S") is None


# --------------------------------------------- guards around the new pattern


def test_cell_ranges_still_win_over_spans():
    e = R("S", "insert_rows", index=1, count=1)
    assert tf("=SUM(A1:B2)", e) == "=SUM(A2:B3)"


def test_time_literal_strings_untouched():
    e = R("S", "insert_rows", index=1, count=5)
    assert tf('=IF(A1>0,"12:30","")', e) == '=IF(A6>0,"12:30","")'


def test_3d_sheet_prefix_not_mangled_by_span_pattern():
    e = R("Sheet1", "insert_rows", index=1, count=1)
    got = refs.transpose_formula("=SUM(Sheet1:Sheet3!A1)", e, "Other")
    assert got == "=SUM(Sheet1:Sheet3!A2)"


def test_malformed_half_spans_never_match_as_spans():
    """'A1:B' and 'A:B2' are not legal Excel refs; the span pattern must not
    half-match them. The CELL half still shifts as a standalone cell, which is
    required behavior for the legitimate range-operator form
    '=SUM(INDEX(...):B5)' where a bare cell legitimately follows ':'."""
    e = R("S", "insert_cols", index=1, count=1)
    assert tf("=A1:B", e) == "=B1:B"     # A1 shifts as a cell; ':B' inert
    assert tf("=SUM(A:B2)", e) == "=SUM(A:C2)"  # B2 shifts as a cell; no span
    # the real range-operator case that behavior exists for:
    assert tf("=SUM(INDEX(C:C,1):B5)", e) == "=SUM(INDEX(D:D,1):C5)"


# ------------------------------------------------- offset (copy) semantics


def test_offset_shifts_relative_spans_only():
    assert refs.offset_formula("=SUM(B:B)", 5, 2) == "=SUM(D:D)"
    assert refs.offset_formula("=SUM($B:$B)", 5, 2) == "=SUM($B:$B)"
    assert refs.offset_formula("=SUM(3:3)", 2, 7) == "=SUM(5:5)"
    assert refs.offset_formula("=SUM($3:$3)", 2, 7) == "=SUM($3:$3)"


def test_offset_span_off_grid_is_ref_error():
    assert refs.offset_formula("=SUM(A:A)", 0, -1) == "=SUM(#REF!)"
    assert refs.offset_formula("=SUM(1:1)", -1, 0) == "=SUM(#REF!)"


# ------------------------------------------------- ArrayFormula rewriting


def test_array_formula_text_and_ref_rewritten_on_insert():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    for r in range(1, 6):
        ws.cell(r, 2, r)
    ws["D1"] = ArrayFormula("D1:D5", "=B1:B5*2")
    rep = refs.rewrite_workbook(wb, R("S", "insert_rows", index=1, count=2))
    val = ws["D1"].value
    assert isinstance(val, ArrayFormula)
    assert val.text == "=B3:B7*2"
    assert val.ref == "D3:D7"
    assert rep.formulas == 1


def test_array_formula_ref_error_counted_on_delete():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    ws["D1"] = ArrayFormula("D1", "=B2*2")
    rep = refs.rewrite_workbook(wb, R("S", "delete_rows", index=2, count=1))
    val = ws["D1"].value
    assert isinstance(val, ArrayFormula)
    assert val.text == "=#REF!*2"
    assert rep.ref_errors == 1


def test_plain_string_formulas_unaffected_by_array_path():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A1"] = "=SUM(B:B)"
    refs.rewrite_workbook(wb, R("S", "insert_cols", index=1, count=1))
    assert ws["A1"].value == "=SUM(C:C)"
