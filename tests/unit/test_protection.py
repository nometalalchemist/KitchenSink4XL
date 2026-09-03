"""set_protection tests: every mutating action round-trips through the full
save pipeline and is re-read from a fresh model load; the option-inversion
semantics (user 'allowed' vs file 'blocked') are pinned; refusal paths
cover the closed vocabulary; status is proven read-only byte-for-byte.
"""

from __future__ import annotations

import hashlib

import openpyxl
import pytest

from xlsx_mcp.core.errors import TargetNotFound, XlMcpError
from xlsx_mcp.ops import lifecycle as _lifecycle
from xlsx_mcp.ops import protection as _protection


def _make(tmp_path, sheets=("S", "T")):
    p = str(tmp_path / "p.xlsx")
    _lifecycle.create_workbook(p, sheets=list(sheets))
    return p


def _md5(path):
    with open(path, "rb") as fh:
        return hashlib.md5(fh.read()).hexdigest()


def test_sheet_protection_round_trip(tmp_path):
    p = _make(tmp_path)
    out = _protection.set_protection(
        p, "sheet", sheet="S", password="pw",
        options={"sort": True, "format_cells": True, "insert_rows": False})
    assert out["ok"] is True and out["verified"] is True
    wb = openpyxl.load_workbook(p)
    prot = wb["S"].protection
    assert prot.sheet is True
    assert prot.password  # the legacy hash landed
    # allowed -> attr False (not blocked); disallowed -> attr True
    assert prot.sort is False
    assert prot.formatCells is False
    assert prot.insertRows is True
    # the other sheet is untouched
    assert wb["T"].protection.sheet is not True
    wb.close()


def test_option_inversion_matches_status_report(tmp_path):
    p = _make(tmp_path)
    _protection.set_protection(p, "sheet", sheet="S",
                               options={"auto_filter": True})
    st = _protection.set_protection(p, "status")
    entry = next(s for s in st["sheets"] if s["sheet"] == "S")
    assert entry["protected"] is True
    assert "auto_filter" in entry["allowed_while_protected"]


def test_workbook_structure_lock_round_trip(tmp_path):
    p = _make(tmp_path)
    out = _protection.set_protection(p, "workbook", password="wpw",
                                     structure=True, windows=True)
    assert out["verified"] is True
    wb = openpyxl.load_workbook(p)
    assert wb.security.lockStructure is True
    assert wb.security.lockWindows is True
    assert wb.security.workbookPassword  # hashed, present
    wb.close()


def test_unlock_ranges_round_trip(tmp_path):
    p = _make(tmp_path)
    out = _protection.set_protection(
        p, "unlock", sheet="S", unlock_ranges=["B2:C3", {"cell": "E5"}])
    assert out["verified"] is True
    wb = openpyxl.load_workbook(p)
    ws = wb["S"]
    assert ws["B2"].protection.locked is False
    assert ws["C3"].protection.locked is False
    assert ws["E5"].protection.locked is False
    assert ws["A1"].protection.locked is True  # untouched default
    wb.close()
    # relock via locked=true
    _protection.set_protection(p, "unlock", sheet="S",
                               unlock_ranges=["B2"], locked=True)
    wb = openpyxl.load_workbook(p)
    assert wb["S"]["B2"].protection.locked is True
    wb.close()


def test_remove_scopes(tmp_path):
    p = _make(tmp_path)
    _protection.set_protection(p, "sheet", sheet="S")
    _protection.set_protection(p, "workbook")
    out = _protection.set_protection(p, "remove", scope="sheet", sheet="S")
    assert out["changed"]["protection"]["removed"] == ["sheet:S"]
    st = _protection.set_protection(p, "status")
    assert st["workbook"]["structure_locked"] is True
    _protection.set_protection(p, "remove", scope="workbook")
    st = _protection.set_protection(p, "status")
    assert st["workbook"]["structure_locked"] is False
    assert all(not s["protected"] for s in st["sheets"])


def test_remove_nothing_refuses(tmp_path):
    p = _make(tmp_path)
    with pytest.raises(TargetNotFound, match="nothing to remove"):
        _protection.set_protection(p, "remove", scope="all")


def test_refusal_paths(tmp_path):
    p = _make(tmp_path)
    with pytest.raises(XlMcpError, match="action must be one of"):
        _protection.set_protection(p, "encrypt")
    with pytest.raises(XlMcpError, match="unknown protection option"):
        _protection.set_protection(p, "sheet", sheet="S",
                                   options={"format_cellz": True})
    with pytest.raises(TargetNotFound, match="no sheet named"):
        _protection.set_protection(p, "sheet", sheet="Nope")
    with pytest.raises(XlMcpError, match="needs unlock_ranges"):
        _protection.set_protection(p, "unlock", sheet="S")
    with pytest.raises(XlMcpError, match="scope must be"):
        _protection.set_protection(p, "remove", scope="everything")


def test_status_is_read_only(tmp_path):
    p = _make(tmp_path)
    _protection.set_protection(p, "sheet", sheet="S", password="x")
    before = _md5(p)
    _protection.set_protection(p, "status")
    assert _md5(p) == before
