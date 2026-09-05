"""Regressions for fix wave 4 -- metamorphic round F2 (MAJOR) + F1
(2026-09-05).

F2: any server edit silently snapped untouched 17-significant-digit
Excel-authored doubles by 1 ulp (openpyxl's %.16g float writer); the save
spine now restores the original numeric text of every cell the operation
did not write. F1: a 17-digit WRITE refuses (verify-after-write, correct)
but the message named neither the cells nor the precision cause; it now
names both and the two remedies.
"""

from __future__ import annotations

import shutil
import struct
import time
import zipfile
from pathlib import Path

import openpyxl
import pytest

from xlsx_mcp.core import calc as core_calc
from xlsx_mcp.core import hazard as core_hazard
from xlsx_mcp.core.errors import (
    HazardRefused,
    UnsupportedStructure,
    ValidationFailed,
    WorkbookProtected,
)
from xlsx_mcp.ops import cells as cells_ops
from xlsx_mcp.ops import lifecycle
from xlsx_mcp.ops import structure as structure_ops

FULL = "0.30000000000000004"          # 0.1 + 0.2 as Excel stores it


def _patch_part(path: Path, part: str, transform) -> None:
    """Rewrite one zip member in place (transform: bytes -> bytes)."""
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(path) as zin,             zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == part:
                data = transform(data)
            zout.writestr(item, data)
    tmp.replace(path)


def _sheet_xml(path: Path) -> str:
    with zipfile.ZipFile(path) as z:
        return z.read("xl/worksheets/sheet1.xml").decode("utf-8")


def _make_fullprec_book(path: Path) -> Path:
    """A1 stored as an Excel-authored 17-significant-digit double,
    A2 = the equality gate the report flagged, C9 an unrelated cell."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A1"] = 0.3
    ws["A2"] = '=IF(A1=0.3,"EQUAL","DIFFERENT")'
    ws["C9"] = "x"
    wb.save(path)
    _patch_part(path, "xl/worksheets/sheet1.xml",
                lambda b: b.replace(b"<v>0.3</v>",
                                    b"<v>" + FULL.encode() + b"</v>", 1))
    assert FULL in _sheet_xml(path)
    return path


# ======================================================================
# Metamorphic F2 (MAJOR): untouched 17-digit doubles survive server edits
# ======================================================================


class TestUntouchedNumberPreservation:
    def test_unrelated_set_cell_preserves_full_precision(self, tmp_path):
        # The report's minimized repro: a write to C9 must not rewrite A1.
        book = _make_fullprec_book(tmp_path / "prec.xlsx")
        r = cells_ops.set_cell(str(book), {"cell": "C9"}, 1)
        assert r["ok"] and r["verified"]
        xml = _sheet_xml(book)
        assert FULL in xml, "the untouched double was snapped by 1 ulp"
        assert r["changed"]["full_precision_cells_preserved"] == 1
        # the model reads the exact original double back, so the =IF
        # equality cannot flip on the next recalc
        wb = openpyxl.load_workbook(book)
        assert wb["S"]["A1"].value == float(FULL)
        assert wb["S"]["A1"].value != 0.3
        wb.close()

    def test_second_edit_still_exact_not_cumulative(self, tmp_path):
        book = _make_fullprec_book(tmp_path / "prec2.xlsx")
        cells_ops.set_cell(str(book), {"cell": "C9"}, 1)
        cells_ops.set_cell(str(book), {"cell": "D9"}, 2)
        assert FULL in _sheet_xml(book)

    def test_written_cell_is_the_operations_to_reserialize(self, tmp_path):
        # Writing A1 itself replaces the value; nothing may "restore" it.
        book = _make_fullprec_book(tmp_path / "prec3.xlsx")
        r = cells_ops.set_cell(str(book), {"cell": "A1"}, 0.25)
        assert r["ok"]
        assert FULL not in _sheet_xml(book)
        wb = openpyxl.load_workbook(book)
        assert wb["S"]["A1"].value == 0.25
        wb.close()

    def test_structural_edit_follows_the_cell(self, tmp_path):
        # insert_rows shifts A1 -> A2; the preserved text must follow it.
        book = _make_fullprec_book(tmp_path / "prec4.xlsx")
        r = structure_ops.modify_grid_structure(str(book), "insert_rows", 1)
        assert r["ok"]
        xml = _sheet_xml(book)
        assert f'r="A2"' in xml and FULL in xml
        wb = openpyxl.load_workbook(book)
        assert wb["S"]["A2"].value == float(FULL)
        wb.close()

    def test_copy_onto_at_risk_cell_is_not_restored(self, tmp_path):
        # A copy destination is a WRITTEN region: the pasted value stands,
        # never the original 17-digit double that used to live there.
        book = _make_fullprec_book(tmp_path / "prec5.xlsx")
        wb = openpyxl.load_workbook(book)
        wb["S"]["B1"] = 7
        wb.save(book)
        wb.close()
        # (the plain openpyxl save above snapped A1; re-patch it)
        _patch_part(book, "xl/worksheets/sheet1.xml",
                    lambda b: b.replace(b"<v>0.3</v>",
                                        b"<v>" + FULL.encode() + b"</v>", 1))
        r = cells_ops.copy_range(str(book), {"cell": "B1"}, {"cell": "A1"},
                                 what="values")
        assert r["ok"]
        wb = openpyxl.load_workbook(book)
        assert wb["S"]["A1"].value == 7
        wb.close()
        assert FULL not in _sheet_xml(book)


# ======================================================================
# Metamorphic F1 (MODERATE): the 17-digit write refusal names cells + cause
# ======================================================================


class TestFullPrecisionWriteRefusal:
    def test_refusal_names_cell_and_precision_cause(self, tmp_path):
        book = _make_fullprec_book(tmp_path / "f1.xlsx")
        before = book.read_bytes()
        with pytest.raises(ValidationFailed) as ei:
            cells_ops.set_cell(str(book), {"cell": "B5"},
                               -1865.0515670299737)
        msg = str(ei.value)
        assert "B5" in msg
        assert "17 significant digits" in msg
        assert "COM" in msg
        assert "NOT modified" in msg
        assert book.read_bytes() == before

    def test_precision_cause_marked_in_mismatch_detail(self):
        from xlsx_mcp.core import verify as core_verify
        reason = core_verify._readback_reason([
            {"sheet": "S", "cell": "B5", "expected": -1865.0515670299737,
             "got": -1865.051567029974, "cause": "precision"}])
        assert "S!B5" in reason and "storage limit" in reason


