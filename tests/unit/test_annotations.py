"""Unit tests for ops/annotations.py: manage_comment, manage_hyperlink."""

from __future__ import annotations

import openpyxl
import pytest

from xlsx_mcp.core.errors import TargetNotFound, XlMcpError
from xlsx_mcp.ops import annotations as an


@pytest.fixture()
def book(tmp_path):
    p = tmp_path / "a.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A1"] = "hello"
    ws["A2"] = "world"
    wb.save(p)
    return p


def test_comment_add_edit_delete_list(book):
    r = an.manage_comment(str(book), "add", location={"cell": "A1"},
                          text="note one", author="tester")
    assert r["ok"] and r["verified"]
    wb = openpyxl.load_workbook(book)
    assert wb["S"]["A1"].comment.text == "note one"
    an.manage_comment(str(book), "edit", location={"cell": "A1"},
                      text="note two")
    wb = openpyxl.load_workbook(book)
    assert wb["S"]["A1"].comment.text == "note two"
    listing = an.manage_comment(str(book), "list")
    assert listing["count"] == 1 and listing["comments"][0]["cell"] == "A1"
    an.manage_comment(str(book), "delete", location={"cell": "A1"})
    wb = openpyxl.load_workbook(book)
    assert wb["S"]["A1"].comment is None


def test_comment_add_needs_text(book):
    with pytest.raises(XlMcpError):
        an.manage_comment(str(book), "add", location={"cell": "A1"})


def test_comment_default_author(book):
    an.manage_comment(str(book), "add", location={"cell": "A1"}, text="x")
    wb = openpyxl.load_workbook(book)
    assert wb["S"]["A1"].comment.author == an._DEFAULT_AUTHOR


def test_comment_edit_missing_refuses(book):
    with pytest.raises(TargetNotFound):
        an.manage_comment(str(book), "edit", location={"cell": "B9"}, text="x")


def test_hyperlink_add_list_remove(book):
    r = an.manage_hyperlink(str(book), "add", location={"cell": "A1"},
                            target="https://example.com", tooltip="site")
    assert r["ok"] and r["verified"]
    wb = openpyxl.load_workbook(book)
    assert wb["S"]["A1"].hyperlink.target == "https://example.com"
    listing = an.manage_hyperlink(str(book), "list")
    assert any(h["kind"] == "hyperlink" for h in listing["hyperlinks"])
    an.manage_hyperlink(str(book), "remove", location={"cell": "A1"})
    wb = openpyxl.load_workbook(book)
    assert wb["S"]["A1"].hyperlink is None


def test_hyperlink_internal_target(book):
    an.manage_hyperlink(str(book), "add", location={"cell": "A2"},
                        target="S!B10", display="jump")
    wb = openpyxl.load_workbook(book)
    hl = wb["S"]["A2"].hyperlink
    assert hl.location == "S!B10"
    assert wb["S"]["A2"].value == "jump"


def test_hyperlink_list_detects_formula(book):
    wb = openpyxl.load_workbook(book)
    wb["S"]["A2"] = '=HYPERLINK("https://x.com","x")'
    wb.save(book)
    listing = an.manage_hyperlink(str(book), "list")
    assert any(h["kind"] == "formula" for h in listing["hyperlinks"])
