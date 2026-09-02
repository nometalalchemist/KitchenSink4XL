"""Unit tests for core/refs.py: the reference-rewriting core.

Formula transposition on insert/delete/move, absolute anchors, cross-sheet
refs, #REF! on deleted targets, string-literal safety, and the workbook
rewriter over formulas / names / conditional formats / data validation /
tables / merges. Hand-computed expected rewrites; correctness is load-bearing.
"""

from __future__ import annotations

import openpyxl
import pytest
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table
from openpyxl.workbook.defined_name import DefinedName

from xlsx_mcp.core import refs

R = refs.RefEdit


def tf(formula, edit, sheet="S"):
    return refs.transpose_formula(formula, edit, sheet)


# ------------------------------------------------------------- insert rows


def test_insert_rows_shifts_refs_at_or_below():
    e = R("S", "insert_rows", index=3, count=2)
    assert tf("=SUM(A1:A10)+B5", e) == "=SUM(A1:A12)+B7"
    assert tf("=A2", e) == "=A2"        # above the insert: unchanged
    assert tf("=A3", e) == "=A5"        # at the insert: shifts


def test_insert_rows_preserves_absolute_anchors():
    e = R("S", "insert_rows", index=1, count=1)
    assert tf("=$A$5+$B5+A$5", e) == "=$A$6+$B6+A$6"


def test_insert_inside_range_expands():
    e = R("S", "insert_rows", index=7, count=2)
    assert tf("=SUM(A5:A10)", e) == "=SUM(A5:A12)"


# ------------------------------------------------------------- delete rows


def test_delete_cell_in_band_becomes_ref_error():
    e = R("S", "delete_rows", index=2, count=2)
    assert tf("=A2", e) == "=#REF!"
    assert tf("=A3", e) == "=#REF!"


def test_delete_cell_below_shifts_up():
    e = R("S", "delete_rows", index=2, count=2)
    assert tf("=A5", e) == "=A3"


def test_delete_range_partial_clamps():
    e = R("S", "delete_rows", index=2, count=2)
    assert tf("=SUM(A1:A6)", e) == "=SUM(A1:A4)"


def test_delete_range_whole_is_ref_error():
    e = R("S", "delete_rows", index=2, count=2)
    assert tf("=SUM(A2:A3)", e) == "=SUM(#REF!)"


# ------------------------------------------------------------- columns


def test_insert_and_delete_columns():
    assert tf("=A1+C1", R("S", "insert_cols", index=2, count=1)) == "=A1+D1"
    assert tf("=B1", R("S", "delete_cols", index=2, count=1)) == "=#REF!"
    assert tf("=SUM(A1:D1)", R("S", "delete_cols", index=2, count=1)) \
        == "=SUM(A1:C1)"


# ------------------------------------------------------------- scope


def test_out_of_scope_sheet_untouched():
    e = R("S", "insert_rows", index=1, count=5)
    # formula lives on 'Other', unqualified refs are Other's, not S's
    assert refs.transpose_formula("=A1+B2", e, "Other") == "=A1+B2"


def test_cross_sheet_qualified_ref_shifts_only_target_sheet():
    e = R("Data", "delete_rows", index=2, count=2)
    got = refs.transpose_formula("=Data!A5+Summary!A5", e, "Summary")
    assert got == "=Data!A3+Summary!A5"


def test_quoted_sheet_name():
    e = R("My Sheet", "insert_rows", index=1, count=1)
    got = refs.transpose_formula("='My Sheet'!A5", e, "Other")
    assert got == "='My Sheet'!A6"


# ------------------------------------------------------------- safety


def test_string_literals_untouched():
    e = R("S", "delete_rows", index=2, count=2)
    assert tf('=A5&"keep A5 literal"', e) == '=A3&"keep A5 literal"'


def test_function_names_not_treated_as_refs():
    e = R("S", "delete_rows", index=2, count=2)
    assert tf("=LOG10(A5)", e) == "=LOG10(A3)"


def test_defined_name_not_matched():
    # a defined name that is not itself a valid cell reference must be left
    # untouched while a real ref beside it shifts
    e = R("S", "insert_rows", index=1, count=1)
    assert tf("=TaxRate*A5", e) == "=TaxRate*A6"
    assert tf("=Tax_2*A5", e) == "=Tax_2*A6"


# ------------------------------------------------------------- move


def test_move_translates_refs_inside_source():
    mv = R("S", "move", src=(1, 1, 2, 2), dst=(10, 10))
    assert tf("=A1", mv) == "=J10"
    assert tf("=B2", mv) == "=K11"
    assert tf("=Z9", mv) == "=Z9"        # outside src, untouched


# ------------------------------------------------------------- bare refs


def test_transpose_ref_range_and_whole_delete():
    assert refs.transpose_ref("A1:C10", R("S", "insert_rows", index=1,
                                          count=2), "S") == "A3:C12"
    assert refs.transpose_ref("A2:A3", R("S", "delete_rows", index=2,
                                         count=2), "S") is None


# ------------------------------------------------------------- workbook rewriter


def _rich_wb():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    for r in range(1, 6):
        ws.cell(r, 1, "h" + str(r))
        ws.cell(r, 2, r * 10)
    ws["D1"] = "=SUM(B1:B5)"
    ws.conditional_formatting.add(
        "B1:B5", CellIsRule(operator="greaterThan", formula=["0"],
                            fill=PatternFill(start_color="FF0000")))
    dv = DataValidation(type="whole", formula1="0")
    dv.add("B2:B5")
    ws.add_data_validation(dv)
    ws.merge_cells("F1:G2")
    wb.defined_names["Grand"] = DefinedName("Grand", attr_text="S!$B$5")
    return wb, ws


def test_rewrite_workbook_insert_rows_touches_every_kind():
    wb, ws = _rich_wb()
    edit = R("S", "insert_rows", index=1, count=1)
    rep = refs.rewrite_workbook(wb, edit)
    # rewrite_workbook repairs REFERENCES only (openpyxl moves the cells); the
    # formula stays in D1 with its refs shifted
    assert ws["D1"].value == "=SUM(B2:B6)"
    cf = list(ws.conditional_formatting)[0]
    assert "B2:B6" in str(cf.sqref)                 # CF range shifted
    dv = ws.data_validations.dataValidation[0]
    assert "B3:B6" in str(dv.sqref)                 # DV range shifted
    assert str(list(ws.merged_cells.ranges)[0]) == "F2:G3"  # merge shifted
    assert wb.defined_names["Grand"].value == "S!$B$6"      # name shifted
    assert rep.names == 1 and rep.conditional_formats == 1
    assert rep.data_validations == 1 and rep.merges == 1


def test_rewrite_workbook_delete_dropping_merge_and_reffing_error():
    wb, ws = _rich_wb()
    edit = R("S", "delete_cols", index=6, count=2)   # delete F,G (the merge)
    rep = refs.rewrite_workbook(wb, edit)
    assert rep.dropped_merges == 1
    assert not ws.merged_cells.ranges


def test_table_ref_rewrites_on_insert():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    for r in range(1, 4):
        ws.cell(r, 1, "c" + str(r))
        ws.cell(r, 2, r)
    ws.add_table(Table(displayName="T", ref="A1:B3"))
    rep = refs.rewrite_workbook(wb, R("S", "insert_rows", index=1, count=2))
    assert ws.tables["T"].ref == "A3:B5"
    assert rep.tables == 1


def test_bad_edit_kind_refuses():
    with pytest.raises(ValueError):
        R("S", "frobnicate", index=1, count=1)
    with pytest.raises(ValueError):
        R("S", "insert_rows", index=0, count=1)
