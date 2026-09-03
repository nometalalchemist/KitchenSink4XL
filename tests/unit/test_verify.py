"""Unit tests for core/verify.py: the verify-after-write gate primitives."""

from __future__ import annotations

import zipfile
from pathlib import Path

import openpyxl

from xlsx_mcp.core import verify


def _save(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A1"] = "hello"
    ws["B2"] = 5
    wb.save(path)
    return path


def test_structural_check_passes_on_valid_file(tmp_path):
    p = _save(tmp_path / "ok.xlsx")
    ok, reasons = verify.structural_check(str(p))
    assert ok and not reasons


def test_structural_check_fails_on_garbage(tmp_path):
    p = tmp_path / "bad.xlsx"
    p.write_bytes(b"not a zip file at all")
    ok, reasons = verify.structural_check(str(p))
    assert not ok and reasons


def test_structural_check_fails_when_sheet_dropped(tmp_path):
    p = _save(tmp_path / "s.xlsx")
    with zipfile.ZipFile(p) as zin:
        infos = zin.infolist()
        datas = {i.filename: zin.read(i.filename) for i in infos}
    with zipfile.ZipFile(p, "w") as zout:
        for i in infos:
            if i.filename == "xl/worksheets/sheet1.xml":
                continue
            zout.writestr(i, datas[i.filename])
    ok, _ = verify.structural_check(str(p))
    assert not ok


def test_part_loss_flags_fragile_drop(tmp_path):
    pre = ["[Content_Types].xml", "xl/workbook.xml",
           "xl/slicers/slicer1.xml", "xl/worksheets/sheet1.xml"]
    p = _save(tmp_path / "p.xlsx")           # produced file has no slicer
    ok, lost = verify.part_loss_check(pre, str(p), allow_loss=False)
    assert not ok
    assert any("slicer" in x for x in lost)


def test_part_loss_ignores_expected_droppable(tmp_path):
    pre = ["[Content_Types].xml", "xl/workbook.xml", "xl/calcChain.xml",
           "xl/worksheets/sheet1.xml"]
    p = _save(tmp_path / "p.xlsx")           # calcChain legitimately gone
    ok, lost = verify.part_loss_check(pre, str(p), allow_loss=False)
    assert ok and not lost


def test_part_loss_allow_loss_permits_drop(tmp_path):
    pre = ["xl/slicers/slicer1.xml", "xl/workbook.xml"]
    p = _save(tmp_path / "p.xlsx")
    ok, lost = verify.part_loss_check(pre, str(p), allow_loss=True)
    assert ok                                # allowed, but still reported
    assert any("slicer" in x for x in lost)


def test_content_readback_detects_mismatch(tmp_path):
    p = _save(tmp_path / "c.xlsx")
    ok, mism = verify.content_readback(
        str(p), {("S", "A1"): ("value", "hello")})
    assert ok and not mism
    ok, mism = verify.content_readback(
        str(p), {("S", "A1"): ("value", "WRONG")})
    assert not ok and mism


def test_content_readback_formula_compare(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A3"] = "=SUM(A1:A2)"
    p = tmp_path / "f.xlsx"
    wb.save(p)
    ok, _ = verify.content_readback(str(p), {("S", "A3"): ("formula", "=SUM(A1:A2)")})
    assert ok
    # leading '=' tolerance both ways
    ok, _ = verify.content_readback(str(p), {("S", "A3"): ("formula", "SUM(A1:A2)")})
    assert ok


def _inventory(path):
    with zipfile.ZipFile(path) as zf:
        return {i.filename: i.file_size for i in zf.infolist()}


def test_inventory_flags_uncataloged_loss(tmp_path):
    """The flipped default: a part TYPE the hazard table never cataloged
    vanishes, and the check fails anyway."""
    p = _save(tmp_path / "p.xlsx")
    pre = _inventory(p)
    pre["xl/futureFeature/part1.xml"] = 123     # present before, gone after
    ok, problems = verify.part_inventory_check(pre, str(p))
    assert not ok
    assert any("xl/futureFeature/part1.xml" in x for x in problems)


def test_inventory_pass_list_losses_are_excused(tmp_path):
    p = _save(tmp_path / "p.xlsx")
    pre = _inventory(p)
    pre["xl/calcChain.xml"] = 50
    pre["xl/sharedStrings.xml"] = 200
    pre["xl/printerSettings/printerSettings1.bin"] = 300
    pre["docProps/thumbnail.jpeg"] = 400
    ok, problems = verify.part_inventory_check(pre, str(p))
    assert ok and not problems


def test_inventory_rels_parts_are_excused(tmp_path):
    """Relationship parts are regenerated derived metadata; their targets are
    independently inventoried, so a vanished rels part is not a loss."""
    p = _save(tmp_path / "p.xlsx")
    pre = _inventory(p)
    pre["xl/worksheets/_rels/sheet1.xml.rels"] = 90
    ok, problems = verify.part_inventory_check(pre, str(p))
    assert ok and not problems


def test_inventory_expected_removal_is_scoped(tmp_path):
    p = _save(tmp_path / "p.xlsx")
    pre = _inventory(p)
    pre["xl/tables/table1.xml"] = 500
    ok, _ = verify.part_inventory_check(pre, str(p))
    assert not ok
    ok, problems = verify.part_inventory_check(
        pre, str(p), expected_removals=("xl/tables/",))
    assert ok and not problems


def test_inventory_worksheet_renumber_is_not_a_loss(tmp_path):
    """openpyxl reassigns sheet part names sequentially (observed sheet5 ->
    sheet1); a same-count rename must not fail."""
    p = _save(tmp_path / "p.xlsx")           # produced file has sheet1.xml
    pre = _inventory(p)
    size = pre.pop("xl/worksheets/sheet1.xml")
    pre["xl/worksheets/sheet5.xml"] = size   # pre-scan knew a different name
    ok, problems = verify.part_inventory_check(pre, str(p))
    assert ok and not problems


def test_inventory_worksheet_deficit_fails_unless_registered(tmp_path):
    p = _save(tmp_path / "p.xlsx")           # one worksheet part
    pre = _inventory(p)
    pre["xl/worksheets/sheet2.xml"] = 800    # a second sheet existed before
    ok, problems = verify.part_inventory_check(pre, str(p))
    assert not ok
    assert any("worksheet part(s) missing" in x for x in problems)
    ok, _ = verify.part_inventory_check(
        pre, str(p), expected_removals=("xl/worksheets/",))
    assert ok


def test_inventory_nonfragile_part_emptied_fails(tmp_path):
    p = _save(tmp_path / "p.xlsx")
    with zipfile.ZipFile(p, "a") as zf:
        zf.writestr("xl/futureFeature/part1.xml", b"")
    pre = _inventory(p)
    pre["xl/futureFeature/part1.xml"] = 123  # had content before the save
    ok, problems = verify.part_inventory_check(pre, str(p))
    assert not ok
    assert any("replaced with empty content" in x for x in problems)


def test_inventory_fragile_parts_are_other_checks_jurisdiction(tmp_path):
    """A lost fragile (hazard-table) part is part_loss_check's job, where
    allow_loss governs; the inventory check must not double-report it."""
    p = _save(tmp_path / "p.xlsx")
    pre = _inventory(p)
    pre["xl/slicers/slicer1.xml"] = 700
    ok, problems = verify.part_inventory_check(pre, str(p))
    assert ok and not problems


def test_verify_after_write_combines_gates(tmp_path):
    p = _save(tmp_path / "v.xlsx")
    pre = list(zipfile.ZipFile(p).namelist())
    r = verify.verify_after_write(
        str(p), pre_parts=pre, intended={("S", "A1"): ("value", "hello")})
    assert r.ok
    r2 = verify.verify_after_write(
        str(p), pre_parts=pre, intended={("S", "A1"): ("value", "nope")})
    assert not r2.ok and r2.mismatches
