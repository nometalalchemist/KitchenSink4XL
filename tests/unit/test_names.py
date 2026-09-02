"""Unit tests for ops/names.py: manage_name."""

from __future__ import annotations

import openpyxl
import pytest

from xlsx_mcp.core.errors import AmbiguousTarget, TargetNotFound, XlMcpError
from xlsx_mcp.ops import names as nm


@pytest.fixture()
def book(tmp_path):
    p = tmp_path / "n.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    wb.create_sheet("T")
    ws["A1"] = 1
    wb.save(p)
    return p


def test_add_list_update_delete(book):
    r = nm.manage_name(str(book), "add", name="Total", refers_to="S!$A$1")
    assert r["ok"] and r["verified"]
    listing = nm.manage_name(str(book), "list")
    assert any(n["name"] == "Total" and n["scope"] == "workbook"
               for n in listing["names"])
    nm.manage_name(str(book), "update", name="Total", refers_to="S!$A$1:$A$5")
    wb = openpyxl.load_workbook(book)
    assert wb.defined_names["Total"].value == "S!$A$1:$A$5"
    nm.manage_name(str(book), "delete", name="Total")
    wb = openpyxl.load_workbook(book)
    assert "Total" not in wb.defined_names


def test_sheet_scope_and_rename(book):
    nm.manage_name(str(book), "add", name="Local", refers_to="T!$A$1",
                   scope="T")
    wb = openpyxl.load_workbook(book)
    assert "Local" in wb["T"].defined_names
    nm.manage_name(str(book), "rename", name="Local", new_name="Renamed",
                   scope="T")
    wb = openpyxl.load_workbook(book)
    assert "Renamed" in wb["T"].defined_names
    assert "Local" not in wb["T"].defined_names


def test_scope_collision_is_ambiguous(book):
    nm.manage_name(str(book), "add", name="Dup", refers_to="S!$A$1")
    nm.manage_name(str(book), "add", name="Dup", refers_to="T!$A$1", scope="T")
    with pytest.raises(AmbiguousTarget):
        nm.manage_name(str(book), "delete", name="Dup")
    # disambiguated by scope, deletes fine
    nm.manage_name(str(book), "delete", name="Dup", scope="T")


def test_reserved_name_protected(book):
    from openpyxl.workbook.defined_name import DefinedName
    wb = openpyxl.load_workbook(book)
    wb.defined_names.add(DefinedName(name="_xlnm.Print_Area",
                                     attr_text="S!$A$1:$B$2"))
    wb.save(book)
    with pytest.raises(XlMcpError):
        nm.manage_name(str(book), "delete", name="_xlnm.Print_Area")


def test_add_duplicate_and_missing(book):
    nm.manage_name(str(book), "add", name="X", refers_to="S!$A$1")
    with pytest.raises(XlMcpError):
        nm.manage_name(str(book), "add", name="X", refers_to="S!$A$2")
    with pytest.raises(TargetNotFound):
        nm.manage_name(str(book), "delete", name="Ghost")
