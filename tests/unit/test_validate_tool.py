"""validate tests: the default battery, each check's pass and findings
paths (structure, references, names, merges, tables, formatting_bloat,
hazards, external_links, calc_staleness), the verbatim underlying-op
shapes, and the refusal paths.
"""

from __future__ import annotations

import shutil
import tempfile
import zipfile

import openpyxl
import pytest
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.cell_range import CellRange

from xlsx_mcp.core.errors import WorkbookNotFound, XlMcpError
from xlsx_mcp.ops import formulas as _formulas
from xlsx_mcp.ops import lifecycle as _lifecycle
from xlsx_mcp.ops import validation as _validation


def _make(tmp_path, name="book.xlsx", sheets=None):
    p = str(tmp_path / name)
    _lifecycle.create_workbook(p, sheets=sheets or ["S"])
    return p


def _patch_cached_value(path: str, formula_snippet: str, cached: str) -> None:
    with zipfile.ZipFile(path) as zf:
        payloads = {n: zf.read(n) for n in zf.namelist()}
    needle = f"<f>{formula_snippet}</f>"
    hit = False
    for n, data in payloads.items():
        if n.startswith("xl/worksheets/") and needle in data.decode(
                "utf-8", "replace"):
            payloads[n] = data.decode("utf-8").replace(
                needle, needle + f"<v>{cached}</v>").encode("utf-8")
            hit = True
    assert hit
    tmp = f"{tempfile.mkdtemp()}/patched.xlsx"
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for n, data in payloads.items():
            zout.writestr(n, data)
    shutil.move(tmp, path)


# ------------------------------------------------------------ the battery


def test_default_battery_on_clean_workbook(tmp_path):
    p = _make(tmp_path)
    out = _validation.validate(p)
    assert out["passed"] is True
    assert out["checks_run"] == ["structure", "references", "calc_staleness"]
    for check in out["checks_run"]:
        assert out["results"][check]["passed"] is True


def test_unknown_and_empty_checks_refuse(tmp_path):
    p = _make(tmp_path)
    with pytest.raises(XlMcpError):
        _validation.validate(p, checks=["spelling"])
    with pytest.raises(XlMcpError):
        _validation.validate(p, checks=[])
    with pytest.raises(WorkbookNotFound):
        _validation.validate(str(tmp_path / "no.xlsx"))


def test_duplicate_checks_run_once(tmp_path):
    p = _make(tmp_path)
    out = _validation.validate(p, checks=["structure", "structure"])
    assert out["checks_run"] == ["structure"]


# ------------------------------------------------------------- per check


def test_structure_clean_and_corrupt(tmp_path):
    p = _make(tmp_path)
    out = _validation.validate(p, checks=["structure"])
    f = out["results"]["structure"]["findings"]
    # opens_clean is Excel's verdict, and nothing asked Excel here: the field
    # says so rather than claiming a pass openpyxl cannot vouch for (insane
    # round M-1 -- ten workbooks Excel refused were reported opens_clean:true
    # because the field was hard-coded).
    assert f["opens_clean"] == _validation.NOT_CHECKED
    assert f["openpyxl_loads"] is True and f["visible_sheets"] == 1
    assert "Excel was NOT asked" in f["opens_clean_source"]
    bad = tmp_path / "bad.xlsx"
    bad.write_bytes(b"this is no zip")
    out = _validation.validate(str(bad), checks=["structure"])
    assert out["passed"] is False
    assert out["results"]["structure"]["findings"]["opens_clean"] is False


def test_references_check_finds_ref_errors(tmp_path):
    p = _make(tmp_path)
    _formulas.set_formula(p, {"cell": "A1"}, "=#REF!+1")
    out = _validation.validate(p, checks=["references"])
    r = out["results"]["references"]
    assert r["passed"] is False
    assert r["findings"]["count"] == 1
    assert r["findings"]["items"][0]["error"] == "#REF!"


def test_calc_staleness_flags_uncached_then_clears(tmp_path):
    p = _make(tmp_path)
    _formulas.set_formula(p, {"cell": "A1"}, "=SUM(1,2)")
    out = _validation.validate(p, checks=["calc_staleness"])
    c = out["results"]["calc_staleness"]
    assert c["passed"] is False and c["findings"]["count"] == 1
    assert "blank" in c["findings"]["note"]
    _patch_cached_value(p, "SUM(1,2)", "3")
    out = _validation.validate(p, checks=["calc_staleness"])
    assert out["results"]["calc_staleness"]["passed"] is True


def test_names_check_finds_broken_defined_names(tmp_path):
    p = _make(tmp_path)
    wb = openpyxl.load_workbook(p)
    wb.defined_names["Good"] = DefinedName("Good", attr_text="S!$A$1")
    wb.defined_names["BadRef"] = DefinedName("BadRef", attr_text="S!#REF!")
    wb.defined_names["Ghosted"] = DefinedName(
        "Ghosted", attr_text="Ghost!$A$1")
    wb.save(p)
    out = _validation.validate(p, checks=["names"])
    f = out["results"]["names"]["findings"]
    assert out["results"]["names"]["passed"] is False
    assert f["defined_names"] == 3 and f["count"] == 2
    broken = {b["name"]: b["reasons"] for b in f["broken"]}
    assert "contains #REF!" in broken["BadRef"]
    assert any("Ghost" in r for r in broken["Ghosted"])


def test_merges_check_overlaps_and_orphans(tmp_path):
    p = _make(tmp_path)
    wb = openpyxl.load_workbook(p)
    ws = wb["S"]
    ws["A1"] = "x"
    ws.merge_cells("A1:B2")
    ws.merged_cells.ranges.add(CellRange("B2:C3"))  # overlaps A1:B2
    ws.merged_cells.ranges.add(CellRange("H10:I11"))  # far outside used
    wb.save(p)
    out = _validation.validate(p, checks=["merges"])
    f = out["results"]["merges"]["findings"]
    assert out["results"]["merges"]["passed"] is False
    assert f["merged_ranges"] == 3
    assert len(f["overlaps"]) >= 1
    assert any(o["range"] == "H10:I11" for o in f["orphans"])


def test_merges_check_clean(tmp_path):
    p = _make(tmp_path)
    wb = openpyxl.load_workbook(p)
    ws = wb["S"]
    ws["A1"] = "header"
    ws.merge_cells("A1:B1")
    wb.save(p)
    out = _validation.validate(p, checks=["merges"])
    assert out["results"]["merges"]["passed"] is True


def test_tables_check_clean_then_forged_overlap(tmp_path):
    from xlsx_mcp.ops import cells as _cells
    from xlsx_mcp.ops import tables as _tables
    p = _make(tmp_path)
    _cells.write_range(p, {"cell": "A1", "sheet": "S"},
                       [["h1", "h2", None, "h3", "h4"],
                        [1, 2, None, 3, 4]])
    _tables.create_table(p, {"range": "A1:B2", "sheet": "S"}, "Tbl")
    _tables.create_table(p, {"range": "D1:E2", "sheet": "S"}, "TblTwo")
    out = _validation.validate(p, checks=["tables"])
    assert out["results"]["tables"]["passed"] is True
    assert out["results"]["tables"]["findings"]["tables"] == 2
    # forge an overlap by zip surgery on the second table part (our own
    # create_table refuses overlaps, but files from other writers can
    # carry them)
    with zipfile.ZipFile(p) as zf:
        payloads = {n: zf.read(n) for n in zf.namelist()}
    hit = False
    for n, data in payloads.items():
        if n.startswith("xl/tables/") and b"D1:E2" in data:
            payloads[n] = data.replace(b"D1:E2", b"A1:B2")
            hit = True
    assert hit
    tmp = f"{tempfile.mkdtemp()}/forged.xlsx"
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for n, data in payloads.items():
            zout.writestr(n, data)
    shutil.move(tmp, p)
    out = _validation.validate(p, checks=["tables"])
    assert out["results"]["tables"]["passed"] is False
    probs = out["results"]["tables"]["findings"]["problems"]
    assert any("overlaps" in pr["problem"] for pr in probs)


def test_formatting_bloat_check_uses_audit_styles_shape(tmp_path):
    p = _make(tmp_path)
    out = _validation.validate(p, checks=["formatting_bloat"])
    f = out["results"]["formatting_bloat"]["findings"]
    assert out["results"]["formatting_bloat"]["passed"] is True
    # verbatim audit_styles keys
    assert {"cell_formats", "ceiling", "risk",
            "heaviest_formats"} <= set(f)


def test_hazards_check_flags_drop_risk_part(tmp_path):
    p = _make(tmp_path)
    out = _validation.validate(p, checks=["hazards"])
    assert out["results"]["hazards"]["passed"] is True
    # inject a slicer part (SEV_DROPS) with a content-type override
    src, dst = tmp_path / "book.xlsx", tmp_path / "rich.xlsx"
    with zipfile.ZipFile(src) as zin, \
            zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            payload = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                payload = payload.replace(
                    b"</Types>",
                    b'<Override PartName="/xl/slicers/slicer1.xml" '
                    b'ContentType="application/vnd.ms-excel.slicer+xml"/>'
                    b"</Types>")
            zout.writestr(item, payload)
        zout.writestr("xl/slicers/slicer1.xml", b"<slicers>real</slicers>")
    out = _validation.validate(str(dst), checks=["hazards"])
    assert out["results"]["hazards"]["passed"] is False
    f = out["results"]["hazards"]["findings"]  # verbatim scan shape
    assert f["would_lose"] is True
    assert any(h["key"] == "slicers" for h in f["hazards"])


def test_external_links_check_clean(tmp_path):
    p = _make(tmp_path)
    out = _validation.validate(p, checks=["external_links"])
    r = out["results"]["external_links"]
    assert r["passed"] is True
    assert {"links", "count", "note"} <= set(r["findings"])  # verbatim


def test_full_battery_all_checks(tmp_path):
    p = _make(tmp_path)
    out = _validation.validate(p, checks=list(_validation.CHECK_NAMES))
    assert set(out["results"]) == set(_validation.CHECK_NAMES)
    assert out["passed"] is True
