"""Unit tests for ops/cells.query_range: the token-shaped read (headline).
Filter, project, sort, paginate, aggregate, group. Read-only."""

from __future__ import annotations

import openpyxl
import pytest

from xlsx_mcp.core.errors import XlMcpError
from xlsx_mcp.ops import cells


@pytest.fixture()
def sales(tmp_path):
    p = tmp_path / "sales.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    rows = [
        ["region", "product", "qty", "price"],
        ["East", "apple", 10, 1.5],
        ["West", "apple", 5, 1.5],
        ["East", "pear", 3, 2.0],
        ["West", "pear", 8, 2.0],
        ["East", "fig", 1, 5.0],
    ]
    for r, row in enumerate(rows, 1):
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v)
    wb.save(p)
    return p


def test_filter_and_project(sales):
    r = cells.query_range(str(sales), {"used_range": "Sales"},
                          where=[{"column": "region", "op": "eq",
                                  "value": "East"}],
                          columns=["product", "qty"])
    assert r["matched"] == 3
    assert r["columns"] == ["product", "qty"]
    assert ["apple", 10] in r["rows"]
    assert all(len(row) == 2 for row in r["rows"])


def test_multi_predicate_all_vs_any(sales):
    r = cells.query_range(str(sales), {"used_range": "Sales"},
                          where=[{"column": "region", "op": "eq",
                                  "value": "East"},
                                 {"column": "qty", "op": "gt", "value": 2}],
                          match="all")
    assert r["matched"] == 2   # East apple(10), East pear(3)
    r2 = cells.query_range(str(sales), {"used_range": "Sales"},
                           where=[{"column": "product", "op": "eq",
                                   "value": "fig"},
                                  {"column": "qty", "op": "gt", "value": 7}],
                           match="any")
    assert r2["matched"] == 3   # fig(1) + the two qty>7


def test_string_ops(sales):
    r = cells.query_range(str(sales), {"used_range": "Sales"},
                          where=[{"column": "product", "op": "startswith",
                                  "value": "ap"}])
    assert r["matched"] == 2


def test_sort_and_limit_paginate(sales):
    r = cells.query_range(str(sales), {"used_range": "Sales"},
                          columns=["qty"],
                          order_by=[{"column": "qty", "dir": "desc"}],
                          limit=2)
    assert r["rows"] == [[10], [8]]
    assert r["truncated"] is True
    r2 = cells.query_range(str(sales), {"used_range": "Sales"},
                           columns=["qty"],
                           order_by=[{"column": "qty", "dir": "asc"}],
                           offset=1, limit=2)
    assert r2["rows"] == [[3], [5]]


def test_distinct(sales):
    r = cells.query_range(str(sales), {"used_range": "Sales"},
                          columns=["region"], distinct=True)
    flat = sorted(x[0] for x in r["rows"])
    assert flat == ["East", "West"]


def test_records_shape(sales):
    r = cells.query_range(str(sales), {"used_range": "Sales"},
                          columns=["product", "qty"],
                          where=[{"column": "product", "op": "eq",
                                  "value": "fig"}],
                          records=True)
    assert r["rows"] == [{"product": "fig", "qty": 1}]


def test_aggregate_no_group(sales):
    r = cells.query_range(str(sales), {"used_range": "Sales"},
                          aggregate=[{"column": "qty", "func": "sum"},
                                     {"column": "qty", "func": "max"},
                                     {"column": "region",
                                      "func": "count_distinct"}])
    assert r["mode"] == "aggregate"
    aggs = r["groups"][0]["aggregates"]
    assert aggs["sum_qty"] == 27
    assert aggs["max_qty"] == 10
    assert aggs["count_distinct_region"] == 2


def test_aggregate_group_by(sales):
    r = cells.query_range(str(sales), {"used_range": "Sales"},
                          aggregate=[{"column": "qty", "func": "sum"}],
                          group_by="region")
    by = {tuple(g["group"].values())[0]: g["aggregates"]["sum_qty"]
          for g in r["groups"]}
    assert by == {"East": 14, "West": 13}


def test_filtered_aggregate(sales):
    r = cells.query_range(str(sales), {"used_range": "Sales"},
                          where=[{"column": "product", "op": "eq",
                                  "value": "pear"}],
                          aggregate=[{"column": "qty", "func": "avg"}])
    assert r["groups"][0]["aggregates"]["avg_qty"] == 5.5


def test_headerless_by_letter(tmp_path):
    p = tmp_path / "h.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A1"], ws["B1"] = 1, 100
    ws["A2"], ws["B2"] = 2, 200
    wb.save(p)
    r = cells.query_range(str(p), {"range": "A1:B2"}, header=False,
                          where=[{"column": "A", "op": "ge", "value": 2}])
    assert r["matched"] == 1
    assert r["rows"] == [[2, 200]]


def test_bad_column_and_op(sales):
    with pytest.raises(XlMcpError):
        cells.query_range(str(sales), {"used_range": "Sales"},
                          where=[{"column": "nope", "op": "eq", "value": 1}])
    with pytest.raises(XlMcpError):
        cells.query_range(str(sales), {"used_range": "Sales"},
                          where=[{"column": "qty", "op": "??", "value": 1}])


def test_empty_sheet_returns_empty(tmp_path):
    p = tmp_path / "e.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "Blank"
    wb.save(p)
    r = cells.query_range(str(p), {"used_range": "Blank"})
    assert r["matched"] == 0 and r["rows"] == []
