"""The style layer (format pack): apply_style (builtin + defined named
styles), copy_format (the format painter), and audit_styles (the read-only
style-bloat audit toward the 64,000-format ceiling).
"""

from __future__ import annotations

import openpyxl
import pytest

from xlsx_mcp.core.errors import TargetNotFound, XlMcpError
from xlsx_mcp.ops import format as _format


def _make(path, rows, sheet="Data"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    for i, row in enumerate(rows, start=1):
        for j, v in enumerate(row, start=1):
            ws.cell(i, j, v)
    wb.save(path)
    wb.close()
    return str(path)


# -------------------------------------------------------------- apply_style


def test_apply_builtin_style(tmp_path):
    p = _make(tmp_path / "b.xlsx", [["ok", "meh"], [1, 2]])
    out = _format.apply_style(p, "Good", location={"range": "A1:B1"})
    assert out["ok"] is True and out["verified"] is True
    assert out["changed"]["style"] == {
        "name": "Good", "defined": False, "range": "A1:B1", "cells": 2}
    wb = openpyxl.load_workbook(p)
    assert "Good" in wb.named_styles
    assert wb["Data"]["A1"].style == "Good"
    wb.close()


def test_define_and_apply_named_style(tmp_path):
    p = _make(tmp_path / "d.xlsx", [[1234.5], [6789.0]])
    out = _format.apply_style(
        p, "Money", location={"range": "A1:A2"},
        define={"number_format": "#,##0.00",
                "font": {"bold": True, "color": "006100"},
                "fill": {"color": "C6EFCE"}})
    assert out["changed"]["style"]["defined"] is True
    wb = openpyxl.load_workbook(p)
    cell = wb["Data"]["A1"]
    assert cell.style == "Money"
    assert cell.number_format == "#,##0.00"
    assert cell.font.bold is True
    assert cell.fill.start_color.rgb == "FFC6EFCE"
    wb.close()
    # a second define under the same name refuses (no in-place redefinition)
    with pytest.raises(XlMcpError, match="already exists"):
        _format.apply_style(p, "Money", location={"cell": "A2"},
                            define={"font": {"italic": True}})
    # but applying the now-registered style by name works
    out = _format.apply_style(p, "Money", location={"cell": "A2"})
    assert out["changed"]["style"]["cells"] == 1


def test_define_only_registers_without_applying(tmp_path):
    p = _make(tmp_path / "reg.xlsx", [[1]])
    out = _format.apply_style(p, "Later",
                              define={"font": {"italic": True}})
    assert out["changed"]["style"] == {
        "name": "Later", "defined": True, "range": None, "cells": 0}
    wb = openpyxl.load_workbook(p)
    assert "Later" in wb.named_styles
    wb.close()


def test_unknown_style_refuses_with_candidates_and_case_hint(tmp_path):
    p = _make(tmp_path / "u.xlsx", [[1]])
    with pytest.raises(TargetNotFound, match="Builtins include"):
        _format.apply_style(p, "NoSuchStyle", location={"cell": "A1"})
    with pytest.raises(TargetNotFound, match="did you mean 'Good'"):
        _format.apply_style(p, "good", location={"cell": "A1"})
    with pytest.raises(XlMcpError, match="location"):
        _format.apply_style(p, "Good")


def test_define_validates_its_shape(tmp_path):
    p = _make(tmp_path / "v.xlsx", [[1]])
    with pytest.raises(XlMcpError, match="unknown key"):
        _format.apply_style(p, "X", define={"fnot": {"bold": True}})
    with pytest.raises(XlMcpError, match="at least one"):
        _format.apply_style(p, "X", define={})
    with pytest.raises(XlMcpError, match="color"):
        _format.apply_style(p, "X", define={"fill": {"pattern": "solid"}})


# -------------------------------------------------------------- copy_format


def test_copy_format_paints_full_format_not_values(tmp_path):
    p = _make(tmp_path / "cf.xlsx", [["src", 1], [2, 3], [4, 5]])
    _format.format_cells(
        p, {"cell": "A1"}, number_format="0.00%",
        font={"bold": True}, fill={"color": "FFFF00"})
    out = _format.copy_format(p, {"cell": "A1"}, {"range": "B1:B3"})
    assert out["changed"]["painted"] == {
        "from": "Data!A1", "to": "Data!B1:B3", "cells": 3}
    wb = openpyxl.load_workbook(p)
    ws = wb["Data"]
    for coord in ("B1", "B2", "B3"):
        assert ws[coord].font.bold is True
        assert ws[coord].number_format == "0.00%"
        assert ws[coord].fill.start_color.rgb == "FFFFFF00"
    # values untouched
    assert ws["B1"].value == 1 and ws["B2"].value == 3
    wb.close()


def test_copy_format_source_must_be_single_cell(tmp_path):
    p = _make(tmp_path / "cfs.xlsx", [[1, 2], [3, 4]])
    with pytest.raises(XlMcpError, match="single cell"):
        _format.copy_format(p, {"range": "A1:A2"}, {"cell": "B1"})


# ------------------------------------------------------------- audit_styles


def test_audit_styles_counts_usage_and_offenders(tmp_path):
    p = _make(tmp_path / "a.xlsx",
              [["plain", "plain2"], ["plain3", "bold"]])
    _format.format_cells(p, {"cell": "B2"}, font={"bold": True},
                         fill={"color": "FF0000"})
    out = _format.audit_styles(p)
    assert out["ceiling"] == 64_000
    assert out["risk"] == "ok"
    assert out["cell_formats"] >= 2
    assert out["distinct_formats_in_use"] == 2
    assert out["per_sheet"]["Data"] == {
        "cells_scanned": 4, "distinct_formats": 2}
    heavy = out["heaviest_formats"]
    assert heavy[0]["cells"] == 3  # the three plain cells share one format
    bold_entry = next(e for e in heavy if e["example"] == "Data!B2")
    assert "bold" in bold_entry["font"]
    assert bold_entry["fill"] == "FFFF0000"
    assert bold_entry["number_format"] == "General"


def test_audit_styles_read_only_and_top_validation(tmp_path):
    p = _make(tmp_path / "ro.xlsx", [[1]])
    before = open(p, "rb").read()
    out = _format.audit_styles(p, top=1)
    assert len(out["heaviest_formats"]) == 1
    assert open(p, "rb").read() == before
    with pytest.raises(XlMcpError, match="top"):
        _format.audit_styles(p, top=0)


def test_named_styles_counted_in_audit(tmp_path):
    p = _make(tmp_path / "ns.xlsx", [[1]])
    _format.apply_style(p, "Money", location={"cell": "A1"},
                        define={"number_format": "#,##0.00"})
    out = _format.audit_styles(p)
    assert out["named_styles"] >= 2  # Normal + Money
