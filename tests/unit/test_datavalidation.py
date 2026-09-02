"""Unit tests for ops/datavalidation.py: manage_data_validation."""

from __future__ import annotations

import openpyxl
import pytest

from xlsx_mcp.core.errors import TargetNotFound, XlMcpError
from xlsx_mcp.ops import datavalidation as dv


@pytest.fixture()
def book(tmp_path):
    p = tmp_path / "dv.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    for r in range(1, 6):
        ws.cell(r, 1, r)
    wb.save(p)
    return p


def _dvs(path):
    wb = openpyxl.load_workbook(path)
    return wb["S"].data_validations.dataValidation


def test_add_list_dropdown_and_list(book):
    r = dv.manage_data_validation(str(book), "add", location={"range": "A1:A5"},
                                  dv_type="list", values=["x", "y", "z"])
    assert r["ok"] and r["verified"]
    got = _dvs(book)
    assert len(got) == 1 and got[0].type == "list"
    assert got[0].formula1 == '"x,y,z"'
    listing = dv.manage_data_validation(str(book), "list")
    assert listing["count"] == 1


def test_add_whole_number_bounds(book):
    dv.manage_data_validation(str(book), "add", location={"range": "A1:A5"},
                              dv_type="whole", operator="between",
                              formula1=1, formula2=100)
    got = _dvs(book)
    assert got[0].type == "whole" and got[0].operator == "between"
    assert got[0].formula1 == "1" and got[0].formula2 == "100"


def test_add_custom(book):
    dv.manage_data_validation(str(book), "add", location={"range": "A1:A5"},
                              dv_type="custom", formula1="=ISNUMBER(A1)")
    assert _dvs(book)[0].type == "custom"


def test_dv_survives_structural_edit(book):
    dv.manage_data_validation(str(book), "add", location={"range": "A2:A5"},
                              dv_type="list", values=["a", "b"])
    from xlsx_mcp.core import refs
    from xlsx_mcp.core.package import WorkbookPackage
    pkg = WorkbookPackage.open(str(book))
    pkg.modify_structure(refs.RefEdit("S", refs.INSERT_ROWS, index=1, count=1))
    pkg.save()
    got = _dvs(book)
    assert str(got[0].sqref) == "A3:A6"


def test_delete(book):
    dv.manage_data_validation(str(book), "add", location={"range": "A1:A5"},
                              dv_type="list", values=["x"])
    dv.manage_data_validation(str(book), "delete", location={"range": "A1:A5"})
    assert len(_dvs(book)) == 0


def test_delete_missing_refuses(book):
    with pytest.raises(TargetNotFound):
        dv.manage_data_validation(str(book), "delete",
                                  location={"range": "Z1:Z5"})


def test_bad_type(book):
    with pytest.raises(XlMcpError):
        dv.manage_data_validation(str(book), "add", location={"range": "A1:A5"},
                                  dv_type="rainbow")
