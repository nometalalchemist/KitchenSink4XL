"""set_merge (Excel merge semantics + the data-loss confirm gate) and the
scatter pair get_cells / set_cells (multi-cell single-call read/write through
the package, validate-first batching).
"""

from __future__ import annotations

import openpyxl
import pytest

from conftest import cell_rows
from xlsx_mcp.core.errors import (
    RangeOutOfBounds,
    TargetNotFound,
    XlMcpError,
)
from xlsx_mcp.ops import cells as _cells


def _make(path, rows, sheet="Data"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    for i, row in enumerate(rows, start=1):
        for j, v in enumerate(row, start=1):
            ws.cell(i, j, v)
    wb.save(path)
    wb.close()
    return str(path)


# ------------------------------------------------------------------- merges


def test_merge_and_list_and_unmerge_roundtrip(tmp_path):
    p = _make(tmp_path / "m.xlsx", [["title", None, None], [1, 2, 3]])
    out = _cells.set_merge(p, "merge", location={"range": "A1:C1"})
    assert out["ok"] is True and out["verified"] is True
    assert out["changed"]["merged"]["absorbed_cleared"] == []
    listed = _cells.set_merge(p, "list")
    assert listed == {"merges": {"Data": ["A1:C1"]}, "count": 1}
    out = _cells.set_merge(p, "unmerge", location={"range": "A1:C1"})
    assert out["changed"]["unmerged"]["range"] == "A1:C1"
    assert _cells.set_merge(p, "list")["count"] == 0
    # the top-left value survives the whole round trip
    wb = openpyxl.load_workbook(p)
    assert wb["Data"]["A1"].value == "title"
    wb.close()


def test_merge_absorbing_values_refuses_without_confirm(tmp_path):
    p = _make(tmp_path / "loss.xlsx", [["keep", "gone1", "gone2"]])
    with pytest.raises(XlMcpError, match="DISCARD") as ei:
        _cells.set_merge(p, "merge", location={"range": "A1:C1"})
    assert getattr(ei.value, "code", None) == "CONFLICT"
    assert "B1" in str(ei.value) and "C1" in str(ei.value)
    # nothing changed
    wb = openpyxl.load_workbook(p)
    assert wb["Data"]["B1"].value == "gone1"
    assert not list(wb["Data"].merged_cells.ranges)
    wb.close()


def test_merge_with_confirm_clears_and_warns(tmp_path):
    p = _make(tmp_path / "confirm.xlsx", [["keep", "gone"]])
    out = _cells.set_merge(p, "merge", location={"range": "A1:B1"},
                           confirm_data_loss=True)
    assert out["changed"]["merged"]["absorbed_cleared"] == ["B1"]
    assert any("discarded" in w for w in out["warnings"])
    wb = openpyxl.load_workbook(p)
    assert wb["Data"]["A1"].value == "keep"
    assert wb["Data"]["B1"].value is None
    assert "A1:B1" in {str(r) for r in wb["Data"].merged_cells.ranges}
    wb.close()


def test_overlapping_merge_refuses(tmp_path):
    p = _make(tmp_path / "ov.xlsx", [[None, None, None]])
    _cells.set_merge(p, "merge", location={"range": "A1:B1"})
    with pytest.raises(XlMcpError, match="overlap"):
        _cells.set_merge(p, "merge", location={"range": "B1:C1"})


def test_merge_single_cell_and_unmerge_miss_refuse(tmp_path):
    p = _make(tmp_path / "one.xlsx", [[None, None], [None, None]])
    with pytest.raises(XlMcpError, match="single cell"):
        _cells.set_merge(p, "merge", location={"cell": "A1"})
    _cells.set_merge(p, "merge", location={"range": "A1:B1"})
    # unmerge needs the EXACT stored range, and the refusal lists the merges
    with pytest.raises(TargetNotFound, match="A1:B1"):
        _cells.set_merge(p, "unmerge", location={"range": "A1:B2"})
    with pytest.raises(XlMcpError, match="action"):
        _cells.set_merge(p, "toggle", location={"range": "A1:B1"})


# ---------------------------------------------------------------- get_cells


def test_get_cells_mixed_addresses_and_labels(tmp_path):
    p = _make(tmp_path / "g.xlsx", [["h", 10], [2, "=A2*2"]])
    out = _cells.get_cells(
        p, ["A1", {"cell": "B1"}, {"cell": "B2"}], values="both")
    assert out["count"] == 3
    rows = cell_rows(out)
    assert rows[0] == {"sheet": "Data", "cell": "A1",
                       "value": "h", "label": "value"}
    assert rows[1]["value"] == 10
    # a never-calculated formula is labelled absent, never passed off as blank
    assert rows[2]["label"] == "absent"
    assert "absent" in out["warning"]
    # formula mode returns the formula string
    out = _cells.get_cells(p, [{"cell": "B2"}], values="formula")
    assert cell_rows(out)[0]["value"] == "=A2*2"


def test_get_cells_refusals(tmp_path):
    p = _make(tmp_path / "gr.xlsx", [[1, 2], [3, 4]])
    with pytest.raises(XlMcpError, match="non-empty"):
        _cells.get_cells(p, [])
    with pytest.raises(XlMcpError, match=r"cells\[1\]"):
        _cells.get_cells(p, ["A1", {"range": "A1:B2"}])
    with pytest.raises(RangeOutOfBounds, match="scatter ceiling"):
        _cells.get_cells(p, ["A1"] * 1_001)
    with pytest.raises(XlMcpError, match="values"):
        _cells.get_cells(p, ["A1"], values="live")


# ---------------------------------------------------------------- set_cells


def test_set_cells_one_batch_one_save(tmp_path):
    p = _make(tmp_path / "s.xlsx", [[None]])
    out = _cells.set_cells(p, [
        {"cell": "A1", "value": "x"},
        {"cell": {"cell": "B2"}, "value": 7},
        {"location": "C3", "value": "=XLOOKUP(1,A:A,B:B)"},
    ])
    assert out["ok"] is True and out["verified"] is True
    assert out["backup"] == "prev"
    assert out["changed"]["cells_written"] == 3
    wb = openpyxl.load_workbook(p)
    ws = wb["Data"]
    assert ws["A1"].value == "x" and ws["B2"].value == 7
    # '=' strings became formulas, normalized through the _xlfn shim
    assert ws["C3"].value == "=_xlfn.XLOOKUP(1,A:A,B:B)"
    wb.close()


def test_set_cells_bad_item_refuses_whole_batch(tmp_path):
    p = _make(tmp_path / "sb.xlsx", [["orig"]])
    before = open(p, "rb").read()
    with pytest.raises(XlMcpError, match=r"cells\[1\]"):
        _cells.set_cells(p, [
            {"cell": "A1", "value": "new"},
            {"cell": {"range": "A1:B2"}, "value": 1},
        ])
    with pytest.raises(XlMcpError, match="value"):
        _cells.set_cells(p, [{"cell": "A1"}])
    with pytest.raises(XlMcpError, match="cell"):
        _cells.set_cells(p, [{"value": 1}])
    assert open(p, "rb").read() == before, (
        "a refused batch must leave the file byte-for-byte unchanged")


def test_set_cells_hazard_gate_inherited(tmp_path):
    """The scatter writer refuses on a slicer-bearing workbook like every
    other mutation (WorkbookPackage pipeline, not a side door)."""
    import shutil
    from pathlib import Path
    from xlsx_mcp.core.errors import HazardRefused
    corpus = Path(__file__).resolve().parents[1] / "fixtures" / "corpus"
    p = str(tmp_path / "sl.xlsx")
    shutil.copy(corpus / "pivot_slicer.xlsx", p)
    with pytest.raises(HazardRefused):
        _cells.set_cells(p, [{"cell": "Z99", "value": 1}])


def test_set_cells_per_item_sheet(tmp_path):
    p = _make(tmp_path / "ms.xlsx", [[1]])
    wb = openpyxl.load_workbook(p)
    wb.create_sheet("Other")
    wb.save(p)
    wb.close()
    _cells.set_cells(p, [
        {"cell": "A1", "value": "d"},
        {"cell": "A1", "sheet": "Other", "value": "o"},
    ])
    wb = openpyxl.load_workbook(p)
    assert wb["Data"]["A1"].value == "d"
    assert wb["Other"]["A1"].value == "o"
    wb.close()
