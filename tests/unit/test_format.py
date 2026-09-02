"""Unit tests for ops/format.py: format_cells and set_dimensions."""

from __future__ import annotations

import openpyxl
import pytest

from xlsx_mcp.core.errors import XlMcpError
from xlsx_mcp.ops import format as fmt


@pytest.fixture()
def book(tmp_path):
    p = tmp_path / "f.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    for r in range(1, 4):
        ws.cell(r, 1, f"row{r}")
        ws.cell(r, 2, r)
    wb.save(p)
    return p


def test_format_font_fill_number(book):
    r = fmt.format_cells(str(book), {"range": "A1:B1"},
                         number_format="0.00",
                         font={"bold": True, "color": "FF0000"},
                         fill={"color": "FFFF00"})
    assert r["ok"] and r["verified"]
    wb = openpyxl.load_workbook(book)
    a1 = wb["S"]["A1"]
    assert a1.font.bold is True
    assert a1.font.color.rgb == "FFFF0000"        # red font
    assert a1.fill.fgColor.rgb == "FFFFFF00"      # yellow fill
    assert wb["S"]["B1"].number_format == "0.00"


def test_format_merge_preserves_unspecified(book):
    # set bold, then set italic; bold must remain
    fmt.format_cells(str(book), {"cell": "A1"}, font={"bold": True})
    fmt.format_cells(str(book), {"cell": "A1"}, font={"italic": True})
    a1 = openpyxl.load_workbook(book)["S"]["A1"]
    assert a1.font.bold is True and a1.font.italic is True


def test_format_border_and_alignment(book):
    fmt.format_cells(str(book), {"cell": "A1"},
                     border={"style": "thin"},
                     alignment={"horizontal": "center", "wrap_text": True})
    a1 = openpyxl.load_workbook(book)["S"]["A1"]
    assert a1.border.left.style == "thin"
    assert a1.alignment.horizontal == "center"
    assert a1.alignment.wrap_text is True


def test_format_requires_something(book):
    with pytest.raises(XlMcpError):
        fmt.format_cells(str(book), {"cell": "A1"})


def test_format_bad_color(book):
    with pytest.raises(XlMcpError):
        fmt.format_cells(str(book), {"cell": "A1"}, font={"color": "zzz"})


def test_format_bad_alignment(book):
    with pytest.raises(XlMcpError):
        fmt.format_cells(str(book), {"cell": "A1"},
                         alignment={"horizontal": "sideways"})


def test_set_dimensions_widths_heights(book):
    r = fmt.set_dimensions(str(book), "S",
                           column_widths={"A": 25, 2: 12},
                           row_heights={1: 30})
    assert r["ok"] and r["verified"]
    wb = openpyxl.load_workbook(book)
    assert wb["S"].column_dimensions["A"].width == 25
    assert wb["S"].column_dimensions["B"].width == 12
    assert wb["S"].row_dimensions[1].height == 30


def test_set_dimensions_autofit_is_approximate(book):
    r = fmt.set_dimensions(str(book), "S", autofit_columns=["A"])
    assert r["changed"]["dimensions"]["autofit_columns"]["approximate"] is True
    assert openpyxl.load_workbook(book)["S"].column_dimensions["A"].width > 0


def test_set_dimensions_hide(book):
    fmt.set_dimensions(str(book), "S", hide_columns=["B"], hide_rows=[2])
    wb = openpyxl.load_workbook(book)
    assert wb["S"].column_dimensions["B"].hidden is True
    assert wb["S"].row_dimensions[2].hidden is True


def test_set_dimensions_requires_something(book):
    with pytest.raises(XlMcpError):
        fmt.set_dimensions(str(book), "S")
