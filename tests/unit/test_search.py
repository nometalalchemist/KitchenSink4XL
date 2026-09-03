"""find_cells / replace_cells: scoped search over values and formulas, the
guarded regex path, dry_run previews, count reporting, and the import_data
injection lint on replacement text.
"""

from __future__ import annotations

import openpyxl
import pytest

from xlsx_mcp.core.errors import XlMcpError
from xlsx_mcp.ops import cells as _cells
from xlsx_mcp.ops import search as _search


def _make(path, rows, sheet="Data"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    for i, row in enumerate(rows, start=1):
        for j, v in enumerate(row, start=1):
            ws.cell(i, j, v)
    wb.save(path)
    wb.close()
    return str(path)


# --------------------------------------------------------------- find_cells


def test_find_contains_default_is_case_insensitive(tmp_path):
    p = _make(tmp_path / "f.xlsx",
              [["Total Q3"], ["subtotal"], ["other"], [42]])
    out = _search.find_cells(p, "total")
    assert out["total"] == 2
    assert [m["cell"] for m in out["matches"]] == ["A1", "A2"]
    out = _search.find_cells(p, "total", match_case=True)
    assert [m["cell"] for m in out["matches"]] == ["A2"]


def test_find_exact_and_numbers(tmp_path):
    p = _make(tmp_path / "fe.xlsx", [["4"], [42], [4]])
    out = _search.find_cells(p, "4", match="exact")
    assert [m["cell"] for m in out["matches"]] == ["A1", "A3"], (
        "exact matches the whole cell text, numbers included")


def test_find_in_formulas_and_both(tmp_path):
    p = _make(tmp_path / "ff.xlsx", [[1], [2], ["=SUM(A1:A2)"]])
    # a never-calculated formula has no cached value: no 'values' hit
    assert _search.find_cells(p, "SUM")["total"] == 0
    out = _search.find_cells(p, "SUM", look_in="formulas")
    assert out["total"] == 1
    assert out["matches"][0] == {"sheet": "Data", "cell": "A3",
                                 "match_in": "formula",
                                 "formula": "=SUM(A1:A2)"}
    assert _search.find_cells(p, "SUM", look_in="both")["total"] == 1


def test_find_regex_scope_and_paging(tmp_path):
    p = str(tmp_path / "fs.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    for r in range(1, 6):
        ws.cell(r, 1, f"item-{r}")
    other = wb.create_sheet("Other")
    other["A1"] = "item-x"
    wb.save(p)
    wb.close()
    out = _search.find_cells(p, r"item-\d", match="regex")
    assert out["total"] == 5, "regex digits must not match item-x"
    out = _search.find_cells(p, "item", sheet="Other")
    assert [m["sheet"] for m in out["matches"]] == ["Other"]
    out = _search.find_cells(p, "item", location={"range": "A2:A4"})
    assert [m["cell"] for m in out["matches"]] == ["A2", "A3", "A4"]
    out = _search.find_cells(p, "item", limit=2, offset=1)
    assert out["total"] == 6
    assert out["returned"] == 2 and out["truncated"] is True
    assert [m["cell"] for m in out["matches"]] == ["A2", "A3"]


def test_find_refusals(tmp_path):
    p = _make(tmp_path / "fr.xlsx", [[1]])
    with pytest.raises(XlMcpError, match="invalid regex"):
        _search.find_cells(p, "(unclosed", match="regex")
    with pytest.raises(XlMcpError, match="look_in"):
        _search.find_cells(p, "x", look_in="everything")
    with pytest.raises(XlMcpError, match="match"):
        _search.find_cells(p, "x", match="fuzzy")
    with pytest.raises(XlMcpError, match="non-empty"):
        _search.find_cells(p, "")


# ------------------------------------------------------------ replace_cells


def test_replace_dry_run_previews_without_writing(tmp_path):
    p = _make(tmp_path / "dr.xlsx", [["foo bar foo"], ["foo"], ["baz"]])
    before = open(p, "rb").read()
    out = _search.replace_cells(p, "foo", "qux", dry_run=True)
    assert out["dry_run"] is True
    assert out["cells_matched"] == 2
    assert out["occurrences"] == 3
    assert out["changes"][0] == {"sheet": "Data", "cell": "A1",
                                 "kind": "value", "from": "foo bar foo",
                                 "to": "qux bar qux"}
    assert open(p, "rb").read() == before


def test_replace_contains_counts_and_coerces_numbers(tmp_path):
    p = _make(tmp_path / "rc.xlsx", [["price 42"], [42], ["x"]])
    out = _search.replace_cells(p, "42", "99")
    assert out["verified"] is True
    assert out["changed"]["replaced"] == {
        "find": "42", "match": "contains", "look_in": "values",
        "cells_changed": 2, "occurrences": 2}
    wb = openpyxl.load_workbook(p)
    assert wb["Data"]["A1"].value == "price 99"
    assert wb["Data"]["A2"].value == 99, (
        "a replaced value that parses as a number is written as a number")
    wb.close()


def test_replace_exact_whole_cell_only(tmp_path):
    p = _make(tmp_path / "re.xlsx", [["cat"], ["catalog"]])
    out = _search.replace_cells(p, "cat", "dog", match="exact")
    assert out["changed"]["replaced"]["cells_changed"] == 1
    wb = openpyxl.load_workbook(p)
    assert wb["Data"]["A1"].value == "dog"
    assert wb["Data"]["A2"].value == "catalog"
    wb.close()


def test_replace_regex_with_backreference(tmp_path):
    p = _make(tmp_path / "rr.xlsx", [["id-123"], ["id-9"]])
    _search.replace_cells(p, r"id-(\d+)", r"\1-id", match="regex")
    wb = openpyxl.load_workbook(p)
    assert wb["Data"]["A1"].value == "123-id"
    assert wb["Data"]["A2"].value == "9-id"
    wb.close()
    with pytest.raises(XlMcpError, match="not valid against"):
        _search.replace_cells(p, r"(\d+)", r"\9", match="regex")


def test_replace_in_formulas_keeps_cell_a_formula(tmp_path):
    p = _make(tmp_path / "rf.xlsx", [[1], [2], ["=SUM(A1:A2)"]])
    out = _search.replace_cells(p, "SUM", "AVERAGE", look_in="formulas")
    assert out["changed"]["replaced"]["cells_changed"] == 1
    wb = openpyxl.load_workbook(p)
    assert wb["Data"]["A3"].value == "=AVERAGE(A1:A2)"
    wb.close()
    # values-mode replacement never rewrites formula text
    out = _search.replace_cells(p, "AVERAGE", "MAX")
    assert out["cells_changed"] == 0 and out["saved"] is False


def test_replacement_that_looks_like_a_formula_is_neutralized(tmp_path):
    p = _make(tmp_path / "inj.xlsx", [["payload"], ["safe"]])
    out = _search.replace_cells(p, "payload", "=CMD()")
    assert any("formula injection" in w for w in out["warnings"])
    wb = openpyxl.load_workbook(p)
    cell = wb["Data"]["A1"]
    assert cell.value == "=CMD()"
    assert cell.data_type == "s", (
        "the replacement must be TEXT, never a live formula")
    wb.close()
    # the explicit override writes it live (and normalized)
    _search.replace_cells(p, "safe", "=XLOOKUP(1,A:A,A:A)", formulas=True)
    wb = openpyxl.load_workbook(p)
    assert wb["Data"]["A2"].value == "=_xlfn.XLOOKUP(1,A:A,A:A)"
    assert wb["Data"]["A2"].data_type == "f"
    wb.close()


def test_replace_scoped_to_range_leaves_rest_alone(tmp_path):
    p = _make(tmp_path / "scope.xlsx", [["x"], ["x"], ["x"]])
    out = _search.replace_cells(p, "x", "y", location={"range": "A2:A3"})
    assert out["changed"]["replaced"]["cells_changed"] == 2
    got = _cells.read_range(p, {"range": "A1:A3"})
    assert [r[0] for r in got["values"]] == ["x", "y", "y"]


def test_replace_no_match_reports_and_skips_save(tmp_path):
    p = _make(tmp_path / "nm.xlsx", [["a"]])
    before = open(p, "rb").read()
    out = _search.replace_cells(p, "zzz", "y")
    assert out["cells_changed"] == 0 and out["saved"] is False
    assert open(p, "rb").read() == before
