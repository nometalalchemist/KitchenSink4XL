"""Regressions for the combined pre-beta fix wave (2026-09-05).

Both final-round reports are remediated here, each test built from the
report's exact repro:
- fresh-eyes round (1H/5M/9L): out_file self-clobber, leading-apostrophe
  sheet names, sort vs filter-hidden rows, aggregate coercion,
  KS4XL_PACK_POLICY fail-open, write_range docstring drift, plus LOWs.
- destroyer round (3H/4M/4L): export overwrite class, create/copy
  overwrite with no backup, restore validating too little, promote-failure
  burning prev, byte-range-lock misdiagnosis, allow_loss dead-end,
  recovery discoverability, plus LOWs.
"""

from __future__ import annotations

import os
import shutil
import zipfile
from pathlib import Path

import openpyxl
import pytest

from xlsx_mcp.core import safesave
from xlsx_mcp.core.errors import (
    ValidationFailed,
    WorkbookCorrupt,
    WorkbookLocked,
    XlMcpError,
)
from xlsx_mcp.ops import backups as backups_ops
from xlsx_mcp.ops import cells as cells_ops
from xlsx_mcp.ops import dataio
from xlsx_mcp.ops import lifecycle
from xlsx_mcp.ops import sortfilter


def _make_book(path, rows, sheet="S"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    for r, row in enumerate(rows, 1):
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v)
    wb.save(path)
    return path


@pytest.fixture()
def victim(tmp_path):
    return _make_book(tmp_path / "victim.xlsx",
                      [["Name", "Amount"], ["a", 10], ["b", 20]])


# ======================================================================
# Convergent headline: fresh-eyes H-1 == destroyer H-1 (export self-clobber)
# ======================================================================


class TestExportOutTargetGuard:
    def test_export_range_refuses_source_workbook(self, victim):
        with pytest.raises(XlMcpError):
            dataio.export_range(str(victim), out_file=str(victim))
        # the workbook is untouched and still opens
        openpyxl.load_workbook(victim).close()

    def test_export_range_refuses_source_even_with_overwrite(self, victim):
        with pytest.raises(XlMcpError):
            dataio.export_range(str(victim), out_file=str(victim),
                                overwrite=True)
        openpyxl.load_workbook(victim).close()

    def test_export_range_refuses_backup_slot(self, victim, tmp_path):
        slot = (tmp_path / safesave.BACKUP_DIR_NAME / "victim.xlsx"
                / "prev.xlsx")
        with pytest.raises(XlMcpError):
            dataio.export_range(str(victim), out_file=str(slot))
        assert not slot.exists()

    def test_export_range_refuses_any_workbook_extension(self, victim,
                                                         tmp_path):
        other = _make_book(tmp_path / "other.xlsx", [["x"]])
        before = other.read_bytes()
        with pytest.raises(XlMcpError):
            dataio.export_range(str(victim), out_file=str(other))
        assert other.read_bytes() == before
        # even a NONEXISTENT workbook-named target refuses (path mix-up)
        with pytest.raises(XlMcpError):
            dataio.export_range(str(victim),
                                out_file=str(tmp_path / "new.xlsm"))

    def test_export_file_json_refuses_source(self, victim):
        with pytest.raises(XlMcpError):
            dataio.export_file(str(victim), fmt="json",
                               out_file=str(victim))
        openpyxl.load_workbook(victim).close()

    def test_export_existing_out_file_refuses_without_overwrite(
            self, victim, tmp_path):
        out = tmp_path / "out.csv"
        out.write_text("precious old data", encoding="utf-8")
        with pytest.raises(XlMcpError) as ei:
            dataio.export_range(str(victim), out_file=str(out))
        assert getattr(ei.value, "code", None) == "CONFLICT"
        assert out.read_text(encoding="utf-8") == "precious old data"

    def test_export_overwrite_keeps_timestamped_bak(self, victim, tmp_path):
        out = tmp_path / "out.csv"
        out.write_text("precious old data", encoding="utf-8")
        r = dataio.export_range(str(victim), out_file=str(out),
                                overwrite=True)
        assert r["overwrote"] is True
        bak = Path(r["backup_of_replaced"])
        assert bak.exists()
        assert bak.read_text(encoding="utf-8") == "precious old data"
        assert "Name" in out.read_text(encoding="utf-8")

    def test_export_file_out_dir_refuses_existing_without_overwrite(
            self, victim, tmp_path):
        outdir = tmp_path / "exports"
        outdir.mkdir()
        clash = outdir / "victim_S.csv"
        clash.write_text("old", encoding="utf-8")
        with pytest.raises(XlMcpError):
            dataio.export_file(str(victim), fmt="csv", out_dir=str(outdir))
        assert clash.read_text(encoding="utf-8") == "old"
        r = dataio.export_file(str(victim), fmt="csv", out_dir=str(outdir),
                               overwrite=True)
        assert r["overwrote"] is True
        assert any(p.endswith(".bak") for p in r["backups_of_replaced"])

    def test_export_file_out_dir_refuses_backup_store(self, victim,
                                                      tmp_path):
        bdir = tmp_path / safesave.BACKUP_DIR_NAME / "victim.xlsx"
        bdir.mkdir(parents=True)
        with pytest.raises(XlMcpError):
            dataio.export_file(str(victim), fmt="csv", out_dir=str(bdir))


# ======================================================================
# Destroyer H-2: create/copy overwrite must rotate-before-clobber
# ======================================================================


class TestOverwriteRotatesBeforeClobber:
    def test_create_workbook_overwrite_backs_up_victim(self, victim):
        before = victim.read_bytes()
        r = lifecycle.create_workbook(str(victim), overwrite=True)
        assert r["backup"] == "prev"
        slot = safesave.slot_dir(victim) / safesave.PREV_SLOT
        assert slot.exists()
        assert slot.read_bytes() == before
        # and the restore round-trip brings the heirloom back
        backups_ops.restore_backup(str(victim), "prev")
        assert victim.read_bytes() == before

    def test_copy_workbook_overwrite_backs_up_victim(self, victim,
                                                     tmp_path):
        src = _make_book(tmp_path / "src.xlsx", [["new"]])
        before = victim.read_bytes()
        r = lifecycle.copy_workbook(str(src), str(victim), overwrite=True)
        assert r["backup"] == "prev"
        slot = safesave.slot_dir(victim) / safesave.PREV_SLOT
        assert slot.read_bytes() == before

    def test_create_workbook_refuses_backup_store(self, tmp_path):
        target = (tmp_path / safesave.BACKUP_DIR_NAME / "x.xlsx"
                  / "prev.xlsx")
        target.parent.mkdir(parents=True)
        with pytest.raises(XlMcpError):
            lifecycle.create_workbook(str(target))

    def test_copy_workbook_refuses_self_copy(self, victim):
        with pytest.raises(XlMcpError):
            lifecycle.copy_workbook(str(victim), str(victim),
                                    overwrite=True)

    def test_no_overwrite_still_refuses(self, victim):
        with pytest.raises(FileExistsError):
            lifecycle.create_workbook(str(victim))


# ======================================================================
# Fresh-eyes M-1: leading-apostrophe sheet name kills the file in Excel
# ======================================================================


class TestSheetNameApostrophe:
    def test_create_workbook_refuses_leading_apostrophe(self, tmp_path):
        with pytest.raises(XlMcpError):
            lifecycle.create_workbook(str(tmp_path / "n.xlsx"),
                                      sheets=["'quoted"])
        assert not (tmp_path / "n.xlsx").exists()

    def test_manage_worksheet_add_refuses(self, victim):
        with pytest.raises(XlMcpError):
            lifecycle.manage_worksheet(str(victim), "add",
                                       new_name="'quoted")

    def test_manage_worksheet_rename_refuses_trailing(self, victim):
        with pytest.raises(XlMcpError):
            lifecycle.manage_worksheet(str(victim), "rename", sheet="S",
                                       new_name="quoted'")

    def test_normal_names_still_fine(self, victim):
        r = lifecycle.manage_worksheet(str(victim), "add",
                                       new_name="It's fine")
        assert r["ok"]  # apostrophe INSIDE a name is legal
