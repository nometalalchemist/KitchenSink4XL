"""Unit tests for ops/tables.py: create_table, get_table, manage_table."""

from __future__ import annotations

import openpyxl
import pytest

from xlsx_mcp.core.errors import TargetNotFound, XlMcpError
from xlsx_mcp.ops import tables as tb


@pytest.fixture()
def book(tmp_path):
    p = tmp_path / "t.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    rows = [["Name", "Amount"], ["a", 10], ["b", 20], ["c", 30]]
    for r, row in enumerate(rows, 1):
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v)
    # a formula BELOW the table, to prove structural edits rewrite it
    ws["A7"] = "=SUM(B2:B4)"
    wb.save(p)
    return p


def test_create_and_get_table(book):
    r = tb.create_table(str(book), {"range": "A1:B4"}, "Sales")
    assert r["ok"] and r["verified"]
    wb = openpyxl.load_workbook(book)
    assert "Sales" in wb["S"].tables
    got = tb.get_table(str(book), "Sales", records=True)
    assert got["columns"] == ["Name", "Amount"]
    assert got["row_count"] == 3
    assert got["rows"][0] == {"Name": "a", "Amount": 10}


def test_create_table_rejects_bad_name(book):
    with pytest.raises(XlMcpError):
        tb.create_table(str(book), {"range": "A1:B4"}, "A1")
    with pytest.raises(XlMcpError):
        tb.create_table(str(book), {"range": "A1:B4"}, "has space")


def test_create_table_duplicate_name(book):
    tb.create_table(str(book), {"range": "A1:B4"}, "Sales")
    with pytest.raises(XlMcpError):
        tb.create_table(str(book), {"range": "A1:B4"}, "sales")


def test_create_table_with_totals(book):
    r = tb.create_table(str(book), {"range": "A1:B4"}, "Sales",
                        totals_row=True, totals={"Amount": "sum"})
    assert r["ok"]
    wb = openpyxl.load_workbook(book)
    t = wb["S"].tables["Sales"]
    assert t.totalsRowCount == 1
    assert t.ref == "A1:B5"


def test_manage_table_add_row_rewrites_refs(book):
    tb.create_table(str(book), {"range": "A1:B4"}, "Sales")
    # A7 formula =SUM(B2:B4) must shift down when a data row is inserted
    r = tb.manage_table(str(book), "Sales", "add_row", values=["d", 40])
    assert r["ok"] and r["verified"]
    wb = openpyxl.load_workbook(book)
    t = wb["S"].tables["Sales"]
    assert t.ref == "A1:B5"
    assert wb["S"]["A5"].value == "d"
    assert wb["S"]["B5"].value == 40
    # the formula moved from A7 to A8 and its range expanded/shifted
    assert wb["S"]["A7"].value is None
    assert wb["S"]["A8"].value is not None
    assert wb["S"]["A8"].value.startswith("=SUM(")


def test_manage_table_delete_row(book):
    tb.create_table(str(book), {"range": "A1:B4"}, "Sales")
    tb.manage_table(str(book), "Sales", "delete_row", index=1)
    wb = openpyxl.load_workbook(book)
    assert wb["S"].tables["Sales"].ref == "A1:B3"
    assert wb["S"]["A2"].value == "b"  # 'a' row removed


def test_manage_table_add_and_delete_column(book):
    tb.create_table(str(book), {"range": "A1:B4"}, "Sales")
    tb.manage_table(str(book), "Sales", "add_column", column="Tax",
                    values=[1, 2, 3])
    wb = openpyxl.load_workbook(book)
    assert wb["S"].tables["Sales"].ref == "A1:C4"
    assert wb["S"]["C1"].value == "Tax"
    assert wb["S"]["C2"].value == 1
    tb.manage_table(str(book), "Sales", "delete_column", column="Tax")
    wb = openpyxl.load_workbook(book)
    assert wb["S"].tables["Sales"].ref == "A1:B4"


def test_manage_table_rename_and_style(book):
    tb.create_table(str(book), {"range": "A1:B4"}, "Sales")
    tb.manage_table(str(book), "Sales", "rename", new_name="Revenue")
    wb = openpyxl.load_workbook(book)
    assert "Revenue" in wb["S"].tables and "Sales" not in wb["S"].tables
    tb.manage_table(str(book), "Revenue", "set_style",
                    style="TableStyleLight1")
    wb = openpyxl.load_workbook(book)
    assert wb["S"].tables["Revenue"].tableStyleInfo.name == "TableStyleLight1"


def test_manage_table_toggle_totals_and_to_range(book):
    tb.create_table(str(book), {"range": "A1:B4"}, "Sales")
    tb.manage_table(str(book), "Sales", "toggle_totals", on=True,
                    totals={"Amount": "sum"})
    wb = openpyxl.load_workbook(book)
    assert wb["S"].tables["Sales"].ref == "A1:B5"
    tb.manage_table(str(book), "Sales", "toggle_totals", on=False)
    wb = openpyxl.load_workbook(book)
    assert wb["S"].tables["Sales"].ref == "A1:B4"
    tb.manage_table(str(book), "Sales", "to_range")
    wb = openpyxl.load_workbook(book)
    assert "Sales" not in wb["S"].tables
    assert wb["S"]["A1"].value == "Name"  # data kept


def test_get_table_not_found(book):
    with pytest.raises(TargetNotFound):
        tb.get_table(str(book), "Nope")


def test_manage_table_bad_action(book):
    tb.create_table(str(book), {"range": "A1:B4"}, "Sales")
    with pytest.raises(XlMcpError):
        tb.manage_table(str(book), "Sales", "explode")
