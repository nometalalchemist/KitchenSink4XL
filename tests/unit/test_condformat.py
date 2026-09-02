"""Unit tests for ops/condformat.py: manage_conditional_format."""

from __future__ import annotations

import openpyxl
import pytest

from xlsx_mcp.core.errors import TargetNotFound, XlMcpError
from xlsx_mcp.ops import condformat as cf


@pytest.fixture()
def book(tmp_path):
    p = tmp_path / "cf.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    for r in range(1, 6):
        ws.cell(r, 1, f"n{r}")
        ws.cell(r, 2, r * 10)
    wb.save(p)
    return p


def _cf_count(path):
    wb = openpyxl.load_workbook(path)
    return sum(len(v) for v in wb["S"].conditional_formatting._cf_rules.values())


def test_add_cell_is_and_list(book):
    r = cf.manage_conditional_format(
        str(book), "add", location={"range": "B1:B5"}, cf_type="cell_is",
        params={"operator": "greaterThan", "formula": 20, "fill": "FF0000"})
    assert r["ok"] and r["verified"]
    listing = cf.manage_conditional_format(str(book), "list")
    assert listing["count"] == 1
    assert listing["rules"][0]["type"] == "cellIs"


def test_add_each_type(book):
    cf.manage_conditional_format(str(book), "add", location={"range": "B1:B5"},
                                 cf_type="color_scale")
    cf.manage_conditional_format(str(book), "add", location={"range": "B1:B5"},
                                 cf_type="data_bar")
    cf.manage_conditional_format(str(book), "add", location={"range": "B1:B5"},
                                 cf_type="icon_set")
    cf.manage_conditional_format(
        str(book), "add", location={"range": "B1:B5"}, cf_type="formula",
        params={"formula": "$B1>15", "fill": "00FF00"})
    cf.manage_conditional_format(
        str(book), "add", location={"range": "B1:B5"}, cf_type="top_bottom",
        params={"rank": 3})
    assert _cf_count(book) == 5


def test_cf_survives_structural_edit(book):
    from xlsx_mcp.ops import cells
    cf.manage_conditional_format(
        str(book), "add", location={"range": "B2:B5"}, cf_type="cell_is",
        params={"operator": "greaterThan", "formula": 20})
    # insert a row above the CF range; the sqref must move down
    cells.write_range  # noqa: reference to confirm import
    from xlsx_mcp.core import refs
    from xlsx_mcp.core.package import WorkbookPackage
    pkg = WorkbookPackage.open(str(book))
    pkg.modify_structure(refs.RefEdit("S", refs.INSERT_ROWS, index=1, count=1))
    pkg.save()
    listing = cf.manage_conditional_format(str(book), "list")
    assert listing["rules"][0]["range"] == "B3:B6"


def test_delete_by_range(book):
    cf.manage_conditional_format(str(book), "add", location={"range": "B1:B5"},
                                 cf_type="data_bar")
    cf.manage_conditional_format(str(book), "delete", location={"range": "B1:B5"})
    assert _cf_count(book) == 0


def test_delete_missing_refuses(book):
    with pytest.raises(TargetNotFound):
        cf.manage_conditional_format(str(book), "delete",
                                     location={"range": "Z1:Z5"})


def test_bad_type(book):
    with pytest.raises(XlMcpError):
        cf.manage_conditional_format(str(book), "add",
                                     location={"range": "B1:B5"},
                                     cf_type="rainbow")
