"""Regressions for the author's insane-mode field-test findings (2026-09-05).

Each test uses the field report's EXACT reproduction so a re-run of that stress
battery would pass. Findings triaged and fixed:

  * empty-string writes (CRITICAL): "" tripped verify-after-write and no write
    tool could store a blank string.
  * query_range order_by direction (MAJOR): a {"order": "desc"} spec sorted
    ascending SILENTLY; only the "dir" key was honored, none was documented,
    and an unknown direction word defaulted quietly to ascending.
  * query_range group_by JSON-string form (MAJOR): a client delivering the
    untyped group_by as '["Region"]' or '"Region"' hit `no column '["Region"]'`.

The COM-tier sparkline mis-binding regression lives in test_comtier.py (needs
Excel) since it is a live-COM finding.
"""

from __future__ import annotations

import openpyxl
import pytest

from xlsx_mcp.core.errors import XlMcpError
from xlsx_mcp.ops import cells


# ------------------------------------------------------- empty-string writes

def _new(tmp_path):
    p = tmp_path / "blank.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "Test"
    wb.save(p)
    return str(p)


def test_set_cell_empty_string_writes(tmp_path):
    p = _new(tmp_path)
    r = cells.set_cell(p, "A1", "", sheet="Test")
    assert r["verified"] is True
    back = cells.read_range(p, "A1", sheet="Test")
    assert back["values"] == [[None]]  # "" serializes to a blank cell


def test_write_range_with_empty_strings(tmp_path):
    p = _new(tmp_path)
    r = cells.write_range(p, "A1", [["x", "", "y"], ["", "z", ""]], sheet="Test")
    assert r["verified"] is True


def test_set_cells_batch_one_empty(tmp_path):
    p = _new(tmp_path)
    r = cells.set_cells(
        p, [{"cell": "B1", "value": "a"}, {"cell": "B2", "value": ""}],
        sheet="Test")
    assert r["verified"] is True


# ------------------------------------------------------ order_by direction

@pytest.fixture()
def q(tmp_path):
    p = tmp_path / "q.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Test"
    for r, row in enumerate([["Product", "Q1", "Region"],
                             ["a", 100, "E"], ["b", 150, "W"],
                             ["c", 80, "E"], ["d", 120, "W"]], 1):
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v)
    wb.save(p)
    return str(p)


@pytest.mark.parametrize("key", ["dir", "order", "direction"])
def test_order_by_desc_honored_for_every_key(q, key):
    r = cells.query_range(q, "A1:C5", sheet="Test",
                          order_by=[{"column": "Q1", key: "desc"}],
                          columns=["Q1"])
    assert [row[0] for row in r["rows"]] == [150, 120, 100, 80]


def test_order_by_asc_still_ascending(q):
    r = cells.query_range(q, "A1:C5", sheet="Test",
                          order_by=[{"column": "Q1", "order": "asc"}],
                          columns=["Q1"])
    assert [row[0] for row in r["rows"]] == [80, 100, 120, 150]


def test_order_by_unknown_direction_refuses(q):
    with pytest.raises(XlMcpError, match="not understood"):
        cells.query_range(q, "A1:C5", sheet="Test",
                          order_by=[{"column": "Q1", "dir": "sideways"}])


# --------------------------------------------------------- group_by strings

@pytest.mark.parametrize("gb", ['["Region"]', '"Region"', ["Region"], "Region"])
def test_group_by_accepts_json_string_forms(q, gb):
    r = cells.query_range(q, "A1:C5", sheet="Test", group_by=gb,
                          aggregate=[{"column": "Q1", "func": "sum"}])
    got = {tuple(g["group"].values())[0]: g["aggregates"]["sum_Q1"]
           for g in r["groups"]}
    assert got == {"E": 180, "W": 270}
