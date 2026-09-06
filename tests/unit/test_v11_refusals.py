"""V1.1 refusal-surface regressions (the triage's Tier 2 mechanical set).

Each test is the triage item's own repro. The theme across all of them is
that a refusal has to name what actually happened: a verdict that agrees
with its payload, a codec name checked before it is used, a rotation range
stated instead of enumerated, a device name identified as a device, an
offset that refuses instead of clamping, and a column reference that reads
as a column reference instead of a Python type error.
"""

from __future__ import annotations

import openpyxl
import pytest

from xlsx_mcp.core import outguard
from xlsx_mcp.core.errors import XlMcpError
from xlsx_mcp.ops import cells as cells_ops
from xlsx_mcp.ops import dataio
from xlsx_mcp.ops import format as format_ops
from xlsx_mcp.ops import validation as validation_ops


def _book(path, rows, sheet="S"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    for r in rows:
        ws.append(r)
    wb.save(path)
    return str(path)


# ------------------------------------------------------------------- V1-6


def test_validate_references_fails_when_it_found_unknown_sheets(tmp_path):
    """A formula pointing at a deleted sheet is a finding, so the verdict
    that carries it cannot read passed:true. An agent gating automation on
    passed used to ship a workbook whose formulas all break to #REF! the
    moment Excel opens it."""
    p = tmp_path / "refs.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Main"
    ws["A1"] = "=Data!A1+1"
    wb.save(p)

    out = validation_ops.validate(str(p), checks=["references"])
    findings = out["results"]["references"]["findings"]
    assert findings.get("unknown_sheet_references"), (
        "fixture did not produce the finding this test is about")
    assert out["results"]["references"]["passed"] is False
    assert out["passed"] is False


def test_validate_references_still_passes_on_a_clean_workbook(tmp_path):
    """The guard above must not turn every workbook red."""
    p = _book(tmp_path / "clean.xlsx", [["a", 1], ["b", 2]])
    out = validation_ops.validate(p, checks=["references"])
    assert out["results"]["references"]["passed"] is True
    assert out["passed"] is True


# ------------------------------------------------------------------- V1-7


def test_import_data_refuses_an_unknown_codec_inline(tmp_path):
    """The file branch discovered a bad codec inside open(); the inline
    branch never looked, so utf-99 was accepted in silence."""
    p = _book(tmp_path / "imp.xlsx", [["x"]])
    with pytest.raises(XlMcpError) as exc:
        dataio.import_data(p, source="a,b\n1,2", fmt="csv",
                           encoding="utf-99")
    assert "utf-99" in str(exc.value)
    assert "encoding" in str(exc.value).lower()


def test_import_data_refuses_an_unknown_codec_from_a_file(tmp_path):
    p = _book(tmp_path / "imp2.xlsx", [["x"]])
    src = tmp_path / "in.csv"
    src.write_text("a,b\n1,2", encoding="utf-8")
    with pytest.raises(XlMcpError) as exc:
        dataio.import_data(p, source_file=str(src), fmt="csv",
                           encoding="utf-99")
    assert "utf-99" in str(exc.value)


def test_import_data_still_accepts_real_codecs(tmp_path):
    p = _book(tmp_path / "imp3.xlsx", [["x"]])
    for enc in ("utf-8", "utf-8-sig", "cp1252", "latin-1", "UTF8"):
        out = dataio.import_data(p, source="a,b\n1,2", fmt="csv",
                                 encoding=enc, location="A1")
        assert out["ok"] is not False


# ------------------------------------------------------- V1-8 (five items)


def test_a_list_column_reference_reads_as_a_column_reference(tmp_path):
    """columns=[[]] used to leak "unhashable type: 'list'" out of a dict
    lookup, which tells the caller nothing about columns."""
    p = _book(tmp_path / "q.xlsx", [["name", "n"], ["a", 1]])
    with pytest.raises(XlMcpError) as exc:
        cells_ops.query_range(p, columns=[[]])
    msg = str(exc.value)
    assert "unhashable" not in msg
    assert "column reference" in msg
    assert "list" in msg


def test_text_rotation_states_its_range_instead_of_listing_it(tmp_path):
    """openpyxl refuses 720 by printing all 182 legal values."""
    p = _book(tmp_path / "rot.xlsx", [["x"]])
    with pytest.raises(XlMcpError) as exc:
        format_ops.format_cells(p, "A1", alignment={"text_rotation": 720})
    msg = str(exc.value)
    assert "0 to 180" in msg and "255" in msg
    assert msg.count(",") < 5, f"refusal still dumps a set: {msg}"


@pytest.mark.parametrize("value", [0, 45, 180, 255])
def test_text_rotation_accepts_every_legal_value(tmp_path, value):
    p = _book(tmp_path / f"rot{value}.xlsx", [["x"]])
    format_ops.format_cells(p, "A1", alignment={"text_rotation": value})
    wb = openpyxl.load_workbook(p)
    assert wb.active["A1"].alignment.text_rotation == value
    wb.close()


def test_a_negative_offset_refuses_instead_of_clamping(tmp_path):
    """offset:-5 used to clamp to 0 and page from the top, returning rows
    nobody asked for with no signal that the argument was ignored. limit
    already refused; offset now matches it."""
    p = _book(tmp_path / "off.xlsx", [["n"], [1], [2], [3]])
    with pytest.raises(XlMcpError) as exc:
        cells_ops.query_range(p, offset=-5)
    assert "offset" in str(exc.value)
    out = cells_ops.query_range(p, offset=1)
    assert out["returned"] == 2


@pytest.mark.parametrize("target", ["CON", "con.csv", "NUL.json", "AUX",
                                    "COM1.csv", "lpt9.tsv", "con. "])
def test_reserved_device_names_refuse_as_output_targets(tmp_path, target):
    """A write to CON reports success and stores nothing, so an export that
    accepted it lost the data and said ok."""
    p = _book(tmp_path / "dev.xlsx", [["a", 1]])
    with pytest.raises(XlMcpError) as exc:
        dataio.export_range(p, out_file=str(tmp_path / target))
    msg = str(exc.value)
    assert "device" in msg.lower()
    assert target.split(".")[0].strip().rstrip(". ").upper() in msg


def test_ordinary_output_names_are_untouched_by_the_device_guard(tmp_path):
    """The guard matches whole stems, so a file that merely contains a
    device name keeps working."""
    p = _book(tmp_path / "dev2.xlsx", [["a", 1]])
    for name in ("console.csv", "control.csv", "auxiliary.json",
                 "com.csv", "prnt.csv"):
        assert outguard._device_name(tmp_path / name) is None
        dataio.export_range(p, out_file=str(tmp_path / name))


# ------------------------------------------------------------------- V1-9


def test_a_read_only_file_is_not_diagnosed_as_open_in_excel(tmp_path):
    """A read-only-attribute file refused with "is open in Excel", which
    sends the user to close a program they never opened. Files off old
    media, cloud restores, and locked shares all arrive this way."""
    import os
    import stat as _stat

    from xlsx_mcp.core.errors import WorkbookLocked

    p = _book(tmp_path / "ro.xlsx", [["a", 1]])
    os.chmod(p, _stat.S_IREAD)
    try:
        with pytest.raises(WorkbookLocked) as exc:
            cells_ops.set_cell(p, "A1", 2)
        msg = str(exc.value)
        assert "read-only" in msg
        assert "open in Excel" not in msg
        assert "No program is holding it" in msg, (
            "the refusal must say nobody has the file open; the old "
            "message sent users to close a program they never opened")
        assert "attribute" in msg
    finally:
        os.chmod(p, _stat.S_IWRITE)
    # and the same file writes once the attribute is cleared
    cells_ops.set_cell(p, "A1", 2)


def test_a_genuinely_held_file_still_says_so(tmp_path):
    """The new branch must not swallow the lock message it split off."""
    from xlsx_mcp.core.package import held_refusal, read_only_refusal

    assert "open in Excel" in held_refusal("b.xlsx")
    assert "read-only" in read_only_refusal("b.xlsx")
    assert "open in Excel" not in read_only_refusal("b.xlsx")


# ---------------------------------------------- the exit-less hazard gate


HAZARD_FIXTURE = "shape.xlsx"


def test_the_hazard_refusal_stops_naming_a_route_com_cannot_take(tmp_path):
    """The refusal offered "enable COM (com pack)" for a refused cell
    write. The com pack's eleven tools recalculate, drive pivots and goal
    seek, export PDF, render, convert, encrypt, set sparklines, autofit,
    and report status. None of them writes a cell, so a caller who
    enabled the pack as instructed found nothing there that could do the
    thing that was refused."""
    import shutil
    from pathlib import Path

    from xlsx_mcp.core.errors import HazardRefused

    src = Path(__file__).parent.parent / "fixtures" / "corpus" / HAZARD_FIXTURE
    if not src.exists():  # pragma: no cover
        pytest.skip(f"{HAZARD_FIXTURE} fixture not built")
    p = tmp_path / HAZARD_FIXTURE
    shutil.copy2(src, p)

    with pytest.raises(HazardRefused) as exc:
        cells_ops.set_cell(str(p), "A1", 1)
    msg = str(exc.value)
    detail = getattr(exc.value, "detail", {}) or {}

    assert "enable COM" not in msg and "enable the COM" not in msg
    assert "com pack" not in msg.lower(), msg
    # the two routes that exist, both of which can actually be taken
    assert "allow_loss=true" in msg, msg
    assert "in Excel itself" in msg, msg
    for route in detail.get("routes", []):
        assert "com pack" not in route.lower(), route
    # the exact costs, per class, are stated rather than named
    assert detail.get("losses"), "the refusal no longer states the losses"
    assert "A future release adds preservation" in msg, (
        "preservation is the real fix and the refusal says so")
    # and the allow_loss route still works, which is what makes it an exit
    out = cells_ops.set_cell(str(p), "A1", 1, allow_loss=True)
    assert out["changed"]


def test_the_loss_costs_come_from_the_hazard_table(tmp_path):
    from xlsx_mcp.core import hazard

    costs = hazard.loss_costs(["slicers", "media"])
    assert len(costs) == 2
    assert costs[0].startswith("slicers: ")
    assert "removed on save" in costs[0]
    assert hazard.loss_costs(["not-a-key"]) == []


# ------------------------------------------------------------------ V1-11


def test_render_sheet_refuses_by_name_when_the_session_has_no_clipboard(
        tmp_path, monkeypatch):
    """Excel's only range-to-bitmap route is CopyPicture, which goes
    through the Windows clipboard. Without one the user got Excel's raw
    "CopyPicture method of Range class failed" and no cause."""
    from xlsx_mcp.ops import comtier

    monkeypatch.setattr(comtier, "_clipboard_available", lambda: False)
    p = _book(tmp_path / "render.xlsx", [["a", 1]])
    with pytest.raises(XlMcpError) as exc:
        comtier.com_render_sheet(p, str(tmp_path / "out.png"))
    msg = str(exc.value)
    assert "clipboard" in msg
    assert "session does not have one" in msg, (
        "the refusal must name the SESSION as the cause; the old one "
        "was Excel's raw CopyPicture error, which named nothing")
    assert "service accounts" in msg, "the refusal names no cause"
    assert "export the range as values" in msg, (
        "the refusal names no working route; exporting values is the "
        "one that does not touch the clipboard")


def test_the_clipboard_probe_never_raises():
    """The probe answers about the environment and must not become a new
    failure mode of its own on a machine with no pywin32."""
    from xlsx_mcp.ops import comtier

    assert comtier._clipboard_available() in (True, False)


# ------------------------------------------------------------------ V1-10


class TestConstructedNamesStayWritable:
    """Every filename this server BUILDS is held inside both per-component
    limits, on every platform.

    Windows caps a component at 255 characters, ext4 and friends cap it at
    255 BYTES, and a Korean character is three bytes in UTF-8. A workbook
    name that is comfortable on Windows therefore produced an unwritable
    snapshot, .bak, or per-sheet export on Linux, and CI dodged it once by
    shortening the test's fixture name instead of bounding the product.
    Now that PyPI serves Linux and macOS, that is a live save bug.
    """

    LONG_KO = "매우" * 60 + "장부"        # 122 chars / 366 bytes

    def test_bound_name_leaves_ordinary_names_alone(self):
        from xlsx_mcp.core import safesave

        for name in ("book.xlsx", "장부.xlsx", "a" * 200 + ".xlsx"):
            assert safesave.bound_name(name) == name

    def test_bound_name_fits_both_ceilings(self):
        from xlsx_mcp.core import safesave

        out = safesave.bound_name(self.LONG_KO + ".xlsx")
        assert safesave.name_fits(out), out
        assert out.endswith(".xlsx"), "the extension was truncated"
        assert len(out.encode("utf-8")) <= safesave.MAX_NAME_BYTES

    def test_bound_name_honors_headroom(self):
        from xlsx_mcp.core import safesave

        out = safesave.bound_name(self.LONG_KO + ".xlsx", headroom=24)
        assert safesave.name_fits(out, headroom=24), out

    def test_two_long_names_sharing_a_prefix_do_not_collide(self):
        from xlsx_mcp.core import safesave

        a = safesave.bound_name(self.LONG_KO + "1.xlsx")
        b = safesave.bound_name(self.LONG_KO + "2.xlsx")
        assert a != b, "the hash suffix is not doing its job"

    def test_a_snapshot_of_a_long_named_workbook_is_writable(self, tmp_path):
        """The snapshot name adds a 14-character DTG and up to 61 more for a
        label, on top of a name that may already be near the ceiling."""
        from xlsx_mcp.core import safesave
        from xlsx_mcp.ops import backups as backups_ops

        p = _book(tmp_path / (self.LONG_KO[:60] + ".xlsx"), [["a", 1]])
        out = backups_ops.create_snapshot(p, label="x" * 60)
        import os
        name = os.path.basename(out["snapshot"])
        assert safesave.name_fits(name), name
        assert os.path.exists(out["snapshot"])

    def test_an_overwrite_bak_of_a_long_named_target_is_writable(
            self, tmp_path):
        from xlsx_mcp.core import safesave
        from xlsx_mcp.core.outguard import _timestamped_backup

        target = tmp_path / (self.LONG_KO[:75] + ".csv")
        target.write_text("a,b", encoding="utf-8")
        bak = _timestamped_backup(str(target))
        import os
        assert safesave.name_fits(os.path.basename(bak)), bak
        assert os.path.exists(bak)

    def test_a_per_sheet_export_of_a_long_named_workbook_is_writable(
            self, tmp_path):
        from xlsx_mcp.core import safesave

        p = tmp_path / (self.LONG_KO[:70] + ".xlsx")
        wb = openpyxl.Workbook()
        wb.active.title = "S1"
        wb.active.append(["a", 1])
        wb.create_sheet("두번째시트이름입니다").append(["b", 2])
        wb.save(p)
        wb.close()

        out_dir = tmp_path / "out"
        out_dir.mkdir()
        out = dataio.export_file(str(p), out_dir=str(out_dir), fmt="csv")
        import os
        files = out["written_to"]
        assert len(files) == 2, out
        for f in files:
            assert safesave.name_fits(os.path.basename(f)), f
            assert os.path.exists(f)
