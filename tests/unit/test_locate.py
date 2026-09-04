"""Unit tests for core/locate.py: the grid location resolver.

Every selector, the ambiguity/not-found refusals, the true-used-range
correctness vs openpyxl's dimension (the incumbent defect-11 regression), and
structured table refs.
"""

from __future__ import annotations

import openpyxl
import pytest
from openpyxl.worksheet.table import Table
from openpyxl.workbook.defined_name import DefinedName

from xlsx_mcp.core import locate
from xlsx_mcp.core.errors import (
    AmbiguousTarget,
    RangeOutOfBounds,
    TargetNotFound,
    UnsupportedStructure,
    XlMcpError,
)


def _wb():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "Name"
    ws["B1"] = "Amount"
    ws["A2"] = "x"
    ws["B2"] = 10
    ws["A3"] = "y"
    ws["B3"] = 20
    ws["D5"] = "Q3 Total"
    wb.create_sheet("Summary")
    wb["Summary"]["B9"] = "Q3 Total"
    return wb


def res(location, **kw):
    return locate.resolve_location(_wb(), location, **kw)


# ------------------------------------------------------------- a1 / cell / range


def test_cell_selector():
    g = res({"cell": "B7"})
    assert (g.sheet, g.min_row, g.min_col, g.max_row, g.max_col) == \
        ("Data", 7, 2, 7, 2)
    assert g.is_single and g.a1 == "B7"


def test_range_selector():
    g = res({"range": "A1:C10"})
    assert g.a1 == "A1:C10"
    assert (g.min_row, g.min_col, g.max_row, g.max_col) == (1, 1, 10, 3)


def test_a1_alias_and_bare_string():
    assert res({"a1": "B7"}).a1 == "B7"
    assert res("A1:B2").a1 == "A1:B2"


def test_sheet_sibling_and_embedded():
    assert res({"cell": "B9", "sheet": "Summary"}).sheet == "Summary"
    assert res({"a1": "Summary!B9"}).sheet == "Summary"


def test_sheet_conflict_refuses():
    with pytest.raises(XlMcpError):
        res({"a1": "Summary!B9", "sheet": "Data"})


def test_out_of_bounds_refuses():
    with pytest.raises(RangeOutOfBounds):
        res({"range": "A10:A1"})              # inverted
    with pytest.raises(XlMcpError):
        res({"cell": "notacell"})


# ------------------------------------------------------------- r1c1


def test_r1c1_single_and_range():
    assert res({"r1c1": "R7C2"}).a1 == "B7"
    g = res({"r1c1": "R1C1:R10C3"})
    assert g.a1 == "A1:C10"


def test_r1c1_rejects_relative_form():
    with pytest.raises(XlMcpError):
        res({"r1c1": "R[1]C2"})


# ------------------------------------------------------------- name


def test_name_selector_resolves_range():
    wb = _wb()
    wb.defined_names["Amounts"] = DefinedName("Amounts", attr_text="Data!$B$2:$B$3")
    g = locate.resolve_location(wb, {"name": "Amounts"})
    assert g.sheet == "Data" and g.a1 == "B2:B3"


def test_name_not_found():
    with pytest.raises(TargetNotFound):
        res({"name": "Nope"})


def test_name_scope_collision_is_ambiguous():
    wb = _wb()
    wb.defined_names["Total"] = DefinedName("Total", attr_text="Data!$B$2")
    wb["Data"].defined_names["Total"] = DefinedName("Total", attr_text="Data!$B$3")
    with pytest.raises(AmbiguousTarget) as ei:
        locate.resolve_location(wb, {"name": "Total"})
    assert ei.value.matches and len(ei.value.matches) == 2
    # a scope modifier disambiguates
    g = locate.resolve_location(wb, {"name": "Total", "scope": "Data"})
    assert g.a1 == "B3"


# ------------------------------------------------------------- table


def _table_wb():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "T"
    headers = ["Region", "Amount"]
    for c, h in enumerate(headers, 1):
        ws.cell(1, c, h)
    for r in range(2, 5):
        ws.cell(r, 1, "r%d" % r)
        ws.cell(r, 2, r * 5)
    ws.add_table(Table(displayName="Sales", ref="A1:B4"))
    return wb


def test_table_parts_and_column():
    wb = _table_wb()
    assert locate.resolve_location(wb, {"table": "Sales"}).a1 == "A1:B4"
    assert locate.resolve_location(
        wb, {"table": "Sales", "part": "headers"}).a1 == "A1:B1"
    assert locate.resolve_location(
        wb, {"table": "Sales", "part": "data"}).a1 == "A2:B4"
    assert locate.resolve_location(
        wb, {"table": "Sales", "column": "Amount", "part": "data"}).a1 == "B2:B4"


def test_table_unknown_column():
    with pytest.raises(TargetNotFound):
        locate.resolve_location(_table_wb(), {"table": "Sales", "column": "Nope"})


def test_table_not_found():
    with pytest.raises(TargetNotFound):
        locate.resolve_location(_table_wb(), {"table": "Ghost"})


# ------------------------------------------------------------- used_range


def test_true_used_range_ignores_formatted_empties():
    # data only in A1:C3, but a cell far away is formatted (no value). openpyxl
    # dimension over-reports; true_used_range must not.
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "U"
    for r in range(1, 4):
        for c in range(1, 4):
            ws.cell(r, c, r * c)
    ws["Z100"].font = ws["Z100"].font.copy(bold=True)  # touch, no value
    g = locate.resolve_location(wb, {"used_range": "U"})
    assert g.a1 == "A1:C3"
    # openpyxl's own dimension disagrees (instantiated cell Z100 pulls it out)
    assert ws.calculate_dimension() != "A1:C3"


def test_used_range_empty_sheet_is_explicit_empty():
    wb = openpyxl.Workbook()
    wb.active.title = "E"
    g = locate.resolve_location(wb, {"used_range": "E"})
    assert g.empty is True


def test_true_used_range_direct_helper():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["B2"] = 1
    ws["D5"] = 2
    assert locate.true_used_range(ws) == (2, 2, 5, 4)
    ws2 = wb.create_sheet("blank")
    assert locate.true_used_range(ws2) is None


# ------------------------------------------------------------- region


def test_region_grows_contiguous_island():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "I"
    for r in range(2, 5):
        for c in range(2, 4):
            ws.cell(r, c, "x")
    ws["Z1"] = "far away"        # separated by blanks, not part of the island
    g = locate.resolve_location(wb, {"region": {"near": "C3"}})
    assert g.a1 == "B2:C4"


# ------------------------------------------------------------- search


def test_search_single_match():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["D5"] = "Q3 Total"
    g = locate.resolve_location(wb, {"search": {"text": "Q3 Total"}})
    assert g.a1 == "D5"


def test_search_ambiguous_lists_every_match():
    with pytest.raises(AmbiguousTarget) as ei:
        res({"search": {"text": "Q3 Total"}})
    assert len(ei.value.matches) == 2
    sheets = {m["sheet"] for m in ei.value.matches}
    assert sheets == {"Data", "Summary"}


def test_search_occurrence_and_sheet_disambiguate():
    assert res({"search": {"text": "Q3 Total", "occurrence": 1}}).sheet == "Data"
    assert res({"search": {"text": "Q3 Total", "sheet": "Summary"}}).sheet \
        == "Summary"


def test_search_not_found_with_case_hint():
    with pytest.raises(TargetNotFound) as ei:
        res({"search": {"text": "q3 total", "match_case": True}})
    assert "ignoring case" in str(ei.value)


def test_search_contains_mode():
    g = res({"search": {"text": "Total", "match": "contains", "sheet": "Data"}})
    assert g.a1 == "D5"


# ------------------------------------------------------------- shape rules


def test_exactly_one_selector_required():
    with pytest.raises(XlMcpError):
        res({"cell": "A1", "range": "A1:B2"})
    with pytest.raises(XlMcpError):
        res({})


def test_unknown_key_refuses():
    with pytest.raises(XlMcpError):
        res({"cell": "A1", "bogus": 1})


def test_anchor_selector_round_trip_and_staleness():
    """The consolidation-phase anchor: make_anchor over a rectangle resolves
    back to the same rectangle while content is unchanged, refuses
    StaleAnchor on any content change or a deleted sheet, and a malformed
    token is BAD_PARAMS (XlMcpError), not a stale refusal."""
    from xlsx_mcp.core.errors import StaleAnchor

    wb = _wb()
    ws = wb["Data"]
    token = locate.make_anchor(ws, 1, 1, 3, 2)  # A1:B3
    grid = locate.resolve_location(wb, {"anchor": token})
    assert (grid.sheet, grid.a1, grid.selector) == ("Data", "A1:B3", "anchor")

    # a content change inside the rectangle goes stale
    ws["B2"] = 999
    with pytest.raises(StaleAnchor):
        locate.resolve_location(wb, {"anchor": token})

    # a change OUTSIDE the rectangle does not
    ws["B2"] = 10
    ws["Z99"] = "elsewhere"
    assert locate.resolve_location(wb, {"anchor": token}).a1 == "A1:B3"

    # single-cell anchors use the single-cell A1 form
    tok1 = locate.make_anchor(ws, 5, 4, 5, 4)
    assert locate.resolve_location(wb, {"anchor": tok1}).a1 == "D5"

    # deleted sheet -> stale, named clearly
    wb2 = _wb()
    tok2 = locate.make_anchor(wb2["Summary"], 9, 2, 9, 2)
    del wb2["Summary"]
    with pytest.raises(StaleAnchor):
        locate.resolve_location(wb2, {"anchor": tok2})

    # malformed tokens refuse as BAD_PARAMS (plain XlMcpError, not a stale
    # refusal), never resolve
    for bogus in ("a3f9", "gv1:only-two-parts", "gv1:!!!:A1:00", 7, None,
                  "gv1::A1:deadbeef00"):
        with pytest.raises(XlMcpError) as exc:
            locate.resolve_location(wb, {"anchor": bogus})
        assert not isinstance(exc.value, StaleAnchor), bogus
