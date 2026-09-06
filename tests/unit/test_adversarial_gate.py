"""Regressions from the adversarial + fidelity gate.

One test per confirmed finding, named for the defect it locks down. The
attacks ran through the real tool surface (an in-process fastmcp client) on
scratch copies; these tests reproduce each one at the layer that owns the fix.

Findings covered here:
  anchor      crafted-rectangle DoS, rectangle-widening forgery, malformed
              token classified as staleness
  package     a concurrent writer clobbered between open and save; a verify
              pass that RAISED skipped the post-promote backup restore
  verify      allow_loss excusing fragile families the caller was never
              warned about
  envelope    corrupt packages and unopenable paths leaving the envelope;
              fastmcp schema failures leaving the envelope
  ops         sort over a merge dying half-way with an openpyxl message;
              defined names Excel refuses to open the workbook for; group_by
              silently ignored; sheet add inventing a name; a list validation
              with no source; a ragged write block
  hazard      part-name spellings that walked past the scan
  fidelity    x14 (in-part) extension blocks openpyxl discards in silence,
              now detected by the scan and REFUSED as drop-class
"""

from __future__ import annotations

import asyncio
import shutil
import time
import zipfile
from pathlib import Path

import openpyxl
import pytest

from fastmcp import Client

from xlsx_mcp import server
from xlsx_mcp.core import hazard, locate, verify
from xlsx_mcp.core.errors import (
    HazardRefused,
    RangeOutOfBounds,
    StaleAnchor,
    UnsupportedStructure,
    ValidationFailed,
    XlMcpError,
)
from xlsx_mcp.core.package import WorkbookPackage
from xlsx_mcp.ops import cells as _cells
from xlsx_mcp.ops import datavalidation as _dv
from xlsx_mcp.ops import gridio as _gridio
from xlsx_mcp.ops import lifecycle as _lifecycle
from xlsx_mcp.ops import names as _names
from xlsx_mcp.ops import sortfilter as _sortfilter
from xlsx_mcp.ops import view as _view

CORPUS = Path(__file__).resolve().parents[1] / "fixtures" / "corpus"


@pytest.fixture
def book(tmp_path):
    p = str(tmp_path / "adv.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Name", "Qty"])
    for n, q in (("b", 3), ("a", 1), ("c", 2)):
        ws.append([n, q])
    wb.save(p)
    wb.close()
    return p


# ----------------------------------------------------------------- anchors


def test_crafted_anchor_rectangle_refuses_instead_of_hanging(book):
    """A hand-widened token naming the whole grid used to send the
    fingerprint on a 17.2-billion-cell walk (20-40 CPU-minutes) inside a
    single-threaded server. It must refuse, and refuse fast."""
    token = _view.get_grid_view(book, sheet="Data")["anchor"]
    tag, b64, rest = token.split(":", 2)
    _a1, fp = rest.rsplit(":", 1)
    crafted = f"{tag}:{b64}:A1:XFD1048576:{fp}"
    t0 = time.perf_counter()
    with pytest.raises(RangeOutOfBounds) as exc:
        _cells.read_range(book, {"anchor": crafted})
    assert time.perf_counter() - t0 < 5.0
    assert "anchor ceiling" in str(exc.value)


def test_anchor_cannot_be_widened_over_empty_space(book):
    """The fingerprint now binds the bounds. Before, an A1:B4 token
    re-pointed at A1:Z2000 still verified (the extra cells were empty and
    empty cells contribute nothing to the digest), so an edit could land on
    cells the view never showed."""
    token = _view.get_grid_view(book, sheet="Data")["anchor"]
    tag, b64, rest = token.split(":", 2)
    a1_ref, fp = rest.rsplit(":", 1)
    assert a1_ref == "A1:B4"
    widened = f"{tag}:{b64}:A1:Z2000:{fp}"
    with pytest.raises(StaleAnchor):
        _cells.read_range(book, {"anchor": widened})
    # the untouched token still resolves
    assert _cells.read_range(book, {"anchor": token})["range"] == "A1:B4"


def test_malformed_anchor_is_bad_params_not_staleness(book):
    """A truncated token (bounds eaten into the fingerprint slot) used to
    answer STALE_ANCHOR, sending the caller hunting for a content change
    that never happened."""
    token = _view.get_grid_view(book, sheet="Data")["anchor"]
    tag, b64, rest = token.split(":", 2)
    a1_ref, _fp = rest.rsplit(":", 1)
    truncated = f"{tag}:{b64}:{a1_ref}"
    with pytest.raises(XlMcpError) as exc:
        _cells.read_range(book, {"anchor": truncated})
    assert not isinstance(exc.value, StaleAnchor)
    assert "get_grid_view" in str(exc.value)


# ------------------------------------------------------- concurrent writer


def test_concurrent_write_between_open_and_save_refuses(book):
    """The write lock serializes this server's writers; it cannot see Excel
    or another tool writing between the read and the save. That write used to
    be silently overwritten."""
    pkg = WorkbookPackage.open(book)
    pkg.set_cell("Data", "A6", "server edit")
    _ = pkg.workbook
    wb = openpyxl.load_workbook(book)
    wb["Data"]["D1"] = "OTHER PROCESS"
    wb.save(book)
    wb.close()

    with pytest.raises(XlMcpError) as exc:
        pkg.save()
    assert getattr(exc.value, "code", None) == "CONFLICT"
    check = openpyxl.load_workbook(book)
    assert check["Data"]["D1"].value == "OTHER PROCESS"
    assert check["Data"]["A6"].value is None
    check.close()


def test_second_save_on_the_same_package_still_works(book):
    """The staleness stamp must re-baseline after a promotion, or a package
    reused for two saves would refuse its own previous write."""
    pkg = WorkbookPackage.open(book)
    pkg.set_cell("Data", "A6", "one")
    pkg.save()
    pkg.set_cell("Data", "A7", "two")
    out = pkg.save()
    assert out["saved"] is True


# ------------------------------------------------------------- verify gate


def test_verify_failure_by_exception_is_a_failure_not_an_escape(tmp_path):
    """A produced package that makes the READER throw (a dangling chart
    relationship raised KeyError out of content_readback) must count as a
    failed verify. In the post-promote position the raise skipped the
    restore-from-backup branch entirely."""
    src = tmp_path / "chart.xlsx"
    if not (CORPUS / "chart.xlsx").exists():
        pytest.skip("corpus not built")
    shutil.copy2(CORPUS / "chart.xlsx", src)

    def saver(tmp):
        pkg._default_saver(tmp)
        with zipfile.ZipFile(tmp) as zin:
            items = [(i.filename, zin.read(i.filename))
                     for i in zin.infolist()]
        with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zo:
            for name, data in items:
                if name.startswith("xl/charts/"):
                    continue
                zo.writestr(name, data)

    pkg = WorkbookPackage.open(str(src))
    pkg.set_cell(None, "H1", "probe")
    before = src.read_bytes()
    with pytest.raises(ValidationFailed):
        pkg.save(allow_loss=True, saver=saver)
    assert src.read_bytes() == before


def test_content_readback_reports_an_unreadable_package(tmp_path):
    junk = tmp_path / "junk.xlsx"
    junk.write_bytes(b"not a zip at all")
    ok, mismatches = verify.content_readback(
        str(junk), {("S", "A1"): ("value", 1)})
    assert ok is False and mismatches


def test_allow_loss_does_not_excuse_an_unwarned_fragile_family(tmp_path):
    """allow_loss is consent to lose the families the gate NAMED, not a
    blanket amnesty: a chart part (warned only as 'may degrade') vanishing
    whole must still fail the write."""
    if not (CORPUS / "chart.xlsx").exists():
        pytest.skip("corpus not built")
    src = tmp_path / "chart.xlsx"
    shutil.copy2(CORPUS / "chart.xlsx", src)
    with zipfile.ZipFile(src) as z:
        pre = list(z.namelist())
        sizes = {i.filename: i.file_size for i in z.infolist()}
    stripped = tmp_path / "stripped.xlsx"
    with zipfile.ZipFile(src) as zin, \
            zipfile.ZipFile(stripped, "w", zipfile.ZIP_DEFLATED) as zo:
        for i in zin.infolist():
            if i.filename.startswith("xl/charts/"):
                continue
            zo.writestr(i.filename, zin.read(i.filename))

    # warned about slicers only: a chart loss is NOT covered
    ok, lost = verify.part_loss_check(
        pre, str(stripped), allow_loss=True, loss_keys={"slicers"})
    assert ok is False and any("chart" in x for x in lost)
    # warned about charts: the same loss is accepted
    ok2, _ = verify.part_loss_check(
        pre, str(stripped), allow_loss=True, loss_keys={"charts"})
    assert ok2 is True
    # and the un-scoped call keeps the old blanket behavior for callers that
    # cannot know the warned set
    ok3, _ = verify.part_loss_check(pre, str(stripped), allow_loss=True)
    assert ok3 is True
    assert sizes  # sanity: the fixture had content


# --------------------------------------------------------------- envelope


def test_corrupt_and_unopenable_paths_refuse_in_envelope(tmp_path):
    junk = tmp_path / "junk.xlsx"
    junk.write_bytes(b"PK\x03\x04not-a-zip")
    with pytest.raises(XlMcpError):
        _gridio.open_wb(str(junk))
    adir = tmp_path / "adir.xlsx"
    adir.mkdir()
    with pytest.raises(XlMcpError) as exc:
        _gridio.open_wb(str(adir))
    assert "directory" in str(exc.value)

    from xlsx_mcp import envelope
    assert envelope.classify(zipfile.BadZipFile("x")) == "UNSUPPORTED_CONTENT"
    assert envelope.classify(PermissionError(13, "denied")) == \
        "WORKBOOK_LOCKED"
    assert envelope.classify(OSError("device")) == "BAD_PARAMS"
    for _etype, code in envelope.CODE_MAP:
        assert code in envelope.CLOSED_CODES


def test_schema_validation_refusal_arrives_in_the_envelope(book):
    """fastmcp validates arguments before the tool body (and so before the
    boundary wrapper): unknown keywords and wrong types used to come back as
    a bare pydantic dump with no code and no hint."""

    async def run():
        async with Client(server.mcp) as c:
            bad_kw = await c.call_tool(
                "read_range",
                {"path": book, "location": {"cell": "A1"}, "nope": 1},
                raise_on_error=False)
            bad_type = await c.call_tool(
                "query_range",
                {"path": book, "where": {"column": "Qty"}},
                raise_on_error=False)
            return bad_kw, bad_type

    for res in asyncio.run(run()):
        assert res.is_error
        payload = res.structured_content
        assert payload["ok"] is False
        assert payload["error"]["code"] == "BAD_PARAMS"
        assert payload["error"]["hint"]
        assert "Traceback" not in payload["error"]["message"]


# ------------------------------------------------------------------- ops


def test_sort_over_a_merge_refuses_before_writing_anything(book):
    _cells.set_cell(book, {"cell": "A2", "sheet": "Data"}, "b")
    from xlsx_mcp.ops import structure  # noqa: F401  (import graph sanity)
    wb = openpyxl.load_workbook(book)
    wb["Data"].merge_cells("A2:A3")
    wb.save(book)
    wb.close()
    before = Path(book).read_bytes()
    with pytest.raises(UnsupportedStructure) as exc:
        _sortfilter.sort_range(book, {"range": "A1:B4", "sheet": "Data"},
                               [{"column": "Name", "order": "asc"}])
    assert "A2:A3" in str(exc.value) and "unmerge" in str(exc.value)
    assert Path(book).read_bytes() == before


@pytest.mark.parametrize("name", ["1bad", "has space", "a-b", "A1", "R1C1",
                                  "R", "=x"])
def test_invalid_defined_names_refuse(book, name):
    """Excel REFUSES TO OPEN a workbook carrying an invalid defined name
    (confirmed against Excel itself in the fidelity pass), and no part-level
    verify can see it."""
    with pytest.raises(XlMcpError):
        _names.manage_name(book, "add", name=name, refers_to="Data!$A$1")


@pytest.mark.parametrize("refers_to", ["(((", "Data!$A$1+", 'SUM("a',
                                       "+++", "'Data!$A$1"])
def test_invalid_refers_to_refuses(book, refers_to):
    with pytest.raises(XlMcpError):
        _names.manage_name(book, "add", name="Ok", refers_to=refers_to)


def test_valid_names_still_pass(book):
    out = _names.manage_name(book, "add", name="Good.Name_1",
                             refers_to="Data!$A$1:$B$4")
    assert out["saved"] is True
    out = _names.manage_name(book, "add", name="Calc",
                             refers_to="=SUM(Data!$B$2:$B$4)")
    assert out["saved"] is True


def test_group_by_without_aggregate_refuses(book):
    with pytest.raises(XlMcpError) as exc:
        _cells.query_range(book, sheet="Data", header=True, group_by="Name")
    assert "aggregate" in str(exc.value)
    # the aggregate form still works and still validates the column
    out = _cells.query_range(book, sheet="Data", header=True,
                             group_by="Name",
                             aggregate=[{"func": "sum", "column": "Qty"}])
    assert out["mode"] == "aggregate"
    with pytest.raises(XlMcpError):
        _cells.query_range(book, sheet="Data", header=True, group_by="Ghost",
                           aggregate=[{"func": "count"}])


def test_worksheet_add_requires_a_name(book):
    with pytest.raises(XlMcpError) as exc:
        _lifecycle.manage_worksheet(book, "add")
    assert "new_name" in str(exc.value)
    assert _lifecycle.manage_worksheet(
        book, "add", new_name="Named")["saved"] is True


def test_list_validation_without_values_refuses(book):
    with pytest.raises(XlMcpError) as exc:
        _dv.manage_data_validation(
            book, "add", location={"range": "A1:A3", "sheet": "Data"},
            dv_type="list")
    assert "values" in str(exc.value)


def test_ragged_write_block_refuses(book):
    with pytest.raises(XlMcpError) as exc:
        _cells.write_range(book, {"cell": "D1", "sheet": "Data"},
                           [[1, 2], [3]])
    assert "rectangular" in str(exc.value)
    with pytest.raises(XlMcpError) as exc2:
        _view.apply_edits(book, [{"op": "write_range",
                                 "location": {"cell": "D1",
                                              "sheet": "Data"},
                                 "data": [[1, 2], [3]]}])
    assert "rectangular" in str(exc2.value)


# ---------------------------------------------------------------- hazard


@pytest.mark.parametrize("member", [
    "./xl/slicers/slicer1.xml",
    "/xl/slicers/slicer1.xml",
    "xl//slicers/slicer1.xml",
    "xl\\slicers\\slicer1.xml",
    "./customXml/item1.xml",
])
def test_part_name_spellings_do_not_evade_the_scan(member):
    rep = hazard.scan_names(["xl/workbook.xml", member])
    assert rep.hazards, f"{member} evaded the hazard scan"
    assert rep.parts == ["xl/workbook.xml", member]  # real names preserved


# --------------------------------------------------------------- fidelity


@pytest.mark.skipif(not (CORPUS / "x14_condformat.xlsx").exists(),
                    reason="COM fixture corpus not built")
def test_x14_extension_loss_refuses(tmp_path):
    """THE FIDELITY-GATE FINDING, now drop-class. An Excel-authored data bar /
    icon set puts its real rule in the worksheet's x14 extLst. The part
    survives the save at full size, so the part-inventory verify sees nothing,
    yet openpyxl discards the rule: the rules VANISH. Under the ratified
    package policy that is drop-risk, not degrade-risk, so the mutation
    refuses rather than warning and proceeding. The scan reads the worksheet
    extLst, so it no longer reports the workbook CLEAN either."""
    src = tmp_path / "x14.xlsx"
    shutil.copy2(CORPUS / "x14_condformat.xlsx", src)
    before = src.read_bytes()
    rep = hazard.scan_path(str(src))
    assert not rep.clean
    assert hazard.IN_PART_EXT_KEY in rep.lossy_keys

    pkg = WorkbookPackage.open(str(src))
    pkg.set_cell(None, "D1", "probe")
    with pytest.raises(HazardRefused) as exc:
        pkg.save()
    msg = str(exc.value)
    assert "x14 conditional formatting" in msg   # names the rules at risk
    # names both outs, and the second one is Excel now: the com pack writes
    # no cells, so it was never an out for a refused cell write.
    assert "allow_loss=true" in msg and "in Excel itself" in msg
    assert src.read_bytes() == before            # refusal touched nothing


@pytest.mark.skipif(not (CORPUS / "x14_condformat.xlsx").exists(),
                    reason="COM fixture corpus not built")
def test_x14_extension_loss_is_announced_under_allow_loss(tmp_path):
    """allow_loss is the sanctioned way through, and it still has to say what
    it destroyed: the itemized announcement from openpyxl's own load-time
    warnings survives the policy flip."""
    src = tmp_path / "x14.xlsx"
    shutil.copy2(CORPUS / "x14_condformat.xlsx", src)
    pkg = WorkbookPackage.open(str(src))
    pkg.set_cell(None, "D1", "probe")
    out = pkg.save(allow_loss=True)
    joined = " ".join(out["warnings"])
    assert "NOT in the saved file" in joined
    assert "Conditional Formatting" in joined
    assert "allow_loss: proceeding despite" in joined
    assert openpyxl.load_workbook(src)["CF"]["D1"].value == "probe"


@pytest.mark.skipif(not CORPUS.exists(), reason="corpus not built")
def test_corpus_fixtures_carry_no_authoring_identity():
    """COM-authored fixtures are written BY EXCEL, which stamps the signed-in
    user into docProps, the threaded-comment persons registry, and absolute
    paths. The corpus lives in the repo, so a rebuild used to commit a real
    name and local paths into a deliberately pseudonymized project;
    build_corpus.scrub_corpus() now runs after every build."""
    import os
    needles = [n for n in (os.environ.get("USERNAME"),
                           os.environ.get("USERPROFILE"),
                           str(Path.home())) if n]
    leaks = []
    for fixture in sorted(CORPUS.glob("*.xls*")):
        with zipfile.ZipFile(fixture) as z:
            for member in z.namelist():
                blob = z.read(member).decode("utf-8", "ignore").lower()
                for needle in needles:
                    if needle.lower() in blob:
                        leaks.append(f"{fixture.name}:{member} -> {needle}")
    assert not leaks, f"authoring identity left in the corpus: {leaks}"


def test_clean_workbook_carries_no_extension_warning(book):
    pkg = WorkbookPackage.open(book)
    pkg.set_cell("Data", "D1", "probe")
    out = pkg.save()
    assert not [w for w in out["warnings"] if "extension" in w]
