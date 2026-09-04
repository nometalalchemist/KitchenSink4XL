"""The grid-view anchor end-to-end: view-stable addressing through the ops.

get_grid_view returns one self-contained token for the rectangle it showed;
{"anchor": token} resolves in every positional tool to exactly that
rectangle while its content is unchanged and refuses STALE_ANCHOR after any
content change inside it. The token is stateless (sheet + A1 bounds + a
content fingerprint), so it works across processes and server restarts, and
the fingerprint lives in the raw (formula-string) value space, so reads that
resolve against cached loads re-open a formula view (gridio.resolve's path
escape hatch) rather than false-staling on formula cells.
"""

from __future__ import annotations

import openpyxl
import pytest

from xlsx_mcp.core.errors import StaleAnchor
from xlsx_mcp.ops import cells as _cells
from xlsx_mcp.ops import dataio as _dataio
from xlsx_mcp.ops import view as _view


@pytest.fixture
def book(tmp_path):
    p = str(tmp_path / "anchored.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Name", "Amount", "Total"])
    ws.append(["x", 10, "=B2*2"])
    ws.append(["y", 20, "=B3*2"])
    wb.create_sheet("Empty")
    wb.save(p)
    return p


def test_view_returns_anchor_and_reads_resolve_it(book):
    view = _view.get_grid_view(book, sheet="Data")
    token = view["anchor"]
    assert token.startswith("gv1:")

    # read_range values='cached' resolves against a data_only load; the
    # path escape hatch keeps the fingerprint in the formula value space.
    got = _cells.read_range(book, {"anchor": token}, values="cached")
    assert got["sheet"] == "Data"
    assert got["range"] == view["used_range"] == "A1:C3"

    # query_range and export_range accept the same token.
    q = _cells.query_range(book, location={"anchor": token}, header=True)
    assert q["matched"] == 2
    ex = _dataio.export_range(book, location={"anchor": token}, fmt="csv")
    assert ex["range"] == "A1:C3"


def test_anchor_goes_stale_after_content_change(book):
    token = _view.get_grid_view(book, sheet="Data")["anchor"]
    _cells.set_cell(book, {"cell": "B2", "sheet": "Data"}, 999)
    with pytest.raises(StaleAnchor):
        _cells.read_range(book, {"anchor": token})
    # a fresh view issues a fresh, working anchor
    fresh = _view.get_grid_view(book, sheet="Data")["anchor"]
    assert fresh != token
    assert _cells.read_range(book, {"anchor": fresh})["range"] == "A1:C3"


def test_change_outside_anchor_keeps_it_fresh(book):
    token = _view.get_grid_view(book, sheet="Data")["anchor"]
    _cells.set_cell(book, {"cell": "H20", "sheet": "Data"}, "far away")
    assert _cells.read_range(book, {"anchor": token})["range"] == "A1:C3"


def test_apply_edits_accepts_anchor_then_staleness_guards(book):
    token = _view.get_grid_view(book, sheet="Data")["anchor"]
    out = _view.apply_edits(book, [
        {"op": "write_range", "location": {"anchor": token},
         "data": [["Name", "Qty", "Total"]]},
    ])
    assert out["changed"]["edits_applied"] == 1
    # the batch's own write changed the anchored region; the token is now
    # stale for a second batch, exactly the edit-what-you-see contract
    with pytest.raises(StaleAnchor):
        _view.apply_edits(book, [
            {"op": "set_value", "location": {"anchor": token}, "value": 1},
        ])


def test_anchor_covers_shown_rectangle_when_truncated(book):
    view = _view.get_grid_view(book, sheet="Data", max_rows=2, max_cols=2)
    assert view["truncated"]["rows"] and view["truncated"]["cols"]
    grid = _cells.read_range(book, {"anchor": view["anchor"]})
    assert grid["range"] == "A1:B2"
    # a change OUTSIDE the shown rectangle (row 3) leaves it fresh
    _cells.set_cell(book, {"cell": "C3", "sheet": "Data"}, 0)
    assert _cells.read_range(
        book, {"anchor": view["anchor"]})["range"] == "A1:B2"


def test_empty_sheet_view_has_no_anchor(book):
    view = _view.get_grid_view(book, sheet="Empty")
    assert view["empty"] is True
    assert "anchor" not in view


def test_anchor_survives_process_boundaries_statelessly(book):
    """No server-side table: a token built from a fresh load (a different
    'process') still resolves."""
    wb = openpyxl.load_workbook(book, data_only=False)
    from xlsx_mcp.core import locate
    token = locate.make_anchor(wb["Data"], 1, 1, 3, 3)
    wb.close()
    assert _cells.read_range(book, {"anchor": token})["range"] == "A1:C3"
