"""modify_grid_structure: the structural-edit flagship over core.refs.

The centerpiece is the end-to-end: insert rows above a region holding a
formula, a table, a conditional format, a data validation, a merge, and a
defined name, then reload the SAVED file and verify every reference shifted
coherently and the save verified. Plus the whole-span semantics the re-audit
fixed (=SUM(B:B) shifting on column edits), #REF! reporting on deletes, the
`at` argument's accepted shapes, and the refusal paths (bad action/count,
axis mismatch, out-of-grid, Excel's push-off-the-edge insert guard).
"""

from __future__ import annotations

import openpyxl
import pytest

from xlsx_mcp.core.errors import RangeOutOfBounds, TargetNotFound, XlMcpError
from xlsx_mcp.ops import (
    cells as _cells,
    condformat as _condformat,
    datavalidation as _datavalidation,
    names as _names,
    structure as _structure,
    tables as _tables,
)


def _make(path, rows, sheet="Data"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    for i, row in enumerate(rows, start=1):
        for j, v in enumerate(row, start=1):
            ws.cell(i, j, v)
    wb.save(path)
    wb.close()
    return str(path)


# --------------------------------------------------------------- end-to-end


def test_insert_rows_shifts_everything_coherently(tmp_path):
    """Insert 2 rows above a formula + table + CF + DV + merge + name region;
    the saved file must reload with every reference shifted."""
    p = _make(tmp_path / "e2e.xlsx", [
        ["h", None, "x", "y", "e", "m"],
        [1, None, 10, 20, 5, "note"],
        [2, None, 30, 40, 6, None],
        ["=SUM(A2:A3)"],
    ])
    _tables.create_table(p, {"range": "C1:D3"}, "T")
    _condformat.manage_conditional_format(
        p, "add", location={"range": "A2:A3"}, cf_type="cell_is",
        params={"operator": "greaterThan", "formula": 1})
    _datavalidation.manage_data_validation(
        p, "add", location={"range": "E2:E3"}, dv_type="whole",
        operator="greaterThan", formula1=0)
    _names.manage_name(p, "add", name="Box", refers_to="Data!$A$2:$A$3")
    _cells.set_merge(p, "merge", location={"range": "F2:G2"},
                     confirm_data_loss=False)

    out = _structure.modify_grid_structure(p, "insert_rows", at=2, count=2)
    assert out["ok"] is True
    assert out["verified"] is True
    assert out["backup"] == "prev"
    rw = out["changed"]["structure"]["rewrites"]
    assert rw["formulas"] >= 1
    assert rw["names"] >= 1
    assert rw["conditional_formats"] >= 1
    assert rw["data_validations"] >= 1
    assert rw["tables"] >= 1
    assert rw["merges"] >= 1
    assert rw["ref_errors"] == 0

    wb = openpyxl.load_workbook(p)
    ws = wb["Data"]
    # values moved down with the insert
    assert ws["A4"].value == 1 and ws["A5"].value == 2
    # the formula moved AND its range was rewritten
    assert ws["A6"].value == "=SUM(A4:A5)"
    # table ref grew with the shift
    assert ws.tables["T"].ref == "C1:D5"
    # conditional format range shifted
    cf_ranges = [str(r) for cf in ws.conditional_formatting
                 for r in cf.sqref.ranges]
    assert "A4:A5" in cf_ranges
    # data validation shifted
    dv_ranges = [str(r) for dv in ws.data_validations.dataValidation
                 for r in dv.sqref.ranges]
    assert "E4:E5" in dv_ranges
    # merge shifted
    assert "F4:G4" in {str(r) for r in ws.merged_cells.ranges}
    # defined name shifted
    assert "$A$4:$A$5" in str(dict(wb.defined_names)["Box"].value)
    wb.close()


# ---------------------------------------------------------- span semantics


def test_whole_column_span_shifts_on_column_insert(tmp_path):
    p = _make(tmp_path / "span.xlsx", [[None, 1], [None, 2]])
    wb = openpyxl.load_workbook(p)
    wb["Data"]["A1"] = "=SUM(B:B)"
    wb.save(p)
    wb.close()
    _structure.modify_grid_structure(p, "insert_cols", at="B", count=1)
    wb = openpyxl.load_workbook(p)
    assert wb["Data"]["A1"].value == "=SUM(C:C)"
    wb.close()


def test_row_insert_leaves_column_span_alone(tmp_path):
    p = _make(tmp_path / "span2.xlsx", [["=SUM(B:B)", 1], [None, 2]])
    _structure.modify_grid_structure(p, "insert_rows", at=1, count=3)
    wb = openpyxl.load_workbook(p)
    assert wb["Data"]["A4"].value == "=SUM(B:B)"
    wb.close()


# ------------------------------------------------------------- delete + REF


def test_delete_rows_reports_ref_errors_and_warns(tmp_path):
    p = _make(tmp_path / "ref.xlsx",
              [["h"], [1], [2], [None], ["=SUM(A2:A3)"]])
    out = _structure.modify_grid_structure(p, "delete_rows", at=2, count=2)
    assert out["changed"]["structure"]["rewrites"]["ref_errors"] == 1
    assert any("#REF!" in w for w in out["warnings"])
    wb = openpyxl.load_workbook(p)
    assert wb["Data"]["A3"].value == "=SUM(#REF!)"
    wb.close()


def test_multi_column_delete_shifts_values_and_formulas(tmp_path):
    p = _make(tmp_path / "cols.xlsx",
              [[1, 2, 3, 4, "=D1*2"]])
    out = _structure.modify_grid_structure(p, "delete_cols", at="B", count=2)
    assert out["ok"] is True
    wb = openpyxl.load_workbook(p)
    ws = wb["Data"]
    assert ws["A1"].value == 1 and ws["B1"].value == 4
    assert ws["C1"].value == "=B1*2"
    wb.close()


# ------------------------------------------------------------ 'at' shapes


def test_at_accepts_cell_string_and_location_object(tmp_path):
    p = _make(tmp_path / "at.xlsx", [["h"], [1], [2]])
    _structure.modify_grid_structure(p, "insert_rows", at="A2")
    wb = openpyxl.load_workbook(p)
    assert wb["Data"]["A3"].value == 1
    wb.close()
    _structure.modify_grid_structure(p, "insert_rows",
                                     at={"search": {"text": "h"}})
    wb = openpyxl.load_workbook(p)
    assert wb["Data"]["A2"].value == "h"
    wb.close()


# ---------------------------------------------------------------- refusals


def test_bad_action_and_count_refuse(tmp_path):
    p = _make(tmp_path / "bad.xlsx", [[1]])
    with pytest.raises(XlMcpError, match="action"):
        _structure.modify_grid_structure(p, "insert_row", at=1)
    with pytest.raises(XlMcpError, match="count"):
        _structure.modify_grid_structure(p, "insert_rows", at=1, count=0)
    with pytest.raises(XlMcpError, match="count"):
        _structure.modify_grid_structure(p, "insert_rows", at=1, count=True)


def test_column_letter_for_row_action_refuses(tmp_path):
    p = _make(tmp_path / "axis.xlsx", [[1]])
    with pytest.raises(XlMcpError, match="works on rows"):
        _structure.modify_grid_structure(p, "delete_rows", at="B")


def test_out_of_grid_positions_refuse(tmp_path):
    p = _make(tmp_path / "oob.xlsx", [[1]])
    with pytest.raises(RangeOutOfBounds):
        _structure.modify_grid_structure(p, "insert_rows", at=0)
    with pytest.raises(RangeOutOfBounds):
        _structure.modify_grid_structure(p, "insert_cols", at=20_000)
    with pytest.raises(RangeOutOfBounds):
        _structure.modify_grid_structure(
            p, "delete_rows", at=1_048_570, count=100)


def test_insert_pushing_data_off_grid_refuses(tmp_path):
    """Excel's own guard: an insert that would push value-bearing cells past
    the grid edge refuses instead of truncating them."""
    p = str(tmp_path / "edge.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.cell(1, 1, "top")
    ws.cell(1_048_576, 1, "bottom")
    wb.save(p)
    wb.close()
    with pytest.raises(RangeOutOfBounds, match="push"):
        _structure.modify_grid_structure(p, "insert_rows", at=1)
    # the file is untouched
    wb = openpyxl.load_workbook(p)
    assert wb["Data"]["A1"].value == "top"
    wb.close()


def test_missing_sheet_refuses_with_locate_hints(tmp_path):
    p = _make(tmp_path / "sheet.xlsx", [[1]])
    with pytest.raises(TargetNotFound, match="data"):
        _structure.modify_grid_structure(p, "insert_rows", at=1, sheet="data")


def test_hazard_gate_inherited_from_the_package(tmp_path):
    """A structural edit of a slicer-bearing workbook refuses like every
    other mutation (the WorkbookPackage pipeline is not bypassed)."""
    import shutil
    from pathlib import Path
    from xlsx_mcp.core.errors import HazardRefused
    corpus = Path(__file__).resolve().parents[1] / "fixtures" / "corpus"
    p = str(tmp_path / "sl.xlsx")
    shutil.copy(corpus / "pivot_slicer.xlsx", p)
    before = open(p, "rb").read()
    with pytest.raises(HazardRefused):
        _structure.modify_grid_structure(p, "insert_rows", at=1)
    assert open(p, "rb").read() == before
