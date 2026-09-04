"""Standing regressions for the INSANE-MODE ADVERSARIAL ROUND (2026-09-05).

One test per finding, each carrying the round's EXACT payload, so none of the
twenty can come back silently. The round's headline was a single theme with
ten routes: the safety core would write a workbook Excel refuses to open,
report ok/saved/verified, and then tell you the file was fine.

Expected values are Excel's own, taken from the COM ground-truth probes run
for this fix wave (recorded next to each constant in core/calc.py,
core/arrays.py and core/limits.py), not from what the code happens to do.
"""

from __future__ import annotations

import json
import zipfile
from pathlib import Path

import openpyxl
import pytest
from openpyxl.worksheet.formula import ArrayFormula

from xlsx_mcp import envelope as _envelope
from xlsx_mcp.core import arrays as core_arrays
from xlsx_mcp.core import calc as core_calc
from xlsx_mcp.core import limits as core_limits
from xlsx_mcp.core import safesave as core_safesave
from xlsx_mcp.core.errors import (
    ExcelWouldRefuse,
    FormulaRejected,
    UnsupportedStructure,
)
from xlsx_mcp.ops import (
    annotations,
    cells,
    datavalidation,
    formulas,
    names,
    objects,
    pagelayout,
    search,
    sortfilter,
    structure,
    validation,
)


# --------------------------------------------------------------- fixtures

def _book(tmp_path, name="t.xlsx", sheets=("Data",)):
    p = tmp_path / name
    wb = openpyxl.Workbook()
    wb.active.title = sheets[0]
    for extra in sheets[1:]:
        wb.create_sheet(extra)
    for i in range(1, 4):
        wb[sheets[0]].cell(i, 1).value = i
    wb.save(p)
    return str(p)


def _cse_book(tmp_path, name="cse.xlsx"):
    """The round's legacy Ctrl+Shift+Enter corpus shape: a multi-cell array
    over B1:B5 and a single-cell one at C1, with no xl/metadata.xml, which is
    why the dynamic-array hazard gate never saw it."""
    p = tmp_path / name
    wb = openpyxl.Workbook()
    ws = wb.active
    for i in range(1, 6):
        ws.cell(i, 1).value = i
        ws.cell(i, 4).value = 6 - i
    ws["B1"] = ArrayFormula("B1:B5", "=A1:A5*2")
    ws["C1"] = ArrayFormula("C1", "=SUM(A1:A5*2)")
    wb.save(p)
    return str(p)


def _stored(path, part="xl/worksheets/sheet1.xml"):
    with zipfile.ZipFile(path) as zf:
        return zf.read(part).decode("utf-8")


# ============================================================== C-1
# An unbalanced '(' inside a quoted sheet name silently disabled the whole
# _xlpm pass. Excel's storage for every one of these was confirmed by COM.

@pytest.mark.parametrize("src,want", [
    ("=LET(x,'Q1 (draft'!A1,x+1)",
     "=_xlfn.LET(_xlpm.x,'Q1 (draft'!A1,_xlpm.x+1)"),
    ("=LET(x,'It''s (fun'!A1,x+1)",
     "=_xlfn.LET(_xlpm.x,'It''s (fun'!A1,_xlpm.x+1)"),
    ("=LAMBDA(x,'a(b'!A1+x)(1)",
     "=_xlfn.LAMBDA(_xlpm.x,'a(b'!A1+_xlpm.x)(1)"),
    ("=LET(x,'((('!A1,x)", "=_xlfn.LET(_xlpm.x,'((('!A1,_xlpm.x)"),
    # a '"' inside a single-quoted sheet name used to start a phantom string
    # that swallowed the rest of the argument list, hiding the y declaration
    ('=LET(x,\'a"b\'!A1,y,"c",x&y)',
     '=_xlfn.LET(_xlpm.x,\'a"b\'!A1,_xlpm.y,"c",_xlpm.x&_xlpm.y)'),
    # the even-paren case that happened to survive, so it stays surviving
    ("=LET(x,'Q1 (draft)'!A1,x+1)",
     "=_xlfn.LET(_xlpm.x,'Q1 (draft)'!A1,_xlpm.x+1)"),
])
def test_c1_quoted_sheet_names_do_not_disable_the_xlpm_pass(src, want):
    assert core_calc.normalize_formula(src)[0] == want


def test_c1_no_bare_parameter_name_can_reach_a_saved_file(tmp_path):
    p = _book(tmp_path, sheets=("Data", "Q1 (draft"))
    formulas.set_formula(p, {"cell": "A1"}, "=LET(x,'Q1 (draft'!A1,x+1)",
                         sheet="Data")
    xml = _stored(p)
    assert "_xlpm.x" in xml
    assert "LET(x," not in xml


def test_c1_an_unparseable_declaration_refuses_loudly(tmp_path):
    """The failure mode must be a refusal, never bare names on disk."""
    p = _book(tmp_path)
    before = Path(p).read_bytes()
    with pytest.raises(FormulaRejected):
        formulas.set_formula(p, {"cell": "A1"}, "=LET(x,(1,x)")
    assert Path(p).read_bytes() == before


def test_c1_literal_spans_track_both_quote_kinds():
    body = "LET(x,'a\"b'!A1,y,\"c\",x&y)"
    kinds = [k for _s, _e, k in core_calc.literal_spans(body)]
    assert kinds == ["sheet", "string"]


# ============================================================== C-2
# Whitespace before a builtin's '(' produced _xlpm.SUM (1,2). COM ground
# truth: Excel's FILE parser refuses "SUM (1,2)" in ANY form, so the write is
# refused rather than stored.

@pytest.mark.parametrize("src", [
    "=LET(sum,2,sum+SUM (1,2))",
    "=LET(sum,2,sum+SUM\t(1,2))",
    "=LET(sum,2,sum+SUM\n(1,2))",
    "=SUM (1,2)",
])
def test_c2_whitespace_before_a_paren_is_refused(src):
    with pytest.raises(FormulaRejected) as exc:
        core_calc.normalize_formula(src)
    assert "whitespace" in str(exc.value)


def test_c2_the_builtin_wins_guard_tolerates_whitespace_internally():
    """The guard itself matches Excel's tolerance even though a spaced call
    is separately refused: the name must never be captured as a parameter."""
    decl = [core_calc.Declaration(0, 80, ["sum"])]
    out = core_calc._prefix_params_scoped("LET(sum,2,sum+SUM (1,2))", decl)
    assert "_xlpm.SUM" not in out
    assert "_xlpm.sum,2,_xlpm.sum" in out


def test_c2_reaches_the_data_validation_custom_surface(tmp_path):
    p = _book(tmp_path)
    with pytest.raises(FormulaRejected):
        datavalidation.manage_data_validation(
            p, "add", location={"range": "A1:A3"}, dv_type="custom",
            formula1="=LET(sum,2,sum+SUM (1,2))", sheet="Data")


# ============================================================== C-3 / H-3
# CSE arrays: the movers treated an ArrayFormula OBJECT as a non-formula.

def test_c3_sort_refuses_over_a_cse_array(tmp_path):
    p = _cse_book(tmp_path)
    before = Path(p).read_bytes()
    with pytest.raises(UnsupportedStructure) as exc:
        sortfilter.sort_range(p, {"range": "A1:D5"},
                              [{"column": 4, "order": "asc"}],
                              has_header=False)
    assert "array" in str(exc.value)
    assert Path(p).read_bytes() == before


def test_c3_copy_translates_the_array_anchor_and_the_formula(tmp_path):
    """Excel's own answer, COM-confirmed: A1:B5 copied to F1 stores
    <c r="G1"><f t="array" ref="G1:G5">F1:F5*2</f>."""
    p = _cse_book(tmp_path)
    cells.copy_range(p, {"range": "A1:B5"}, {"cell": "F1"}, what="all")
    xml = _stored(p)
    assert 'ref="G1:G5"' in xml and "F1:F5*2" in xml
    # the untranslated anchor is the shape Excel refuses to open
    assert xml.count('ref="B1:B5"') == 1


def test_c3_a_values_copy_never_pastes_a_live_array_formula(tmp_path):
    p = _cse_book(tmp_path)
    cells.copy_range(p, {"range": "A1:B5"}, {"cell": "F1"}, what="values")
    xml = _stored(p)
    assert 'ref="G1:G5"' not in xml
    assert "F1:F5*2" not in xml


@pytest.mark.parametrize("action,at", [
    ("delete_rows", 1), ("delete_rows", 3), ("insert_rows", 2),
    ("insert_rows", 5),
])
def test_h3_structural_edits_refuse_to_split_an_array(tmp_path, action, at):
    p = _cse_book(tmp_path)
    before = Path(p).read_bytes()
    with pytest.raises(UnsupportedStructure):
        structure.modify_grid_structure(p, action, at, 1)
    assert Path(p).read_bytes() == before


def test_h3_removing_an_array_whole_is_allowed(tmp_path):
    """Excel's rule is about changing PART of an array; deleting the column
    the whole array lives in is an ordinary delete."""
    p = _cse_book(tmp_path)
    structure.modify_grid_structure(p, "delete_cols", 2, 1)
    assert 'ref="B1:B5"' not in _stored(p)


def test_h3_an_edit_clear_of_the_array_still_goes_through(tmp_path):
    """The guard is Excel's ("you can't change PART of an array"), not a
    blanket ban: an edit outside every array is unaffected."""
    p = _cse_book(tmp_path)
    structure.modify_grid_structure(p, "insert_rows", 7, 1)
    assert 'ref="B1:B5"' in _stored(p)


def test_h3_replace_cells_keeps_the_array_and_changes_only_the_text(tmp_path):
    """The silent 30-becomes-2 corruption: stripping t="array" made Excel
    resolve the formula with implicit intersection. Clean open, wrong number,
    no warning -- the most dangerous class in the report."""
    p = _cse_book(tmp_path)
    out = search.replace_cells(p, "A1:A5", "A1:A4", look_in="formulas")
    xml = _stored(p)
    assert 't="array" ref="C1"' in xml
    assert "SUM(A1:A4*2)" in xml
    assert out["changed"]["replaced"]["array_formulas_preserved"] == 2


def test_c3_partial_array_operations_refuse(tmp_path):
    p = _cse_book(tmp_path)
    for call in (
        lambda: cells.clear_range(p, {"range": "B1:B3"}),
        lambda: cells.copy_range(p, {"range": "A1:B3"}, {"cell": "F1"}),
        lambda: cells.move_range(p, {"range": "A1:B3"}, {"cell": "F1"}),
    ):
        with pytest.raises(UnsupportedStructure):
            call()


def test_c3_array_ranges_reads_the_anchor_and_the_owned_rectangle(tmp_path):
    p = _cse_book(tmp_path)
    wb = openpyxl.load_workbook(p)
    found = {a.anchor: a.ref for a in core_arrays.array_ranges(wb.active)}
    assert found == {"B1": "B1:B5", "C1": "C1"}


# ============================================================== C-4
# Five single-call routes to an unopenable workbook, each reported verified.
# Every threshold below is the MEASURED Excel boundary (core/limits.py).

def test_c4_oversize_comment_refuses(tmp_path):
    p = _book(tmp_path)
    before = Path(p).read_bytes()
    with pytest.raises(ExcelWouldRefuse):
        annotations.manage_comment(p, "add", location={"cell": "A1"},
                                   text="z" * 60000, sheet="Data")
    assert Path(p).read_bytes() == before
    # the boundary itself: 32,767 opens in Excel, 32,768 does not
    annotations.manage_comment(p, "add", location={"cell": "A1"},
                               text="z" * core_limits.MAX_COMMENT_CHARS,
                               sheet="Data")


def test_c4_oversize_data_validation_list_refuses(tmp_path):
    p = _book(tmp_path)
    with pytest.raises(ExcelWouldRefuse) as exc:
        datavalidation.manage_data_validation(
            p, "add", location={"range": "A1:A3"}, dv_type="list",
            values=[f"x{i}" for i in range(10000)], sheet="Data")
    assert "range" in str(exc.value)


def test_c4_oversize_header_refuses_on_the_assembled_string(tmp_path):
    p = _book(tmp_path)
    with pytest.raises(ExcelWouldRefuse):
        pagelayout.set_header_footer(p, sheet="Data",
                                     header={"left": "&Z&bad" * 500})
    # the limit is on the WHOLE &L/&C/&R string, so three legal-looking
    # sections can still cross it
    with pytest.raises(ExcelWouldRefuse):
        pagelayout.set_header_footer(
            p, sheet="Data",
            header={"left": "a" * 90, "center": "b" * 90, "right": "c" * 90},
            raw=True)
    pagelayout.set_header_footer(p, sheet="Data",
                                 header={"left": "a" * 80}, raw=True)


def test_c4_invalid_hyperlink_authority_refuses(tmp_path):
    p = _book(tmp_path)
    with pytest.raises(ExcelWouldRefuse) as exc:
        annotations.manage_hyperlink(p, "add", location={"cell": "A1"},
                                     target='http://x"/><evil a="',
                                     sheet="Data")
    assert "host" in str(exc.value)
    # MEASURED: the same characters in the PATH are fine and must not refuse
    annotations.manage_hyperlink(p, "add", location={"cell": "A2"},
                                 target='http://example.com/x"y',
                                 sheet="Data")


def test_c4_negative_image_extents_refuse(tmp_path):
    with pytest.raises(ExcelWouldRefuse):
        core_limits.check_image_extents(-5, -5)
    core_limits.check_image_extents(0, 0)
    core_limits.check_image_extents(10 ** 9, 10)


# ============================================================== H-1
# denormalize_formula corrupted string literals on every read.

@pytest.mark.parametrize("src,want", [
    # Excel stores this verbatim and the cell displays "_xlpm.total"
    ('=CONCATENATE("_xlpm.","total")', '=CONCATENATE("_xlpm.","total")'),
    ('="_xlfn._xlws.FILTER is the prefix"',
     '="_xlfn._xlws.FILTER is the prefix"'),
    ('=_xlfn.LET(_xlpm.msg,"_xlfn.LAMBDA / _xlpm.p / _xlop.q",_xlpm.msg)',
     '=LET(msg,"_xlfn.LAMBDA / _xlpm.p / _xlop.q",msg)'),
    ("='_xlpm.sheet'!A1", "='_xlpm.sheet'!A1"),
])
def test_h1_denormalize_never_edits_a_literal(src, want):
    assert core_calc.denormalize_formula(src) == want


def test_h1_a_read_write_cycle_does_not_corrupt_the_users_text(tmp_path):
    p = _book(tmp_path)
    formulas.set_formula(p, {"cell": "A1"}, '=CONCATENATE("_xlpm.","total")',
                         sheet="Data")
    got = cells.read_range(p, {"cell": "A1"}, values="formula",
                           sheet="Data")["values"][0][0]
    assert got == '=CONCATENATE("_xlpm.","total")'
    formulas.set_formula(p, {"cell": "A2"}, got, sheet="Data")
    assert '_xlpm.' in _stored(p)


# ============================================================== H-2
# Prefixing was global; Excel scopes it to the declaring LET/LAMBDA.

@pytest.mark.parametrize("src,want", [
    ("=LET(Rate,1,Rate)+Rate", "=_xlfn.LET(_xlpm.Rate,1,_xlpm.Rate)+Rate"),
    ("=LAMBDA(Factor,Factor*2)(3)+Factor",
     "=_xlfn.LAMBDA(_xlpm.Factor,_xlpm.Factor*2)(3)+Factor"),
    ("=LET(a,1,a)+LET(b,2,b)+a",
     "=_xlfn.LET(_xlpm.a,1,_xlpm.a)+_xlfn.LET(_xlpm.b,2,_xlpm.b)+a"),
])
def test_h2_prefixing_is_scoped_to_the_declaring_call(src, want):
    assert core_calc.normalize_formula(src)[0] == want


def test_h2_normalization_is_idempotent_under_scoping():
    once, _ = core_calc.normalize_formula("=LET(Rate,1,Rate)+Rate")
    twice, _ = core_calc.normalize_formula(once)
    assert once == twice


# ============================================================== H-4
# The pooled COM worker was never terminated and never reaped.

class _FakeWorker:
    def __init__(self, pid):
        self.pid = pid
        self.app = object()
        self.busy = False


def test_h4_the_journal_marks_a_dead_owners_pids_as_abandoned(tmp_path):
    from xlsx_mcp.com import instances as com_instances

    j = com_instances.PidJournal(tmp_path / "j.json")
    j.record(4242)
    rec = json.loads((tmp_path / "j.json").read_text())["4242"]
    assert rec["owner_pid"] > 0
    # a record owned by THIS live process is a concurrent session, not litter
    assert j.abandoned_pids() == set()
    rec["owner_pid"] = 999999          # a server that is certainly gone
    (tmp_path / "j.json").write_text(json.dumps({"4242": rec}))
    assert j.abandoned_pids() == {4242}


def test_h4_shutdown_reclaims_owned_pids_by_pid(tmp_path, monkeypatch):
    """Quit alone was never enough: it was issued from the wrong COM
    apartment at interpreter exit, raised, and was swallowed. Shutdown must
    end with a reclaim BY OWNED PID."""
    from xlsx_mcp.com import session as com_session

    ex = com_session.ComExecutor(journal_path=tmp_path / "j.json")
    reclaimed = []

    class _Mgr:
        def __init__(self):
            self._workers = [_FakeWorker(1234)]
            self._pool = None

        def force_reclaim(self):
            reclaimed.append(True)

    ex._manager = _Mgr()
    ex.shutdown()
    assert reclaimed == [True]
    assert ex._manager is None


def test_h4_status_reports_journaled_pids_without_a_live_worker(tmp_path):
    """A fresh server reported journaled_pids: [] next to a stranded Excel it
    had every record of, because status read the journal through the live
    manager."""
    from xlsx_mcp.com import instances as com_instances
    from xlsx_mcp.com import session as com_session

    com_instances.PidJournal(tmp_path / "j.json").record(5150)
    ex = com_session.ComExecutor(journal_path=tmp_path / "j.json")
    assert 5150 in ex.status()["journaled_pids"]


# ============================================================== H-5
# The cross-process write lockfile leaked under contention.

def test_h5_the_lockfile_is_published_complete(tmp_path):
    """The leak began with a two-syscall create: a concurrent acquirer read
    the empty file, saw no pid, called it stale and deleted a LIVE lock."""
    lock = tmp_path / "write.lock"
    assert core_safesave._publish_lockfile(lock) is True
    info = json.loads(lock.read_text())
    assert info["pid"] and info["token"] and info["time"]


def test_h5_a_second_acquisition_by_this_process_is_re_entrant(tmp_path):
    lock = tmp_path / "write.lock"
    assert core_safesave._acquire_lockfile(lock, "t.xlsx") is True
    assert core_safesave._acquire_lockfile(lock, "t.xlsx") is False


def test_h5_a_recycled_pid_under_a_foreign_token_is_stale(tmp_path):
    """Our own PID with someone else's token means the number was recycled.
    The old re-entrancy amnesty read that as 'mine, already held' and never
    cleaned it up, forever."""
    import os

    lock = tmp_path / "write.lock"
    lock.write_text(json.dumps({"pid": os.getpid(), "token": "someone-else",
                                "time": __import__("time").time()}))
    assert core_safesave._acquire_lockfile(lock, "t.xlsx") is True
    assert json.loads(lock.read_text())["token"] == core_safesave._OWNER_TOKEN


def test_h5_release_never_removes_someone_elses_lock(tmp_path):
    lock = tmp_path / "write.lock"
    lock.write_text(json.dumps({"pid": 999999, "token": "not-ours",
                                "time": __import__("time").time()}))
    core_safesave._release_lockfile(lock)
    assert lock.exists()
    lock.unlink()
    core_safesave._publish_lockfile(lock)
    core_safesave._release_lockfile(lock)
    assert not lock.exists()


def test_h5_a_live_foreign_holder_is_reported_by_its_own_pid(tmp_path,
                                                             monkeypatch):
    import time as _time

    lock = tmp_path / "write.lock"
    lock.write_text(json.dumps({"pid": 4321, "token": "other",
                                "time": _time.time()}))
    monkeypatch.setattr(core_safesave, "_pid_alive", lambda pid: True)
    monkeypatch.setattr(core_safesave, "LOCK_WAIT_SECONDS", 0.2)
    with pytest.raises(core_safesave.MutationLockTimeout) as exc:
        core_safesave._acquire_lockfile(lock, "t.xlsx")
    assert "PID 4321" in str(exc.value)


def test_h5_the_write_lock_leaves_nothing_behind(tmp_path):
    p = _book(tmp_path)
    with core_safesave.write_lock(p):
        pass
    lock = core_safesave.slot_dir(p) / core_safesave.LOCK_FILE_NAME
    assert not lock.exists()


# ============================================================== M-1
# validate claimed opens_clean: true for workbooks Excel refuses to open.

def test_m1_opens_clean_is_not_claimed_without_an_excel_verdict(tmp_path):
    p = _book(tmp_path)
    f = validation.validate(p, checks=["structure"])["results"]["structure"][
        "findings"]
    assert f["opens_clean"] == validation.NOT_CHECKED
    assert f["openpyxl_loads"] is True
    assert "com_validate_opens_clean" in f["opens_clean_source"]


# ============================================================== M-2 / M-3
# openpyxl's IllegalCharacterError escaped the envelope; a lone surrogate got
# no response at all.

@pytest.mark.parametrize("value", ["a\x01b", "a\x0bc", "a\x00b"])
def test_m2_control_characters_refuse_in_envelope(tmp_path, value):
    p = _book(tmp_path)
    with pytest.raises(ExcelWouldRefuse):
        cells.set_cell(p, {"cell": "A1"}, value, sheet="Data")


def test_m2_the_openpyxl_error_is_classified_as_a_backstop():
    from openpyxl.utils.exceptions import IllegalCharacterError

    assert _envelope.classify(IllegalCharacterError("x")) == "BAD_PARAMS"


def test_m2_a_refusal_never_echoes_raw_control_bytes():
    payload = _envelope.refusal(ExcelWouldRefuse("bad \x01 value"))
    assert "\x01" not in payload["error"]["message"]
    assert "<U+0001>" in payload["error"]["message"]


def test_m3_a_lone_surrogate_refuses_instead_of_hanging(tmp_path):
    p = _book(tmp_path)
    with pytest.raises(ExcelWouldRefuse) as exc:
        cells.set_cell(p, {"cell": "A1"}, "lone \ud800 sur", sheet="Data")
    assert "surrogate" in str(exc.value)
    # the refusal itself must be encodable, which is the whole point
    json.dumps(_envelope.refusal(exc.value)).encode("utf-8")


def test_m3_a_paired_surrogate_is_ordinary_text(tmp_path):
    p = _book(tmp_path)
    cells.set_cell(p, {"cell": "A1"}, "ok \U0001f600 emoji", sheet="Data")


# ============================================================== M-4
# A file-tier save dropped ca="1" and the cached value, re-creating the
# stuck-#VALUE! failure under iterative calculation.

def test_m4_the_always_calc_flag_and_cache_survive_an_unrelated_edit(
        tmp_path):
    p = _book(tmp_path)
    # hand-build the Excel shape openpyxl cannot model
    raw = _stored(p)
    patched = raw.replace(
        "</sheetData>",
        '<row r="9"><c r="B9"><f ca="1">B9+A1</f><v>2.5</v></c></row>'
        "</sheetData>")
    tmp = tmp_path / "patched.xlsx"
    with zipfile.ZipFile(p) as zin, zipfile.ZipFile(tmp, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = patched.encode("utf-8")
            zout.writestr(item, data)
    tmp.replace(p)

    cells.set_cell(p, {"cell": "A1"}, 99, sheet="Data")
    xml = _stored(p)
    assert 'ca="1"' in xml
    assert "<v>2.5</v>" in xml


def test_m4_an_edited_formula_never_keeps_a_stale_cached_value(tmp_path):
    """The restore is keyed on the formula text being UNCHANGED, so a value
    can never be pinned to a formula the edit rewrote."""
    p = _book(tmp_path)
    raw = _stored(p)
    patched = raw.replace(
        "</sheetData>",
        '<row r="9"><c r="B9"><f ca="1">B9+A1</f><v>2.5</v></c></row>'
        "</sheetData>")
    tmp = tmp_path / "patched.xlsx"
    with zipfile.ZipFile(p) as zin, zipfile.ZipFile(tmp, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = patched.encode("utf-8")
            zout.writestr(item, data)
    tmp.replace(p)

    formulas.set_formula(p, {"cell": "B9"}, "=B9+A1+7", sheet="Data")
    xml = _stored(p)
    assert "B9+A1+7" in xml
    assert "<v>2.5</v>" not in xml


# ============================================================== M-5
# sort_range erased every cached value in the sorted range without warning.

def test_m5_sort_warns_when_it_drops_a_populated_cache(tmp_path):
    p = tmp_path / "cached.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    for i in range(1, 4):
        ws.cell(i, 1).value = 4 - i
    wb.save(p)
    # give the formulas a real cached value, the way Excel would
    raw = _stored(str(p))
    patched = raw.replace(
        "</sheetData>",
        '<row r="1"><c r="B1"><f>A1*2</f><v>6</v></c></row>'
        '<row r="2"><c r="B2"><f>A2*2</f><v>4</v></c></row>'
        "</sheetData>")
    tmp = tmp_path / "p.xlsx"
    with zipfile.ZipFile(p) as zin, zipfile.ZipFile(tmp, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = patched.encode("utf-8")
            zout.writestr(item, data)
    tmp.replace(p)

    out = sortfilter.sort_range(str(p), {"range": "A1:B3"},
                                [{"column": 1, "order": "asc"}],
                                has_header=False)
    assert any("cached value before this sort" in w
               for w in out.get("warnings", []))


# ============================================================== M-6
# audit_formulas and read_range disagreed about the same cell.

def test_m6_the_audit_carries_both_the_stored_and_the_display_form(tmp_path):
    p = _book(tmp_path)
    formulas.set_formula(p, {"cell": "A1"}, "=LAMBDA(x,[y],x)", sheet="Data")
    item = formulas.audit_formulas(p)["formulas"]["items"][0]
    read = cells.read_range(p, {"cell": "A1"}, values="formula",
                            sheet="Data")["values"][0][0]
    assert item["formula"].startswith("=_xlfn.LAMBDA(")
    assert item["formula_display"] == read


# ============================================================== M-7
# manage_hyperlink wrote any URI scheme with no lint.

@pytest.mark.parametrize("target", [
    "javascript:alert(1)", "vbscript:msgbox(1)", "ms-msdt:/id PCWDiagnostic",
    "data:text/html;base64,PHNjcmlwdD4=",
])
def test_m7_unsafe_hyperlink_schemes_refuse(tmp_path, target):
    p = _book(tmp_path)
    with pytest.raises(ExcelWouldRefuse):
        annotations.manage_hyperlink(p, "add", location={"cell": "A1"},
                                     target=target, sheet="Data")


def test_m7_a_unc_target_is_written_but_warned_about(tmp_path):
    p = _book(tmp_path)
    out = annotations.manage_hyperlink(
        p, "add", location={"cell": "A1"},
        target=r"\\evil\share\payload.exe", sheet="Data")
    assert any("network share" in w for w in out.get("warnings", []))


def test_m7_ordinary_targets_are_untouched(tmp_path):
    p = _book(tmp_path)
    out = annotations.manage_hyperlink(
        p, "add", location={"cell": "A1"},
        target="https://example.com/report.pdf", sheet="Data")
    assert not [w for w in out.get("warnings", []) if "share" in w]


# ============================================================== L-1
# manage_name refused legal Unicode defined names.

@pytest.mark.parametrize("name", ["환율", "データ", "Ünïcode", "_x", "Rate"])
def test_l1_unicode_defined_names_are_accepted(name):
    assert names._validate_name(name) == name


@pytest.mark.parametrize("name", ["Tbl1", "Bad1", "Sc1", "N1", "R", "C",
                                  "1bad", "a b"])
def test_l1_the_correct_refusals_still_refuse(name):
    with pytest.raises(Exception):
        names._validate_name(name)


# ============================================================== L-2
# One refusal path returned a numeric error code.

def test_l2_an_expat_parse_error_does_not_leak_its_integer_code():
    from xml.etree.ElementTree import ParseError

    exc = ParseError("no element found: line 1, column 273")
    exc.code = 3                      # what expat sets
    payload = _envelope.refusal(exc)
    assert payload["error"]["code"] in _envelope.CLOSED_CODES
    assert isinstance(payload["error"]["code"], str)


def test_l2_a_missing_package_part_gets_a_sentence_not_a_repr():
    payload = _envelope.refusal(KeyError("xl/workbook.xml"))
    assert payload["error"]["code"] == "UNSUPPORTED_CONTENT"
    assert "missing xl/workbook.xml" in payload["error"]["message"]


# ============================================================== L-3
# com_goal_seek leaked a raw HRESULT.

def test_l3_an_hresult_becomes_plain_english():
    from xlsx_mcp.com import session as com_session

    class _ComError(Exception):
        hresult = -2147352567
        excepinfo = (0, None, None, None, 0, -2146827284)

    msg = com_session.excel_error_message(_ComError())
    assert "HRESULT" not in msg
    assert "Excel rejected the call" in msg


def test_l3_an_unknown_hresult_is_hex_not_a_pywin32_repr():
    from xlsx_mcp.com import session as com_session

    class _ComError(Exception):
        hresult = -2147000000
        excepinfo = None

    msg = com_session.excel_error_message(_ComError())
    assert "0x" in msg and "com_error" not in msg


# ============================================================== L-4
# Cosmetic normalize drift: documented and pinned, not "fixed".

@pytest.mark.parametrize("src,stored", [
    ("=LAMBDA(x, [ y ], x+y)",
     "=_xlfn.LAMBDA(_xlpm.x, _xlop.y, _xlpm.x+_xlpm.y)"),
    ("=LET(_xlpm.x,1,_xlpm.x+1)", "=_xlfn.LET(_xlpm.x,1,_xlpm.x+1)"),
])
def test_l4_the_drift_is_cosmetic_and_the_stored_form_is_stable(src, stored):
    """Excel stores the name without brackets or spaces, so the bracketed
    display form is re-derived on read. What must hold is that the STORED
    string is stable under a round trip."""
    out, _ = core_calc.normalize_formula(src)
    assert out == stored
    again, _ = core_calc.normalize_formula(
        core_calc.denormalize_formula(out))
    assert again == stored


# ============================================================== the gap
# The architectural finding: three layers, each independently checkable.

def test_the_refuse_class_table_is_published():
    table = core_limits.describe()
    assert table["comment_text_chars"] == 32767
    assert table["header_footer_string_chars"] == 255
    assert "2026-09-05" in table["measured_against"]


def test_the_deep_verification_default_is_off_but_switchable(monkeypatch):
    from xlsx_mcp.core import package as core_package

    monkeypatch.delenv(core_package.VERIFY_COM_ENV, raising=False)
    assert core_package.com_verify_default() is False
    monkeypatch.setenv(core_package.VERIFY_COM_ENV, "1")
    assert core_package.com_verify_default() is True
