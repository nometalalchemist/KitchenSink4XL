"""set_page_layout and set_header_footer tests: full-field round-trips
through the save pipeline and a fresh reload, the whole-row/col print-title
spans, the fit-to/scale exclusion, the & escaping and placeholder expansion
in headers/footers, and the refusal paths.
"""

from __future__ import annotations

import openpyxl
import pytest

from xlsx_mcp.core.errors import TargetNotFound, XlMcpError
from xlsx_mcp.ops import lifecycle as _lifecycle
from xlsx_mcp.ops import pagelayout as _pagelayout


def _make(tmp_path):
    p = str(tmp_path / "pl.xlsx")
    _lifecycle.create_workbook(p, sheets=["S"])
    return p


# -------------------------------------------------------------- page layout


def test_full_layout_round_trip(tmp_path):
    p = _make(tmp_path)
    out = _pagelayout.set_page_layout(
        p, sheet="S", orientation="landscape", paper_size="a4",
        margins={"left": 0.5, "top": 1.25, "header": 0.2},
        fit_to_width=1, fit_to_height=0, print_area="A1:D10",
        print_title_rows="1:2", print_title_cols="A:B",
        gridlines=True, headings=True)
    assert out["ok"] is True and out["verified"] is True
    wb = openpyxl.load_workbook(p)
    ws = wb["S"]
    assert ws.page_setup.orientation == "landscape"
    assert ws.page_setup.paperSize == 9
    assert ws.page_margins.left == 0.5
    assert ws.page_margins.top == 1.25
    assert ws.page_margins.header == 0.2
    assert int(ws.page_setup.fitToWidth) == 1
    assert int(ws.page_setup.fitToHeight) == 0
    assert ws.sheet_properties.pageSetUpPr.fitToPage is True
    assert ws.print_area == "'S'!$A$1:$D$10"
    assert ws.print_title_rows == "$1:$2"
    assert ws.print_title_cols == "$A:$B"
    assert ws.print_options.gridLines is True
    assert ws.print_options.headings is True
    wb.close()


def test_scale_and_paper_code(tmp_path):
    p = _make(tmp_path)
    _pagelayout.set_page_layout(p, sheet="S", scale=85, paper_size=5)
    wb = openpyxl.load_workbook(p)
    assert int(wb["S"].page_setup.scale) == 85
    assert wb["S"].page_setup.paperSize == 5
    wb.close()


def test_clear_print_area_and_titles(tmp_path):
    p = _make(tmp_path)
    _pagelayout.set_page_layout(p, sheet="S", print_area="A1:B2",
                                print_title_rows="1:1")
    out = _pagelayout.set_page_layout(p, sheet="S", print_area="clear",
                                      print_title_rows="clear")
    assert out["changed"]["page_layout"]["print_area"] == "cleared"
    wb = openpyxl.load_workbook(p)
    assert not wb["S"].print_area
    assert wb["S"].print_title_rows is None
    wb.close()


def test_layout_refusal_paths(tmp_path):
    p = _make(tmp_path)
    with pytest.raises(XlMcpError, match="nothing to set"):
        _pagelayout.set_page_layout(p, sheet="S")
    with pytest.raises(XlMcpError, match="mutually exclusive"):
        _pagelayout.set_page_layout(p, sheet="S", scale=80, fit_to_width=1)
    with pytest.raises(XlMcpError, match="orientation must be"):
        _pagelayout.set_page_layout(p, sheet="S", orientation="sideways")
    with pytest.raises(XlMcpError, match="unknown paper size"):
        _pagelayout.set_page_layout(p, sheet="S", paper_size="a11")
    with pytest.raises(XlMcpError, match="outside the ECMA-376 range"):
        _pagelayout.set_page_layout(p, sheet="S", paper_size=500)
    with pytest.raises(XlMcpError, match="unknown margin key"):
        _pagelayout.set_page_layout(p, sheet="S", margins={"middle": 1})
    with pytest.raises(XlMcpError, match="outside 0-10"):
        _pagelayout.set_page_layout(p, sheet="S", margins={"left": 99})
    with pytest.raises(XlMcpError, match="scale must be 10-400"):
        _pagelayout.set_page_layout(p, sheet="S", scale=5)
    with pytest.raises(XlMcpError, match="whole-row span"):
        _pagelayout.set_page_layout(p, sheet="S", print_title_rows="A:B")
    with pytest.raises(XlMcpError, match="whole-column span"):
        _pagelayout.set_page_layout(p, sheet="S", print_title_cols="1:2")
    with pytest.raises(TargetNotFound, match="no sheet named"):
        _pagelayout.set_page_layout(p, sheet="Nope", scale=100)


# ------------------------------------------------------------ header/footer


def test_header_footer_escaping_and_placeholders(tmp_path):
    p = _make(tmp_path)
    out = _pagelayout.set_header_footer(
        p, sheet="S",
        header={"center": "R&D report, page {page} of {pages}"},
        footer={"left": "{file} / {sheet}", "right": "{date} {time}"})
    assert out["verified"] is True
    wb = openpyxl.load_workbook(p)
    ws = wb["S"]
    assert ws.oddHeader.center.text == "R&&D report, page &P of &N"
    assert ws.oddFooter.left.text == "&F / &A"
    assert ws.oddFooter.right.text == "&D &T"
    wb.close()


def test_header_footer_raw_passthrough(tmp_path):
    p = _make(tmp_path)
    _pagelayout.set_header_footer(
        p, sheet="S", header={"right": "&P of &N"}, raw=True)
    wb = openpyxl.load_workbook(p)
    assert wb["S"].oddHeader.right.text == "&P of &N"
    wb.close()


def test_header_footer_even_first_flags(tmp_path):
    p = _make(tmp_path)
    _pagelayout.set_header_footer(p, sheet="S", apply_to="even",
                                  header={"center": "even"})
    _pagelayout.set_header_footer(p, sheet="S", apply_to="first",
                                  footer={"center": "first"})
    wb = openpyxl.load_workbook(p)
    ws = wb["S"]
    assert ws.evenHeader.center.text == "even"
    assert ws.firstFooter.center.text == "first"
    assert ws.HeaderFooter.differentOddEven is True
    assert ws.HeaderFooter.differentFirst is True
    wb.close()


def test_header_footer_clear_section(tmp_path):
    p = _make(tmp_path)
    _pagelayout.set_header_footer(p, sheet="S",
                                  header={"center": "x", "left": "y"})
    out = _pagelayout.set_header_footer(p, sheet="S", header={"center": ""})
    assert out["changed"]["header_footer"]["header"]["center"] == "cleared"
    wb = openpyxl.load_workbook(p)
    assert not wb["S"].oddHeader.center.text
    assert wb["S"].oddHeader.left.text == "y"
    wb.close()


def test_header_footer_refusals(tmp_path):
    p = _make(tmp_path)
    with pytest.raises(XlMcpError, match="nothing to set"):
        _pagelayout.set_header_footer(p, sheet="S")
    with pytest.raises(XlMcpError, match="apply_to must be"):
        _pagelayout.set_header_footer(p, sheet="S", header={"center": "x"},
                                      apply_to="both")
    with pytest.raises(XlMcpError, match="unknown header section"):
        _pagelayout.set_header_footer(p, sheet="S", header={"middle": "x"})
