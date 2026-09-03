"""Re-audit tests for the ops layer: each test pins a correctness trap the
audit found in the shipped Phase 3a/3b tools.

query_range: blank cells slipped through lt/le ordered filters (blank coerced
to "" and "" <= "5"); order_by on a column excluded from the projection was
silently ignored. condformat/datavalidation delete matched ranges by
SUBSTRING, so a delete aimed at A1 could remove the rule on A10:A20.
manage_name stored a leading '=' verbatim, which broke destinations parsing
and made the name unresolvable. manage_worksheet delete checked sheet count
but not visibility, contradicting its own docstring. Table names shaped like
R1C1 refs passed validation. CSV import coerced "nan"/"inf" into float
specials Excel cannot store. sort_range silently ordered uncached formula
cells by their formula text.
"""

from __future__ import annotations

import openpyxl
import pytest

from xlsx_mcp.core.errors import TargetNotFound, XlMcpError
from xlsx_mcp.ops import (
    cells as _cells,
    condformat as _condformat,
    dataio as _dataio,
    datavalidation as _datavalidation,
    lifecycle as _lifecycle,
    names as _names,
    sortfilter as _sortfilter,
    tables as _tables,
)


def _make(path, rows, sheet="Data"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    for i, row in enumerate(rows, start=1):
        for j, v in enumerate(row, start=1):
            ws.cell(i, j, v)
    wb.save(path)
    wb.close()
    return str(path)


# ------------------------------------------------------------ query_range


def test_blank_cells_fail_ordered_filters(tmp_path):
    p = _make(tmp_path / "q.xlsx",
              [["name", "qty"], ["a", 3], ["b", None], ["c", 10]])
    out = _cells.query_range(p, where=[{"column": "qty", "op": "le",
                                        "value": 5}])
    assert [r[0] for r in out["rows"]] == ["a"], (
        "a blank qty must not satisfy 'le 5'")
    out = _cells.query_range(p, where=[{"column": "qty", "op": "gt",
                                        "value": 0}])
    assert [r[0] for r in out["rows"]] == ["a", "c"]
    # blanks are still addressable explicitly
    out = _cells.query_range(p, where=[{"column": "qty", "op": "is_blank"}])
    assert [r[0] for r in out["rows"]] == ["b"]


def test_order_by_unprojected_column_sorts(tmp_path):
    p = _make(tmp_path / "q.xlsx",
              [["name", "revenue"], ["a", 5], ["b", 50], ["c", 20]])
    out = _cells.query_range(
        p, columns=["name"],
        order_by=[{"column": "revenue", "dir": "desc"}])
    assert out["columns"] == ["name"]
    assert [r[0] for r in out["rows"]] == ["b", "c", "a"], (
        "order_by must apply even when the sort column is not projected")


def test_regex_filter_is_search_not_fullmatch(tmp_path):
    p = _make(tmp_path / "q.xlsx",
              [["name"], ["alpha"], ["beta"], ["gamma"]])
    out = _cells.query_range(p, where=[{"column": "name", "op": "regex",
                                        "value": "a$"}])
    assert [r[0] for r in out["rows"]] == ["alpha", "beta", "gamma"]
    out = _cells.query_range(p, where=[{"column": "name", "op": "regex",
                                        "value": "^beta$"}])
    assert [r[0] for r in out["rows"]] == ["beta"]


# ------------------------------------- condformat / datavalidation delete


def test_cf_delete_no_substring_false_match(tmp_path):
    p = _make(tmp_path / "cf.xlsx", [[1], [2], [3]])
    _condformat.manage_conditional_format(
        p, "add", location={"range": "A10:A20"}, cf_type="cell_is",
        params={"operator": "greaterThan", "formula": 1})
    with pytest.raises(TargetNotFound):
        _condformat.manage_conditional_format(
            p, "delete", location={"cell": "A1"})
    # the rule on A10:A20 is still there
    out = _condformat.manage_conditional_format(p, "list")
    assert out["count"] == 1
    # exact address deletes it
    _condformat.manage_conditional_format(
        p, "delete", location={"range": "A10:A20"})
    assert _condformat.manage_conditional_format(p, "list")["count"] == 0


def test_dv_delete_no_substring_false_match(tmp_path):
    p = _make(tmp_path / "dv.xlsx", [[1], [2], [3]])
    _datavalidation.manage_data_validation(
        p, "add", location={"range": "A10:A20"}, dv_type="whole",
        formula1=0, operator="greaterThan")
    with pytest.raises(TargetNotFound):
        _datavalidation.manage_data_validation(
            p, "delete", location={"cell": "A1"})
    assert _datavalidation.manage_data_validation(p, "list")["count"] == 1


# ------------------------------------------------------------ manage_name


def test_refers_to_leading_equals_stripped_and_resolvable(tmp_path):
    p = _make(tmp_path / "n.xlsx", [["h", 1], ["x", 2]])
    _names.manage_name(p, "add", name="Box", refers_to="=Data!$A$1:$B$2")
    out = _names.manage_name(p, "list")
    entry = next(n for n in out["names"] if n["name"] == "Box")
    assert not str(entry["refers_to"]).startswith("="), (
        "definitions are stored without '='; storing it breaks destinations")
    # and the name actually resolves through the locate layer
    got = _cells.read_range(p, {"name": "Box"})
    assert got["range"] == "A1:B2"


# ------------------------------------------------------- manage_worksheet


def test_delete_only_visible_sheet_refuses_upfront(tmp_path):
    p = str(tmp_path / "w.xlsx")
    _lifecycle.create_workbook(p, sheets=["Main", "Aux"])
    _lifecycle.manage_worksheet(p, "hide", sheet="Aux")
    with pytest.raises(XlMcpError, match="VISIBLE"):
        _lifecycle.manage_worksheet(p, "delete", sheet="Main")
    # deleting the hidden one is fine
    out = _lifecycle.manage_worksheet(p, "delete", sheet="Aux")
    assert out["ok"] is True


# ------------------------------------------------------------ table names


def test_r1c1_shaped_table_names_refused(tmp_path):
    p = _make(tmp_path / "t.xlsx", [["h1", "h2"], [1, 2]])
    for bad in ("R1C1", "R2", "C33", "r", "c"):
        with pytest.raises(XlMcpError):
            _tables.create_table(p, {"range": "A1:B2"}, bad)
    out = _tables.create_table(p, {"range": "A1:B2"}, "RealName")
    assert out["ok"] is True


# ------------------------------------------------------------ import coercion


def test_nan_inf_stay_text_on_import(tmp_path):
    p = _make(tmp_path / "i.xlsx", [[None]])
    _dataio.import_data(p, source="v\nnan\ninf\n-1.5\n", fmt="csv")
    out = _cells.read_range(p, {"range": "A1:A4"})
    vals = [r[0] for r in out["values"]]
    assert vals == ["v", "nan", "inf", -1.5]


# ------------------------------------------------------------ sort warning


def test_sort_on_uncached_formula_keys_warns(tmp_path):
    p = _make(tmp_path / "s.xlsx", [["k"], [3], [1]])
    wb = openpyxl.load_workbook(p)
    wb["Data"]["A4"] = "=1+1"  # formula, never calculated: no cached value
    wb.save(p)
    wb.close()
    out = _sortfilter.sort_range(p, {"range": "A1:A4"},
                                 keys=[{"column": "k"}])
    assert any("formula TEXT" in w for w in out.get("warnings", [])), (
        "sorting by a never-calculated formula must warn, not silently order "
        "by formula text")


# ------------------------------------------------------- table overlap/names


def test_overlapping_tables_refused(tmp_path):
    p = _make(tmp_path / "ov.xlsx",
              [["a", "b", "c"], [1, 2, 3], [4, 5, 6], [7, 8, 9]])
    _tables.create_table(p, {"range": "A1:B3"}, "First")
    with pytest.raises(XlMcpError, match="overlap"):
        _tables.create_table(p, {"range": "B2:C4"}, "Second")
    # a disjoint table on the same sheet is fine
    out = _tables.create_table(p, {"range": "C1:C4"}, "Second")
    assert out["ok"] is True


def test_table_name_defined_name_collision_refused(tmp_path):
    p = _make(tmp_path / "nc.xlsx", [["h", 1], [2, 3]])
    _names.manage_name(p, "add", name="Shared", refers_to="Data!$A$1")
    with pytest.raises(XlMcpError, match="defined name"):
        _tables.create_table(p, {"range": "A1:B2"}, "Shared")


# ------------------------------------- TableList.items() latent crashes


def test_metadata_lists_tables_without_crashing(tmp_path):
    """openpyxl's TableList.items() yields (name, ref-STRING) pairs, not
    Table objects. get_workbook_metadata iterated items() and read .ref off
    the string, so ANY workbook containing a table crashed the discovery
    call. Same trap fixed in both _find_table case-insensitive fallbacks."""
    p = _make(tmp_path / "m.xlsx", [["h1", "h2"], [1, 2]])
    _tables.create_table(p, {"range": "A1:B2"}, "Inv")
    meta = _lifecycle.get_workbook_metadata(p)
    assert meta["tables"] == [
        {"name": "Inv", "sheet": "Data", "ref": "A1:B2"}]


def test_case_insensitive_table_lookup_returns_table(tmp_path):
    p = _make(tmp_path / "ci.xlsx", [["h1", "h2"], [1, 2]])
    _tables.create_table(p, {"range": "A1:B2"}, "Sales")
    out = _tables.get_table(p, "sales")  # case-insensitive fallback path
    assert out["ref"] == "A1:B2"
    got = _cells.read_range(p, {"table": "SALES", "part": "data"})
    assert got["range"] == "A2:B2"
