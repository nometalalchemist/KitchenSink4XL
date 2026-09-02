"""Unit tests for ops/sortfilter.py: sort_range, set_filter, clear_filter."""

from __future__ import annotations

import openpyxl
import pytest

from xlsx_mcp.core.errors import XlMcpError
from xlsx_mcp.ops import sortfilter as sf


@pytest.fixture()
def book(tmp_path):
    p = tmp_path / "sf.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    rows = [["Name", "Score"], ["c", 30], ["a", 10], ["b", 20]]
    for r, row in enumerate(rows, 1):
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v)
    wb.save(p)
    return p


def test_sort_ascending_by_header(book):
    r = sf.sort_range(str(book), {"range": "A1:B4"},
                      keys=[{"column": "Score", "order": "asc"}])
    assert r["ok"] and r["verified"]
    wb = openpyxl.load_workbook(book)
    col_b = [wb["S"].cell(r, 2).value for r in range(2, 5)]
    assert col_b == [10, 20, 30]
    assert wb["S"]["A1"].value == "Name"  # header stays


def test_sort_descending_and_multikey(book):
    # add a tie in Score to exercise the second key
    wb = openpyxl.load_workbook(book)
    wb["S"]["B4"] = 10  # b now ties a at 10
    wb.save(book)
    sf.sort_range(str(book), {"range": "A1:B4"},
                  keys=[{"column": "Score", "order": "asc"},
                        {"column": "Name", "order": "desc"}])
    wb = openpyxl.load_workbook(book)
    names = [wb["S"].cell(r, 1).value for r in range(2, 5)]
    # scores 10,10,30 -> among the two 10s, Name desc puts b before a
    assert names == ["b", "a", "c"]


def test_sort_shifts_relative_formula(tmp_path):
    p = tmp_path / "f.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A1"] = "K"
    ws["B1"] = "Calc"
    ws["A2"] = 3
    ws["B2"] = "=A2*10"
    ws["A3"] = 1
    ws["B3"] = "=A3*10"
    wb.save(p)
    sf.sort_range(str(p), {"range": "A1:B3"}, keys=[{"column": "K"}])
    wb = openpyxl.load_workbook(p)
    # after sort A2=1, A3=3; the formula in each row still points at its own row
    assert wb["S"]["A2"].value == 1
    assert wb["S"]["B2"].value == "=A2*10"
    assert wb["S"]["A3"].value == 3
    assert wb["S"]["B3"].value == "=A3*10"


def test_set_filter_hides_nonmatching(book):
    r = sf.set_filter(str(book), {"range": "A1:B4"},
                      criteria=[{"column": "Score", "op": "ge", "value": 20}])
    assert r["ok"] and r["verified"]
    assert r["changed"]["filter"]["rows_shown"] == 2
    wb = openpyxl.load_workbook(book)
    # row with Score 10 (row 3, 'a') is hidden; 30 and 20 visible
    assert wb["S"].row_dimensions[3].hidden is True
    assert wb["S"].row_dimensions[2].hidden is False
    assert wb["S"].auto_filter.ref == "A1:B4"


def test_clear_filter_unhides(book):
    sf.set_filter(str(book), {"range": "A1:B4"},
                  criteria=[{"column": "Score", "op": "ge", "value": 25}])
    r = sf.clear_filter(str(book), sheet="S")
    assert r["ok"]
    wb = openpyxl.load_workbook(book)
    assert wb["S"].auto_filter.ref is None
    for row in range(2, 5):
        assert not wb["S"].row_dimensions[row].hidden


def test_clear_filter_no_filter_refuses(book):
    with pytest.raises(XlMcpError):
        sf.clear_filter(str(book), sheet="S")


def test_sort_bad_keys(book):
    with pytest.raises(XlMcpError):
        sf.sort_range(str(book), {"range": "A1:B4"}, keys=[])
