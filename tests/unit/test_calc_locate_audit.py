"""Re-audit tests for core/calc.py and core/locate.py.

calc: the _xlfn future-function list was expanded from ~34 names to the full
documented set (a missing name means a silently broken formula write, proven
by the Phase 1 spike's #NAME?/repair-only findings); the _xlws set was
corrected to the documented {FILTER, SORT} pair; string literals are now
skipped by the shim; formula-only true-used-range semantics are pinned.

locate: the used_range selector honors the sibling 'sheet' key; region
probing no longer instantiates cells into the model.
"""

from __future__ import annotations

import openpyxl
import pytest

from xlsx_mcp.core import calc, locate


# ------------------------------------------------------- _xlfn completeness


@pytest.mark.parametrize("name", [
    # one representative per era the pre-audit list missed entirely
    "NORM.DIST", "STDEV.P", "RANK.EQ", "NETWORKDAYS.INTL", "WORKDAY.INTL",
    "IFNA", "BITAND", "SHEET", "ISFORMULA", "FORECAST.LINEAR",
    "FORECAST.ETS", "MAP", "REDUCE", "BYROW", "MAKEARRAY", "STOCKHISTORY",
    "REGEXEXTRACT",
])
def test_previously_missing_future_functions_are_prefixed(name):
    out, prefixed = calc.normalize_formula(f"={name}(A1)")
    assert out == f"=_xlfn.{name}(A1)"
    assert prefixed == [name]


def test_classic_functions_never_prefixed():
    for fn in ("SUM", "VLOOKUP", "IF", "INDEX", "MATCH", "COUNTIF",
               "SUMPRODUCT", "RANDBETWEEN", "FORECAST"):
        out, prefixed = calc.normalize_formula(f"={fn}(A1)")
        assert out == f"={fn}(A1)", fn
        assert prefixed == []


def test_xlws_set_is_filter_and_sort_only():
    """Documented storage uses _xlfn._xlws. for FILTER and SORT only. The
    pre-audit set also routed SORTBY/UNIQUE/SEQUENCE/RANDARRAY/ANCHORARRAY/
    SINGLE through _xlws, risking the repair-only failure the spike proved
    for a mis-prefixed dynamic array. COM-tier verification is flagged."""
    assert calc.XLFN_XLWS_FUNCS == frozenset({"FILTER", "SORT"})
    out, _ = calc.normalize_formula("=SORT(FILTER(A1:A9,B1:B9>0))")
    assert out == "=_xlfn._xlws.SORT(_xlfn._xlws.FILTER(A1:A9,B1:B9>0))"
    for fn in ("SORTBY", "UNIQUE", "SEQUENCE", "RANDARRAY"):
        out, _ = calc.normalize_formula(f"={fn}(A1:A9)")
        assert out == f"=_xlfn.{fn}(A1:A9)", fn


def test_normalize_skips_string_literals():
    out, prefixed = calc.normalize_formula(
        '=IF(A1>0,"call XLOOKUP(never)",XLOOKUP(D1,A:A,B:B))')
    assert '"call XLOOKUP(never)"' in out
    assert out.count("_xlfn.XLOOKUP") == 1
    assert prefixed == ["XLOOKUP"]


def test_already_prefixed_not_double_prefixed():
    out, prefixed = calc.normalize_formula("=_xlfn.XLOOKUP(D1,A:A,B:B)")
    assert out == "=_xlfn.XLOOKUP(D1,A:A,B:B)"
    assert prefixed == []


def test_denormalize_round_trip():
    src = "=_xlfn._xlws.SORT(_xlfn.UNIQUE(A1:A9))"
    assert calc.denormalize_formula(src) == "=SORT(UNIQUE(A1:A9))"


# --------------------------------------------------- true used range edges


def _wb():
    wb = openpyxl.Workbook()
    return wb, wb.active


def test_true_used_range_fully_empty_sheet():
    _, ws = _wb()
    assert locate.true_used_range(ws) is None


def test_true_used_range_single_cell():
    _, ws = _wb()
    ws["C7"] = "x"
    assert locate.true_used_range(ws) == (7, 3, 7, 3)


def test_formula_only_cell_counts_as_value_bearing():
    """A formula cell with NO cached value must still count as used: on a
    data_only=False load its value is the formula string. (On a data_only
    load it would read None; the read paths that drive used_range load with
    data_only=False for exactly this reason.)"""
    _, ws = _wb()
    ws["B2"] = "=SUM(A1:A9)"
    assert locate.true_used_range(ws) == (2, 2, 2, 2)


def test_formatted_but_empty_cells_not_value_bearing():
    _, ws = _wb()
    from openpyxl.styles import Font
    ws["Z99"].font = Font(bold=True)  # instantiated, styled, no value
    ws["A1"] = 1
    assert locate.true_used_range(ws) == (1, 1, 1, 1)


# --------------------------------------------------- locate re-audit fixes


def test_used_range_honors_sibling_sheet_key():
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "First"
    ws1["A1"] = 1
    ws2 = wb.create_sheet("Second")
    ws2["D4"] = 9
    grid = locate.resolve_location(wb, {"used_range": True, "sheet": "Second"})
    assert grid.sheet == "Second"
    assert (grid.min_row, grid.min_col) == (4, 4)
    # the string form still overrides the sibling
    grid2 = locate.resolve_location(
        wb, {"used_range": "First", "sheet": "Second"})
    assert grid2.sheet == "First"


def test_region_probe_does_not_instantiate_cells():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["B2"] = 1
    ws["B3"] = 2
    ws["C2"] = 3
    before = set(ws._cells)
    grid = locate.resolve_location(wb, {"region": {"near": "B2"}})
    assert (grid.min_row, grid.min_col, grid.max_row, grid.max_col) \
        == (2, 2, 3, 3)
    assert set(ws._cells) == before, (
        "region probing must not create empty cells in the model")
