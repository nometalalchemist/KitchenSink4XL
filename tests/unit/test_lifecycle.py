"""Unit tests for ops/lifecycle.py: workbook + worksheet lifecycle, metadata,
and diagnose. Happy paths and refusal paths; every mutation is verified and
backed up through WorkbookPackage."""

from __future__ import annotations

import hashlib
import shutil
from pathlib import Path

import openpyxl
import pytest

from xlsx_mcp.core.errors import (
    HazardRefused, TargetNotFound, XlMcpError)
from xlsx_mcp.ops import lifecycle

CORPUS = Path(__file__).resolve().parents[1] / "fixtures" / "corpus"


def _md5(p) -> str:
    return hashlib.md5(Path(p).read_bytes()).hexdigest()


# ------------------------------------------------------------- create / copy


def test_create_workbook(tmp_path):
    p = tmp_path / "new.xlsx"
    r = lifecycle.create_workbook(str(p), ["Alpha", "Beta"])
    assert r["created"] and r["sheets"] == ["Alpha", "Beta"]
    wb = openpyxl.load_workbook(p)
    assert wb.sheetnames == ["Alpha", "Beta"]


def test_create_refuses_existing_without_overwrite(tmp_path):
    p = tmp_path / "new.xlsx"
    lifecycle.create_workbook(str(p))
    with pytest.raises(FileExistsError):
        lifecycle.create_workbook(str(p))
    r = lifecycle.create_workbook(str(p), ["Z"], overwrite=True)
    assert r["sheets"] == ["Z"]


def test_create_rejects_duplicate_and_long_names(tmp_path):
    with pytest.raises(XlMcpError):
        lifecycle.create_workbook(str(tmp_path / "a.xlsx"), ["S", "s"])
    with pytest.raises(XlMcpError):
        lifecycle.create_workbook(str(tmp_path / "b.xlsx"), ["x" * 32])


def test_copy_workbook_is_byte_identical(tmp_path):
    src = tmp_path / "src.xlsx"
    lifecycle.create_workbook(str(src), ["Data"])
    dst = tmp_path / "dst.xlsx"
    r = lifecycle.copy_workbook(str(src), str(dst))
    assert r["copied_from"]
    assert _md5(src) == _md5(dst)
    with pytest.raises(FileExistsError):
        lifecycle.copy_workbook(str(src), str(dst))


def test_copy_missing_source_refuses(tmp_path):
    with pytest.raises(TargetNotFound):
        lifecycle.copy_workbook(str(tmp_path / "nope.xlsx"),
                                str(tmp_path / "out.xlsx"))


# --------------------------------------------------------------- metadata


def test_get_workbook_metadata(tmp_path):
    p = tmp_path / "m.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "h"
    ws["A2"] = 5
    ws["C3"] = "x"
    wb.create_sheet("Hidden").sheet_state = "hidden"
    wb.save(p)
    md = lifecycle.get_workbook_metadata(str(p))
    assert md["sheet_count"] == 2
    data = next(s for s in md["sheets"] if s["name"] == "Data")
    assert data["used_range"] == "A1:C3"
    assert data["state"] == "visible"
    hidden = next(s for s in md["sheets"] if s["name"] == "Hidden")
    assert hidden["state"] == "hidden"
    assert md["hazards"]["clean"] is True


def test_diagnose_clean_and_hazardous(tmp_path):
    p = tmp_path / "d.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A1"] = "=SUM(1,2)"
    wb.save(p)
    d = lifecycle.diagnose_workbook(str(p))
    assert d["readable"] and d["hazards"]["clean"] is True
    assert d["health"]["formula_cells"] == 1
    assert d["routing"]["surgical_edit"]["route"] == "openpyxl"

    src = tmp_path / "sl.xlsx"
    shutil.copy(CORPUS / "pivot_slicer.xlsx", src)
    d2 = lifecycle.diagnose_workbook(str(src))
    assert d2["hazards"]["would_lose"] is True
    assert "slicers" in [h["key"] for h in d2["hazards"]["hazards"]]


# --------------------------------------------------------- manage_worksheet


@pytest.fixture()
def book(tmp_path):
    p = tmp_path / "wb.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "One"
    wb.create_sheet("Two")
    wb.save(p)
    return p


def test_ws_add_rename_reorder_copy(book):
    r = lifecycle.manage_worksheet(str(book), "add", new_name="Three")
    assert r["ok"] and r["verified"]
    assert openpyxl.load_workbook(book).sheetnames == ["One", "Two", "Three"]

    lifecycle.manage_worksheet(str(book), "rename", sheet="Two",
                               new_name="Middle")
    assert "Middle" in openpyxl.load_workbook(book).sheetnames

    lifecycle.manage_worksheet(str(book), "reorder", sheet="Three", index=0)
    assert openpyxl.load_workbook(book).sheetnames[0] == "Three"

    lifecycle.manage_worksheet(str(book), "copy", sheet="One",
                               new_name="OneCopy")
    assert "OneCopy" in openpyxl.load_workbook(book).sheetnames


def test_ws_hide_unhide(book):
    lifecycle.manage_worksheet(str(book), "hide", sheet="Two")
    assert openpyxl.load_workbook(book)["Two"].sheet_state == "hidden"
    lifecycle.manage_worksheet(str(book), "hide", sheet="Two",
                               state="very_hidden")
    assert openpyxl.load_workbook(book)["Two"].sheet_state == "veryHidden"
    lifecycle.manage_worksheet(str(book), "unhide", sheet="Two")
    assert openpyxl.load_workbook(book)["Two"].sheet_state == "visible"


def test_ws_delete_and_last_sheet_guard(book):
    lifecycle.manage_worksheet(str(book), "delete", sheet="Two")
    assert openpyxl.load_workbook(book).sheetnames == ["One"]
    with pytest.raises(XlMcpError):
        lifecycle.manage_worksheet(str(book), "delete", sheet="One")


def test_ws_cannot_hide_last_visible(book):
    lifecycle.manage_worksheet(str(book), "hide", sheet="Two")
    with pytest.raises(XlMcpError):
        lifecycle.manage_worksheet(str(book), "hide", sheet="One")


def test_ws_bad_action_and_missing_sheet(book):
    with pytest.raises(XlMcpError):
        lifecycle.manage_worksheet(str(book), "frobnicate")
    with pytest.raises(TargetNotFound):
        lifecycle.manage_worksheet(str(book), "rename", sheet="Nope",
                                   new_name="X")


def test_ws_on_hazardous_refuses_without_allow_loss(tmp_path):
    src = tmp_path / "sl.xlsx"
    shutil.copy(CORPUS / "pivot_slicer.xlsx", src)
    before = _md5(src)
    with pytest.raises(HazardRefused):
        lifecycle.manage_worksheet(str(src), "add", new_name="New")
    assert _md5(src) == before
    r = lifecycle.manage_worksheet(str(src), "add", new_name="New",
                                   allow_loss=True)
    assert r["ok"] and r["warnings"]
