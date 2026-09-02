"""Unit tests for ops/dataio.py: import_data, export_range."""

from __future__ import annotations

import json

import openpyxl
import pytest

from xlsx_mcp.core.errors import RangeOutOfBounds, XlMcpError
from xlsx_mcp.ops import dataio as io


@pytest.fixture()
def book(tmp_path):
    p = tmp_path / "io.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    rows = [["Name", "Amount"], ["a", 10], ["b", 20]]
    for r, row in enumerate(rows, 1):
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v)
    wb.save(p)
    return p


def test_import_csv(book):
    r = io.import_data(str(book), source="x,y\n1,2\n3,4\n",
                       location={"cell": "D1"})
    assert r["ok"] and r["verified"]
    wb = openpyxl.load_workbook(book)
    assert wb["S"]["D1"].value == "x"
    assert wb["S"]["E2"].value == 2
    assert wb["S"]["D3"].value == 3


def test_import_json_records(book):
    payload = json.dumps([{"a": 1, "b": 2}, {"a": 3, "b": 4}])
    io.import_data(str(book), source=payload, fmt="json",
                   location={"cell": "D1"})
    wb = openpyxl.load_workbook(book)
    assert wb["S"]["D1"].value == "a"
    assert wb["S"]["E3"].value == 4


def test_import_neutralizes_injection(book):
    r = io.import_data(str(book), source="=1+1\n@cmd\nsafe\n",
                       location={"cell": "D1"})
    wb = openpyxl.load_workbook(book)
    d1 = wb["S"]["D1"]
    assert d1.value == "=1+1" and d1.data_type == "s"  # text, not a formula
    assert any("formula injection" in w for w in r["warnings"])


def test_import_formulas_true_writes_formula(book):
    io.import_data(str(book), source="=1+1\n", location={"cell": "D1"},
                   formulas=True)
    wb = openpyxl.load_workbook(book)
    assert wb["S"]["D1"].data_type == "f"


def test_import_bad_format(book):
    with pytest.raises(XlMcpError):
        io.import_data(str(book), source="x", fmt="xml")


def test_import_needs_source(book):
    with pytest.raises(XlMcpError):
        io.import_data(str(book), location={"cell": "A1"})


def test_export_csv_roundtrip(book):
    r = io.export_range(str(book), {"range": "A1:B3"}, fmt="csv")
    assert r["format"] == "csv"
    assert r["content"].splitlines()[0] == "Name,Amount"
    assert "a,10" in r["content"]


def test_export_json_records(book):
    r = io.export_range(str(book), {"range": "A1:B3"}, fmt="json",
                        records=True)
    data = json.loads(r["content"])
    assert data[0] == {"Name": "a", "Amount": 10}


def test_export_used_range_default(book):
    r = io.export_range(str(book), sheet="S", fmt="tsv")
    assert r["range"] == "A1:B3"
    assert "\t" in r["content"]


def test_export_to_file(book, tmp_path):
    out = tmp_path / "out.csv"
    r = io.export_range(str(book), {"range": "A1:B3"}, fmt="csv",
                        out_file=str(out))
    assert "written_to" in r
    assert out.read_text().startswith("Name,Amount")


def test_import_export_roundtrip(book):
    exported = io.export_range(str(book), {"range": "A1:B3"}, fmt="csv")
    io.import_data(str(book), source=exported["content"],
                   location={"cell": "D1"})
    reexported = io.export_range(str(book), {"range": "D1:E3"}, fmt="csv")
    assert reexported["content"].strip() == exported["content"].strip()
