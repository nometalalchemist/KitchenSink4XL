"""THE NUMBERS-SAFETY GATE (PLAN Addendum 2026-09-03, author ruling).

The author's bar for shipping: calculations, formulas, and all the numbers
must be safe for the end-user. This gate is the proof, run against real Excel
with Excel CLOSED, through the shipped tool bodies only.

  1. GROUND TRUTH per structural-edit class. One workbook of interlocking
     formulas (relative / absolute / mixed refs, cross-sheet, ranges, a
     defined name, whole-column spans), expectations derived BY HAND in this
     file, then insert/delete rows, insert/delete cols, sort, move, merge and
     an insert+delete composition, each verified by asking Excel to
     recalculate and comparing its answers to the hand-derived ones. A
     divergence means the reference rewriter lied about the numbers.
  2. FORMULA-FAMILY WRITE ROUND-TRIPS across every family the normalize shim
     documents (classic, the 2010 dot-statisticals, the 2013 wave, 2016
     IFS/TEXTJOIN, the IFNA/XLOOKUP era, dynamic arrays, LAMBDA helpers,
     volatiles), written through set_formula / set_cell / write_range /
     apply_edits: zero repair prompts, zero #NAME?, correct computed values.
  3. STALENESS LABELS across every read surface, including the neutralized
     injection-text seam (a cell storing "=cmd" as TEXT must never be
     labelled a formula).
  4. SORT / FILTER SEMANTICS against Excel's own Sort and AutoFilter on
     identical mixed-type data.
  5. CALC SETTINGS: manual mode, fullCalcOnLoad, iterative calc, volatiles.
  6. verify_com deep verification wired into a real structural-edit save.
  7. THE REFUSE-CLASS BATTERY (added after the insane-mode adversarial round,
     2026-09-05): every route that round found from a single tool call to a
     workbook Excel will not open, replayed through the shipped tool bodies.
     Two outcomes are acceptable and no third: a loud refusal that leaves the
     file byte-identical, or a file Excel itself opens. The boundary values
     are exercised too, so the refusals cannot be a blunt over-refusal.

PID DISCIPLINE: private DispatchEx workers only, journaled; SKIPS honestly if
a foreign EXCEL.EXE is present at start; never touches a foreign PID.
"""

from __future__ import annotations

import shutil
import sys
import tempfile
import time
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(ROOT / "src"))

import openpyxl  # noqa: E402

from xlsx_mcp.com import session as com_session  # noqa: E402
from xlsx_mcp.com.instances import (  # noqa: E402
    list_excel_pids, pid_alive, taskkill,
)
from xlsx_mcp.core.package import WorkbookPackage  # noqa: E402
from xlsx_mcp.ops import (  # noqa: E402
    cells, comtier, dataio, formulas, properties, sortfilter, structure,
    tables, view,
)

FAILS: list[str] = []
NOTES: list[str] = []


def check(cond: bool, label: str) -> bool:
    if cond:
        print(f"PASS {label}")
    else:
        print(f"FAIL {label}")
        FAILS.append(label)
    return bool(cond)


def note(text: str) -> None:
    print(f"NOTE {text}")
    NOTES.append(text)


# ---------------------------------------------------------------- helpers

def cached(path, sheet, coord):
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        return wb[sheet][coord].value
    finally:
        wb.close()


def cached_many(path, wants: dict[str, list[str]]) -> dict[str, object]:
    """{'Sheet!A1': value} for every requested address, from the saved cache."""
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        return {f"{s}!{c}": wb[s][c].value
                for s, coords in wants.items() for c in coords}
    finally:
        wb.close()


def formulas_of(path, wants: dict[str, list[str]]) -> dict[str, object]:
    wb = openpyxl.load_workbook(path, data_only=False)
    try:
        return {f"{s}!{c}": wb[s][c].value
                for s, coords in wants.items() for c in coords}
    finally:
        wb.close()


def approx(a, b) -> bool:
    if isinstance(a, (int, float)) and isinstance(b, (int, float)) \
            and not isinstance(a, bool) and not isinstance(b, bool):
        return abs(float(a) - float(b)) < 1e-9
    return a == b


def verify_numbers(path, expected: dict[str, object], label: str,
                   wants: dict[str, list[str]]) -> None:
    """Recalculate in Excel, then compare every address to the hand-derived
    expectation. Prints the whole delta so a failure is diagnosable."""
    comtier.recalculate(str(path))
    got = cached_many(path, wants)
    bad = {k: (expected[k], got.get(k)) for k in expected
           if not approx(expected[k], got.get(k))}
    if bad:
        print(f"  formulas: {formulas_of(path, wants)}")
        print(f"  MISMATCH (expected, excel): {bad}")
    check(not bad, label)


# ---------------------------------------------------------- the fixture
#
#   Sheet "D"                          Sheet "S"
#   A       B              C              D            A1 =D!B5
# 1 hdr     Qty            Cost           Line         A2 =SUM(D!B2:B4)
# 2 a       2              10             =B2*C2       A3 =D!B6+D!C6
# 3 b       3              20             =B3*C3       A4 =SUM(D!C:C)
# 4 c       4              30             =B4*C4       A5 =SUM(D!D2:D4)
# 5 tot     =SUM(B2:B4)    =SUM(C2:C4)    =SUM(D2:D4)
# 6 rel     =B2*C2         =$B$3*C$3
# 7 named   =SUM(Qty)      =B$2+$C3
#
# Defined name Qty -> D!$B$2:$B$4.

WANTS = {"D": ["B5", "C5", "D5", "B6", "C6", "B7", "C7"],
         "S": ["A1", "A2", "A3", "A4", "A5"]}

BASE = {
    "D!B5": 9, "D!C5": 60, "D!D5": 200, "D!B6": 20, "D!C6": 60,
    "D!B7": 9, "D!C7": 22,
    "S!A1": 9, "S!A2": 9, "S!A3": 80, "S!A4": 202, "S!A5": 200,
}


def build_fixture(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "D"
    ws.append(["hdr", "Qty", "Cost", "Line"])
    for r, (name, qty, cost) in enumerate(
            (("a", 2, 10), ("b", 3, 20), ("c", 4, 30)), start=2):
        ws.cell(r, 1, name)
        ws.cell(r, 2, qty)
        ws.cell(r, 3, cost)
        ws.cell(r, 4, f"=B{r}*C{r}")
    ws["A5"] = "tot"
    ws["B5"] = "=SUM(B2:B4)"
    ws["C5"] = "=SUM(C2:C4)"
    ws["D5"] = "=SUM(D2:D4)"
    ws["A6"] = "rel"
    ws["B6"] = "=B2*C2"
    ws["C6"] = "=$B$3*C$3"
    ws["A7"] = "named"
    ws["B7"] = "=SUM(Qty)"
    ws["C7"] = "=B$2+$C3"
    s = wb.create_sheet("S")
    s["A1"] = "=D!B5"
    s["A2"] = "=SUM(D!B2:B4)"
    s["A3"] = "=D!B6+D!C6"
    s["A4"] = "=SUM(D!C:C)"
    s["A5"] = "=SUM(D!D2:D4)"
    from openpyxl.workbook.defined_name import DefinedName
    wb.defined_names.add(DefinedName("Qty", attr_text="D!$B$2:$B$4"))
    wb.save(path)
    wb.close()


def item1(scratch: Path) -> None:
    print("\n=== ITEM 1: ground-truth computation per structural-edit class ===")
    src = scratch / "_fixture.xlsx"
    build_fixture(src)

    def fresh(tag: str) -> Path:
        p = scratch / f"gt_{tag}.xlsx"
        shutil.copy2(src, p)
        return p

    # ---- 1.0 the baseline itself must compute to the hand-derived numbers
    p = fresh("base")
    verify_numbers(p, BASE, "1.0 baseline: Excel agrees with the hand math",
                   WANTS)

    # ---- 1.1 insert_rows INSIDE the summed block: every value invariant,
    #          at the SHIFTED addresses (rows 5-7 become 6-8).
    p = fresh("insrow")
    structure.modify_grid_structure(str(p), "insert_rows", 3, count=1,
                                    sheet="D")
    wants = {"D": ["B6", "C6", "D6", "B7", "C7", "B8", "C8"],
             "S": ["A1", "A2", "A3", "A4", "A5"]}
    exp = {"D!B6": 9, "D!C6": 60, "D!D6": 200, "D!B7": 20, "D!C7": 60,
           "D!B8": 9, "D!C8": 22,
           "S!A1": 9, "S!A2": 9, "S!A3": 80, "S!A4": 202, "S!A5": 200}
    verify_numbers(p, exp, "1.1 insert_rows: ranges expanded, refs shifted, "
                           "every number invariant", wants)

    # ---- 1.2 delete_rows removing the 'c' row: ranges contract, the name
    #          contracts, cross-sheet refs follow.
    p = fresh("delrow")
    structure.modify_grid_structure(str(p), "delete_rows", 4, count=1,
                                    sheet="D")
    wants = {"D": ["B4", "C4", "D4", "B5", "C5", "B6", "C6"],
             "S": ["A1", "A2", "A3", "A4", "A5"]}
    exp = {"D!B4": 5, "D!C4": 30, "D!D4": 80, "D!B5": 20, "D!C5": 60,
           "D!B6": 5, "D!C6": 22,
           "S!A1": 5, "S!A2": 5, "S!A3": 80, "S!A4": 142, "S!A5": 80}
    verify_numbers(p, exp, "1.2 delete_rows: SUMs contract, defined name "
                           "contracts, cross-sheet refs follow", wants)

    # ---- 1.3 delete_rows destroying an absolute target -> Excel's #REF!
    p = fresh("delref")
    r = structure.modify_grid_structure(str(p), "delete_rows", 3, count=1,
                                        sheet="D")
    rewrites = r["changed"]["structure"]["rewrites"]
    comtier.recalculate(str(p))
    c5 = cached(p, "D", "C5")
    check(rewrites.get("ref_errors", 0) >= 1 and c5 == "#REF!",
          f"1.3 delete_rows: a wholly-deleted target becomes #REF! in Excel "
          f"and the count is reported (ref_errors={rewrites.get('ref_errors')},"
          f" C5={c5!r})")

    # ---- 1.4 insert_cols: column refs and the whole-column span transpose
    p = fresh("inscol")
    structure.modify_grid_structure(str(p), "insert_cols", "C", count=1,
                                    sheet="D")
    wants = {"D": ["B5", "D5", "E5", "B6", "D6", "B7", "D7"],
             "S": ["A1", "A2", "A3", "A4", "A5"]}
    exp = {"D!B5": 9, "D!D5": 60, "D!E5": 200, "D!B6": 20, "D!D6": 60,
           "D!B7": 9, "D!D7": 22,
           "S!A1": 9, "S!A2": 9, "S!A3": 80, "S!A4": 202, "S!A5": 200}
    verify_numbers(p, exp, "1.4 insert_cols: refs shift right and the "
                           "whole-column span C:C transposes to D:D", wants)

    # ---- 1.5 delete_cols: the label column goes, everything slides left
    p = fresh("delcol")
    structure.modify_grid_structure(str(p), "delete_cols", "A", count=1,
                                    sheet="D")
    wants = {"D": ["A5", "B5", "C5", "A6", "B6", "A7", "B7"],
             "S": ["A1", "A2", "A3", "A4", "A5"]}
    exp = {"D!A5": 9, "D!B5": 60, "D!C5": 200, "D!A6": 20, "D!B6": 60,
           "D!A7": 9, "D!B7": 22,
           "S!A1": 9, "S!A2": 9, "S!A3": 80, "S!A4": 202, "S!A5": 200}
    verify_numbers(p, exp, "1.5 delete_cols: refs, the defined name and the "
                           "whole-column span all slide left", wants)

    # ---- 1.6 sort_range: moved formulas keep their own row, the numbers
    #          OUTSIDE the sorted block see the reordered data.
    p = fresh("sort")
    sortfilter.sort_range(str(p), "A2:D4", keys=[{"column": "B",
                                                  "order": "desc"}],
                          has_header=False, sheet="D")
    wants = {"D": ["B5", "C5", "D5", "B6", "C6", "B7", "C7",
                   "D2", "D3", "D4"],
             "S": ["A1", "A2", "A3", "A4", "A5"]}
    exp = {"D!D2": 120, "D!D3": 60, "D!D4": 20,
           "D!B5": 9, "D!C5": 60, "D!D5": 200, "D!B6": 120, "D!C6": 60,
           "D!B7": 9, "D!C7": 24,
           "S!A1": 9, "S!A2": 9, "S!A3": 180, "S!A4": 204, "S!A5": 200}
    verify_numbers(p, exp, "1.6 sort_range: each moved formula still points "
                           "at its own row and the dependents follow", wants)

    # ---- 1.7 move_range: references INTO the moved rectangle follow it
    p = fresh("move")
    cells.move_range(str(p), "A2:D2", "A10", sheet="D")
    wants = {"D": ["B5", "C5", "D5", "B6", "C6", "B7", "C7", "D10"],
             "S": ["A1", "A2", "A3", "A4", "A5"]}
    exp = {"D!B5": 7, "D!C5": 50, "D!D5": 180, "D!B6": 20, "D!C6": 60,
           "D!B7": 7, "D!C7": 22, "D!D10": 20,
           "S!A1": 7, "S!A2": 7, "S!A3": 80, "S!A4": 192, "S!A5": 180}
    verify_numbers(p, exp, "1.7 move_range: refs into the moved rectangle "
                           "follow the cells, ranges left behind shrink",
                   wants)

    # ---- 1.8 set_merge must not disturb a number, and must refuse to
    #          discard one silently.
    p = fresh("merge")
    cells.set_merge(str(p), "merge", "F1:G2", sheet="D")
    refused = False
    try:
        cells.set_merge(str(p), "merge", "B4:B5", sheet="D")
    except Exception as exc:  # noqa: BLE001
        refused = "confirm_data_loss" in str(exc)
    check(refused, "1.8a set_merge refuses to swallow the values under a "
                   "merge without confirm_data_loss")
    verify_numbers(p, BASE, "1.8b set_merge: every number unchanged after a "
                            "merge beside the formulas", WANTS)

    # ---- 1.9 COMPOSITION: insert then delete the same row/col returns the
    #          workbook to the baseline numbers exactly.
    p = fresh("compose")
    structure.modify_grid_structure(str(p), "insert_rows", 3, count=2,
                                    sheet="D")
    structure.modify_grid_structure(str(p), "delete_rows", 3, count=2,
                                    sheet="D")
    structure.modify_grid_structure(str(p), "insert_cols", "B", count=1,
                                    sheet="D")
    structure.modify_grid_structure(str(p), "delete_cols", "B", count=1,
                                    sheet="D")
    verify_numbers(p, BASE, "1.9 composition: insert+delete rows and cols "
                            "round-trips to the baseline numbers", WANTS)


# ------------------------------------------------------- item 2: families

#: (address, formula, expected value). Written through the shipped writers.
CORE_FAMILIES = [
    # classic
    ("E1", "=SUM(B2:B4)", 9),
    ("E2", "=AVERAGE(C2:C4)", 20),
    ("E3", "=IF(B2>1,10,20)", 10),
    ("E4", "=COUNTIF(B2:B4,\">2\")", 2),
    ("E5", "=SUMIF(B2:B4,\">2\",C2:C4)", 50),
    ("E6", "=VLOOKUP(\"b\",A2:C4,3,FALSE)", 20),
    ("E7", "=INDEX(C2:C4,MATCH(4,B2:B4,0))", 30),
    ("E8", "=ROUND(10/3,2)", 3.33),
    # 2010 dot-statisticals
    ("F1", "=STDEV.S(B2:B4)", 1),
    ("F2", "=VAR.P(B2:B4)", 0.6666666666666666),
    ("F3", "=NORM.S.DIST(0,TRUE)", 0.5),
    ("F4", "=PERCENTILE.INC(B2:B4,0.5)", 3),
    ("F5", "=RANK.EQ(4,B2:B4)", 1),
    ("F6", "=MODE.SNGL({1,2,2,3})", 2),
    ("F7", "=QUARTILE.INC(B2:B4,2)", 3),
    ("F8", "=COVARIANCE.P(B2:B4,C2:C4)", 6.666666666666667),
    ("F9", "=CHISQ.DIST.RT(0,1)", 1),
    ("F10", "=T.DIST.2T(1,10)", 0.3408931323),
    # 2013 wave
    ("G1", "=IFNA(NA(),42)", 42),
    ("G2", "=ISFORMULA(B5)", True),
    ("G3", "=DAYS(DATE(2026,1,31),DATE(2026,1,1))", 30),
    ("G4", "=BASE(255,16)", "FF"),
    ("G5", "=ARABIC(\"XIV\")", 14),
    ("G6", "=CEILING.MATH(4.2)", 5),
    ("G7", "=FLOOR.MATH(4.8)", 4),
    ("G8", "=COMBINA(4,2)", 10),
    ("G9", "=ISOWEEKNUM(DATE(2026,1,8))", 2),
    ("G10", "=XOR(TRUE,FALSE)", True),
    ("G11", "=BITAND(12,10)", 8),
    ("G12", "=DECIMAL(\"FF\",16)", 255),
    ("G13", "=UNICHAR(65)", "A"),
    ("G14", "=NUMBERVALUE(\"1.5\")", 1.5),
    ("G15", "=PERMUTATIONA(3,2)", 9),
    ("G16", "=SHEETS()", 2),
    # 2016
    ("H1", "=CONCAT(\"a\",\"b\")", "ab"),
    ("H2", "=IFS(B2=2,\"two\",TRUE,\"other\")", "two"),
    ("H3", "=MAXIFS(C2:C4,B2:B4,\">2\")", 30),
    ("H4", "=MINIFS(C2:C4,B2:B4,\">2\")", 20),
    ("H5", "=SWITCH(2,1,\"one\",2,\"two\",\"none\")", "two"),
    ("H6", "=TEXTJOIN(\"-\",TRUE,A2:A4)", "a-b-c"),
    ("H7", "=FORECAST.LINEAR(5,C2:C4,B2:B4)", 40),
    # 365 lookups / LET
    ("I1", "=XLOOKUP(\"b\",A2:A4,C2:C4)", 20),
    ("I2", "=XMATCH(4,B2:B4)", 3),
    ("I3", "=LET(x,B2,y,C2,x*y)", 20),
    # dynamic arrays (first spilled cell)
    ("J1", "=FILTER(C2:C4,B2:B4>2)", 20),
    ("K1", "=SORT(C2:C4,1,-1)", 30),
    ("L1", "=SORTBY(C2:C4,B2:B4,-1)", 30),
    ("M1", "=UNIQUE(B2:B4)", 2),
    ("N1", "=SEQUENCE(3)", 1),
    ("O1", "=SUM(RANDARRAY(2)*0)", 0),
    # 365 text / array shaping
    ("P1", "=TEXTBEFORE(\"a-b\",\"-\")", "a"),
    ("P2", "=TEXTAFTER(\"a-b\",\"-\")", "b"),
    ("P3", "=TAKE(C2:C4,1)", 10),
    ("P4", "=DROP(C2:C4,2)", 30),
    ("P5", "=CHOOSEROWS(C2:C4,2)", 20),
    ("P6", "=INDEX(CHOOSECOLS(B2:C4,2),1,1)", 10),
    ("P7", "=INDEX(TOCOL(B2:C2),1,1)", 2),
    ("P8", "=INDEX(TOROW(B2:B4),1,1)", 2),
    ("P9", "=INDEX(VSTACK(B2:B2,C2:C2),1,1)", 2),
    ("P10", "=INDEX(HSTACK(B2:B2,C2:C2),1,1)", 2),
    ("P11", "=SUM(EXPAND(B2:B2,1,2,0))", 2),
    ("P12", "=ARRAYTOTEXT(B2:B3)", "2, 3"),
    ("P13", "=VALUETOTEXT(B2)", "2"),
    ("P14", "=INDEX(TEXTSPLIT(\"a,b\",\",\"),1,2)", "b"),
    ("P15", "=INDEX(WRAPROWS(SEQUENCE(4),2),2,1)", 3),
    # LAMBDA and its helpers (the _xlpm parameter-name seam)
    ("Q1", "=LAMBDA(x,x*2)(21)", 42),
    ("Q2", "=SUM(MAP(B2:B4,LAMBDA(v,v*2)))", 18),
    ("Q3", "=REDUCE(0,B2:B4,LAMBDA(acc,v,acc+v))", 9),
    ("Q4", "=INDEX(SCAN(0,B2:B4,LAMBDA(acc,v,acc+v)),3,1)", 9),
    ("Q5", "=INDEX(BYROW(B2:C4,LAMBDA(r,SUM(r))),1,1)", 12),
    ("Q6", "=INDEX(BYCOL(B2:C4,LAMBDA(c,SUM(c))),1,1)", 9),
    ("Q7", "=INDEX(MAKEARRAY(2,2,LAMBDA(r,c,r*c)),2,2)", 4),
]

#: Volatile functions: assert only that they compute to something sane.
VOLATILE_FAMILY = [
    ("R1", "=IF(NOW()>0,1,0)", 1),
    ("R2", "=IF(TODAY()>0,1,0)", 1),
    ("R3", "=IF(AND(RAND()>=0,RAND()<=1),1,0)", 1),
    ("R4", "=RANDBETWEEN(7,7)", 7),
    ("R5", "=OFFSET(B2,1,0)", 3),
    ("R6", "=INDIRECT(\"B4\")", 4),
    ("R7", "=CELL(\"col\",C2)", 3),
    ("R8", "=IF(LEN(INFO(\"numfile\"))>0,1,0)", 1),
]


def item2(scratch: Path) -> None:
    print("\n=== ITEM 2: formula-family write round-trips ===")
    p = scratch / "families.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "F"
    ws.append(["hdr", "Qty", "Cost"])
    for r, (n, q, c) in enumerate(
            (("a", 2, 10), ("b", 3, 20), ("c", 4, 30)), start=2):
        ws.cell(r, 1, n)
        ws.cell(r, 2, q)
        ws.cell(r, 3, c)
    ws["B5"] = "=SUM(B2:B4)"
    wb.create_sheet("Other")
    wb.save(p)
    wb.close()

    battery = CORE_FAMILIES + VOLATILE_FAMILY
    # Spread the battery across all four shipped write paths so the shim is
    # proven on each of them, not just on set_formula.
    quarter = len(battery) // 4 + 1
    via_set_formula = battery[:quarter]
    via_set_cell = battery[quarter:quarter * 2]
    via_write_range = battery[quarter * 2:quarter * 3]
    via_apply_edits = battery[quarter * 3:]

    for addr, f, _v in via_set_formula:
        formulas.set_formula(str(p), addr, f, sheet="F", backup=False)
    for addr, f, _v in via_set_cell:
        cells.set_cell(str(p), addr, f, sheet="F", backup=False)
    for addr, f, _v in via_write_range:
        cells.write_range(str(p), addr, [[f]], sheet="F", backup=False)
    view.apply_edits(str(p), [
        {"op": "set_formula", "location": addr, "sheet": "F", "formula": f}
        for addr, f, _v in via_apply_edits], backup=False)

    oc = comtier.com_validate_opens_clean(str(p))
    check(oc["opens_clean"] is True,
          f"2.1 every shim family opens with ZERO repair prompts "
          f"({oc.get('excel_says', '')[:70]})")

    comtier.recalculate(str(p))
    got = cached_many(p, {"F": [a for a, _f, _v in battery]})
    names = {}
    wrong = {}
    for addr, f, want in battery:
        v = got.get(f"F!{addr}")
        if v == "#NAME?":
            names[addr] = f
        elif not approx(want, v):
            wrong[addr] = (f, want, v)
    if names:
        print(f"  #NAME? cells: {names}")
    if wrong:
        print(f"  wrong values: {wrong}")
    check(not names, f"2.2 zero #NAME? across {len(battery)} formulas in "
                     f"every documented family")
    check(not wrong, f"2.3 every hand-derivable value computed correctly "
                     f"({len(battery) - len(wrong)}/{len(battery)})")

    # The four write paths must all be represented in what Excel accepted.
    check(all(got.get(f"F!{a}") not in (None, "#NAME?")
              for a, _f, _v in (via_set_formula[:1] + via_set_cell[:1]
                                + via_write_range[:1] + via_apply_edits[:1])),
          "2.4 set_formula / set_cell / write_range / apply_edits all "
          "produced live, computing formulas")


# ------------------------------------------------- item 3: staleness labels

def item3(scratch: Path) -> None:
    print("\n=== ITEM 3: staleness labels across every read surface ===")
    p = scratch / "stale.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "D"
    ws.append(["Item", "Amount"])
    for r, (n, v) in enumerate((("a", 10), ("b", 20), ("c", 30)), start=2):
        ws.cell(r, 1, n)
        ws.cell(r, 2, v)
    ws["A5"] = "tot"
    ws["B5"] = "=SUM(B2:B4)"
    wb.save(p)
    wb.close()
    tables.create_table(str(p), "A1:B5", name="Tbl", sheet="D")

    # ---- 3.1 BEFORE any calc: every surface must say 'absent', never blank
    rr = cells.read_range(str(p), "A1:B5", values="both", sheet="D")
    gc = cells.get_cells(str(p), ["B5"], values="both", sheet="D")
    qr = cells.query_range(str(p), "A1:B5", sheet="D",
                           aggregate=[{"column": "Amount", "func": "sum"}])
    gv = view.get_grid_view(str(p), sheet="D")
    gt = tables.get_table(str(p), "Tbl", values="cached")
    ex = dataio.export_range(str(p), "A1:B5", sheet="D")
    ef = dataio.export_file(str(p), fmt="csv")
    au = formulas.audit_formulas(str(p), sheet="D")
    check(rr["labels"][4][1] == "absent" and "warning" in rr,
          "3.1a read_range(both): the uncalculated SUM is labelled 'absent' "
          "with the warning")
    check(gc["cells"][0]["label"] == "absent" and "warning" in gc,
          "3.1b get_cells: 'absent' label + warning")
    check("warning" in qr and qr.get("uncalculated_cells") == ["B5"],
          "3.1c query_range(aggregate): names the uncalculated cell instead "
          "of quietly summing around it")
    check("B5" in gv["formula_cells"],
          "3.1d get_grid_view: the formula cell is marked, not shown blank")
    check(gt.get("absent_cells") == ["B5"] and "warning" in gt,
          "3.1e get_table: names the uncalculated body cell")
    check(ex.get("absent_cells") == ["B5"] and "warning" in ex,
          "3.1f export_range: the export says which cells left as blanks")
    check("warning" in ef and ef["sheets"][0].get("absent_cells") == 1,
          "3.1g export_file: same honesty per sheet")
    check(au["missing_cached_values"]["count"] == 1 and "warning" in au,
          "3.1h audit_formulas: reports the missing cached value")

    # ---- 3.2 the DEFAULT read mode (cached) must label too
    rr_c = cells.read_range(str(p), "A1:B5", values="cached", sheet="D")
    check(rr_c.get("labels") and rr_c["labels"][4][1] == "absent"
          and "warning" in rr_c,
          "3.2 read_range values='cached' (the DEFAULT) labels the absent "
          "formula instead of returning a silent blank")

    # ---- 3.3 AFTER an Excel recalc: 'cached', real numbers, no warning
    comtier.recalculate(str(p))
    rr2 = cells.read_range(str(p), "A1:B5", values="both", sheet="D")
    qr2 = cells.query_range(str(p), "A1:B5", sheet="D",
                            aggregate=[{"column": "Amount", "func": "sum"}])
    ex2 = dataio.export_range(str(p), "A1:B5", sheet="D")
    gt2 = tables.get_table(str(p), "Tbl", values="cached")
    check(rr2["values"][4][1] == 60 and rr2["labels"][4][1] == "cached"
          and "warning" not in rr2,
          "3.3a read_range: after recalculate the value is 60, labelled "
          "'cached', warning gone")
    check(qr2["groups"][0]["aggregates"]["sum_Amount"] == 120
          and "warning" not in qr2,
          "3.3b query_range: the total row now counts (10+20+30+60=120)")
    check("absent_cells" not in ex2 and "60" in ex2["content"],
          "3.3c export_range: the computed number is exported")
    check("absent_cells" not in gt2, "3.3d get_table: clean after recalc")
    check(comtier.recalculate(str(p))["changed"]
          ["formula_cells_without_cached_value_after"] == 0,
          "3.3e recalculate reports zero absent-cache cells remaining")

    # ---- 3.4 an edit AFTER the recalc invalidates the cache: say so
    cells.set_cell(str(p), "B2", 100, sheet="D")
    rr3 = cells.read_range(str(p), "A1:B5", values="both", sheet="D")
    check(rr3["labels"][4][1] == "absent" and "warning" in rr3,
          "3.4 an edit drops the caches and every read says 'absent' again "
          "rather than serving the pre-edit number as current")

    # ---- 3.5 THE INJECTION SEAM: text that merely starts with '='
    p2 = scratch / "neutral.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "S"
    wb.save(p2)
    wb.close()
    imp = dataio.import_data(
        str(p2), source="label\n=cmd|'/c calc'!A1\n=1+1\n", header=False,
        location={"cell": "A1"}, sheet="S")
    check(any("TEXT" in w for w in imp["warnings"]),
          "3.5a import_data neutralized the formula-shaped text")
    rr4 = cells.read_range(str(p2), "A1:A3", values="both", sheet="S")
    gc4 = cells.get_cells(str(p2), ["A2", "A3"], values="both", sheet="S")
    au4 = formulas.audit_formulas(str(p2), sheet="S")
    labs4 = rr4.get("labels") or [["value"]]
    check(all(lab == "value" for row in labs4 for lab in row)
          and "warning" not in rr4,
          f"3.5b read_range treats neutralized text as a VALUE, not a "
          f"formula ({rr4.get('labels')})")
    check(all(c["label"] == "value" for c in gc4["cells"]),
          "3.5c get_cells labels neutralized text 'value'")
    check(au4["formulas"]["count"] == 0,
          f"3.5d audit_formulas does not report neutralized text as a "
          f"formula ({au4['formulas']['count']})")
    # and no structural edit, copy, move or sort may re-arm it
    cells.copy_range(str(p2), "A2", "C2", sheet="S")
    cells.move_range(str(p2), "A3", "E3", sheet="S")
    structure.modify_grid_structure(str(p2), "insert_rows", 1, count=1,
                                    sheet="S")
    wbx = openpyxl.load_workbook(p2)
    types = {c: wbx["S"][c].data_type for c in ("A3", "C3", "E4")}
    vals = {c: wbx["S"][c].value for c in ("A3", "C3", "E4")}
    wbx.close()
    check(all(t == "s" for t in types.values()),
          f"3.5e copy, move and an insert leave neutralized text as TEXT, "
          f"never re-armed into a live formula ({types}, {vals})")
    oc = comtier.com_validate_opens_clean(str(p2))
    check(oc["opens_clean"] is True,
          "3.5f Excel opens the neutralized workbook clean")
    comtier.recalculate(str(p2))
    survived = cached_many(p2, {"S": ["A3", "C3", "E4"]})
    check(survived == {"S!A3": "=cmd|'/c calc'!A1",
                       "S!C3": "=cmd|'/c calc'!A1", "S!E4": "=1+1"},
          f"3.5g Excel itself holds the neutralized cells as TEXT after a "
          f"real recalculation, computing nothing ({survived})")


# --------------------------------------------- item 4: sort/filter vs Excel

MIXED = [42, "apple", None, True, -7, "10", False, "Zebra", 1.5, "banana"]


def _excel_sort(path: Path, sheet: str, rng: str, key_col: str,
                order: int) -> list:
    """Sort a range with EXCEL's own Sort and read the key column back."""
    def body(manager):
        w = manager.acquire()
        wb = None
        try:
            wb = com_session.open_workbook(w.app, str(path))
            ws = wb.Worksheets(sheet)
            rg = ws.Range(rng)
            rg.Sort(Key1=ws.Range(key_col), Order1=order, Header=2)
            wb.Save()
            n = rg.Rows.Count
            first = rg.Cells(1, 1).Row
            col = rg.Cells(1, 1).Column
            return [ws.Cells(first + i, col).Value for i in range(n)]
        finally:
            if wb is not None:
                wb.Close(SaveChanges=False)
            manager.release_to_pool(w)
    return com_session.run_com("excel sort", body)


def _excel_autofilter(path: Path, sheet: str, rng: str, field: int,
                      criteria: str) -> list[int]:
    """Apply Excel's own AutoFilter and report which data rows it hides."""
    def body(manager):
        w = manager.acquire()
        wb = None
        try:
            wb = com_session.open_workbook(w.app, str(path))
            ws = wb.Worksheets(sheet)
            rg = ws.Range(rng)
            rg.AutoFilter(Field=field, Criteria1=criteria)
            first = rg.Cells(1, 1).Row
            n = rg.Rows.Count
            hidden = [first + i for i in range(1, n)
                      if ws.Rows(first + i).Hidden]
            return hidden
        finally:
            if wb is not None:
                wb.Close(SaveChanges=False)
            manager.release_to_pool(w)
    return com_session.run_com("excel autofilter", body)


def item4(scratch: Path) -> None:
    print("\n=== ITEM 4: sort / filter semantics against Excel itself ===")

    def build(p: Path) -> None:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "S"
        ws["A1"] = "K"
        ws["B1"] = "tag"
        for i, v in enumerate(MIXED, start=2):
            ws.cell(i, 1, v)
            ws.cell(i, 2, i)
        wb.save(p)
        wb.close()

    for order, label, com_order in (("asc", "ascending", 1),
                                    ("desc", "descending", 2)):
        mine = scratch / f"sort_mine_{order}.xlsx"
        theirs = scratch / f"sort_excel_{order}.xlsx"
        build(mine)
        build(theirs)
        sortfilter.sort_range(str(mine), "A1:B11",
                              keys=[{"column": "K", "order": order}],
                              sheet="S")
        wbx = openpyxl.load_workbook(mine)
        ours = [wbx["S"].cell(r, 1).value for r in range(2, 12)]
        wbx.close()
        excel = _excel_sort(theirs, "S", "A2:B11", "A2", com_order)
        same = [(a == b) or (a is None and b is None) for a, b in
                zip(ours, excel)]
        print(f"  ours  ({order}): {ours}")
        print(f"  excel ({order}): {excel}")
        check(all(same), f"4.1 sort_range {label} matches Excel's own Sort "
                         f"on a mixed-type column")

    # ---- blanks: Excel puts them last in BOTH directions
    wbx = openpyxl.load_workbook(scratch / "sort_mine_desc.xlsx")
    last = wbx["S"].cell(11, 1).value
    wbx.close()
    check(last is None,
          f"4.2 blanks sort LAST on a descending sort, as in Excel "
          f"(got {last!r})")

    # ---- numeric filter criteria vs Excel's AutoFilter
    fmine = scratch / "filter_mine.xlsx"
    fthem = scratch / "filter_excel.xlsx"

    def build_num(p: Path) -> None:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "S"
        ws["A1"] = "N"
        for i, v in enumerate([5, 15, None, 25, 35], start=2):
            ws.cell(i, 1, v)
        wb.save(p)
        wb.close()

    build_num(fmine)
    build_num(fthem)
    sortfilter.set_filter(str(fmine), "A1:A6",
                          criteria=[{"column": "N", "op": "gt", "value": 15}],
                          sheet="S")
    wbx = openpyxl.load_workbook(fmine)
    ours_hidden = [r for r in range(2, 7)
                   if wbx["S"].row_dimensions[r].hidden]
    wbx.close()
    excel_hidden = _excel_autofilter(fthem, "S", "A1:A6", 1, ">15")
    print(f"  ours hidden: {ours_hidden} | excel hidden: {excel_hidden}")
    check(ours_hidden == excel_hidden,
          f"4.3 set_filter 'greater than 15' hides exactly the rows Excel's "
          f"AutoFilter hides (ours {ours_hidden}, excel {excel_hidden})")


# ------------------------------------------------- item 5: calc settings

def item5(scratch: Path) -> None:
    print("\n=== ITEM 5: calc-settings battery ===")
    # ---- 5.1 manual mode: an explicit CalculateFull must still populate
    p = scratch / "manual.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "M"
    ws["A1"] = 7
    ws["A2"] = "=A1*3"
    wb.calculation.calcMode = "manual"
    wb.save(p)
    wb.close()
    comtier.recalculate(str(p))
    check(cached(p, "M", "A2") == 21,
          "5.1 manual-mode workbook: explicit CalculateFull populated the "
          "cache anyway")
    rep = properties.set_workbook_properties(str(p))
    check(rep["calc"]["calc_mode"] == "manual" and "caveat" in rep["calc"],
          f"5.2 set_workbook_properties reports calc_mode honestly with the "
          f"manual caveat ({rep['calc']['calc_mode']})")

    # ---- 5.3 the mode survives the server's own save
    cells.set_cell(str(p), "A1", 8, sheet="M")
    rep = properties.set_workbook_properties(str(p))
    check(rep["calc"]["calc_mode"] == "manual",
          "5.3 a file-tier edit does not silently flip the calc mode")

    # ---- 5.4 fullCalcOnLoad is set by a formula write and honored by Excel
    p2 = scratch / "fullcalc.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "F"
    ws["A1"] = 6
    wb.save(p2)
    wb.close()
    formulas.set_formula(str(p2), "A2", "=A1*7", sheet="F")
    wbx = openpyxl.load_workbook(p2)
    flag = bool(wbx.calculation.fullCalcOnLoad)
    wbx.close()
    check(flag, "5.4a a formula write flags fullCalcOnLoad for the next open")

    def open_and_save_untouched(path):
        """Open in Excel and save WITHOUT asking for a calculation: only the
        fullCalcOnLoad flag can populate the cache here."""
        def body(manager):
            w = manager.acquire()
            wbx = None
            try:
                wbx = com_session.open_workbook(w.app, str(path))
                wbx.Save()
                return True
            finally:
                if wbx is not None:
                    wbx.Close(SaveChanges=False)
                manager.release_to_pool(w)
        return com_session.run_com("open+save", body)

    open_and_save_untouched(p2)
    check(cached(p2, "F", "A2") == 42,
          "5.4b Excel recalculated on open because of the flag (no explicit "
          "calc was issued)")

    # ---- 5.5 iterative calc round-trips and actually converges in Excel
    p3 = scratch / "iter.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "I"
    ws["A1"] = 1
    wb.save(p3)
    wb.close()
    properties.set_workbook_properties(
        str(p3), iterative_calc=True, max_iterations=120, max_change=0.0001)
    formulas.set_formula(str(p3), "B1", "=B1+1", sheet="I")
    rep = properties.set_workbook_properties(str(p3))
    check(rep["calc"]["iterative_calc"] is True
          and rep["calc"]["max_iterations"] == 120
          and abs(rep["calc"]["max_change"] - 0.0001) < 1e-12,
          f"5.5a iterative-calc settings round-trip ({rep['calc']})")
    r = comtier.recalculate(str(p3))
    v = cached(p3, "I", "B1")
    check(isinstance(v, (int, float)) and not isinstance(v, bool) and v > 1,
          f"5.5b Excel converged the circular reference instead of erroring "
          f"(B1={v!r})")
    check(r["changed"].get("iterative_cells_reseeded", 0) >= 1
          and any("seed" in w for w in r.get("warnings", [])),
          f"5.5c recalculate says it re-seeded the unseeded circular cell "
          f"({r['changed'].get('iterative_cells_reseeded')})")

    # ---- 5.6 volatile freshness: a volatile recomputes on every calc
    p4 = scratch / "volatile.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "V"
    wb.save(p4)
    wb.close()
    formulas.set_formula(str(p4), "A1", "=RAND()", sheet="V")
    formulas.set_formula(str(p4), "A2", "=NOW()", sheet="V")
    comtier.recalculate(str(p4))
    first = (cached(p4, "V", "A1"), cached(p4, "V", "A2"))
    time.sleep(1.2)
    comtier.recalculate(str(p4))
    second = (cached(p4, "V", "A1"), cached(p4, "V", "A2"))
    check(first[0] != second[0] and first[1] != second[1],
          f"5.6a volatiles recompute on every recalculate ({first} -> "
          f"{second})")
    au = formulas.audit_formulas(str(p4), sheet="V")
    check(au["volatile"]["count"] == 2,
          f"5.6b audit_formulas names both volatile cells so a caller knows "
          f"the cached number is a snapshot ({au['volatile']['count']})")


def item2b(scratch: Path) -> None:
    """The two defects this gate uncovered, as standing proofs."""
    print("\n=== ITEM 2b: the _xlpm and empty-cache defects ===")
    # ---- LAMBDA / LET parameter names: bare names made a workbook Excel
    #      REFUSED TO OPEN AT ALL, not merely a #NAME?.
    p = scratch / "lambda.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "L"
    for r, v in enumerate((2, 3, 4), start=1):
        ws.cell(r, 1, v)
    wb.save(p)
    wb.close()
    for addr, f in (("C1", "=LET(x,A1,y,A2,x*y)"),
                    ("C2", "=LAMBDA(v,v*2)(21)"),
                    ("C3", "=SUM(MAP(A1:A3,LAMBDA(v,v*2)))"),
                    ("C4", "=REDUCE(0,A1:A3,LAMBDA(acc,v,acc+v))"),
                    ("C5", "=INDEX(MAKEARRAY(2,2,LAMBDA(r,c,r*c)),2,2)")):
        formulas.set_formula(str(p), addr, f, sheet="L", backup=False)
    stored = formulas_of(p, {"L": ["C1"]})["L!C1"]
    check("_xlpm.x" in str(stored),
          f"2b.1 LET / LAMBDA parameter names are stored with the _xlpm "
          f"prefix Excel requires ({stored})")
    oc = comtier.com_validate_opens_clean(str(p))
    check(oc["opens_clean"] is True,
          f"2b.2 Excel OPENS the LAMBDA-family workbook "
          f"({oc.get('excel_says', '')[:60]})")
    comtier.recalculate(str(p))
    got = cached_many(p, {"L": ["C1", "C2", "C3", "C4", "C5"]})
    check(got == {"L!C1": 6, "L!C2": 42, "L!C3": 18, "L!C4": 9, "L!C5": 4},
          f"2b.3 every LAMBDA-family formula computes correctly ({got})")

    # ---- the empty <v></v> artifact: harmless under a normal recalc, fatal
    #      under iterative calculation.
    import zipfile
    with zipfile.ZipFile(p) as zf:
        sheet_xml = zf.read("xl/worksheets/sheet1.xml").decode("utf-8")
    check("<v></v>" not in sheet_xml and "<v/>" not in sheet_xml,
          "2b.4 no empty cached-value element is left on a formula cell")


def item2c(scratch: Path) -> None:
    """The edge-audit follow-up defects (2026-09-04), against real Excel."""
    print("\n=== ITEM 2c: optional LAMBDA params, dynamic-array reseed, "
          "error sort ===")

    # ---- C1: OPTIONAL lambda parameters. Excel stores the declaration under
    #      a THIRD prefix, _xlop., with the brackets removed, and every use
    #      site under _xlpm. A bare [y] used to be dropped entirely, which is
    #      the class where Excel REFUSES TO OPEN the workbook.
    p = scratch / "optlambda.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "L"
    wb.save(p)
    wb.close()
    want = {"L!A1": 3, "L!A2": 1, "L!A3": 6, "L!A4": 2, "L!A5": 6, "L!A6": 5}
    for addr, f in (
            ("A1", "=LAMBDA(x,[y],IF(ISOMITTED(y),x,x+y))(1,2)"),
            ("A2", "=LAMBDA(x,[y],x)(1)"),
            ("A3", "=LAMBDA(x,[f],f(x))(2,LAMBDA(a,a*3))"),
            ("A4", "=LET(q,1,LAMBDA(y,[z],y+q)(q,2))"),
            ("A5", "=LAMBDA(x,[y],[z],IF(ISOMITTED(z),x+y,x+y+z))(1,2,3)"),
            ("A6", "=LAMBDA(x,[y],IF(ISOMITTED(y),x,x+y))(5)")):
        formulas.set_formula(str(p), addr, f, sheet="L", backup=False)
    stored = str(formulas_of(p, {"L": ["A1"]})["L!A1"])
    check("_xlop.y" in stored and "_xlpm.y" in stored and "[y]" not in stored,
          f"2c.1 an optional LAMBDA parameter is stored _xlop. at the "
          f"declaration and _xlpm. at every use ({stored})")
    oc = comtier.com_validate_opens_clean(str(p))
    check(oc["opens_clean"] is True,
          f"2c.2 Excel OPENS the optional-LAMBDA workbook "
          f"({oc.get('excel_says', '')[:60]})")
    comtier.recalculate(str(p))
    got = cached_many(p, {"L": ["A1", "A2", "A3", "A4", "A5", "A6"]})
    check(got == want,
          f"2c.3 every optional-LAMBDA formula EVALUATES correctly "
          f"(got {got}, want {want})")
    shown = cells.read_range(str(p), "A1:A6", values="formula",
                             sheet="L")["values"]
    check(all("_xl" not in str(r[0]) for r in shown)
          and str(shown[0][0]) == "=LAMBDA(x,[y],IF(ISOMITTED(y),x,x+y))(1,2)",
          f"2c.4 the read-back displays the formula the way Excel's formula "
          f"bar does, brackets restored ({shown[0][0]!r})")

    # ---- S1: the iterative reseed must not convert a DYNAMIC array. HasArray
    #      is False for one, so the CSE skip never covered it, and assigning
    #      legacy .Formula re-enters in implicit-intersection mode.
    def _reseed_body(manager):
        w = manager.acquire()
        wb2 = None
        try:
            app = w.app
            wb2 = app.Workbooks.Add()
            s = wb2.Worksheets(1)
            for r, v in enumerate((1, 2, 3), start=1):
                s.Cells(r, 1).Value = v
            s.Range("C1").Formula2 = "=FILTER(A1:A3,A1:A3>99)"   # #CALC!
            app.CalculateFull()
            before = str(s.Range("C1").Formula2)
            touched = comtier._reseed_iterative_errors(app, wb2)
            after = str(s.Range("C1").Formula2)
            return before, after, touched
        finally:
            if wb2 is not None:
                try:
                    wb2.Close(SaveChanges=False)
                except Exception:  # noqa: BLE001
                    pass
            manager.release_to_pool(w)

    before, after, touched = com_session.run_com("reseed probe", _reseed_body)
    check(after == before and "@" not in after,
          f"2c.5 the iterative reseed leaves a DYNAMIC-array formula "
          f"unchanged (before={before!r} after={after!r}, "
          f"{touched} cell(s) reseeded)")
    note(f"S1 confirmed and fixed: legacy .Formula assignment turned "
         f"{before!r} into '=@FILTER(A1:A3,A1:A3>99)' and dropped the "
         f"t=\"array\" spill attributes in a pre-fix probe; .Formula2 "
         f"round-trips it ({after!r})")

    # ---- S2: does Excel order error values among themselves in a sort?
    def _errsort_body(manager):
        w = manager.acquire()
        wb3 = None
        try:
            app = w.app
            wb3 = app.Workbooks.Add()
            s = wb3.Worksheets(1)
            # deliberately NOT alphabetical, so a sort that reorders the
            # error run is distinguishable from one that does not
            for r, f in enumerate(("=\"a\"+1", "=A99+#REF!", "=NA()",
                                   "=1/0"), start=1):
                try:
                    s.Cells(r, 1).Formula = f
                except Exception:  # noqa: BLE001
                    s.Cells(r, 1).Value = f
            app.CalculateFull()
            pre = [str(s.Cells(r, 1).Text) for r in range(1, 5)]
            s.Range("A1:A4").Sort(Key1=s.Range("A1"), Order1=1,
                                  Header=2)
            post = [str(s.Cells(r, 1).Text) for r in range(1, 5)]
            return pre, post
        finally:
            if wb3 is not None:
                try:
                    wb3.Close(SaveChanges=False)
                except Exception:  # noqa: BLE001
                    pass
            manager.release_to_pool(w)

    pre, post = com_session.run_com("error sort probe", _errsort_body)
    ours = sorted(pre, key=cells._sort_key)
    check(post == pre,
          f"2c.6 Excel preserves the original order of error values in a "
          f"sort (before={pre}, after={post})")
    check(ours == post,
          f"2c.7 _sort_key reproduces that: every error keys equal, so the "
          f"stable sort leaves them alone (ours={ours}, excel={post})")


# --------------------------------------- item 6: verify_com end to end

def item6(scratch: Path) -> None:
    print("\n=== ITEM 6: verify_com deep verification, end to end ===")
    from xlsx_mcp.core import refs as _refs
    corpus = ROOT / "tests" / "fixtures" / "corpus"
    for name in ("datavalidation.xlsx", "condformat.xlsx", "chart.xlsx",
                 "clean.xlsx"):
        src = corpus / name
        if not src.exists():
            continue
        p = scratch / f"v6_{name}"
        shutil.copy2(src, p)
        try:
            pkg = WorkbookPackage.open(str(p))
            sheet = pkg.workbook.sheetnames[0]
            ws = pkg.workbook[sheet]
            ws["H1"] = 6
            ws["H2"] = 7
            pkg.set_formula(sheet, "H3", "=H1*H2")
            r1 = pkg.save(verify_com=True)

            pkg = WorkbookPackage.open(str(p))
            pkg.modify_structure(_refs.RefEdit(sheet, _refs.INSERT_ROWS,
                                               index=1, count=2))
            r2 = pkg.save(verify_com=True)
        except Exception as exc:  # noqa: BLE001
            note(f"item 6: {name} is not editable on the file tier "
                 f"({type(exc).__name__}: {str(exc)[:80]}); trying the next "
                 "fixture")
            continue
        check(r1["ok"] and r1.get("verified_com") is True,
              f"6.1 formula write on the {name} corpus fixture passed "
              f"Excel's own open check (verify_com)")
        check(r2["ok"] and r2.get("verified_com") is True,
              f"6.2 structural edit on {name} + save(verify_com=True) "
              f"round trip")
        comtier.recalculate(str(p))
        v = cached(p, sheet, "H5")
        check(v == 42,
              f"6.3 the number survived the structural edit and Excel "
              f"computes it at its new address (H5={v!r})")
        return
    check(False, "6.0 a rich corpus fixture survived a file-tier edit")


def item7(scratch: Path) -> None:
    """ITEM 7: THE REFUSE-CLASS BATTERY (insane round, 2026-09-05).

    The round's headline was ten single tool calls that returned
    ok/saved/verified and produced a workbook Excel answers "Open method of
    Workbooks class failed" for. Each route is replayed here through the
    SHIPPED tool bodies, and the gate asserts the only two acceptable
    outcomes: a loud refusal that leaves the file byte-identical, or a file
    EXCEL ITSELF OPENS. Nothing in between."""
    print("\n=== ITEM 7: refuse-class battery (Excel is the judge) ===")
    import hashlib

    from xlsx_mcp.core.errors import (
        ExcelWouldRefuse, FormulaRejected, UnsupportedStructure,
    )
    from xlsx_mcp.ops import (
        annotations, datavalidation, objects, pagelayout, search, sortfilter,
    )
    from openpyxl.worksheet.formula import ArrayFormula

    def md5(p) -> str:
        return hashlib.md5(Path(p).read_bytes()).hexdigest()

    def fresh(name: str, sheets=("Data",)) -> Path:
        p = scratch / f"r7_{name}.xlsx"
        wb = openpyxl.Workbook()
        wb.active.title = sheets[0]
        for extra in sheets[1:]:
            wb.create_sheet(extra)
        for i in range(1, 4):
            wb[sheets[0]].cell(i, 1).value = i
        wb.save(p)
        return p

    def excel_opens(p) -> bool:
        try:
            return bool(com_session.opens_clean(str(p)).get("opens_clean"))
        except Exception as exc:  # noqa: BLE001
            note(f"item 7: opens_clean raised on {Path(p).name}: {exc}")
            return False

    REFUSALS = (ExcelWouldRefuse, FormulaRejected, UnsupportedStructure)

    def route(label: str, name: str, call, sheets=("Data",)):
        """Run one route; assert refusal-with-file-untouched OR Excel opens."""
        p = fresh(name, sheets)
        before = md5(p)
        try:
            call(str(p))
        except REFUSALS as exc:
            ok = md5(p) == before
            check(ok, f"7.{label} refused ({type(exc).__name__}) and left the "
                      f"file byte-identical")
            return
        check(excel_opens(p),
              f"7.{label} produced a file EXCEL OPENS (no refusal was raised)")

    # --- C-1: quoted sheet names that used to disable the _xlpm pass -------
    for i, formula in enumerate((
            "=LET(x,'Q1 (draft'!A1,x+1)",
            "=LET(x,'It''s (fun'!A1,x+1)",
            "=LAMBDA(x,'Q1 (draft'!A1+x)(1)",
    ), start=1):
        route(f"1{i} C-1 {formula}", f"c1_{i}",
              lambda pp, f=formula: formulas.set_formula(
                  pp, {"cell": "A1"}, f, sheet="Data"),
              sheets=("Data", "Q1 (draft", "It's (fun"))

    # --- C-2: whitespace before a builtin's paren -------------------------
    route("2 C-2 spaced builtin call", "c2",
          lambda pp: formulas.set_formula(
              pp, {"cell": "A1"}, "=LET(sum,2,sum+SUM (1,2))", sheet="Data"))

    # --- H-2: a defined name colliding with a LET parameter ---------------
    p = fresh("h2")
    from xlsx_mcp.ops import names as _names
    _names.manage_name(str(p), "add", name="Rate", refers_to="=Data!$A$1")
    formulas.set_formula(str(p), {"cell": "B1"}, "=LET(Rate,1,Rate)+Rate",
                         sheet="Data")
    stored = formulas_of(p, {"Data": ["B1"]})["Data!B1"]
    check(stored == "=_xlfn.LET(_xlpm.Rate,1,_xlpm.Rate)+Rate",
          f"7.3a H-2 the out-of-scope name stays bare, as Excel stores it "
          f"({stored!r})")
    check(excel_opens(p), "7.3b H-2 the scoped output opens in Excel")
    comtier.recalculate(str(p))
    v = cached(p, "Data", "B1")
    check(v == 2, f"7.3c H-2 Excel computes LET(Rate,1,Rate)+Rate = 2 "
                  f"(got {v!r})")

    # --- C-3 / H-3: legacy CSE arrays --------------------------------------
    def cse(name: str) -> Path:
        p = scratch / f"r7_{name}.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Data"
        for i in range(1, 6):
            ws.cell(i, 1).value = i
            ws.cell(i, 4).value = 6 - i
        ws["B1"] = ArrayFormula("B1:B5", "=A1:A5*2")
        ws["C1"] = ArrayFormula("C1", "=SUM(A1:A5*2)")
        wb.save(p)
        return p

    for label, name, call in (
        ("4a sort over a CSE array", "cse_sort",
         lambda pp: sortfilter.sort_range(pp, {"range": "A1:D5"},
                                          [{"column": 4, "order": "asc"}],
                                          has_header=False)),
        ("4b delete_rows across a CSE array", "cse_del",
         lambda pp: structure.modify_grid_structure(pp, "delete_rows", 1, 1)),
    ):
        p = cse(name)
        before = md5(p)
        try:
            call(str(p))
            check(excel_opens(p), f"7.{label}: no refusal, so Excel must open "
                                  "the result")
        except REFUSALS:
            check(md5(p) == before,
                  f"7.{label}: refused and left the file byte-identical")

    p = cse("cse_copy")
    cells.copy_range(str(p), {"range": "A1:B5"}, {"cell": "F1"}, what="all")
    check(excel_opens(p),
          "7.4c a CSE array copied with its anchor rebased opens in Excel")
    comtier.recalculate(str(p))
    got = cached_many(p, {"Data": ["B1", "G1", "G5"]})
    check(approx(got["Data!B1"], 2) and approx(got["Data!G1"], 2),
          f"7.4d the copied array computes at its new address ({got})")

    p = cse("cse_replace")
    search.replace_cells(str(p), "A1:A5", "A1:A4", look_in="formulas")
    check(excel_opens(p), "7.4e replace_cells kept the array and Excel opens "
                          "the result")
    comtier.recalculate(str(p))
    v = cached(p, "Data", "C1")
    check(approx(v, 20),
          f"7.4f the rewritten array still computes as an ARRAY (C1={v!r}; "
          "the de-arrayed form gives 2 by implicit intersection)")

    # --- C-4: the five parameter routes ------------------------------------
    route("5a oversize comment", "c4_comment",
          lambda pp: annotations.manage_comment(
              pp, "add", location={"cell": "A1"}, text="z" * 60000,
              sheet="Data"))
    route("5b oversize DV list", "c4_dv",
          lambda pp: datavalidation.manage_data_validation(
              pp, "add", location={"range": "A1:A3"}, dv_type="list",
              values=[f"x{i}" for i in range(10000)], sheet="Data"))
    route("5c oversize header", "c4_hf",
          lambda pp: pagelayout.set_header_footer(
              pp, sheet="Data", header={"left": "&Z&bad" * 500}))
    route("5d invalid hyperlink authority", "c4_link",
          lambda pp: annotations.manage_hyperlink(
              pp, "add", location={"cell": "A1"},
              target='http://x"/><evil a="', sheet="Data"))

    png = scratch / "r7.png"
    if not png.exists():
        png.write_bytes(bytes.fromhex(
            "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c489"
            "0000000a49444154789c636000000200010005fe02fe0000000049454e44ae426082"))
    route("5e negative image extents", "c4_img",
          lambda pp: objects.manage_image(
              pp, "insert", location={"cell": "B2"}, image_file=str(png),
              width=-5, height=-5, sheet="Data"))

    # --- the boundary must still be usable ---------------------------------
    p = fresh("c4_ok")
    annotations.manage_comment(str(p), "add", location={"cell": "A1"},
                               text="z" * 32767, sheet="Data")
    check(excel_opens(p),
          "7.6a a comment at Excel's exact 32,767 limit still opens")
    p = fresh("c4_ok2")
    annotations.manage_hyperlink(str(p), "add", location={"cell": "A1"},
                                 target="https://example.com/a b",
                                 sheet="Data")
    check(excel_opens(p),
          "7.6b a legal hyperlink with a space in the PATH still opens")

    # --- M-1: the file tier no longer vouches for what it did not check ----
    from xlsx_mcp.ops import validation as _validation
    f = _validation.validate(str(fresh("m1")), checks=["structure"])[
        "results"]["structure"]["findings"]
    check(f["opens_clean"] == _validation.NOT_CHECKED,
          "7.7 validate reports opens_clean as 'not checked' rather than "
          "claiming an Excel verdict it never asked for")


# --------------------------------------------------------------- driver

def main() -> int:
    baseline = list_excel_pids()
    if baseline:
        print(f"SKIPPED: foreign EXCEL.EXE present at start {sorted(baseline)}")
        return 0

    scratch = Path(tempfile.mkdtemp(prefix="ks4xl_numbers_"))
    print(f"scratch: {scratch}")
    com_session._EXECUTOR = com_session.ComExecutor(
        journal_path=scratch / "_gate_journal.json")
    ex = com_session.get_executor()
    owned: list[int] = []

    def _excel_version():
        def body(manager):
            w = manager.acquire()
            try:
                return f"{w.app.Version} build {w.app.Build}"
            finally:
                manager.release_to_pool(w)
        try:
            return com_session.run_com("excel version", body)
        except Exception as exc:  # noqa: BLE001
            return f"unknown ({exc})"

    print(f"excel: {_excel_version()}")

    try:
        for fn in (item1, item2, item2b, item2c, item3, item4, item5, item6,
                   item7):
            try:
                fn(scratch)
            except Exception as exc:  # noqa: BLE001
                import traceback
                traceback.print_exc()
                check(False, f"{fn.__name__} raised "
                             f"{type(exc).__name__}: {exc}")
        st = comtier.com_status()
        print(f"\n  com status: pooled_pid={st['pooled_instance_pid']} "
              f"ops={st['ops_completed']} timeouts={st['timeouts']}")
        check(st["timeouts"] == 0, "gate: zero COM timeouts")
    finally:
        # Snapshot the journal BEFORE shutdown. Shutdown now reaps the pooled
        # worker by owned PID and forgets each PID once it is confirmed gone
        # (the H-4 fix), so reading the journal afterwards would report
        # owned=[] and then count the still-terminating process as FOREIGN.
        # The accounting this gate exists to do needs the pre-shutdown list.
        mgr = com_session.ExcelInstanceManager(
            journal_path=scratch / "_gate_journal.json")
        owned = sorted(mgr.journal.owned_pids())
        ex.shutdown()
        owned = sorted(set(owned) | set(mgr.journal.owned_pids()))
        for pid in owned:
            if pid_alive(pid):
                taskkill(pid)
        deadline = time.monotonic() + 60
        while time.monotonic() < deadline and any(
                pid_alive(p) for p in owned):
            time.sleep(1.0)

    alive_mine = [p for p in owned if pid_alive(p)]
    foreign = sorted(list_excel_pids() - set(owned))
    check(not alive_mine, f"orphan check: no owned PID alive (owned={owned})")
    check(not foreign,
          f"orphan check: no foreign EXCEL.EXE touched/left ({foreign})")

    if NOTES:
        print("\nNOTES:")
        for n in NOTES:
            print(f"  - {n}")
    if FAILS:
        print(f"\nVERDICT numbers_safety_gate FAIL ({len(FAILS)}):")
        for f in FAILS:
            print(f"  - {f}")
        return 1
    print(f"\nVERDICT numbers_safety_gate PASS ({len(owned)} PIDs journaled, "
          "0 orphans by owned PID, 0 foreign touched)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
