"""COM gate: the Phase 5 COM tier proof against real Excel.

Everything runs through the shipped tool bodies (ops/comtier over
com/session's serialized executor), not ad-hoc COM. Proves, on this machine:

  1.  recalculate: openpyxl-written formulas (absent cache) gain Excel-computed
      cached values; explicit CalculateFull fires even under MANUAL calc mode.
  2.  PRE-COM CHECKLIST (a): the _xlfn/_xlws shim verification in Excel.
      FILTER and SORT written as _xlfn._xlws.*, SORTBY/UNIQUE/SEQUENCE/
      RANDARRAY as plain _xlfn.* (the audit's correction): the file must open
      with zero repair prompts and every formula must compute (no #NAME?).
  3.  com_manage_pivot: create a REAL pivot from a source range, verify via
      list and by reading the pivot's computed cells through Excel, edit the
      source, refresh, see the numbers move, delete.
  4.  com_export_pdf (workbook + range), com_render_sheet (PNG), and
      com_convert_format (xlsb round-trip, csv values).
  5.  com_save_with_password: real encryption (plain zip open fails), reopen
      with password, opens_clean without password refuses honestly, decrypt.
  6.  com_autofit widens a squeezed column; com_goal_seek converges;
      com_set_sparkline creates and lists a real sparkline group.
  7.  com_validate_opens_clean: clean file passes; a corrupted file fails
      HONESTLY with Excel's own message; verify_com wiring in package.save.
  8.  PRE-COM CHECKLIST (b): the stale ~$ lockfile experiment. A stale
      lockfile (no real hold) degrades to a warning on both the file tier and
      the COM tier; a REAL hold (a worker holding the book open) refuses with
      WORKBOOK_LOCKED semantics.
  9.  Volume: enough pooled operations to exceed 50 COM ops total, then the
      zero-orphan accounting (owned PIDs only, no foreign EXCEL.EXE touched).

PID DISCIPLINE: private DispatchEx workers only, journaled; SKIPS honestly if
a foreign EXCEL.EXE is present at start; never touches a foreign PID.
"""

from __future__ import annotations

import sys
import time
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(ROOT / "src"))

import openpyxl  # noqa: E402

from xlsx_mcp.com import session as com_session  # noqa: E402
from xlsx_mcp.com.instances import (  # noqa: E402
    list_excel_pids, pid_alive, taskkill,
)
from xlsx_mcp.core import calc as core_calc  # noqa: E402
from xlsx_mcp.core.errors import WorkbookLocked, XlMcpError  # noqa: E402
from xlsx_mcp.core.package import WorkbookPackage  # noqa: E402
from xlsx_mcp.ops import comtier  # noqa: E402

FAILS: list[str] = []


def check(cond: bool, label: str) -> None:
    if cond:
        print(f"PASS {label}")
    else:
        print(f"FAIL {label}")
        FAILS.append(label)


def _data_only(path, sheet, coord):
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        return wb[sheet][coord].value
    finally:
        wb.close()


def main() -> int:  # noqa: PLR0915
    baseline = list_excel_pids()
    if baseline:
        print(f"SKIPPED: foreign EXCEL.EXE present at start {sorted(baseline)}")
        return 0

    scratch = Path(__import__("tempfile").mkdtemp(prefix="ks4xl_comtier_"))
    # A gate-scoped executor + journal so the accounting below owns exactly
    # this run's PIDs.
    com_session._EXECUTOR = com_session.ComExecutor(
        journal_path=scratch / "_gate_journal.json")
    ex = com_session.get_executor()

    try:
        # ============ 1. recalculate (absent cache -> computed) ============
        p_calc = scratch / "calc.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Data"
        ws["A1"] = 2
        ws["A2"] = 40
        ws["A3"] = "=A1+A2"
        wb.create_sheet("Summary")["B1"] = "=Data!A1*10"
        wb.save(p_calc)
        wb.close()
        check(_data_only(p_calc, "Data", "A3") is None,
              "1 setup: openpyxl formula has NO cached value")
        r = comtier.recalculate(str(p_calc))
        check(r["ok"] and r["recalculated"] and r["engine"] == "com",
              "1 recalculate: COM engine ran and saved")
        check(_data_only(p_calc, "Data", "A3") == 42,
              "1 recalculate: A3 cache = 42 (Excel-computed)")
        check(_data_only(p_calc, "Summary", "B1") == 20,
              "1 recalculate: cross-sheet cache = 20")
        check(r["changed"]["formula_cells_without_cached_value_after"] == 0,
              "1 recalculate: zero absent-cache cells remain")

        # ---- manual calc mode: explicit CalculateFull must still fire ----
        p_man = scratch / "manual.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "M"
        ws["A1"] = 7
        ws["A2"] = "=A1*3"
        wb.calculation.calcMode = "manual"
        wb.save(p_man)
        wb.close()
        comtier.recalculate(str(p_man))
        check(_data_only(p_man, "M", "A2") == 21,
              "1 manual mode: explicit CalculateFull populated the cache")

        # ============ 2. CHECKLIST (a): _xlfn/_xlws in-Excel proof ========
        p_x = scratch / "xlws.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "F"
        for i, v in enumerate((10, 250, 30), start=1):
            ws.cell(row=i + 1, column=1, value=f"k{i}")
            ws.cell(row=i + 1, column=2, value=v)
        ws["D1"] = "k2"
        shim_cases = {
            "E1": ("=XLOOKUP(D1,A2:A4,B2:B4)", "_xlfn.XLOOKUP("),
            "F1": ("=FILTER(B2:B4,B2:B4>100)", "_xlfn._xlws.FILTER("),
            "F5": ("=SORT(B2:B4)", "_xlfn._xlws.SORT("),
            "G1": ("=SORTBY(B2:B4,B2:B4)", "_xlfn.SORTBY("),
            "H1": ("=UNIQUE(B2:B4)", "_xlfn.UNIQUE("),
            "I1": ("=SEQUENCE(2)", "_xlfn.SEQUENCE("),
            "J1": ("=SUM(RANDARRAY(2))", "_xlfn.RANDARRAY("),
        }
        stored_ok = True
        for coord, (formula, want) in shim_cases.items():
            normalized, _ = core_calc.normalize_formula(formula)
            if want not in normalized:
                stored_ok = False
                print(f"  shim MISS: {formula} -> {normalized}")
            ws[coord] = normalized
        wb.save(p_x)
        wb.close()
        check(stored_ok, "2 _xlws: shim stores FILTER/SORT as _xlfn._xlws "
                         "and SORTBY/UNIQUE/SEQUENCE/RANDARRAY as plain _xlfn")
        oc = comtier.com_validate_opens_clean(str(p_x))
        check(oc["opens_clean"] is True,
              "2 _xlws: Excel opens the shimmed workbook with ZERO repair "
              "prompts")
        comtier.recalculate(str(p_x))
        vals = {c: _data_only(p_x, "F", c) for c in
                ("E1", "F1", "F5", "G1", "H1", "I1", "J1")}
        print(f"  computed: {vals}")
        check(vals["E1"] == 250, "2 _xlws: XLOOKUP computed (250)")
        check(vals["F1"] == 250, "2 _xlws: FILTER spilled and computed (250)")
        check(vals["F5"] == 10, "2 _xlws: SORT computed (10 first)")
        check(vals["G1"] == 10, "2 _xlws: SORTBY computed (10 first)")
        check(vals["H1"] == 10, "2 _xlws: UNIQUE computed (10 first)")
        check(vals["I1"] == 1, "2 _xlws: SEQUENCE computed (1 first)")
        check(isinstance(vals["J1"], (int, float)) and 0 <= vals["J1"] <= 2,
              "2 _xlws: RANDARRAY computed (numeric)")
        check(all(v != "#NAME?" for v in vals.values()),
              "2 _xlws: zero #NAME? errors across the shim families")

        # ============ 3. REAL pivot round-trip ============================
        p_piv = scratch / "pivot.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sales"
        ws.append(["Region", "Product", "Amount"])
        data = [("East", "A", 100), ("East", "B", 50), ("West", "A", 70),
                ("West", "B", 30), ("East", "A", 25)]
        for row in data:
            ws.append(row)
        wb.save(p_piv)
        wb.close()
        r = comtier.com_manage_pivot(
            str(p_piv), "create", source_sheet="Sales",
            source_range="A1:C6", dest_sheet="Pivot", dest_cell="A3",
            name="GatePivot", rows=["Region"],
            values=[{"field": "Amount", "func": "sum"}])
        check(r["ok"] and r["changed"]["pivot_created"] == "GatePivot",
              "3 pivot: created a REAL pivot table through Excel")
        listing = comtier.com_manage_pivot(str(p_piv), "list")
        check(any(pt["name"] == "GatePivot"
                  for pt in listing["pivot_tables"]),
              "3 pivot: list sees the pivot with source and location")
        oc = comtier.com_validate_opens_clean(str(p_piv))
        check(oc["opens_clean"] is True, "3 pivot: file opens clean after "
                                         "pivot creation")

        def read_pivot_total(manager):
            w = manager.acquire()
            wbx = None
            try:
                wbx = com_session.open_workbook(w.app, str(p_piv),
                                                read_only=True)
                pt = wbx.Worksheets("Pivot").PivotTables("GatePivot")
                vals = pt.DataBodyRange.Value
                if isinstance(vals, tuple):
                    return float(vals[-1][-1])  # grand total, last cell
                return float(vals)
            finally:
                if wbx is not None:
                    wbx.Close(False)
                manager.release_to_pool(w)

        total = com_session.run_com("read pivot total", read_pivot_total)
        check(total == 275.0,
              f"3 pivot: grand total computed by the pivot = 275 (got {total})")

        # Edit the SOURCE with the file tier, then refresh the pivot.
        pkg = WorkbookPackage.open(str(p_piv))
        pkg.set_cell("Sales", "C2", 1100)  # was 100 -> +1000
        pkg.save()
        r = comtier.com_manage_pivot(str(p_piv), "refresh", name="GatePivot")
        check(r["ok"] and r["changed"]["refreshed"] == ["GatePivot"],
              "3 pivot: refresh after a file-tier source edit")
        total2 = com_session.run_com("read pivot total 2", read_pivot_total)
        check(total2 == 1275.0,
              f"3 pivot: refreshed total = 1275 (got {total2})")
        r = comtier.com_manage_pivot(str(p_piv), "delete", name="GatePivot")
        check(r["ok"] and r["changed"]["pivot_deleted"] == "GatePivot",
              "3 pivot: delete cleared the pivot")
        listing = comtier.com_manage_pivot(str(p_piv), "list")
        check(listing["pivot_tables"] == [],
              "3 pivot: list is empty after delete")

        # ============ 4. export / render / convert ========================
        pdf1 = scratch / "book.pdf"
        r = comtier.com_export_pdf(str(p_calc), str(pdf1))
        check(r["ok"] and pdf1.read_bytes()[:4] == b"%PDF",
              "4 export_pdf: workbook scope produced a real PDF")
        pdf2 = scratch / "range.pdf"
        r = comtier.com_export_pdf(str(p_calc), str(pdf2), scope="range",
                                   sheet="Data", range_a1="A1:A3")
        check(r["ok"] and pdf2.read_bytes()[:4] == b"%PDF",
              "4 export_pdf: range scope produced a real PDF")
        png = scratch / "render.png"
        r = comtier.com_render_sheet(str(p_calc), str(png), sheet="Data")
        check(r["ok"] and png.read_bytes()[:8] == b"\x89PNG\r\n\x1a\n",
              "4 render_sheet: produced a real PNG of the used range")
        xlsb = scratch / "conv.xlsb"
        r = comtier.com_convert_format(str(p_calc), str(xlsb))
        check(r["ok"] and r["format"] == "xlsb",
              "4 convert: xlsx -> xlsb via SaveAs")
        oc = comtier.com_validate_opens_clean(str(xlsb))
        check(oc["opens_clean"] is True, "4 convert: the xlsb opens clean")
        csvp = scratch / "conv.csv"
        r = comtier.com_convert_format(str(p_calc), str(csvp))
        text = csvp.read_text(encoding="utf-8-sig")
        check("42" in text, "4 convert: csv carries the computed values")

        # ============ 5. real encryption ==================================
        p_enc = scratch / "secret.xlsx"
        wb = openpyxl.Workbook()
        wb.active["A1"] = "classified"
        wb.save(p_enc)
        wb.close()
        r = comtier.com_save_with_password(str(p_enc), "hunter2")
        check(r["ok"] and r["changed"]["encryption"] == "applied",
              "5 encrypt: password save succeeded")
        import zipfile
        try:
            zipfile.ZipFile(p_enc).close()
            plain = True
        except Exception:
            plain = False
        check(not plain, "5 encrypt: the file is no longer a plain zip "
                         "(real encryption, not advisory)")
        oc = comtier.com_validate_opens_clean(str(p_enc))
        check(oc["opens_clean"] is False,
              "5 encrypt: opens_clean WITHOUT the password refuses honestly")
        oc = comtier.com_validate_opens_clean(str(p_enc),
                                              password="hunter2")
        check(oc["opens_clean"] is True,
              "5 encrypt: opens_clean WITH the password passes")
        r = comtier.com_save_with_password(str(p_enc), "",
                                           current_password="hunter2")
        check(r["ok"] and r["changed"]["encryption"] == "removed",
              "5 encrypt: decryption restored a verifiable plain package")
        check(_read_a1(p_enc) == "classified",
              "5 encrypt: content survived the encrypt/decrypt round-trip")

        # ============ 6. autofit / goal seek / sparklines =================
        p_fit = scratch / "fit.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "W"
        ws["A1"] = "a very long piece of text that needs a wide column"
        ws.column_dimensions["A"].width = 4
        wb.save(p_fit)
        wb.close()
        r = comtier.com_autofit(str(p_fit))
        wb = openpyxl.load_workbook(p_fit)
        width = wb["W"].column_dimensions["A"].width
        wb.close()
        check(r["ok"] and width and width > 20,
              f"6 autofit: column A widened by real metrics (width={width})")

        p_goal = scratch / "goal.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "G"
        ws["A1"] = 1
        ws["A2"] = "=A1*2"
        wb.save(p_goal)
        wb.close()
        r = comtier.com_goal_seek(str(p_goal), "A2", 42, "A1", sheet="G")
        check(r["ok"] and r["changed"]["converged"]
              and abs(r["changed"]["solution"] - 21) < 1e-6,
              "6 goal_seek: converged, A1 = 21 for A2 -> 42")
        check(_data_only(p_goal, "G", "A2") == 42,
              "6 goal_seek: saved file carries the achieved value")
        try:
            comtier.com_goal_seek(str(p_goal), "A1", 5, "A2", sheet="G")
            check(False, "6 goal_seek: non-formula target refused")
        except XlMcpError:
            check(True, "6 goal_seek: non-formula target refused")

        p_spark = scratch / "spark.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "S"
        for r_i, row in enumerate(((1, 5, 3, 8, 2), (4, 1, 6, 2, 9)),
                                  start=1):
            for c_i, v in enumerate(row, start=1):
                ws.cell(row=r_i, column=c_i, value=v)
        wb.save(p_spark)
        wb.close()
        r = comtier.com_set_sparkline(str(p_spark), "create",
                                      location="G1:G2", source="A1:E2",
                                      type="line", sheet="S")
        check(r["ok"] and r["changed"]["sparklines_created"] == "G1:G2",
              "6 sparkline: created a line sparkline group through Excel")
        listing = comtier.com_set_sparkline(str(p_spark), "list")
        groups = listing["sparkline_groups"]
        check(len(groups) >= 1 and groups[0]["type"] == "line",
              f"6 sparkline: list sees the group ({groups})")
        oc = comtier.com_validate_opens_clean(str(p_spark))
        check(oc["opens_clean"] is True, "6 sparkline: file opens clean")

        # ============ 7. opens_clean + verify_com wiring ==================
        p_bad = scratch / "corrupt.xlsx"
        good = p_calc.read_bytes()
        p_bad.write_bytes(good[: len(good) // 2])  # truncated package
        oc = comtier.com_validate_opens_clean(str(p_bad))
        check(oc["opens_clean"] is False and oc.get("excel_says"),
              f"7 opens_clean: corrupt file fails HONESTLY with Excel's "
              f"message ({oc.get('excel_says', '')[:60]}...)")

        pkg = WorkbookPackage.open(str(p_calc))
        pkg.set_cell("Data", "B1", 123)
        result = pkg.save(verify_com=True)
        check(result["ok"] and result.get("verified_com") is True,
              "7 verify_com: package.save(verify_com=True) deep-verified "
              "through a real Excel open")

        # ============ 8. CHECKLIST (b): stale lockfile experiment =========
        stale = scratch / ("~$" + p_calc.name)
        stale.write_bytes(b"stale owner record from a crashed session")
        pkg = WorkbookPackage.open(str(p_calc))
        pkg.set_cell("Data", "B2", 5)
        result = pkg.save()
        check(result["ok"] and any("stale" in w for w in result["warnings"]),
              "8 lockfile: FILE tier proceeds past a stale ~$ with a warning")
        r = comtier.com_autofit(str(p_calc))
        check(r["ok"] and any("stale" in w for w in r["warnings"]),
              "8 lockfile: COM tier open+Save proceeds past a stale ~$ "
              "with a warning (SaveAs/Save unaffected by the stale record)")
        stale.unlink(missing_ok=True)

        # A REAL hold: a worker holds the book open; com ops must refuse.
        def hold_open(manager):
            w = manager.acquire()
            wbx = com_session.open_workbook(w.app, str(p_calc))
            return (w, wbx)

        w_hold, wb_hold = com_session.run_com("hold open", hold_open)
        try:
            try:
                comtier.com_autofit(str(p_calc))
                check(False, "8 lockfile: REAL hold refuses WORKBOOK_LOCKED")
            except WorkbookLocked:
                check(True, "8 lockfile: REAL hold refuses WORKBOOK_LOCKED")
        finally:
            def release(manager):
                try:
                    wb_hold.Close(False)
                finally:
                    manager.release_to_pool(w_hold)
            com_session.run_com("release hold", release)

        # ============ 9. volume + status + zero orphans ===================
        st = comtier.com_status()
        base_ops = st["ops_completed"]
        for _ in range(25):
            com_session.opens_clean(str(p_calc))
        st = comtier.com_status()
        check(st["ops_completed"] >= base_ops + 25 and st["ops_completed"] > 50,
              f"9 volume: {st['ops_completed']} serialized COM ops completed")
        check(st["busy"] is False and st["last_op"] is not None
              and st["last_op_seconds"] is not None,
              "9 status: honest idle report with real last-op timing")
        check(st["timeouts"] == 0, "9 status: zero timeouts across the gate")
        print(f"  status: pooled_pid={st['pooled_instance_pid']} "
              f"ops={st['ops_completed']} "
              f"contention_waits={st['contention_waits']}")

    except Exception as exc:  # noqa: BLE001
        import traceback
        traceback.print_exc()
        check(False, f"unhandled exception: {type(exc).__name__}: {exc}")
    finally:
        # PID-precise teardown: reclaim ONLY journal-owned PIDs.
        ex.shutdown()
        mgr = com_session.ExcelInstanceManager(
            journal_path=scratch / "_gate_journal.json")
        owned = sorted(mgr.journal.owned_pids())
        for pid in owned:
            if pid_alive(pid):
                taskkill(pid)
        deadline = time.monotonic() + 60
        while time.monotonic() < deadline and any(
                pid_alive(p) for p in owned):
            time.sleep(1.0)

    alive_mine = [p for p in owned if pid_alive(p)]
    foreign = sorted(list_excel_pids() - set(owned))
    check(not alive_mine,
          f"orphan check: no owned PID alive (owned={owned})")
    check(not foreign,
          f"orphan check: no foreign EXCEL.EXE touched/left ({foreign})")

    if FAILS:
        print(f"VERDICT com_tier_gate FAIL: {FAILS}")
        return 1
    print(f"VERDICT com_tier_gate PASS ({len(owned)} PIDs journaled, "
          "0 orphans by owned PID, 0 foreign touched)")
    return 0


def _read_a1(path) -> str:
    wb = openpyxl.load_workbook(path)
    try:
        return wb.active["A1"].value
    finally:
        wb.close()


if __name__ == "__main__":
    sys.exit(main())
