"""scripts/calc_spike.py: the formula / calc round-trip proof (SPIKE 3).

Proves the honest calc story end to end on real files, in ONE Excel session:
  A. openpyxl writes formulas with NO cache -> data_only reads None (absent).
  B. COM opens, calculates, saves -> openpyxl data_only reads the computed
     value. Measured on a multi-formula sheet with a cross-sheet reference and
     a volatile function (clean.xlsx from the corpus).
  C. Manual-calc-mode probe (com_ground_truth deferred item): with
     Application.Calculation = xlManual, does open-time calc still populate the
     cache, and does an explicit CalculateFull cover it?
  D. The _xlfn shim end to end: openpyxl writes XLOOKUP/FILTER with and without
     the _xlfn prefix; Excel calculates; do the prefixed ones avoid #NAME?

Plus a no-Excel probe of the `formulas` library coverage on the same sheet.

COM-gated; spawns ONE DispatchEx worker, journals its PID, reclaims owned PIDs
at the end (zero orphans by PID). Run only with the user's Excel CLOSED.
"""

from __future__ import annotations

import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
sys.path.insert(0, str(ROOT / "src"))
CORPUS = ROOT / "tests" / "fixtures" / "corpus"

from xlsx_mcp.com.instances import ExcelInstanceManager, list_excel_pids  # noqa: E402
from xlsx_mcp.core import calc  # noqa: E402

XL_MANUAL = -4135
XL_AUTOMATIC = -4105
CELL_ERROR_NAME = -2146826259  # xlErrName (#NAME?) as a Range.Value error code


def _openpyxl_data_only(path: Path, cell: str, sheet: str | None = None):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet else wb.active
    return ws[cell].value


def probe_formulas_library() -> dict:
    """No-Excel coverage probe of the `formulas` library on clean.xlsx."""
    out = {"available": False, "ok": False, "note": "", "computed": {}}
    try:
        import formulas  # type: ignore
    except Exception as exc:  # noqa: BLE001
        out["note"] = f"not installed: {exc}"
        return out
    out["available"] = True
    try:
        xl = formulas.ExcelModel().loads(str(CORPUS / "clean.xlsx")).finish()
        sol = xl.calculate()
        picked = {}
        for k, v in sol.items():
            ku = str(k).upper()
            if "A9" in ku or "B1" in ku or "B2" in ku:
                try:
                    picked[str(k)] = v.value[0, 0]
                except Exception:
                    picked[str(k)] = str(v)
        out["ok"] = True
        out["computed"] = picked
        out["note"] = "best-effort pure-Python recalc succeeded"
    except Exception as exc:  # noqa: BLE001
        out["note"] = f"{type(exc).__name__}: {exc}"
    return out


def _write_modern_fn_file(path: Path, formula: str, prefixed: bool) -> None:
    """openpyxl writes a lookup table + ONE modern-function formula in E1.
    When prefixed, the formula goes through the _xlfn shim first."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"], ws["B1"] = "key", "val"
    for i, (k, v) in enumerate([("x", 100), ("y", 200), ("z", 300)], start=2):
        ws[f"A{i}"], ws[f"B{i}"] = k, v
    ws["D1"] = "y"
    if prefixed:
        formula, _ = calc.normalize_formula(formula)
    ws["E1"] = formula
    wb.save(path)


def main() -> int:
    foreign = list_excel_pids()
    if foreign:
        print(f"SKIPPED: foreign EXCEL.EXE present {sorted(foreign)}")
        return 0

    scratch = ROOT / "tests" / "fixtures" / "corpus"
    tmpdir = Path(sys.argv[1]) if len(sys.argv) > 1 else \
        Path(__import__("tempfile").mkdtemp(prefix="ks4xl_calc_"))
    tmpdir.mkdir(parents=True, exist_ok=True)

    results: dict = {}

    # --- A: openpyxl formulas have no cache -------------------------------
    import shutil
    calc_file = tmpdir / "calc.xlsx"
    shutil.copy(CORPUS / "clean.xlsx", calc_file)
    a_before = _openpyxl_data_only(calc_file, "A9", "Data")
    b_cross_before = _openpyxl_data_only(calc_file, "B1", "Summary")
    results["A_absent_A9_before"] = a_before
    results["A_absent_crosssheet_before"] = b_cross_before

    mgr = ExcelInstanceManager(journal_path=tmpdir / "_calc_journal.json")
    w = mgr.acquire()  # becomes the pool; recalc_via_com reuses this instance
    print(f"pooled worker PID {w.pid}")
    try:
        # --- B: COM recalc round-trip (reuses the pooled instance) ---------
        rr = calc.recalc_via_com(str(calc_file), mgr, calculate=True, full=True)
        results["B_recalc_ok"] = rr.ok
        results["B_recalc_note"] = rr.note
        results["B_A9_after"] = _openpyxl_data_only(calc_file, "A9", "Data")
        results["B_crosssheet_after"] = _openpyxl_data_only(
            calc_file, "B1", "Summary")
        results["B_volatile_now_after"] = str(_openpyxl_data_only(
            calc_file, "B2", "Summary"))

        # --- C: manual-calc-mode probe ------------------------------------
        # Application.Calculation can only be set while a workbook is open, so
        # open a keeper first, switch to manual, THEN open the formula file to
        # observe open-time behavior under manual mode.
        manual_file = tmpdir / "manual.xlsx"
        shutil.copy(CORPUS / "clean.xlsx", manual_file)
        app = w.app
        keeper = app.Workbooks.Add()
        prior = app.Calculation
        app.Calculation = XL_MANUAL
        wb = app.Workbooks.Open(str(manual_file.resolve()))
        # Read A9 immediately after open under manual mode (before any calc):
        open_time_val = wb.Worksheets("Data").Range("A9").Value
        results["C_manual_opentime_A9"] = open_time_val
        app.CalculateFull()
        after_calc_val = wb.Worksheets("Data").Range("A9").Value
        results["C_manual_after_calcfull_A9"] = after_calc_val
        wb.Save()
        wb.Close(SaveChanges=False)
        try:
            app.Calculation = XL_AUTOMATIC
        except Exception:
            pass
        keeper.Close(SaveChanges=False)
        results["C_manual_cache_on_disk_A9"] = _openpyxl_data_only(
            manual_file, "A9", "Data")

        # --- D: the _xlfn shim end to end ---------------------------------
        # Each formula in its own file, opened independently, fault-isolated:
        # a bare modern function may make Excel demand a repair on open, which
        # throws under DisplayAlerts=False. That is itself a finding.
        for fn_name, formula in (("XLOOKUP", "=XLOOKUP(D1,A2:A4,B2:B4)"),
                                 ("FILTER", "=FILTER(B2:B4,B2:B4>150)")):
            for prefixed in (False, True):
                tag = "prefixed" if prefixed else "bare"
                key = f"D_{fn_name}_{tag}"
                xf = tmpdir / f"{fn_name}_{tag}.xlsx"
                _write_modern_fn_file(xf, formula, prefixed)
                try:
                    wb = app.Workbooks.Open(str(xf.resolve()))
                    app.CalculateFull()
                    e1 = wb.Worksheets(1).Range("E1").Value
                    wb.Close(SaveChanges=False)
                    results[key] = repr(e1)
                except Exception as exc:  # noqa: BLE001
                    results[key] = f"OPEN/CALC FAILED: {type(exc).__name__}"

        mgr.quit(w)
    finally:
        res = mgr.force_reclaim()
        print(f"reclaim: killed={res.killed} exited={res.exited_on_own} "
              f"checked={res.checked}")

    # --- formulas library probe (no Excel) --------------------------------
    results["formulas_probe"] = probe_formulas_library()

    print("\n=== CALC SPIKE RESULTS ===")
    for k, v in results.items():
        print(f"{k}: {v}")
    leftover = list_excel_pids()
    print(f"\nEXCEL.EXE after run: {sorted(leftover) if leftover else 'NONE'}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
